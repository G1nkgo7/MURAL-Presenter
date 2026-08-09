import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import attachments, engine


class _Upload:
    def __init__(self, name: str, data: bytes):
        self.filename = name
        self.file = io.BytesIO(data)


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 6), (80, 120, 200)).save(buf, format="PNG")
    return buf.getvalue()


class AttachmentModeTests(unittest.TestCase):
    def test_default_attachment_mode_is_web_parse(self):
        with patch.dict(os.environ, {}, clear=True):
            self.assertEqual(attachments.normalize_mode(""), "web_parse")

    def test_env_can_switch_attachment_mode_to_pipeline_agent(self):
        with patch.dict(os.environ, {"STUDIO_ATTACHMENT_MODE": "pipeline_agent"}, clear=True):
            self.assertEqual(attachments.normalize_mode(""), "pipeline_agent")

    def test_store_raw_attachments_does_not_extract_text(self):
        old_uploads = engine.UPLOADS_DIR
        with tempfile.TemporaryDirectory() as tmp:
            engine.UPLOADS_DIR = Path(tmp)
            try:
                records = attachments.store_attachments(
                    7,
                    [_Upload("notes.txt", b"secret text that web should not parse")],
                    parse=False,
                )
            finally:
                engine.UPLOADS_DIR = old_uploads

        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["name"], "notes.txt")
        self.assertEqual(records[0]["text"], "")
        self.assertEqual(records[0]["images"], [])
        raw = attachments.raw_attachments(records)
        self.assertEqual(len(raw), 1)
        self.assertEqual(raw[0]["workspace_rel"], "attachments/raw/notes.txt")

    def test_pipeline_agent_prompt_mentions_manifest_without_file_text(self):
        records = [{
            "name": "notes.txt",
            "stored_name": "notes.txt",
            "path": "/tmp/notes.txt",
            "size": 11,
            "type": "txt",
            "text": "this must not appear",
            "images": [],
            "notes": [],
        }]
        block = attachments.pipeline_agent_query_block(records)
        self.assertIn("attachments/manifest.json", block)
        self.assertIn("attachment_reader", block)
        self.assertIn("notes.txt", block)
        self.assertNotIn("this must not appear", block)

    def test_docx_embedded_images_are_extracted_for_model_input(self):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            img = Path(tmp) / "inline.png"
            img.write_bytes(_png_bytes())
            doc = Document()
            doc.add_paragraph("Doc heading")
            doc.add_picture(str(img))
            buf = io.BytesIO()
            doc.save(buf)

            old_uploads = engine.UPLOADS_DIR
            engine.UPLOADS_DIR = Path(tmp) / "uploads"
            try:
                records = attachments.store_attachments(
                    8,
                    [_Upload("brief.docx", buf.getvalue())],
                    parse=True,
                )
            finally:
                engine.UPLOADS_DIR = old_uploads

        self.assertIn("Doc heading", records[0]["text"])
        images = attachments.attachment_images(records)
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["kind"], "docx_image")
        self.assertIn("docx_image", images[0]["workspace_rel"])

    def test_pptx_embedded_images_are_extracted_for_model_input(self):
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as tmp:
            img = Path(tmp) / "slide.png"
            img.write_bytes(_png_bytes())
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Slide title"
            slide.shapes.add_picture(str(img), Inches(1), Inches(2), Inches(2), Inches(1.5))
            buf = io.BytesIO()
            prs.save(buf)

            old_uploads = engine.UPLOADS_DIR
            engine.UPLOADS_DIR = Path(tmp) / "uploads"
            try:
                records = attachments.store_attachments(
                    9,
                    [_Upload("deck.pptx", buf.getvalue())],
                    parse=True,
                )
            finally:
                engine.UPLOADS_DIR = old_uploads

        self.assertIn("Slide title", records[0]["text"])
        images = attachments.attachment_images(records)
        self.assertGreaterEqual(len(images), 1)
        self.assertIn("pptx_image", {img["kind"] for img in images})

    def test_xlsx_embedded_images_are_extracted_for_model_input(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as SheetImage

        with tempfile.TemporaryDirectory() as tmp:
            img = Path(tmp) / "sheet.png"
            img.write_bytes(_png_bytes())
            wb = Workbook()
            ws = wb.active
            ws.title = "Summary"
            ws["A1"] = "Metric"
            ws["B1"] = "Value"
            ws.add_image(SheetImage(str(img)), "C3")
            buf = io.BytesIO()
            wb.save(buf)

            old_uploads = engine.UPLOADS_DIR
            engine.UPLOADS_DIR = Path(tmp) / "uploads"
            try:
                records = attachments.store_attachments(
                    10,
                    [_Upload("sheet.xlsx", buf.getvalue())],
                    parse=True,
                )
            finally:
                engine.UPLOADS_DIR = old_uploads

        self.assertIn("[Sheet Summary]", records[0]["text"])
        images = attachments.attachment_images(records)
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["kind"], "xlsx_image")


if __name__ == "__main__":
    unittest.main()
