#!/usr/bin/env python3
"""Build the compact two-sheet XLSX for the intermediate-artifact rubric."""

from __future__ import annotations

import tempfile
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RUBRIC_PATH = ROOT / "rubric" / "long_horizon_agentic_intermediate_rubric_v2.yaml"
OUTPUT_PATH = ROOT / "rubric" / "long_horizon_agentic_intermediate_rubric_v2.xlsx"

NAVY = "17365D"
BLUE = "4472C4"
PALE_BLUE = "EAF2F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRAY = "666666"
GRID = "B4C6E7"
THIN = Side(style="thin", color=GRID)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def join(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return "\n".join(f"{key}: {join(item)}" for key, item in value.items())
    if isinstance(value, list):
        return "\n".join(f"• {join(item)}" for item in value)
    return str(value)


def set_title(ws, title: str, subtitle: str, width: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    ws.cell(1, 1).fill = fill(NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(name="Aptos", size=10, italic=True, color=GRAY)
    ws.cell(2, 1).fill = fill(PALE_BLUE)
    ws.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28


def style_header(row) -> None:
    for cell in row:
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = fill(BLUE)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_body(ws, start_row: int, end_row: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Aptos", size=9.5)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def configure(ws, widths: list[float], freeze: str, filter_ref: str) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = freeze
    ws.auto_filter.ref = filter_ref
    ws.sheet_view.showGridLines = False


def flatten_points(rubric: dict) -> list[dict]:
    return [
        {"dimension_id": dim["id"], "dimension_name": dim["name"], "dimension_weight": dim["weight"], **point}
        for dim in rubric["dimensions"]
        for point in dim["scoring_points"]
    ]


def dimension_paths(dim: dict) -> tuple[list[str], list[str]]:
    required: list[str] = []
    conditional: list[str] = []
    for point in dim["scoring_points"]:
        for path in point.get("required_intermediate_artifacts", []):
            if path not in required:
                required.append(path)
        for path in point.get("conditional_intermediate_artifacts", []):
            if path not in conditional:
                conditional.append(path)
    return required, conditional


def build_dimensions(ws, rubric: dict) -> None:
    set_title(ws, "评分维度", "保留完整评分锚点与硬边界；路径均相对于每个run目录。", 11)
    headers = ["维度ID", "能力维度", "权重", "适用性", "定义", "评分点数", "代码/Judge", "所需实际产物路径", "Judge关注点", "评分锚点", "硬边界"]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    style_header(ws[header_row])
    conditional_dims = set(rubric["applicability"]["conditional_dimensions"])
    for dim in rubric["dimensions"]:
        required, conditional = dimension_paths(dim)
        paths = list(required)
        paths.extend(f"{path}（条件）" for path in conditional)
        det = sum(point["evaluator"] == "deterministic" for point in dim["scoring_points"])
        judge_binary = sum(
            point["evaluator"] == "judge" and point["scale"] == "binary"
            for point in dim["scoring_points"]
        )
        judge_graded = len(dim["scoring_points"]) - det - judge_binary
        ws.append([
            dim["id"], dim["name"], dim["weight"], "条件适用" if dim["id"] in conditional_dims else "默认适用",
            dim["definition"], len(dim["scoring_points"]),
            f"代码0/1：{det}\nJudge 0/1：{judge_binary}\nJudge分档：{judge_graded}", join(paths),
            join(dim["judge_focus"]), join(dim["anchors"]), join(dim["hard_boundaries"]),
        ])
    style_body(ws, header_row + 1, ws.max_row, 11)
    for row in range(header_row + 1, ws.max_row + 1):
        ws.row_dimensions[row].height = 210
    configure(ws, [9, 34, 9, 12, 58, 10, 18, 58, 52, 66, 48], f"A{header_row+1}", f"A{header_row}:K{ws.max_row}")


def build_points(ws, rubric: dict, points: list[dict]) -> None:
    set_title(ws, "评分细则", "12个通用评分点；保留评分点适用性、允许分数与各档位含义。", 14)
    headers = ["维度", "维度名称", "评分点ID", "评测方式", "评分点适用性", "点权重", "档位", "允许分数", "档位含义", "评什么/通过条件", "内部观察项", "所需实际产物路径", "条件实际产物路径", "外部参照（非中间产物）"]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    style_header(ws[header_row])
    profiles = rubric["score_scale"]["profiles"]
    for point in points:
        criterion = point.get("pass_if") or point.get("evaluate") or ""
        scale_spec = profiles[point["scale"]]
        ws.append([
            point["dimension_id"], point["dimension_name"], point["id"],
            (
                "代码0/1"
                if point["evaluator"] == "deterministic"
                else "LLM Judge 0/1"
                if point["scale"] == "binary"
                else "LLM Judge 分档"
            ),
            point.get("applicability", "随所属维度适用"), point["point_weight"], point["scale"],
            " / ".join(str(value) for value in scale_spec["allowed_scores"]), join(scale_spec["semantics"]),
            criterion, join(point.get("judge_subcriteria")), join(point["required_intermediate_artifacts"]),
            join(point.get("conditional_intermediate_artifacts")), join(point["required_reference_context"]),
        ])
    style_body(ws, header_row + 1, ws.max_row, 14)
    for row in range(header_row + 1, ws.max_row + 1):
        ws.cell(row, 4).fill = fill(GREEN if ws.cell(row, 4).value == "代码0/1" else YELLOW)
        ws.row_dimensions[row].height = 175
    configure(ws, [8, 30, 44, 13, 34, 10, 13, 24, 54, 62, 48, 58, 38, 48], f"A{header_row+1}", f"A{header_row}:N{ws.max_row}")


def verify(path: Path, rubric: dict, points: list[dict]) -> None:
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["评分维度", "评分细则"]
    assert len(rubric["dimensions"]) == 6
    assert len(points) == 12
    assert wb["评分维度"].max_row == 10
    assert wb["评分细则"].max_row == 16
    assert all(point.get("required_intermediate_artifacts") for point in points)
    assert all(sum(point["point_weight"] for point in dim["scoring_points"]) == 100 for dim in rubric["dimensions"])


def main() -> None:
    rubric = yaml.safe_load(RUBRIC_PATH.read_text(encoding="utf-8"))
    points = flatten_points(rubric)
    wb = Workbook()
    wb.remove(wb.active)
    build_dimensions(wb.create_sheet("评分维度"), rubric)
    build_points(wb.create_sheet("评分细则"), rubric, points)
    wb.properties.title = rubric["title"]
    wb.properties.subject = "PPT intermediate artifact and agentic capability rubric"
    wb.properties.creator = "NovaPresent-bench"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=OUTPUT_PATH.parent, delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        wb.save(temp_path)
        verify(temp_path, rubric, points)
        temp_path.replace(OUTPUT_PATH)
        verify(OUTPUT_PATH, rubric, points)
    finally:
        temp_path.unlink(missing_ok=True)
    print(f"wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
