#!/usr/bin/env python3
"""Build the v2 rubric XLSX and Markdown from the canonical YAML.

Requires openpyxl >= 3.1.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parents[2]
YAML_PATH = ROOT / "rubric" / "multi_page_image_rubric_v2.yaml"
XLSX_PATH = ROOT / "rubric" / "multi_page_image_rubric_v2.xlsx"
MARKDOWN_PATH = ROOT / "rubric" / "multi_page_image_rubric_v2.md"

NAVY = "17365D"
BLUE = "4472C4"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF2F8"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
WHITE = "FFFFFF"
GRID = "B4C6E7"
GRAY = "666666"
RED = "F4CCCC"
ORANGE = "FCE5CD"
DARK_GREEN = "B6D7A8"


def solid(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


THIN_BORDER = Border(
    left=Side(style="thin", color=GRID),
    right=Side(style="thin", color=GRID),
    top=Side(style="thin", color=GRID),
    bottom=Side(style="thin", color=GRID),
)


def style_title(cell) -> None:
    cell.font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
    cell.fill = solid(NAVY)
    cell.alignment = Alignment(vertical="center")


def style_header(cell) -> None:
    cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
    cell.fill = solid(BLUE)
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_section(cell) -> None:
    cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)
    cell.fill = solid(LIGHT_BLUE)
    cell.border = THIN_BORDER
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def style_body(cell, *, center: bool = False, muted: bool = False) -> None:
    cell.font = Font(name="Aptos", size=9 if muted else 10, color=GRAY if muted else "000000")
    cell.border = THIN_BORDER
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center" if center else "top",
        wrap_text=True,
    )


def style_input(cell) -> None:
    style_body(cell, center=True)
    cell.fill = solid(YELLOW)


def style_formula(cell) -> None:
    style_body(cell, center=True)
    cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)
    cell.fill = solid(GREEN)
    cell.number_format = "0.0"


def set_widths(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def build_detail_sheet(wb: Workbook, rubric: dict) -> None:
    ws = wb.create_sheet("v2评分细则")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.merge_cells("A1:M1")
    ws["A1"] = rubric["title"]
    style_title(ws["A1"])
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:M2")
    ws["A2"] = "仅评内容与美学；5 个 Deck-level 维度采用严格的 0–5 六档整数分。0=未实现/根本失效，3=最低完整合格，5=标杆。"
    ws["A2"].font = Font(name="Aptos", size=9, color=GRAY)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    summary = [
        ("评分范围", "Deck-level 全页巡检"),
        ("维度数", 5),
        ("内容/美学", "3 / 2"),
        ("维度分", "0–5 整数"),
        ("最低合格", 3),
        ("最终分", "五项平均"),
    ]
    col = 1
    for label, value in summary:
        style_section(ws.cell(3, col, label))
        style_body(ws.cell(3, col + 1, value), center=True)
        col += 2
    ws.merge_cells("L3:M3")
    ws["L3"] = "最终分保留一位小数"
    style_body(ws["L3"], center=True)
    ws.row_dimensions[3].height = 28

    ws.merge_cells("A4:M4")
    ws["A4"] = "严格原则：先检查 0 分触发条件；选择全部条件均满足的最高档。核心要求未完整实现最高 2 分；存在实质问题最高 3 分；5 分必须全页核查且无实质缺陷。"
    style_body(ws["A4"], muted=True)
    ws["A4"].fill = solid(PALE_BLUE)
    ws.row_dimensions[4].height = 42

    headers = [
        "模块", "编号", "评分维度", "评什么 / 定义", "主要证据",
        "0｜未实现/根本失效", "1｜严重不足", "2｜部分实现",
        "3｜基本合格", "4｜良好", "5｜标杆", "硬门槛 / 边界", "关键性",
    ]
    for col, value in enumerate(headers, 1):
        style_header(ws.cell(6, col, value))
    ws.row_dimensions[6].height = 40

    module_names = {"content": "Content", "aesthetic": "Aesthetic"}
    for row, dim in enumerate(rubric["dimensions"], 7):
        values = [
            module_names[dim["module"]],
            dim["id"].split("_", 1)[0],
            dim["name"],
            dim["definition"],
            "\n".join(f"• {item}" for item in dim["evidence_sources"]),
            *[dim["anchors"][str(score)] for score in range(6)],
            "\n".join(f"• {item}" for item in dim["hard_boundaries"]),
            "关键" if dim["criticality"] == "critical" else "标准",
        ]
        for col, value in enumerate(values, 1):
            c = ws.cell(row, col, value)
            if col in {1, 3}:
                style_section(c)
            elif col in {2, 13}:
                style_body(c, center=True)
            else:
                style_body(c)
        ws.row_dimensions[row].height = 210

    ws.merge_cells("A13:M13")
    ws["A13"] = "统一分值语义"
    style_section(ws["A13"])
    for row, score in enumerate(range(6), 14):
        ws.cell(row, 1, score)
        style_formula(ws.cell(row, 1))
        ws.cell(row, 1).number_format = "0"
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=13)
        ws.cell(row, 2, rubric["dimension_scale"]["score_semantics"][str(score)])
        style_body(ws.cell(row, 2))
        ws.row_dimensions[row].height = 34

    ws.merge_cells("A21:M21")
    ws["A21"] = "最终平均分解释区间（仅用于报告，不重算分数）"
    style_section(ws["A21"])
    for row, band in enumerate(rubric["scoring"]["reporting_bands"], 22):
        ws.cell(row, 1, band["range"])
        style_formula(ws.cell(row, 1))
        ws.cell(row, 1).number_format = "General"
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=13)
        ws.cell(row, 2, band["label"])
        style_body(ws.cell(row, 2))
        ws.row_dimensions[row].height = 27

    ws.auto_filter.ref = "A6:M11"
    set_widths(ws, [11, 9, 24, 36, 28, 44, 44, 48, 50, 52, 56, 46, 10])
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"


def build_template_sheet(wb: Workbook, rubric: dict) -> None:
    ws = wb.create_sheet("v2评分模板")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A7"
    ws.merge_cells("A1:I1")
    ws["A1"] = "多页 Deck v2 评分模板｜0–5 六档严格评分"
    style_title(ws["A1"])
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:I2")
    ws["A2"] = "黄色分数格只能选择 0–5 整数。0=未实现/根本失效，3=最低完整合格，4=无重大问题，5=全页核查后仍无实质缺陷。低于 5 必须填写阻止进入下一档的证据。"
    ws["A2"].font = Font(name="Aptos", size=9, color=GRAY)
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 38

    metadata = ["样本 / Case ID", "", "Deck 页数", "", "评审者", "", "评审日期", "", ""]
    for col, value in enumerate(metadata, 1):
        c = ws.cell(4, col, value)
        if col in {1, 3, 5, 7}:
            style_section(c)
        else:
            style_body(c)
    ws.row_dimensions[4].height = 25

    headers = [
        "模块", "编号", "评分维度", "分数", "证据定位（页码/章节/相邻页对）",
        "评分理由", "阻止进入下一档的缺陷 / 5分依据", "已核查范围", "关键性",
    ]
    for col, value in enumerate(headers, 1):
        style_header(ws.cell(6, col, value))
    ws.row_dimensions[6].height = 40

    module_names = {"content": "Content", "aesthetic": "Aesthetic"}
    for row, dim in enumerate(rubric["dimensions"], 7):
        values = [
            module_names[dim["module"]],
            dim["id"].split("_", 1)[0],
            dim["name"],
            None, None, None, None, None,
            "关键" if dim["criticality"] == "critical" else "标准",
        ]
        for col, value in enumerate(values, 1):
            c = ws.cell(row, col, value)
            if col in {1, 3}:
                style_section(c)
            elif col in {2, 9}:
                style_body(c, center=True)
            elif col == 4:
                style_input(c)
                c.number_format = "0"
            else:
                style_body(c)
        ws.row_dimensions[row].height = 68

    summaries = [
        (13, "Content Mean", "=IF(COUNT(D7:D9)<3,\"\",ROUND(AVERAGE(D7:D9),1))"),
        (14, "Aesthetic Mean", "=IF(COUNT(D10:D11)<2,\"\",ROUND(AVERAGE(D10:D11),1))"),
        (15, "Final Score（五维等权平均）", "=IF(COUNT(D7:D11)<5,\"\",ROUND(AVERAGE(D7:D11),1))"),
        (16, "Critical Failure", '=IF(COUNT(D7:D11)<5,"",IF(OR(D7=0,D9=0,D10=0),"true","false"))'),
        (17, "评测状态", None),
    ]
    for row, label, value in summaries:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1, label)
        style_section(ws.cell(row, 1))
        ws.cell(row, 4, value)
        if row == 17:
            style_input(ws.cell(row, 4))
            ws.cell(row, 4).number_format = "General"
        else:
            style_formula(ws.cell(row, 4))
            if row == 16:
                ws.cell(row, 4).number_format = "General"
        ws.row_dimensions[row].height = 28

    ws["A19"] = "最终分公式"
    style_section(ws["A19"])
    ws.merge_cells("B19:I19")
    ws["B19"] = "Final Score = (C1 + C2 + C3 + A1 + A2) / 5，保留一位小数；不转换百分制，不做第二次离散化。"
    style_body(ws["B19"], muted=True)
    ws.row_dimensions[19].height = 34
    ws["A20"] = "0 与 invalid"
    style_section(ws["A20"])
    ws.merge_cells("B20:I20")
    ws["B20"] = "Deck 可观察但某维度未实现/根本失效，该维度记 0；输入本身不足以评价整套 Deck 时状态选 invalid，不填写分数。"
    style_body(ws["B20"], muted=True)
    ws.row_dimensions[20].height = 38

    score_validation = DataValidation(type="list", formula1="=_score_values", allow_blank=False)
    score_validation.errorTitle = "无效分数"
    score_validation.error = "只能选择 0、1、2、3、4 或 5"
    score_validation.errorStyle = "stop"
    score_validation.showErrorMessage = True
    score_validation.promptTitle = "维度分"
    score_validation.prompt = "请选择严格满足的最高整数档"
    score_validation.showInputMessage = True
    ws.add_data_validation(score_validation)
    score_validation.add("D7:D11")

    status_validation = DataValidation(type="list", formula1="=_status_values", allow_blank=False)
    status_validation.errorTitle = "无效状态"
    status_validation.error = "只能选择 valid 或 invalid"
    status_validation.errorStyle = "stop"
    status_validation.showErrorMessage = True
    ws.add_data_validation(status_validation)
    status_validation.add("D17")

    score_cells = "D7:D11"
    ws.conditional_formatting.add(score_cells, CellIsRule(operator="equal", formula=["0"], fill=solid(RED)))
    ws.conditional_formatting.add(score_cells, CellIsRule(operator="between", formula=["1", "2"], fill=solid(ORANGE)))
    ws.conditional_formatting.add(score_cells, CellIsRule(operator="equal", formula=["3"], fill=solid(YELLOW)))
    ws.conditional_formatting.add(score_cells, CellIsRule(operator="equal", formula=["4"], fill=solid(GREEN)))
    ws.conditional_formatting.add(score_cells, CellIsRule(operator="equal", formula=["5"], fill=solid(DARK_GREEN)))

    set_widths(ws, [13, 10, 28, 13, 42, 48, 48, 34, 11])
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:6"


def build_lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("_lists")
    for row, value in enumerate(range(6), 1):
        ws.cell(row, 1, value)
    for row, value in enumerate(("valid", "invalid"), 1):
        ws.cell(row, 2, value)
    ws.sheet_state = "hidden"
    wb.defined_names.add(DefinedName("_score_values", attr_text="'_lists'!$A$1:$A$6"))
    wb.defined_names.add(DefinedName("_status_values", attr_text="'_lists'!$B$1:$B$2"))


def verify_workbook(path: Path, rubric: dict) -> None:
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["v2评分细则", "v2评分模板", "_lists"]
    assert wb["_lists"].sheet_state == "hidden"
    detail = wb["v2评分细则"]
    template = wb["v2评分模板"]
    assert [detail.cell(row, 3).value for row in range(7, 12)] == [d["name"] for d in rubric["dimensions"]]
    assert [template.cell(row, 2).value for row in range(7, 12)] == ["C1", "C2", "C3", "A1", "A2"]
    assert template["D13"].value == '=IF(COUNT(D7:D9)<3,"",ROUND(AVERAGE(D7:D9),1))'
    assert template["D15"].value == '=IF(COUNT(D7:D11)<5,"",ROUND(AVERAGE(D7:D11),1))'
    assert "OR(D7=0,D9=0,D10=0)" in template["D16"].value
    assert len(template.data_validations.dataValidation) == 2
    assert [wb["_lists"].cell(row, 1).value for row in range(1, 7)] == list(range(6))


def markdown_table_row(values: list[object]) -> str:
    return "| " + " | ".join(str(value).replace("|", "\\|").replace("\n", "<br>") for value in values) + " |"


def build_markdown(rubric: dict) -> str:
    lines = [
        f"# {rubric['title']}",
        "",
        "## 1. 设计结论",
        "",
        "本 Rubric 保留 5 个压缩后的 Deck-level 维度，采用严格的 `0–5` 六档整数分。六档并不过多：`0` 专门表示未实现或根本失效，`1–5` 分别表示严重不足、部分实现、基本合格、良好和标杆。维度禁止小数，最终分为五个维度的算术平均值。",
        "",
        "- `3` 是完整可用的最低合格档，不是高分。",
        "- `4` 要求不存在重大问题，不能作为默认分。",
        "- `5` 必须全页核查、证据充分且没有实质缺陷。",
        "- 没有实现、核心目标完全缺失或根本失效时直接为 `0`。",
        "",
        "## 2. 适用范围与有效性",
        "",
        "只评价任务上下文与全 Deck 渲染图中可观察的内容和美学质量，不评价 HTML、DOM、CSS、离线运行、加载、控制台、导航或隐藏状态等工程属性。",
        "",
        "有完整可观察的 Deck，但某个维度未实现或根本失效时，该维度记 `0`；只有输入本身不足以评价整套 Deck 时才记 `invalid`，不填写维度分和最终分。",
        "",
        "## 3. 统一评分尺度",
        "",
        "| 分数 | 等级 | 严格定义 |",
        "|---:|---|---|",
    ]
    labels = ["未实现/根本失效", "严重不足", "部分实现", "基本合格", "良好", "标杆"]
    for score, label in enumerate(labels):
        lines.append(markdown_table_row([f"`{score}`", label, rubric["dimension_scale"]["score_semantics"][str(score)]]))
    lines.extend(["", "### 评分纪律", ""])
    lines.extend(f"- {rule}" for rule in rubric["strict_scoring_rules"])
    lines.extend([
        "",
        "## 4. 五个评分维度",
        "",
        "| 编号 | 模块 | 维度 | 关键性 | 核心问题 |",
        "|---|---|---|---|---|",
    ])
    module_names = {"content": "Content", "aesthetic": "Aesthetic"}
    for dim in rubric["dimensions"]:
        lines.append(markdown_table_row([
            dim["id"].split("_", 1)[0], module_names[dim["module"]], dim["name"],
            "关键" if dim["criticality"] == "critical" else "标准", dim["definition"],
        ]))
    for dim in rubric["dimensions"]:
        short_id = dim["id"].split("_", 1)[0]
        lines.extend([
            "",
            f"### {short_id} {dim['name']}",
            "",
            f"**定义：** {dim['definition']}",
            "",
            f"**评价方式：** {dim['evaluation']}",
            "",
            "**主要证据：** " + "、".join(f"`{item}`" for item in dim["evidence_sources"]),
            "",
            "| 分数 | 严格判定标准 |",
            "|---:|---|",
        ])
        for score in range(6):
            lines.append(markdown_table_row([f"`{score}`", dim["anchors"][str(score)]]))
        lines.extend(["", "**硬门槛：**", ""])
        lines.extend(f"- {item}" for item in dim["hard_boundaries"])
    lines.extend([
        "",
        "## 5. 分数计算",
        "",
        "```text",
        "Content Mean   = (C1 + C2 + C3) / 3",
        "Aesthetic Mean = (A1 + A2) / 2",
        "Final Score    = (C1 + C2 + C3 + A1 + A2) / 5",
        "```",
        "",
        "`Final Score` 范围为 0–5，保留一位小数。由于五个维度都是整数，最终分天然以 `0.2` 为步长；不转换为百分制，也不做第二次离散化。",
        "",
        "C1、C3 或 A1 任一为 `0` 时，另外输出 `critical_failure: true` 和触发维度，但不改变算术平均分。",
        "",
        "| 最终平均分 | 报告标签 |",
        "|---:|---|",
    ])
    for band in rubric["scoring"]["reporting_bands"]:
        lines.append(markdown_table_row([band["range"], band["label"]]))
    lines.extend([
        "",
        "## 6. 输出要求",
        "",
    ])
    lines.extend(f"- {item}" for item in rubric["judge_output_requirements"]["final_output"])
    lines.extend(["", "### 证据要求", ""])
    lines.extend(f"- {item}" for item in rubric["judge_output_requirements"]["evidence_minimum"])
    lines.append("")
    return "\n".join(lines)


def build() -> None:
    rubric = yaml.safe_load(YAML_PATH.read_text(encoding="utf-8"))
    assert len(rubric["dimensions"]) == 5
    assert rubric["dimension_scale"]["allowed_scores"] == list(range(6))
    assert all(set(dim["anchors"]) == {str(score) for score in range(6)} for dim in rubric["dimensions"])

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = rubric["title"]
    wb.properties.creator = "NovaPresent-bench"
    wb.properties.description = "Deck-level content and aesthetic rubric with strict integer scores from 0 to 5."
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    build_detail_sheet(wb, rubric)
    build_template_sheet(wb, rubric)
    build_lists_sheet(wb)

    with tempfile.NamedTemporaryFile(dir=XLSX_PATH.parent, suffix=".xlsx", delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        wb.save(temp_path)
        verify_workbook(temp_path, rubric)
        reopened = load_workbook(temp_path, data_only=False)
        reopened.save(XLSX_PATH)
        verify_workbook(XLSX_PATH, rubric)
    finally:
        temp_path.unlink(missing_ok=True)

    MARKDOWN_PATH.write_text(build_markdown(rubric), encoding="utf-8")
    os.chmod(XLSX_PATH, 0o664)
    os.chmod(MARKDOWN_PATH, 0o664)


if __name__ == "__main__":
    build()
