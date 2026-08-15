from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pymupdf
from unittest import mock

HARNESS = Path(__file__).resolve().parents[1]
REPO = HARNESS.parents[1]
sys.path.insert(0, str(HARNESS))

from core import tools  # noqa: E402
from core import model_call  # noqa: E402
from core.agent_loop import (  # noqa: E402
    Agent,
    _accept,
    _contract_fields,
    _delegate_task,
    _finish_gap,
    _review_can_complete_needs_improvement,
    _review_required_view_gap,
    _revision_delegation_error,
    _revision_prompt,
    _slide_repair_issue,
)
from core.run_batch import _extract_material, _material_blocked, stage_materials  # noqa: E402
from core.runtime_capabilities import detect_runtime_capabilities  # noqa: E402


def test_runtime_capabilities_omit_optional_stages_without_inputs_or_keys(
    tmp_path: Path,
) -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {"model": "test", "model_base_url": "http://model", "enable_image_gen": True},
            [],
            plan_only=True,
        )
    assert profile["active_roles"] == ["slide", "review"]
    assert profile["omitted_roles"] == ["material", "research", "image"]
    assert profile["workflow"]["image"] == "code_only"
    assert profile["tools"]["web_search"] is False
    assert profile["tools"]["image_generate"] is False
    assert "api_key" not in json.dumps(profile).lower()


def test_runtime_capabilities_route_material_directly_without_search() -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "image-secret", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {"model": "test", "model_base_url": "http://model", "enable_image_gen": True},
            ["inputs/report.pdf.md"],
            plan_only=True,
        )
    assert profile["active_roles"] == ["material", "image", "slide", "review"]
    assert profile["workflow"]["material_handoff"] == "orchestrator"
    assert profile["workflow"]["research"] == "omitted"


def test_runtime_capabilities_inject_only_available_roles_and_tools(tmp_path: Path) -> None:
    profile = {
        "active_roles": ["slide", "review"],
        "omitted_roles": ["material", "research", "image"],
        "workflow": {
            "material": "omitted",
            "research": "omitted",
            "material_handoff": "none",
            "image": "code_only",
        },
        "tools": {
            "web_search": False,
            "web_extract": False,
            "image_generate": False,
            "vision_analyze": True,
        },
    }
    agent = Agent(
        "runtime-capability-test",
        str(tmp_path),
        "制作一套演示",
        {
            "skill_name": "mural-presenter-v0-2-zh",
            "skill_language": "zh",
            "_runtime_capabilities": profile,
        },
    )
    tool_names = {schema["name"] for schema in agent.tool_schemas}
    assert agent.research_required is False
    assert "Research：omitted" in agent.system
    assert "可委派角色：slide, review" in agent.system
    assert "web_search" not in tool_names
    assert "image_generate" not in tool_names
    rejected = _delegate_task(
        agent,
        {"tasks": [{"goal": "Research: collect sources"}]},
    )
    assert "运行时能力检查已省略" in rejected


def test_agent_tools_remove_search_generation_and_vision() -> None:
    names = {
        schema["name"]
        for schema in tools.agent_tools(
            "image",
            enable_web_search=False,
            enable_image_gen=False,
            enable_vision=False,
        )
    }
    assert "web_search" not in names
    assert "web_extract" not in names
    assert "image_generate" not in names
    assert "vision_analyze" not in names


def test_slide_group_is_rejected() -> None:
    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-2-zh"
        response_language = "zh"

    result = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "SlideGroup hero [01,02]: build both pages"}]},
    )
    assert "只接受一页一个" in result


def test_grouped_variant_accepts_only_complete_slide_groups() -> None:
    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-2-grouped-zh"
        response_language = "zh"
        query_language_hint = "zh"

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(
            Parent(),
            {"tasks": [{"goal": "SlideGroup mechanism [03,04,05]: build group"}]},
        )
    assert result == "delegated"
    specs = delegate.call_args.args[1]["tasks"]
    assert specs == [{
        "role": "slide",
        "label": "slide_group_mechanism",
        "group_id": "mechanism",
        "pages": [3, 4, 5],
        "task": "SlideGroup mechanism [03,04,05]: build group",
    }]

    rejected = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "Slide 03: build one page"}]},
    )
    assert "不接受 `Slide NN:`" in rejected


def test_grouped_v02_keeps_research_and_revision_contracts(tmp_path: Path) -> None:
    agent = Agent(
        "grouped-contract-test",
        str(tmp_path),
        "制作八页演示",
        {
            "skill_name": "mural-presenter-v0-2-grouped-zh",
            "skill_language": "zh",
        },
    )
    assert agent.research_required is True

    prompt = _revision_prompt(
        {"query": "介绍台风白海豚", "lang": "zh"},
        {"instruction": "把主题纠正为台风白海豚"},
        grouped=True,
    )
    assert "Review: mode=simple_edit" in prompt
    assert "plan/revision-impact.md" in prompt
    assert "SlideGroup GROUP [NN,NN]" in prompt
    assert "不得把受影响页组拆成单页" in prompt

    class RevisionOrchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-2-grouped-zh"
        revision_mode = True

    error = tools._model_write_error(
        RevisionOrchestrator(), "slides/slide_02.html"
    )
    assert error is not None
    assert "不允许 Orchestrator 直接修改" in error


def test_image_and_slide_must_not_share_a_synchronous_wave() -> None:
    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-2-zh"
        response_language = "zh"
        child_outcomes = {}

    result = _delegate_task(
        Parent(),
        {
            "tasks": [
                {"goal": "Image: resolve all raster assets"},
                {"goal": "Slide 03: complete the code-only page"},
            ]
        },
    )
    assert "不把 Image 与 Slide 放在同一批" in result


def test_one_shot_shared_env_is_explicit_same_model_native() -> None:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(HARNESS)
    env["VISION_BACKEND"] = "one_shot"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "from core import config; "
            "print(config.VISION_BACKEND_REQUESTED, config.VISION_BACKEND)",
        ],
        capture_output=True,
        text=True,
        env=env,
        check=True,
    )
    assert result.stdout.strip() == "one_shot native"


def test_deepseek_no_thinking_is_explicitly_disabled() -> None:
    captured = {}

    class Response:
        usage = None
        content = []

    class Messages:
        def create(self, **kwargs):
            captured.update(kwargs)
            return Response()

    class Client:
        messages = Messages()

    with mock.patch.object(model_call, "_client", return_value=Client()):
        result = model_call.call_with_tools(
            model="bailian/deepseek-v4-flash-0731",
            system="test",
            messages=[{"role": "user", "content": "test"}],
            tools=None,
            max_tokens=16,
            thinking=False,
            effort="high",
            max_retries=1,
        )

    assert result is not None
    assert captured["thinking"] == {"type": "disabled"}
    assert "output_config" not in captured


def test_research_receives_verbatim_raw_user_query_from_harness(tmp_path: Path) -> None:
    raw_query = '做个 8 页 PPT，介绍台风“白海豚”'
    child = Agent(
        "raw-query-test",
        str(tmp_path),
        "Research: investigate the Chinese white dolphin species",
        {
            "skill_name": "mural-presenter-v0-2-zh",
            "skill_language": "zh",
            "_raw_user_query": raw_query,
        },
        role="research",
        label="research",
    )

    assert json.dumps(raw_query, ensure_ascii=False) in child.system
    assert "委派 goal 只是 Orchestrator 的研究假设" in child.system


def test_research_receives_latest_revision_as_separate_user_context(
    tmp_path: Path,
) -> None:
    child = Agent(
        "revision-query-test",
        str(tmp_path),
        "Research: verify the proposed topic",
        {
            "skill_name": "mural-presenter-v0-2-zh",
            "skill_language": "zh",
            "_raw_user_query": "介绍台风白海豚",
            "_revision_instruction": "不是动物，改回台风主题",
        },
        role="research",
        label="research",
    )

    assert "介绍台风白海豚" in child.system
    assert "不是动物，改回台风主题" in child.system
    assert "明确的最新修改优先" in child.system


def test_revision_prompt_routes_simple_and_complex_edits() -> None:
    prompt = _revision_prompt(
        {"query": "介绍台风白海豚", "lang": "zh"},
        {"instruction": "把主题纠正为台风白海豚"},
    )
    assert "Review: mode=simple_edit" in prompt
    assert "plan/revision-impact.md" in prompt
    assert "主题/实体纠正" in prompt
    assert "不能当成局部换字" in prompt


def test_revision_route_is_locked_and_complex_requires_impact_map(
    tmp_path: Path,
) -> None:
    class Parent:
        revision_mode = True
        revision_route = ""
        ws = str(tmp_path)

    parent = Parent()
    error = _revision_delegation_error(
        parent,
        [{"role": "review", "task": "Review: mode=simple_edit; page 2 title"}],
    )
    assert error == ""
    assert parent.revision_route == "simple_edit"
    assert "只能使用唯一" in _revision_delegation_error(
        parent,
        [{"role": "slide", "task": "Slide 02: rewrite page"}],
    )

    complex_parent = Parent()
    assert "revision-impact.md" in _revision_delegation_error(
        complex_parent,
        [{"role": "research", "task": "Research: resolve the named storm"}],
    )
    impact = tmp_path / "plan/revision-impact.md"
    impact.parent.mkdir(parents=True, exist_ok=True)
    impact.write_text("# Impact\n- facts: all pages\n", encoding="utf-8")
    assert _revision_delegation_error(
        complex_parent,
        [{"role": "research", "task": "Research: resolve the named storm"}],
    ) == ""
    assert complex_parent.revision_route == "complex_edit"


def test_revision_orchestrator_cannot_patch_slide_html_directly() -> None:
    class RevisionOrchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-2-zh"
        revision_mode = True

    error = tools._model_write_error(
        RevisionOrchestrator(), "slides/slide_02.html"
    )
    assert error is not None
    assert "不允许 Orchestrator 直接修改" in error


def test_initial_page_plans_use_standard_write_and_scripted_batch_split(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- page_count: 8\n",
        encoding="utf-8",
    )

    class Orchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-2-zh"
        revision_mode = False
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    agent = Orchestrator()
    single = tools.write_file(agent, "plan/slide_01.md", "# slide 1\n")
    assert "不得逐页串行" in single

    manifest = {
        "files": [
            {"path": f"plan/slide_{page:02d}.md", "content": f"# slide {page}\n"}
            for page in range(1, 5)
        ]
    }
    written = tools.write_file(
        agent,
        "plan/plan-batch.json",
        json.dumps(manifest, ensure_ascii=False),
    )
    assert written.startswith("已写入")
    script = (
        REPO
        / "skills/mural-presenter-v0.2/mural-presenter-v0-2-zh/scripts/deck.py"
    )
    applied = subprocess.run(
        [sys.executable, str(script), "apply-plan-batch", str(tmp_path)],
        capture_output=True,
        text=True,
        check=True,
    )
    assert "pages=01,02,03,04" in applied.stdout
    assert all((plan / f"slide_{page:02d}.md").is_file() for page in range(1, 5))
    assert not (plan / "plan-batch.json").exists()

    bad_manifest = {
        "files": [
            {"path": "plan/slide_05.md", "content": "# slide 5\n"},
            {"path": "plan/slide_07.md", "content": "# slide 7\n"},
        ]
    }
    tools.write_file(
        agent,
        "plan/plan-batch.json",
        json.dumps(bad_manifest, ensure_ascii=False),
    )
    nonconsecutive = subprocess.run(
        [sys.executable, str(script), "apply-plan-batch", str(tmp_path)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert nonconsecutive.returncode != 0
    assert "consecutive" in nonconsecutive.stderr
    assert not (plan / "slide_05.md").exists()
    assert not (plan / "slide_07.md").exists()

    tool_names = {
        schema["name"]
        for schema in tools.agent_tools(
            "orchestrator", skill_name="mural-presenter-v0-2-zh"
        )
    }
    assert "write_file" in tool_names
    assert "write_files" not in tool_names


def test_final_single_missing_plan_can_resume_with_write_file(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- page_count: 3\n",
        encoding="utf-8",
    )
    for page in (1, 2):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# slide {page}\n", encoding="utf-8"
        )

    class Orchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-2-zh"
        revision_mode = False
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.write_file(
        Orchestrator(), "plan/slide_03.md", "# recovered slide 3\n"
    )
    assert result.startswith("已写入")
    assert (plan / "slide_03.md").is_file()


def test_slide_repair_required_is_a_structured_nonfatal_handoff(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace/slide-render-states"
    trace.mkdir(parents=True)
    (trace / "page_07.json").write_text(
        json.dumps({"limit": 8, "hashes": [str(i) for i in range(8)]}),
        encoding="utf-8",
    )

    class Slide:
        role = "slide"
        ws = str(tmp_path)
        assigned_slide_pages = (7,)
        repair_required_reason = ""
        final_text = """status: repair_required
pages: 07
issue_type: shared_system
evidence: inherited two-column grid squeezes the primary content
proposed_fix: override the page grid to one column
final_pixels_inspected: yes
"""

    issue = _slide_repair_issue(Slide())
    assert issue is not None
    assert issue["status"] == "repair_required"
    assert issue["pages"] == [7]
    assert issue["issue_type"] == "shared_system"
    assert issue["render_budget_exhausted_pages"] == [7]


def test_review_must_open_flagged_page_current_pixels(tmp_path: Path) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    png = renders / "slide_07.png"
    png.write_bytes(b"current-pixels")

    class Review:
        role = "review"
        ws = str(tmp_path)
        required_review_pages = (7,)
        review_viewed_page_hashes = {}

    assert "P07" in _review_required_view_gap(Review())
    Review.review_viewed_page_hashes = {
        7: __import__("hashlib").sha256(png.read_bytes()).hexdigest()
    }
    assert _review_required_view_gap(Review()) == ""


def test_contract_parser_ignores_free_prose() -> None:
    fields = _contract_fields(
        "The page may be ready later.\nstatus: repair_required\npages: 07\n"
    )
    assert fields == {"status": "repair_required", "pages": "07"}


def test_third_slide_vision_is_routed_to_review() -> None:
    class Slide:
        role = "slide"
        slide_pixel_inspections = 2
        repair_required_reason = ""

    slide = Slide()
    result = tools.vision_analyze(slide, "renders/slide_07.png")
    assert "repair_required" in result
    assert slide.repair_required_reason


def test_bounded_nonblocking_review_is_not_a_whole_deck_failure() -> None:
    class Parent:
        finalize_succeeded = True

    result = {
        "ok": False,
        "status": "needs_orchestrator",
        "blocking": "no",
        "final_pixels_inspected": True,
        "final_view_after_review": True,
        "review_changed": False,
        "attempt": 2,
    }
    assert _review_can_complete_needs_improvement(Parent(), result)
    result["blocking"] = "yes"
    assert not _review_can_complete_needs_improvement(Parent(), result)


def test_finish_gap_nudges_only_incomplete_single_pages(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_01.md").write_text("# P1\n", encoding="utf-8")
    (plan / "slide_02.md").write_text("# P2\n", encoding="utf-8")

    class Agent:
        role = "orchestrator"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = True
        delegated_roles = {"slide"}
        child_outcomes = {
            "slide_01": {"ok": True},
            "slide_02": {"ok": False},
        }
        ws = str(tmp_path)

    gap = _finish_gap(Agent())
    assert "Slide 02:" in gap
    assert "Slide 01:" not in gap
    assert "不要合并为 SlideGroup" in gap


def test_finish_gap_retries_only_incomplete_complete_group(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    for page, group in ((1, "opening"), (2, "mechanism"), (3, "mechanism")):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n- production_group: {group}\n",
            encoding="utf-8",
        )

    class GroupedAgent:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-2-grouped-zh"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = True
        revision_mode = False
        delegated_roles = {"slide"}
        child_outcomes = {
            "slide_group_opening": {
                "ok": True,
                "completed_pages": [1],
            },
            "slide_group_mechanism": {
                "ok": False,
                "completed_pages": [],
                "incomplete_pages": [2, 3],
            },
        }
        ws = str(tmp_path)

    gap = _finish_gap(GroupedAgent())
    assert "SlideGroup mechanism [02,03]:" in gap
    assert "SlideGroup opening" not in gap
    assert "不要拆成单页" in gap


def test_grouped_skill_diff_is_limited_to_ownership_topology() -> None:
    topology_files = {
        "SKILL.md",
        "agents/openai.yaml",
        "references/plan-contract.md",
        "roles/orchestrator.md",
        "roles/slide.md",
        "scripts/deck.py",
    }
    for language in ("zh", "en"):
        single_name = f"mural-presenter-v0-2-{language}"
        grouped_name = f"mural-presenter-v0-2-grouped-{language}"
        single = REPO / "skills/mural-presenter-v0.2" / single_name
        grouped = REPO / "skills/mural-presenter-v0.2" / grouped_name
        single_files = {
            path.relative_to(single).as_posix()
            for path in single.rglob("*")
            if path.is_file()
        }
        grouped_files = {
            path.relative_to(grouped).as_posix()
            for path in grouped.rglob("*")
            if path.is_file()
        }
        assert grouped_files == single_files
        for relative in sorted(single_files - topology_files):
            left = (single / relative).read_bytes()
            right = (grouped / relative).read_bytes().replace(
                grouped_name.encode(), single_name.encode()
            )
            assert right == left, relative


def test_revision_finish_gap_does_not_require_unchanged_pages(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    for page in range(1, 4):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n", encoding="utf-8"
        )

    class RevisionAgent:
        role = "orchestrator"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = False
        revision_mode = True
        revision_route = "complex_edit"
        delegated_roles = {"slide"}
        child_outcomes = {"slide_02": {"ok": True}}
        review_completed = False
        finalize_attempted = False
        finalize_succeeded = False
        ws = str(tmp_path)

    gap = _finish_gap(RevisionAgent())
    assert "Review 尚未完成" in gap
    assert "Slide 01:" not in gap
    assert "Slide 03:" not in gap


def test_pdf_ingestion_writes_complete_text_and_page_derivatives(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    source = inputs / "paper.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Principal finding: attachment evidence is preserved.")
    document.save(source)
    document.close()

    companion = inputs / "paper.pdf.md"
    assert _extract_material(source, companion)
    text = companion.read_text(encoding="utf-8")
    assert "Principal finding" in text
    assert "text_mode: native" in text
    assert (inputs / "paper.pdf.pages/page_001.png").is_file()
    layout = json.loads(
        (inputs / "paper.pdf.pages/page_001.json").read_text(encoding="utf-8")
    )
    assert layout["text_mode"] == "native"
    assert layout["image_pixels"][0] >= 1400


def test_office_ingestion_preserves_text_formulas_and_embedded_media(
    tmp_path: Path,
) -> None:
    from PIL import Image
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorksheetImage
    from pptx import Presentation
    from pptx.util import Inches

    inputs = tmp_path / "inputs"
    inputs.mkdir()
    picture = inputs / "evidence.png"
    Image.new("RGB", (320, 180), "#3c78d8").save(picture)

    docx_path = inputs / "brief.docx"
    doc = Document()
    doc.add_heading("Policy brief", level=1)
    doc.add_paragraph("Target adoption is 42 percent.")
    doc.add_picture(str(picture))
    doc.save(docx_path)
    assert _extract_material(docx_path, inputs / "brief.docx.md")
    docx_text = (inputs / "brief.docx.md").read_text(encoding="utf-8")
    assert "Target adoption is 42 percent" in docx_text
    assert "embedded_media:" in docx_text
    assert list((inputs / "brief.docx.pages").glob("embedded_*.png"))

    pptx_path = inputs / "roadmap.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
    textbox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    textbox.text = "Launch timeline: pilot in Q3"
    deck.save(pptx_path)
    assert _extract_material(pptx_path, inputs / "roadmap.pptx.md")
    pptx_text = (inputs / "roadmap.pptx.md").read_text(encoding="utf-8")
    assert "Launch timeline: pilot in Q3" in pptx_text
    assert list((inputs / "roadmap.pptx.pages").glob("embedded_*.png"))

    xlsx_path = inputs / "portfolio.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(["Project", "Budget", "Double"])
    sheet.append(["Aurora", 125, "=B2*2"])
    sheet.add_image(WorksheetImage(str(picture)), "E2")
    workbook.save(xlsx_path)
    assert _extract_material(xlsx_path, inputs / "portfolio.xlsx.md")
    xlsx_text = (inputs / "portfolio.xlsx.md").read_text(encoding="utf-8")
    assert "Project | Budget | Double" in xlsx_text
    assert "Aurora | 125 | =B2*2" in xlsx_text
    assert list((inputs / "portfolio.xlsx.pages").glob("embedded_*.png"))


def test_material_staging_failure_is_structured_and_blocks_planning(
    tmp_path: Path,
) -> None:
    source = tmp_path / "unknown.bin"
    source.write_bytes(b"not a supported attachment")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)]}
    staged = stage_materials(workspace, seed)
    assert staged == ["inputs/01_unknown.bin"]
    failures = seed["_material_ingestion_failures"]
    assert failures[0]["error"] == "unsupported attachment type"
    _material_blocked(str(workspace), failures)
    blocked = (workspace / "research/material.md").read_text(encoding="utf-8")
    assert "status: material_blocked" in blocked
    assert "unknown.bin" in blocked


def test_standalone_image_attachment_is_material_input_not_parse_failure(
    tmp_path: Path,
) -> None:
    from PIL import Image

    source = tmp_path / "diagram.png"
    Image.new("RGB", (640, 360), "white").save(source)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)]}
    assert stage_materials(workspace, seed) == ["inputs/01_diagram.png"]
    assert seed["_material_ingestion_failures"] == []


def test_scanned_pdf_blocks_before_agent_when_main_model_is_text_only(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(source)
    document.close()
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"materials": [str(source)]}
    with mock.patch("core.run_batch.acfg.VISION_BACKEND", "disabled"):
        stage_materials(workspace, seed)
    assert seed["_material_ingestion_failures"]
    assert "text-only" in seed["_material_ingestion_failures"][0]["error"]


def test_attachment_only_is_a_tool_level_research_scope(tmp_path: Path) -> None:
    tool_names = {
        schema["name"]
        for schema in tools.agent_tools(
            "research", evidence_scope="attachment_only"
        )
    }
    assert "web_search" not in tool_names
    assert "web_extract" not in tool_names

    class Orchestrator:
        role = "orchestrator"
        material_required = True
        material_completed = False

    class Research:
        role = "research"

    assert "不得直接读取" in tools.read_file(Orchestrator(), "inputs/report.pdf")
    assert "不直接重读" in tools.read_file(Research(), "inputs/report.pdf.md")


def test_plan_only_acceptance_requires_material_research_and_plans(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "plan").mkdir()
    (tmp_path / "research/material.md").write_text("status: ready\n", encoding="utf-8")
    (tmp_path / "research/knowledge-brief.md").write_text("# Evidence\n", encoding="utf-8")
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    for page in (1, 2):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n", encoding="utf-8"
        )

    class PlanOnlyAgent:
        child_outcomes = {}
        tool_policy_violations = []
        workspace_policy_violations = []

    with mock.patch.dict(os.environ, {"CLEAN_PLAN_ONLY_EVAL": "1"}):
        ok, reason, detail = _accept(PlanOnlyAgent(), str(tmp_path))
    assert ok is True
    assert reason == "plan-only deliverables complete"
    assert detail["n_slide_plans"] == 2


def test_material_blocked_is_a_terminal_structured_acceptance_state(
    tmp_path: Path,
) -> None:
    class BlockedAgent:
        role = "orchestrator"
        skill_language = "zh"
        material_blocked = True
        material_required = True
        material_completed = False
        child_outcomes = {"material": {"ok": False, "status": "material_blocked"}}
        tool_policy_violations = []
        workspace_policy_violations = []

    assert _finish_gap(BlockedAgent()) == ""
    ok, reason, detail = _accept(BlockedAgent(), str(tmp_path))
    assert ok is False
    assert reason == "material_blocked"
    assert detail["material_blocked"] is True


def test_attachment_pixels_are_rejected_by_crop_and_catalog(tmp_path: Path) -> None:
    from PIL import Image

    page_dir = tmp_path / "inputs/paper.pdf.pages"
    page_dir.mkdir(parents=True)
    source = page_dir / "page_001.png"
    Image.new("RGB", (1600, 2000), "white").save(source)
    source.with_suffix(".json").write_text(
        json.dumps({"page_points": [800, 1000], "text_blocks": []}),
        encoding="utf-8",
    )
    script = (
        REPO
        / "skills/mural-presenter-v0.2/mural-presenter-v0-2-en/scripts/deck.py"
    )
    blocked_crop = subprocess.run(
        [
            sys.executable,
            str(script),
            "material-figure",
            str(tmp_path),
            "--source",
            "inputs/paper.pdf.pages/page_001.png",
            "--output",
            "assets/figure.png",
            "--box",
            "0.10,0.10,0.90,0.35",
        ],
        capture_output=True,
        text=True,
    )
    assert blocked_crop.returncode != 0
    assert "material-figure is disabled" in blocked_crop.stderr
    assert not (tmp_path / "assets/figure.png").exists()

    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "figure.png").write_bytes(source.read_bytes())
    (assets / "catalog.md").write_text(
        """# Asset catalog

## forbidden-attachment
- slides: 1
- kind: material
- path: assets/figure.png
- purpose: copied from attachment
""",
        encoding="utf-8",
    )
    blocked_catalog = subprocess.run(
        [sys.executable, str(script), "fetch-images", str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert blocked_catalog.returncode != 0
    assert "kind must be one of" in blocked_catalog.stderr
    assert "generated" in blocked_catalog.stderr
    assert "real" in blocked_catalog.stderr


def test_image_role_cannot_read_or_view_attachment_pixels(tmp_path: Path) -> None:
    class ImageAgent:
        role = "image"

    read_result = tools.read_file(
        ImageAgent(),
        "inputs/paper.pdf.pages/page_001.json",
    )
    assert "Image 不读取原始附件" in read_result

    view_result = tools.vision_analyze(
        ImageAgent(),
        "inputs/paper.pdf.pages/page_001.png",
    )
    assert "Image 不查看或裁切附件像素" in view_result


def test_image_terminal_does_not_expose_material_figure(tmp_path: Path) -> None:
    script = (
        REPO
        / "skills/mural-presenter-v0.2/mural-presenter-v0-2-en/scripts/deck.py"
    )

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        render_script = str(script)
        bash_relaxed = False

    result = tools.terminal(
        ImageAgent(),
        f"python {script} material-figure . --source inputs/page.png "
        "--output assets/figure.png --box 0.1,0.1,0.9,0.9",
    )
    assert "不允许的 deck.py 动作 `material-figure`" in result


def test_v02_portable_font_bundle_is_self_contained(tmp_path: Path) -> None:
    edition = REPO / "skills/mural-presenter-v0.2/mural-presenter-v0-2-en"
    shutil.copyfile(edition / "assets/base.css", tmp_path / "base.css")
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "slide_01.html").write_text(
        '<section class="slide"><h1 class="type-heavy">RAVE 渲染一致</h1></section>',
        encoding="utf-8",
    )

    result = subprocess.run(
        [sys.executable, str(edition / "scripts/font_bundle.py"), str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr
    css = (tmp_path / "base.css").read_text(encoding="utf-8")
    assert "DECK_FONT_BUNDLE_START" in css
    assert '@font-face {' in css
    manifest = json.loads(
        (tmp_path / "assets/fonts/manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["faces"]
    assert all((tmp_path / face["path"]).is_file() for face in manifest["faces"])
