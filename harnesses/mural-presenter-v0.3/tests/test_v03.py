from __future__ import annotations

import json
import hashlib
import os
import shutil
import subprocess
import sys
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from unittest import mock

HARNESS = Path(__file__).resolve().parents[1]
REPO = HARNESS.parents[1]
sys.path.insert(0, str(HARNESS))
V03_SCRIPTS = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts"
sys.path.insert(0, str(V03_SCRIPTS))

from core import tools  # noqa: E402
from core import config  # noqa: E402
from core import model_call  # noqa: E402
from core.agent_loop import (  # noqa: E402
    Agent,
    _accept,
    _contract_fields,
    _delegate,
    _delegate_task,
    _direct_text_injection,
    _finish_gap,
    _image_required_fullres_gap,
    _mandatory_fullres_review_pages,
    _open_vision_issues,
    _persisted_trace_attempts,
    _review_can_complete_needs_improvement,
    _review_required_view_gap,
    _render_input_digest,
    _review_delivery_fingerprint,
    _revision_delegation_error,
    _revision_prompt,
    _slide_authoring_stop_line,
    _slide_deliverable_gap,
    _slide_render_quality_actions,
    _slide_repair_issue,
    _style_lock_injection,
    _tool_results,
    _update_vision_issue_ledger,
)
from core.run_batch import _extract_material, _material_blocked, stage_materials  # noqa: E402
from core.runtime_capabilities import (  # noqa: E402
    classify_research_mode,
    detect_runtime_capabilities,
    user_requires_bitmap,
)
from core.trace_mode import write_trace  # noqa: E402
from _internal import deck_core  # noqa: E402


def test_soft_asset_quality_marks_blurry_full_bleed_without_rejecting(
    tmp_path: Path,
) -> None:
    from PIL import Image

    path = tmp_path / "archive.png"
    Image.new("RGB", (1600, 900), (92, 92, 92)).save(path)
    report = deck_core._asset_visual_quality(path, "full-bleed")
    assert report["width"] == 1600
    assert report["height"] == 900
    assert not report["full_bleed_ready"]
    assert any("sharpness" in reason for reason in report["reasons"])


def test_high_confidence_layout_defects_are_structured_but_nonblocking(
    tmp_path: Path,
) -> None:
    geometry = {
        "page_type": "content",
        "frame": "content",
        "page_family": "editorial",
        "canvas_variant": "base",
        "text_boxes": [
            {
                "tag": "p",
                "class_name": "card-copy",
                "ancestor_chain": ["div.card", "main.page-body"],
                "text": "clipped body copy",
                "x": 100,
                "y": 100,
                "width": 300,
                "height": 80,
                "overflow_px_x": 0,
                "overflow_px_y": 12,
            },
            {
                "tag": "li",
                "class_name": "field-item",
                "ancestor_chain": ["ul.field-list", "main.page-body"],
                "text": "colliding list item",
                "x": 120,
                "y": 120,
                "width": 260,
                "height": 70,
                "overflow_px_x": 0,
                "overflow_px_y": 0,
            },
        ],
        "layout_blocks": [
            {
                "tag": "div",
                "class_name": "composition composition--sequence",
                "parent_key": "div.content-stage",
                "ancestor_chain": ["div.content-stage", "main.page-body"],
                "x": 80,
                "y": 180,
                "width": 1440,
                "height": 650,
                "overflow_px_x": 0,
                "overflow_px_y": 24,
                "overflow_mode_x": "visible",
                "overflow_mode_y": "visible",
            },
            {
                "tag": "article",
                "class_name": "card card-a",
                "parent_key": "div.grid-2",
                "ancestor_chain": ["div.grid-2", "div.content-stage"],
                "x": 100,
                "y": 300,
                "width": 500,
                "height": 300,
                "overflow_px_x": 0,
                "overflow_px_y": 0,
                "overflow_mode_x": "visible",
                "overflow_mode_y": "visible",
            },
            {
                "tag": "article",
                "class_name": "card card-b",
                "parent_key": "div.grid-2",
                "ancestor_chain": ["div.grid-2", "div.content-stage"],
                "x": 200,
                "y": 320,
                "width": 500,
                "height": 300,
                "overflow_px_x": 0,
                "overflow_px_y": 0,
                "overflow_mode_x": "visible",
                "overflow_mode_y": "visible",
            },
        ],
        "images": [],
        "boxes": {
            "header": {"x": 60, "y": 40, "width": 1480, "height": 100},
            "footer": {"x": 60, "y": 830, "width": 1480, "height": 30},
        },
    }
    defects = deck_core._page_layout_defects(1, geometry)
    assert {item["type"] for item in defects} == {
        "text_overflow",
        "text_collision",
        "block_overflow",
        "block_collision",
    }
    assert any("page-body" in item.get("ancestor_hint", "") for item in defects)
    assert all(item["id"].startswith("LAYOUT-") for item in defects)
    assert defects == deck_core._page_layout_defects(1, geometry)

    rendered = {"pages": [{"page": 1, "geometry": geometry}]}
    plans = [{
        "number": 1,
        "page_type": "content",
        "page_family": "editorial",
        "canvas_variant": "base",
        "show_footer": True,
        "needs_bitmap": False,
    }]
    audit = deck_core._geometry_audit(tmp_path, plans, rendered)
    assert audit["status"] == "PASS"
    assert audit["errors"] == []
    assert audit["layout_defects"]["01"]

    with mock.patch.dict(os.environ, {"MURAL_LAYOUT_DEFECT": "off"}):
        assert deck_core._page_layout_defects(1, geometry) == []


def test_layout_defect_action_is_shared_by_single_and_grouped_v03(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    state_path = tmp_path / "_trace/slide-render-states/page_01.json"
    state_path.write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "attempt_hashes": {"slide_group_intro": ["first"]},
            "authoring_hashes": ["first"],
            "layout_defects": [{"type": "text_overflow"}],
            "quality_action_count": 1,
        }),
        encoding="utf-8",
    )

    class Slide:
        role = "slide"
        skill_name = "mural-presenter-v0-3"
        trace_label = "slide_group_intro"
        ws = str(tmp_path)

    for topology in ("single", "grouped"):
        (tmp_path / "plan/deck.md").write_text(
            f"# Deck\n## Resolved deck brief\n- ownership_topology: {topology}\n",
            encoding="utf-8",
        )
        assert _slide_render_quality_actions(Slide(), 1) == 1

    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["attempt_hashes"]["slide_group_intro"].append("second")
    state["authoring_hashes"].append("second")
    state_path.write_text(json.dumps(state), encoding="utf-8")
    assert _slide_render_quality_actions(Slide(), 1) == 1
    state["attempt_hashes"]["slide_group_intro"].append("third")
    state["authoring_hashes"].append("third")
    state_path.write_text(json.dumps(state), encoding="utf-8")
    assert _slide_render_quality_actions(Slide(), 1) == 0


def test_layout_defects_do_not_turn_usable_delivery_into_rejection(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "assets").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- language: zh\n",
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_01.md").write_text("# Slide 01\n", encoding="utf-8")
    (tmp_path / "assets/catalog.md").write_text("# Catalog\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "present.html").write_text("<main>" + "x" * 600 + "</main>", encoding="utf-8")
    (tmp_path / "renders/render.json").write_text(
        json.dumps({
            "n_pages": 1,
            "blank_pages": [],
            "console_errors": [],
            "static_pages": [1],
            "special_page_geometry": {"status": "PASS"},
            "layout_defects": {
                "01": [{"id": "LAYOUT-TEST", "type": "text_collision"}],
            },
        }),
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        child_outcomes={},
        tool_policy_violations=[],
        workspace_policy_violations=[],
        material_blocked=False,
        n_renders=1,
        deck_language="zh",
        review_completed=True,
        final_render_after_review=True,
        final_view_after_review=True,
        finalize_attempted=True,
        finalize_succeeded=True,
        finalize_failure="",
        quality_status="needs_improvement",
        research_required=False,
        research_completed=False,
        material_required=False,
        material_completed=False,
        image_required=False,
        image_completed=False,
    )

    ok, reason, detail = _accept(agent, str(tmp_path))
    assert ok is True
    assert reason == "ok"
    assert detail["quality_status"] == "needs_improvement"


def test_anti_slop_lint_reports_repeated_dark_glow_and_card_geometry(
    tmp_path: Path,
) -> None:
    (tmp_path / "slides").mkdir()
    html = (
        '<section class="slide"><div class="card"></div><div class="card"></div>'
        '<div class="card"></div><style>.slide{background:radial-gradient(circle,#111,#000)}'
        '</style></section>'
    )
    plans = []
    for page in range(1, 5):
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(html, encoding="utf-8")
        plans.append({
            "number": page,
            "page_type": "content",
            "canvas_variant": "base",
            "needs_bitmap": False,
            "special_layout": "",
        })
    deck = {
        "sections": {"visual": "generic clean interface"},
        "tokens": {
            "--content-canvas": "#11143a",
            "--special-bg": "#11143a",
            "--accent": "#6366f1",
        },
    }
    warnings = deck_core._anti_slop_warnings(tmp_path, deck, plans)
    joined = "\n".join(warnings)
    assert "default indigo" in joined
    assert "radial-gradient" in joined
    assert "card/tile/panel" in joined
    assert "uniform solid background" in joined


def test_anti_slop_reads_localized_visual_contract_and_flags_flat_navy(
    tmp_path: Path,
) -> None:
    plans = [
        {
            "number": page,
            "page_type": "content",
            "canvas_variant": "base",
            "needs_bitmap": False,
            "special_layout": "",
        }
        for page in range(1, 5)
    ]
    deck = {
        "title": "多模态人工智能通识课",
        "resolved": "",
        "sections": {
            "视觉契约": "背景配方：基础画布使用深靛蓝纯色，全册统一背景。",
        },
        "tokens": {
            "--content-canvas": "#14182C",
            "--special-bg": "#0B0E1E",
        },
    }
    joined = "\n".join(deck_core._anti_slop_warnings(tmp_path, deck, plans))
    assert "generic dark navy/indigo" in joined
    assert "uniform solid background" in joined


def test_body_mono_is_an_advisory_audit_warning(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "base.css").write_text(
        ".body-copy { font-family: var(--font-mono); }\n"
        ".tech-id { font-family: var(--font-mono); }\n",
        encoding="utf-8",
    )
    warnings = deck_core._typography_quality_warnings(tmp_path)
    assert len(warnings) == 1
    assert "body-copy" in warnings[0]
    assert "tech-id" not in warnings[0]


def test_divider_with_one_follower_emits_soft_actionable_warning(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck = {
        "ownership_topology": "single",
        "bitmap_strategy": "none",
        "language": "en",
        "warnings": [],
    }
    plans = [
        {"number": 1, "production_group": "p1", "needs_bitmap": False,
         "page_type": "cover", "composition_explicit": True,
         "composition_blueprint": "hero", "composition": "freeform"},
        {"number": 2, "production_group": "p2", "needs_bitmap": False,
         "page_type": "section-divider", "composition_explicit": True,
         "composition_blueprint": "pause", "composition": "freeform"},
        {"number": 3, "production_group": "p3", "needs_bitmap": False,
         "page_type": "content", "composition_explicit": True,
         "composition_blueprint": "evidence", "composition": "editorial"},
        {"number": 4, "production_group": "p4", "needs_bitmap": False,
         "page_type": "closing", "composition_explicit": True,
         "composition_blueprint": "close", "composition": "freeform"},
    ]
    with mock.patch.object(deck_core, "_load_plans", return_value=(deck, plans)):
        assert deck_core.validate_plans(tmp_path, 4) == plans
    output = capsys.readouterr().out
    assert "status:PASS" in output
    assert "leads only 1 content page" in output
    assert "remove the standalone divider" in output


def test_theme_tokens_reject_css_self_reference() -> None:
    with pytest.raises(ValueError, match="invalid self-references: --font-mono"):
        deck_core._theme_tokens(
            ":root { --content-canvas: #f2eee4; --font-mono: var(--font-mono); }"
        )


def test_slide_write_guard_preserves_scaffold_owned_root_metadata(
    tmp_path: Path,
) -> None:
    page = tmp_path / "slide_03.html"
    opening = (
        '<section class="slide" id="slide-03" data-slide="03" '
        'data-page-type="content" data-page-family="editorial" '
        'data-frame="content" data-canvas-variant="base">'
    )
    page.write_text(opening + "<main>scaffold</main></section>", encoding="utf-8")
    agent = SimpleNamespace(role="slide")
    broken = '<section class="slide" id="slide-03"><main>done</main></section>'
    with mock.patch.object(tools, "_is_current_variant_skill", return_value=True):
        error = tools._slide_root_contract_error(
            agent, "slides/slide_03.html", str(page), broken
        )
        assert error and "data-slide='03'" in error
        assert tools._slide_root_contract_error(
            agent,
            "slides/slide_03.html",
            str(page),
            opening + "<main>done</main></section>",
        ) is None


def test_scaffold_force_cannot_destroy_authored_or_rendered_pages(
    tmp_path: Path,
) -> None:
    slides = tmp_path / "slides"
    renders = tmp_path / "renders"
    slides.mkdir()
    renders.mkdir()
    page = slides / "slide_01.html"
    original = '<section class="slide"><main>completed page</main></section>\n'
    page.write_text(original, encoding="utf-8")
    (renders / "slide_01.png").write_bytes(b"stable-render")

    with mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ):
        with pytest.raises(ValueError, match="refusing scaffold --force") as error:
            deck_core.scaffold_from_plans(tmp_path, 1, force=True)

    assert "authored HTML: slide_01.html" in str(error.value)
    assert "render checkpoints: slide_01.png" in str(error.value)
    assert page.read_text(encoding="utf-8") == original


def test_trace_sidecar_distinguishes_real_user_from_tool_result_wrappers(
    tmp_path: Path,
) -> None:
    messages = [
        {"role": "user", "content": "real query"},
        {"role": "assistant", "content": [{"type": "tool_use", "id": "t1"}]},
        {
            "role": "user",
            "content": [
                {"type": "tool_result", "tool_use_id": "t1", "content": "ok"}
            ],
        },
        {"role": "user", "content": "[系统] continue"},
    ]
    write_trace(tmp_path, messages, [], {}, "inference")
    origins = json.loads(
        (tmp_path / "message-origins.json").read_text(encoding="utf-8")
    )["messages"]
    assert [row["semantic_origin"] for row in origins] == [
        "end_user",
        "model_response",
        "harness_tool_result",
        "harness_control",
    ]
    assert json.loads((tmp_path / "messages.json").read_text(encoding="utf-8")) == messages


def test_workspace_audit_allows_empty_harness_tmp_and_result_json(
    tmp_path: Path,
) -> None:
    (tmp_path / "tmp").mkdir()
    (tmp_path / "result.json").write_text('{"status":"completed"}', encoding="utf-8")
    assert deck_core.audit_workspace(tmp_path) == []
    (tmp_path / "tmp" / "residue.tmp").write_text("x", encoding="utf-8")
    violations = deck_core.audit_workspace(tmp_path)
    assert any("retained files" in item for item in violations)


def test_render_state_budget_is_scoped_to_page_lifecycle(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "_trace" / "slide-render-states").mkdir(parents=True)
    (tmp_path / "base.css").write_text("body { color: black; }", encoding="utf-8")
    (tmp_path / "slides" / "slide_01.html").write_text(
        '<section class="slide"><main>completed page</main></section>',
        encoding="utf-8",
    )
    state = tmp_path / "_trace" / "slide-render-states" / "page_01.json"
    state.write_text(json.dumps({
        "limit": 8,
        "authoring_attempt_limit": 3,
        "hashes": ["first", "second"],
        "authoring_hashes": ["first", "second"],
        "attempt_hashes": {"slide_01": ["first", "second"]},
        "published_hash": "second",
    }), encoding="utf-8")

    def publish(_root, _source, target):
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"png")

    with mock.patch.dict(os.environ, {"MURAL_RENDER_ATTEMPT_ID": "slide_01_r2"}), mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ), mock.patch.object(
        deck_core, "_skeleton_is_unfilled", return_value=False
    ), mock.patch.object(
        deck_core, "_validate_fragment", return_value=None
    ), mock.patch.object(
        deck_core, "_run_renderer", return_value=None
    ), mock.patch.object(
        deck_core, "atomic_copy", side_effect=publish
    ):
        deck_core.render(tmp_path, 1, 1)

    payload = json.loads(state.read_text(encoding="utf-8"))
    assert payload["limit_scope"] == "page_lifecycle"
    assert payload["authoring_attempt_limit"] == 3
    assert len(payload["authoring_hashes"]) == 3
    assert len(payload["attempt_hashes"]["slide_01_r2"]) == 1
    assert (tmp_path / "renders" / "slide_01.png").is_file()

    (tmp_path / "slides" / "slide_01.html").write_text(
        '<section class="slide"><main>fourth state</main></section>',
        encoding="utf-8",
    )
    with mock.patch.dict(os.environ, {"MURAL_RENDER_ATTEMPT_ID": "slide_01_r3"}), mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ), mock.patch.object(
        deck_core, "_skeleton_is_unfilled", return_value=False
    ), mock.patch.object(
        deck_core, "_validate_fragment", return_value=None
    ):
        with pytest.raises(ValueError, match="page lifecycle"):
            deck_core.render(tmp_path, 1, 1)


def test_invalid_fragment_does_not_consume_render_state_budget(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "base.css").write_text("body { color: black; }", encoding="utf-8")
    (tmp_path / "slides" / "slide_01.html").write_text(
        '<section class="slide"><main>completed page</main></section>',
        encoding="utf-8",
    )
    with mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ), mock.patch.object(
        deck_core, "_skeleton_is_unfilled", return_value=False
    ), mock.patch.object(
        deck_core, "_validate_fragment", side_effect=ValueError("missing data-slide")
    ):
        with pytest.raises(ValueError, match="missing data-slide"):
            deck_core.render(tmp_path, 1, 1)

    assert not (tmp_path / "_trace" / "slide-render-states" / "page_01.json").exists()


def test_incremental_finalize_checkpoint_tracks_published_pixels(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders/.page_01").mkdir(parents=True)
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    (tmp_path / "base.css").write_text(".slide{color:#111}\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide" data-slide="01">A</section>\n', encoding="utf-8"
    )
    (tmp_path / "renders/slide_01.png").write_bytes(b"png")
    digest = deck_core._page_digest(tmp_path, 1)
    deck_core._publish_render_state(
        tmp_path,
        1,
        tmp_path / "_trace/slide-render-states/page_01.json",
        [digest],
        {"slide_01": [digest]},
        digest,
    )
    isolated = {
        "pages": [{"page": 1, "png": "old", "blank": False, "geometry": {}}],
        "console_errors": [],
    }
    (tmp_path / "renders/.page_01/render.json").write_text(
        json.dumps(isolated), encoding="utf-8"
    )
    manifest = deck_core._incremental_manifest_seed(tmp_path, [1])
    assert manifest["pages"][0]["page"] == 1
    assert deck_core._render_checkpoint_current(tmp_path, 1, {1})

    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide" data-slide="01">B</section>\n', encoding="utf-8"
    )
    assert not deck_core._render_checkpoint_current(tmp_path, 1, {1})


def test_page_render_fingerprint_changes_when_referenced_asset_changes(
    tmp_path: Path,
) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "assets").mkdir()
    (tmp_path / "base.css").write_text(".slide{color:#111}\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide"><img src="assets/hero.png"></section>',
        encoding="utf-8",
    )
    asset = tmp_path / "assets/hero.png"
    asset.write_bytes(b"before")
    before = deck_core._page_digest(tmp_path, 1)
    asset.write_bytes(b"after")
    after = deck_core._page_digest(tmp_path, 1)
    assert before != after


def test_runtime_capabilities_omit_optional_stages_without_inputs_or_keys(
    tmp_path: Path,
) -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {"model": "test", "model_base_url": "http://model", "enable_image_gen": True},
            [],
            plan_only=True,
        )
    assert profile["active_roles"] == ["slide", "review"]
    assert profile["omitted_roles"] == ["material", "research", "image"]
    assert profile["workflow"]["image"] == "bitmap_unavailable"
    assert profile["workflow"]["research_mode"] == "open_research"
    assert profile["tools"]["web_search"] is False
    assert profile["tools"]["image_generate"] is False
    assert "api_key" not in json.dumps(profile).lower()


def test_runtime_capabilities_route_material_directly_without_search() -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "image-secret", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {"model": "test", "model_base_url": "http://model", "enable_image_gen": True},
            ["inputs/report.pdf.md"],
            plan_only=True,
        )
    assert profile["active_roles"] == ["material", "image", "slide", "review"]
    assert profile["workflow"]["material_handoff"] == "orchestrator"
    assert profile["workflow"]["research"] == "omitted"


def test_attachment_can_route_material_figures_without_search_or_generation() -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "enable_image_gen": True,
                "_raw_user_query": "仅使用附件制作演示",
            },
            ["inputs/report.pdf.md"],
            plan_only=True,
        )
    assert profile["active_roles"] == ["material", "image", "slide", "review"]
    assert profile["workflow"]["research_mode"] == "attachment_only"
    assert profile["workflow"]["image"] == "user_or_material_bitmap_only"
    assert profile["tools"]["web_search"] is False
    assert profile["tools"]["image_generate"] is False


@pytest.mark.parametrize(
    ("query", "has_material", "expected"),
    [
        ("把以下文字润色成三页演示", False, "off"),
        ("仅使用附件内容，不要联网", True, "attachment_only"),
        ("基于附件制作并核验最新政策", True, "verify_external"),
        ("介绍台风白海豚的形成和影响", False, "open_research"),
    ],
)
def test_research_mode_is_classified_before_model_call(
    query: str,
    has_material: bool,
    expected: str,
) -> None:
    mode, _reason = classify_research_mode(query, has_material=has_material)
    assert mode == expected


def test_web_search_sets_locale_and_retries_one_simplified_query() -> None:
    calls = []

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    def fake_post(_url, **kwargs):
        calls.append(kwargs["json"])
        return Response({"organic": []} if len(calls) == 1 else {
            "organic": [{"title": "官方资料", "link": "https://example.com", "snippet": "证据"}]
        })

    class Research:
        search_enabled = True
        evidence_scope = "open_research"

    with mock.patch.dict(os.environ, {"SERPER_API_KEY": "test-key"}), mock.patch(
        "core.tools.requests.post", side_effect=fake_post
    ):
        payload = json.loads(
            tools.web_search(Research(), '"台风白海豚" 形成 路径 site:example.invalid')
        )

    assert len(calls) == 2
    assert calls[0]["hl"] == "zh-cn" and calls[0]["gl"] == "cn"
    assert "site:" not in calls[1]["q"] and '"' not in calls[1]["q"]
    assert payload["retried_with_simplified_query"] is True
    assert payload["results"][0]["url"] == "https://example.com"


def test_verify_external_keeps_factual_gate_but_allows_image_asset_search(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "research" / "material.md").write_text(
        "# Material\n- unresolved_items: []\n", encoding="utf-8"
    )

    class AgentStub:
        search_enabled = True
        evidence_scope = "verify_external"
        ws = str(tmp_path)

        def __init__(self, role: str):
            self.role = role

    assert "当前没有可核验项" in tools._external_evidence_gate(AgentStub("research"))
    assert tools._external_evidence_gate(
        AgentStub("image"), purpose="visual_asset_acquisition"
    ) == ""


def test_explicit_bitmap_requirement_is_persisted_in_runtime_contract() -> None:
    assert user_requires_bitmap("须含高质量相关图片") is True
    assert user_requires_bitmap("Images: Include relevant images. Images must be high quality.") is True
    assert user_requires_bitmap("不要使用图片，只做数据图") is False
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "search", "IMAGE_API_KEY": "image", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "enable_image_gen": True,
                "_raw_user_query": "制作演示，须含高质量相关图片",
            },
            [],
            plan_only=True,
        )
    assert profile["requirements"] == {
        "bitmap_required": True,
        "source": "raw_user_query",
        "failure_policy": "keep_needs_bitmap_and_return_image_blocked",
    }


def test_image_generation_provider_failure_preserves_requested_model(
    tmp_path: Path,
) -> None:
    class Response:
        ok = False
        status_code = 429
        text = ""

        @staticmethod
        def json():
            return {
                "error": {
                    "message": "All channels for model gpt-image-2 are unavailable"
                }
            }

    class ImageAgent:
        enable_image_gen = True
        img_base = "http://tokenhub"
        img_key = "secret"
        image_model = "gpt-image-2-adobe-2"
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    with mock.patch("core.tools.requests.post", return_value=Response()):
        result = tools.image_generate(ImageAgent(), "editorial portrait")
    assert "[provider_unavailable]" in result
    assert "requested_model=gpt-image-2-adobe-2" in result
    assert "http_status=429" in result
    assert "不得把 needs_bitmap:true 静默改为 false" in result


def test_runtime_capabilities_inject_only_available_roles_and_tools(tmp_path: Path) -> None:
    profile = {
        "active_roles": ["slide", "review"],
        "omitted_roles": ["material", "research", "image"],
        "workflow": {
            "material": "omitted",
            "research": "omitted",
            "material_handoff": "none",
            "image": "bitmap_unavailable",
            "research_mode": "off",
            "research_reason": "test",
            "material_stage": "omitted",
        },
        "tools": {
            "web_search": False,
            "web_extract": False,
            "image_generate": False,
            "vision_analyze": True,
        },
    }
    agent = Agent(
        "runtime-capability-test",
        str(tmp_path),
        "制作一套演示",
        {
            "skill_name": "mural-presenter-v0-3",
            "skill_language": "zh",
            "_runtime_capabilities": profile,
        },
    )
    tool_names = {schema["name"] for schema in agent.tool_schemas}
    assert agent.research_required is False
    assert agent.role_card_injected is True
    assert '<injected_role_card path="roles/orchestrator.md">' in agent.system
    assert "Orchestrator 角色卡已由 Harness 完整注入" in agent.system
    assert "Research：omitted" in agent.system
    assert "可委派角色：slide, review" in agent.system
    assert "web_search" not in tool_names
    assert "image_generate" not in tool_names
    rejected = _delegate_task(
        agent,
        {"tasks": [{"goal": "Research: collect sources"}]},
    )
    assert "运行时能力检查已省略" in rejected


def test_agent_tools_remove_search_generation_and_vision() -> None:
    names = {
        schema["name"]
        for schema in tools.agent_tools(
            "image",
            enable_web_search=False,
            enable_image_gen=False,
            enable_vision=False,
        )
    }
    assert "web_search" not in names
    assert "web_extract" not in names
    assert "image_generate" not in names
    assert "vision_analyze" not in names


def test_single_topology_rejects_slide_group(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-3"
        response_language = "zh"
        ws = str(tmp_path)

    result = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "SlideGroup hero [01,02]: build both pages"}]},
    )
    assert "ownership_topology: grouped" in result


def test_grouped_topology_accepts_only_complete_slide_groups(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-3"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(
            Parent(),
            {"tasks": [{"goal": "SlideGroup mechanism [03,04,05]: build group"}]},
        )
    assert result == "delegated"
    specs = delegate.call_args.args[1]["tasks"]
    assert specs == [{
        "role": "slide",
        "label": "slide_group_mechanism",
        "group_id": "mechanism",
        "pages": [3, 4, 5],
        "task": "SlideGroup mechanism [03,04,05]: build group",
    }]

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        isolated = _delegate_task(
            Parent(),
            {"tasks": [{"goal": "SlideGroup closing [13]: build isolated closing"}]},
        )
    assert isolated == "delegated"
    assert delegate.call_args.args[1]["tasks"][0]["pages"] == [13]

    duplicate = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "SlideGroup closing [13,13]: build closing"}]},
    )
    assert "use [13], never [13,13]" in duplicate

    rejected = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "Slide 03: build one page"}]},
    )
    assert "ownership_topology: single" in rejected


def test_structured_delegation_derives_labels_without_goal_syntax(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-3"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [{
            "role": "slide",
            "group_id": "mechanism",
            "pages": [3, 4, 5],
            "goal": "兑现跨页机制编码",
        }]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"] == [{
        "role": "slide",
        "label": "slide_group_mechanism",
        "group_id": "mechanism",
        "pages": [3, 4, 5],
        "task": "兑现跨页机制编码",
    }]


def test_style_lock_is_extracted_once_for_visual_workers(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n\n## 视觉契约\n档案纸纹；禁通用深蓝金。\n\n"
        "## 特殊页\n封面使用真实档案主体。\n\n"
        "## 主题变量\n```css\n:root { --content-canvas: #eee; }\n```\n",
        encoding="utf-8",
    )
    lock = _style_lock_injection(str(tmp_path), "zh")
    assert "<style_lock" in lock
    assert "档案纸纹" in lock
    assert "--content-canvas" in lock
    assert "通用暗底模板" in lock


def test_vision_issue_ledger_requires_changed_pixel_closure(tmp_path: Path) -> None:
    agent = SimpleNamespace(ws=str(tmp_path), trace_label="slide_03", role="slide")
    scan = {
        "visible_subjects": ["one portrait"],
        "text_regions": ["title readable"],
        "regions": ["left image, right copy"],
        "edges": {edge: "clear" for edge in ("top", "right", "bottom", "left")},
    }
    issue = {
        "severity": "major",
        "type": "text_overlap",
        "location": "lower right",
        "evidence": "two labels visibly overlap",
        "suggested_fix": "separate the labels",
    }
    _update_vision_issue_ledger(agent, "renders/slide_03.png", "hash-a", {
        "verdict": "repair_required", "summary": "overlap", "scan": scan,
        "issues": [issue],
    })
    assert len(_open_vision_issues(agent)) == 1

    _update_vision_issue_ledger(agent, "renders/slide_03.png", "hash-a", {
        "verdict": "ready", "summary": "generic ready", "scan": scan, "issues": [],
    })
    assert len(_open_vision_issues(agent)) == 1

    _update_vision_issue_ledger(agent, "renders/slide_03.png", "hash-b", {
        "verdict": "ready", "summary": "overlap is visibly gone", "scan": scan,
        "issues": [],
    })
    assert _open_vision_issues(agent) == []
    ledger = json.loads((tmp_path / "_trace/vision-issues.json").read_text())
    assert ledger["issues"][0]["opened_pixel_sha256"] == "hash-a"
    assert ledger["issues"][0]["closed_pixel_sha256"] == "hash-b"


def test_ready_minor_vision_observation_does_not_open_blocking_ledger(
    tmp_path: Path,
) -> None:
    agent = SimpleNamespace(ws=str(tmp_path), trace_label="image")
    scan = {
        "visible_subjects": [], "text_regions": [], "regions": [],
        "edges": {"top": "ok", "right": "ok", "bottom": "ok", "left": "ok"},
    }
    _update_vision_issue_ledger(agent, "assets/contact-sheet.png", "hash-a", {
        "verdict": "ready",
        "summary": "usable set",
        "scan": scan,
        "issues": [{
            "severity": "minor",
            "type": "contact_sheet_balance",
            "location": "row 2",
            "evidence": "preview grid has an empty cell",
        }],
    })
    payload = json.loads(
        (tmp_path / "_trace/vision-issues.json").read_text(encoding="utf-8")
    )
    assert payload["issues"] == []


def test_delegation_reloads_canonical_topology_instead_of_cached_hint(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    deck = tmp_path / "plan/deck.md"

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-3"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)
        ownership_topology = "grouped"

    parent = Parent()
    deck.write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    with mock.patch("core.agent_loop._delegate", return_value="delegated"):
        result = _delegate_task(
            parent,
            {"tasks": [{"goal": "Slide 01: build cover"}]},
        )
    assert result == "delegated"
    assert parent.ownership_topology == "single"
    assert tools._is_grouped_skill(parent) is False

    deck.write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    with mock.patch("core.agent_loop._delegate", return_value="delegated"):
        result = _delegate_task(
            parent,
            {"tasks": [{"goal": "SlideGroup cover [01]: build cover"}]},
        )
    assert result == "delegated"
    assert parent.ownership_topology == "grouped"
    assert tools._is_grouped_skill(parent) is True


def test_v03_keeps_research_and_revision_contracts(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    agent = Agent(
        "grouped-contract-test",
        str(tmp_path),
        "制作八页演示",
        {
            "skill_name": "mural-presenter-v0-3",
            "skill_language": "zh",
        },
    )
    assert agent.research_required is True

    prompt = _revision_prompt(
        {"query": "介绍台风白海豚", "lang": "zh"},
        {"instruction": "把主题纠正为台风白海豚"},
        grouped=True,
    )
    assert "Review: mode=simple_edit" in prompt
    assert "plan/revision-impact.md" in prompt
    assert "group_id: GROUP" in prompt
    assert "不得把受影响页组拆成单页" in prompt

    class RevisionOrchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        revision_mode = True

    error = tools._model_write_error(
        RevisionOrchestrator(), "slides/slide_02.html"
    )
    assert error is not None
    assert "不允许 Orchestrator 直接修改" in error


def test_image_and_slide_share_dependency_aware_wave(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-3"
        response_language = "zh"
        child_outcomes = {}
        ws = str(tmp_path)

    parent = Parent()
    with mock.patch("core.agent_loop._delegate", return_value="dependency-aware") as delegate:
        result = _delegate_task(
            parent,
            {
                "tasks": [
                    {"goal": "Image: resolve all raster assets"},
                    {"goal": "Slide 03: complete the bitmap-independent page"},
                ]
            },
        )
    assert result == "dependency-aware"
    assert [spec["role"] for spec in delegate.call_args.args[1]["tasks"]] == ["image", "slide"]


def test_bitmap_slide_releases_when_image_finishes_without_waiting_for_other_slides(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("speech\n", encoding="utf-8")
    for page, needs_bitmap in ((1, False), (2, True)):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- production_group: page-{page:02d}\n"
            f"- needs_bitmap: {str(needs_bitmap).lower()}\n",
            encoding="utf-8",
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            "<section></section>\n", encoding="utf-8"
        )

    profile = {
        "active_roles": ["image", "slide", "review"],
        "workflow": {
            "research": "omitted",
            "research_mode": "off",
            "material_stage": "omitted",
            "image": "bitmap_available",
        },
        "tools": {
            "web_search": False,
            "web_extract": False,
            "image_generate": True,
            "vision_analyze": True,
        },
    }
    parent = Agent(
        "dependency-scheduler",
        str(tmp_path),
        "build",
        {
            "skill_name": "mural-presenter-v0-3",
            "skill_language": "zh",
            "_runtime_capabilities": profile,
        },
    )
    parent.research_required = False
    events: list[str] = []
    lock = threading.Lock()

    def fake_child(_parent, _index, spec):
        role = spec["role"]
        label = spec["label"]
        if role == "image":
            time.sleep(0.03)
            with lock:
                events.append("image_done")
        elif label == "slide_01":
            time.sleep(0.15)
            with lock:
                events.append("plain_done")
        else:
            with lock:
                events.append("bitmap_started")
        return {
            "label": label,
            "role": role,
            "ok": True,
            "exit_reason": "completed",
            "completed_pages": list(spec.get("pages") or []),
            "incomplete_pages": [],
            "renders": 0,
            "views": 0,
            "summary": "done",
        }

    specs = [
        {"role": "image", "label": "image", "task": "resolve"},
        {"role": "slide", "label": "slide_01", "pages": [1], "task": "plain"},
        {"role": "slide", "label": "slide_02", "pages": [2], "task": "bitmap"},
    ]
    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        result = json.loads(_delegate(parent, {"tasks": specs}))

    assert result["failed"] == 0
    assert events.index("image_done") < events.index("bitmap_started")
    assert events.index("bitmap_started") < events.index("plain_done")


def test_completed_slide_cannot_be_reopened_in_normal_production(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_01.md").write_text(
        "# P1\n- needs_bitmap: false\n", encoding="utf-8"
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    parent = Agent(
        "repair-reopen",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-3", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.child_outcomes = {"slide_01": {"ok": True}}

    ordinary = json.loads(_delegate(parent, {"tasks": [{
        "role": "slide", "label": "slide_01", "pages": [1], "task": "retry"
    }]}))
    assert ordinary["status"] == "already_completed"

    with mock.patch("core.agent_loop._run_child") as run_child:
        repaired = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide",
            "label": "slide_01",
            "pages": [1],
            "task": "fix confirmed duplicate credit",
            "repair": True,
        }]}))
    assert run_child.call_count == 0
    assert repaired["status"] == "already_completed"
    assert repaired["executed"] == 0


def test_repair_required_slide_is_handed_to_review_without_r2(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_01.md").write_text(
        "# P1\n- needs_bitmap: false\n", encoding="utf-8"
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section>authored</section>\n", encoding="utf-8"
    )
    parent = Agent(
        "repair-handoff",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-3", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.child_outcomes = {"slide_01": {
        "ok": True,
        "status": "repair_required",
        "repair_issue": {
            "status": "repair_required",
            "pages": [1],
            "issue_type": "page_authoring",
        },
    }}
    with mock.patch("core.agent_loop._run_child") as run_child:
        result = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide", "label": "slide_01", "pages": [1], "repair": True,
        }]}))
    assert result["status"] == "already_completed"
    assert result["executed"] == 0
    assert run_child.call_count == 0


def test_one_shot_shared_env_maps_to_isolated_same_model_aux() -> None:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(HARNESS)
    env["VISION_BACKEND"] = "one_shot"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "from core import config; "
            "print(config.VISION_BACKEND_REQUESTED, config.VISION_BACKEND)",
        ],
        capture_output=True,
        text=True,
        env=env,
        check=True,
    )
    assert result.stdout.strip() == "one_shot same_model_aux"


def test_gemini_backend_alias_maps_to_external_model() -> None:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(HARNESS)
    env["VISION_BACKEND"] = "gemini"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "from core import config; print(config.VISION_BACKEND)",
        ],
        capture_output=True,
        text=True,
        env=env,
        check=True,
    )
    assert result.stdout.strip() == "external_model"


def test_same_model_vision_critic_uses_fresh_history_free_request() -> None:
    captured = {}

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[SimpleNamespace(
            type="text",
            text='{"schema":"mural.vision-critic.v1","verdict":"ready",'
                 '"summary":"clear","observations":[],"issues":[]}',
        )])

    agent = SimpleNamespace(
        model="main-multimodal-model",
        cfg={"model_base_url": "http://main-model"},
        log=lambda *_args, **_kwargs: None,
        deadline_monotonic=time.monotonic() + 60,
    )
    with mock.patch("core.model_call.call_with_tools", side_effect=fake_call):
        result = model_call.call_vision_same_model_auxiliary(
            agent=agent,
            image_bytes=b"pixels",
            media_type="image/png",
            query="check clipping",
            system_prompt="isolated critic system",
        )

    assert '"verdict":"ready"' in result
    assert captured["model"] == agent.model
    assert captured["system"] == "isolated critic system"
    assert captured["tools"] is None
    assert captured["thinking"] is False
    assert len(captured["messages"]) == 1
    assert captured["messages"][0]["role"] == "user"
    assert captured["messages"][0]["content"][1]["text"] == "check clipping"


def test_vision_tool_returns_structured_text_from_same_model_aux(tmp_path: Path) -> None:
    from PIL import Image

    renders = tmp_path / "renders"
    renders.mkdir()
    Image.new("RGB", (1600, 900), "white").save(renders / "slide_01.png")
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-3",
        ownership_topology="single",
        slide_pixel_inspections=0,
        assigned_slide_pages=(1,),
        slide_group_id="",
        n_views=0,
        max_vision_edge=1600,
        render_script="skills/mural-presenter-v0-3/scripts/slide.py",
        model="main-multimodal-model",
        cfg={"model_base_url": "http://main-model"},
        deadline_monotonic=time.monotonic() + 60,
        ws=str(tmp_path),
        read_path=lambda path: str(tmp_path / path),
        log=lambda *_args, **_kwargs: None,
    )
    critic = json.dumps({
        "schema": "mural.vision-critic.v1",
        "verdict": "repair_required",
        "summary": "右下图片孤立且过小",
        "observations": ["主标题位于左上"],
        "issues": [{
            "severity": "major",
            "type": "isolated_small_image",
            "location": "right-bottom",
            "evidence": "图片仅占画布约 8%",
            "suggested_fix": "扩大为右侧主视觉区",
        }],
    }, ensure_ascii=False)
    with mock.patch("core.tools.config.VISION_BACKEND", "same_model_aux"), mock.patch(
        "core.tools.model_call.call_vision_same_model_auxiliary",
        return_value=critic,
    ) as auxiliary:
        result = tools.vision_analyze(
            agent,
            image="renders/slide_01.png",
            query="检查孤立小图与四边安全区",
        )

    assert result["vision_verdict"] == "repair_required"
    assert result["vision_result"]["issues"][0]["type"] == "isolated_small_image"
    assert result["vision_backend"] == "same_model_aux:main-multimodal-model"
    assert result["source_sha256"]
    assert auxiliary.call_args.kwargs["system_prompt"].startswith(
        "You are an isolated visual agent."
    )
    assert auxiliary.call_args.kwargs["query"].startswith("FIXED OPEN SCAN")
    assert "检查孤立小图与四边安全区" in auxiliary.call_args.kwargs["query"]


def test_vision_tool_result_keeps_pixels_out_of_main_agent_history(tmp_path: Path) -> None:
    trace_dir = tmp_path / "trace"
    (trace_dir / "images").mkdir(parents=True)
    agent = SimpleNamespace(
        role="image",
        started=time.time(),
        log=lambda *_args, **_kwargs: None,
        remote_tool_concurrency=1,
        ws=str(tmp_path),
        trace_dir=trace_dir,
        n_views=0,
        image_by_tool={},
        vision_critic_results={},
    )
    value = {
        "vision_analysis": '{"verdict":"ready","issues":[]}',
        "vision_result": {"verdict": "ready", "issues": []},
        "vision_verdict": "ready",
        "vision_summary": "asset is usable",
        "image_b64": __import__("base64").b64encode(b"pixels").decode(),
        "media_type": "image/png",
        "path": "assets/item.png",
        "source_sha256": __import__("hashlib").sha256(b"pixels").hexdigest(),
        "vision_backend": "same_model_aux:test",
    }
    use = SimpleNamespace(
        id="tool-vision-1",
        name="vision_analyze",
        input={"image": "assets/item.png", "query": "inspect asset"},
    )
    with mock.patch("core.agent_loop.tools.dispatch", return_value=value):
        results = _tool_results(agent, [use], 1, [])

    assert isinstance(results[0]["content"], str)
    assert '"verdict":"ready"' in results[0]["content"]
    assert "image" not in results[0]["content"]
    assert (trace_dir / "images/view_001.png").read_bytes() == b"pixels"


def test_deepseek_no_thinking_is_explicitly_disabled() -> None:
    captured = {}

    class Response:
        usage = None
        content = []

    class Messages:
        def create(self, **kwargs):
            captured.update(kwargs)
            return Response()

    class Client:
        messages = Messages()

    with mock.patch.object(model_call, "_client", return_value=Client()):
        result = model_call.call_with_tools(
            model="bailian/deepseek-v4-flash-0731",
            system="test",
            messages=[{"role": "user", "content": "test"}],
            tools=None,
            max_tokens=16,
            thinking=False,
            effort="high",
            max_retries=1,
        )

    assert result is not None
    assert captured["thinking"] == {"type": "disabled"}
    assert "output_config" not in captured


def test_research_receives_verbatim_raw_user_query_from_harness(tmp_path: Path) -> None:
    raw_query = '做个 8 页 PPT，介绍台风“白海豚”'
    child = Agent(
        "raw-query-test",
        str(tmp_path),
        "Research: investigate the Chinese white dolphin species",
        {
            "skill_name": "mural-presenter-v0-3",
            "skill_language": "zh",
            "_raw_user_query": raw_query,
        },
        role="research",
        label="research",
    )

    assert json.dumps(raw_query, ensure_ascii=False) in child.system
    assert "委派 goal 只是 Orchestrator 的研究假设" in child.system


def test_research_receives_latest_revision_as_separate_user_context(
    tmp_path: Path,
) -> None:
    child = Agent(
        "revision-query-test",
        str(tmp_path),
        "Research: verify the proposed topic",
        {
            "skill_name": "mural-presenter-v0-3",
            "skill_language": "zh",
            "_raw_user_query": "介绍台风白海豚",
            "_revision_instruction": "不是动物，改回台风主题",
        },
        role="research",
        label="research",
    )

    assert "介绍台风白海豚" in child.system
    assert "不是动物，改回台风主题" in child.system
    assert "明确的最新修改优先" in child.system


def test_revision_prompt_routes_simple_and_complex_edits() -> None:
    prompt = _revision_prompt(
        {"query": "介绍台风白海豚", "lang": "zh"},
        {"instruction": "把主题纠正为台风白海豚"},
    )
    assert "Review: mode=simple_edit" in prompt
    assert "plan/revision-impact.md" in prompt
    assert "主题/实体纠正" in prompt
    assert "不能当成局部换字" in prompt


def test_revision_route_is_locked_and_complex_requires_impact_map(
    tmp_path: Path,
) -> None:
    class Parent:
        revision_mode = True
        revision_route = ""
        ws = str(tmp_path)

    parent = Parent()
    error = _revision_delegation_error(
        parent,
        [{"role": "review", "task": "Review: mode=simple_edit; page 2 title"}],
    )
    assert error == ""
    assert parent.revision_route == "simple_edit"
    assert "只能使用唯一" in _revision_delegation_error(
        parent,
        [{"role": "slide", "task": "Slide 02: rewrite page"}],
    )

    complex_parent = Parent()
    assert "revision-impact.md" in _revision_delegation_error(
        complex_parent,
        [{"role": "research", "task": "Research: resolve the named storm"}],
    )
    impact = tmp_path / "plan/revision-impact.md"
    impact.parent.mkdir(parents=True, exist_ok=True)
    impact.write_text("# Impact\n- facts: all pages\n", encoding="utf-8")
    assert _revision_delegation_error(
        complex_parent,
        [{"role": "research", "task": "Research: resolve the named storm"}],
    ) == ""
    assert complex_parent.revision_route == "complex_edit"


def test_revision_orchestrator_cannot_patch_slide_html_directly() -> None:
    class RevisionOrchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        revision_mode = True

    error = tools._model_write_error(
        RevisionOrchestrator(), "slides/slide_02.html"
    )
    assert error is not None
    assert "不允许 Orchestrator 直接修改" in error


def test_initial_page_plans_use_structured_batch_tool(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- page_count: 8\n",
        encoding="utf-8",
    )

    class Orchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        revision_mode = False
        ws = str(tmp_path)
        render_script = str(
            REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts/orchestrator.py"
        )

        def safe(self, path):
            return str(tmp_path / path)

    agent = Orchestrator()
    single = tools.write_file(agent, "plan/slide_01.md", "# slide 1\n")
    assert "不得逐页串行" in single
    direct_manifest = tools.write_file(agent, "plan/plan-batch.json", "{}")
    assert "write_plan_batch" in direct_manifest

    files = [
        {"path": f"plan/slide_{page:02d}.md", "content": f"# slide {page}\n"}
        for page in range(1, 5)
    ]
    applied = tools.write_plan_batch(agent, files)
    assert "pages=01,02,03,04" in applied
    assert all((plan / f"slide_{page:02d}.md").is_file() for page in range(1, 5))
    assert not (plan / "plan-batch.json").exists()

    stringified = [
        {"path": "plan/slide_05.md", "content": "# slide 5\n"},
        {"path": "plan/slide_06.md", "content": "# slide 6\n"},
    ]
    decoded = tools.write_plan_batch(
        agent,
        json.dumps(stringified, ensure_ascii=False),
    )
    assert "pages=05,06" in decoded

    corrected = [
        {
            "path": f"plan/slide_{page:02d}.md",
            "content": f"# corrected slide {page}\n",
        }
        for page in range(1, 5)
    ]
    rejected_overwrite = tools.write_plan_batch(agent, corrected)
    assert "replace_existing=true" in rejected_overwrite
    assert (plan / "slide_01.md").read_text(encoding="utf-8") == "# slide 1\n"
    assert not (plan / "plan-batch.json").exists()

    recovered = tools.write_plan_batch(
        agent,
        corrected,
        replace_existing=True,
    )
    assert "replace_existing=true" in recovered
    assert (plan / "slide_01.md").read_text(encoding="utf-8") == "# corrected slide 1\n"

    bad_files = [
        {"path": "plan/slide_07.md", "content": "# slide 7\n"},
        {"path": "plan/slide_09.md", "content": "# slide 9\n"},
    ]
    nonconsecutive = tools.write_plan_batch(agent, bad_files)
    assert "[non_contiguous]" in nonconsecutive
    assert not (plan / "slide_07.md").exists()
    assert not (plan / "slide_09.md").exists()

    tool_names = {
        schema["name"]
        for schema in tools.agent_tools(
            "orchestrator", skill_name="mural-presenter-v0-3"
        )
    }
    assert "write_file" in tool_names
    assert "write_plan_batch" in tool_names
    assert "write_files" not in tool_names

    batch_schema = next(
        schema
        for schema in tools.agent_tools(
            "orchestrator", skill_name="mural-presenter-v0-3"
        )
        if schema["name"] == "write_plan_batch"
    )
    assert batch_schema["input_schema"]["properties"]["replace_existing"]["type"] == "boolean"


def test_plan_batch_overwrite_is_blocked_after_page_production_starts(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide"></section>',
        encoding="utf-8",
    )

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        revision_mode = False
        delegated_roles = set()
        ws = str(tmp_path)

    files = [
        {"path": "plan/slide_01.md", "content": "# corrected 1"},
        {"path": "plan/slide_02.md", "content": "# corrected 2"},
    ]
    result = tools.write_plan_batch(
        Orchestrator(),
        files,
        replace_existing=True,
    )
    assert "[production_started]" in result


def test_write_plan_batch_reports_actionable_error_codes(tmp_path: Path) -> None:
    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        revision_mode = False
        ws = str(tmp_path)

    agent = Orchestrator()
    one = [{"path": "plan/slide_01.md", "content": "# one"}]
    assert "[too_few]" in tools.write_plan_batch(agent, one)
    assert "write_file" in tools.write_plan_batch(agent, one)
    seven = [
        {"path": f"plan/slide_{page:02d}.md", "content": f"# {page}"}
        for page in range(1, 8)
    ]
    assert "[too_many]" in tools.write_plan_batch(agent, seven)
    duplicate = [
        {"path": "plan/slide_01.md", "content": "# one"},
        {"path": "plan/slide_01.md", "content": "# duplicate"},
    ]
    assert "[duplicate]" in tools.write_plan_batch(agent, duplicate)
    reversed_pages = [
        {"path": "plan/slide_02.md", "content": "# two"},
        {"path": "plan/slide_01.md", "content": "# one"},
    ]
    assert "[out_of_order]" in tools.write_plan_batch(agent, reversed_pages)


def test_pending_read_only_accepts_exact_continuation_and_clears_at_eof(
    tmp_path: Path,
) -> None:
    path = tmp_path / "research" / "material.md"
    path.parent.mkdir()
    path.write_text(
        "\n".join(f"line {index:04d} " + ("x" * 90) for index in range(300)),
        encoding="utf-8",
    )

    class Research:
        role = "research"
        skill_name = "mural-presenter-v0-3"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Research()
    first = tools.read_file(agent, "research/material.md")
    assert "续读 offset=" in first
    required = tools.pending_read_requirement(agent)
    assert required is not None
    wrong = tools.read_file(agent, "research/material.md", offset=required[1] + 1)
    assert "pending_read_only" in wrong
    while (required := tools.pending_read_requirement(agent)) is not None:
        result = tools.read_file(agent, required[0], offset=required[1])
    assert "pending_read 已解除" in result
    assert tools.pending_read_error(agent) is None


def test_orchestrator_cannot_read_material_when_research_is_enabled(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "research" / "material.md").write_text("raw", encoding="utf-8")

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        research_required = True
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    result = tools.read_file(Orchestrator(), "research/material.md")
    assert "research_route_violation" in result
    assert "knowledge-brief.md" in result


def test_final_single_missing_plan_can_resume_with_write_file(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- page_count: 3\n",
        encoding="utf-8",
    )
    for page in (1, 2):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# slide {page}\n", encoding="utf-8"
        )

    class Orchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        revision_mode = False
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.write_file(
        Orchestrator(), "plan/slide_03.md", "# recovered slide 3\n"
    )
    assert result.startswith("已写入")
    assert (plan / "slide_03.md").is_file()


def test_slide_repair_required_is_a_structured_nonfatal_handoff(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace/slide-render-states"
    trace.mkdir(parents=True)
    (trace / "page_07.json").write_text(
        json.dumps({
            "limit": 8,
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "hashes": ["a", "b", "c"],
            "attempt_hashes": {"slide_07": ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )

    class Slide:
        role = "slide"
        ws = str(tmp_path)
        assigned_slide_pages = (7,)
        repair_required_reason = ""
        final_text = """status: repair_required
pages: 07
issue_type: shared_system
evidence: inherited two-column grid squeezes the primary content
proposed_fix: override the page grid to one column
final_pixels_inspected: yes
"""

    issue = _slide_repair_issue(Slide())
    assert issue is not None
    assert issue["status"] == "repair_required"
    assert issue["pages"] == [7]
    assert issue["issue_type"] == "shared_system"
    assert issue["render_budget_exhausted_pages"] == [7]


def test_review_must_open_flagged_page_current_pixels(tmp_path: Path) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    png = renders / "slide_07.png"
    png.write_bytes(b"current-pixels")

    class Review:
        role = "review"
        ws = str(tmp_path)
        review_contact_sheet_inspected = True
        required_review_pages = (7,)
        review_viewed_page_hashes = {}

    assert "P07" in _review_required_view_gap(Review())
    Review.review_viewed_page_hashes = {
        7: __import__("hashlib").sha256(png.read_bytes()).hexdigest()
    }
    assert _review_required_view_gap(Review()) == ""


def test_review_gets_compact_render_summary_without_pending_pagination(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(
        json.dumps(
            {
                "mode": "incremental-finalize",
                "n_pages": 1,
                "console_errors": [],
                "anti_slop_warnings": ["generic dark canvas"],
                "special_page_geometry": {"warnings": []},
                "pages": [
                    {
                        "page": 1,
                        "png": str(renders / "slide_01.png"),
                        "blank": False,
                        "static": False,
                        "geometry": {
                            "page_type": "content",
                            "text_boxes": [
                                {
                                    "text": "too small",
                                    "font_size": 14,
                                    "overflow_x": False,
                                    "overflow_y": False,
                                }
                            ],
                            "images": [],
                        },
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    class Review:
        role = "review"
        skill_name = "mural-presenter-v0-3"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Review()
    result = tools.read_file(agent, "renders/render.json")
    assert "Review compact render summary" in result
    assert "generic dark canvas" in result
    assert "too small" in result
    assert tools.pending_read_requirement(agent) is None


def test_image_and_dense_pages_are_selected_for_full_resolution_review(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_01.md").write_text(
        "# slide_01\n- needs_bitmap: true\n- composition: visual-split\n",
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_02.md").write_text(
        "# slide_02\n- needs_bitmap: false\n- composition: matrix\n",
        encoding="utf-8",
    )
    (tmp_path / "renders/render.json").write_text(
        json.dumps({
            "pages": [],
            "special_page_geometry": {"warnings": []},
            "layout_defects": {"03": [{"type": "text_collision"}]},
        }),
        encoding="utf-8",
    )
    parent = SimpleNamespace(ws=str(tmp_path))
    assert _mandatory_fullres_review_pages(parent) == {1, 2, 3}


def test_review_cannot_override_current_vision_critic_repair_required(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    png = renders / "slide_07.png"
    png.write_bytes(b"defective-current-pixels")
    digest = __import__("hashlib").sha256(png.read_bytes()).hexdigest()

    class Review:
        role = "review"
        ws = str(tmp_path)
        review_contact_sheet_inspected = True
        required_review_pages = ()
        review_viewed_page_hashes = {}
        vision_critic_results = {
            "renders/slide_07.png": {
                "source_sha256": digest,
                "verdict": "repair_required",
                "summary": "isolated small image remains",
                "issues": [],
            }
        }

    assert "Vision Critic" in _review_required_view_gap(Review())
    png.write_bytes(b"new-pixels-not-yet-criticized")
    assert _review_required_view_gap(Review()) == ""


def test_contract_parser_ignores_free_prose() -> None:
    fields = _contract_fields(
        "The page may be ready later.\nstatus: repair_required\npages: 07\n"
    )
    assert fields == {"status": "repair_required", "pages": "07"}


def test_page_vision_budget_is_three_states_across_redispatch(tmp_path: Path) -> None:
    from PIL import Image

    renders = tmp_path / "renders"
    renders.mkdir()
    image_path = renders / "slide_07.png"

    def slide_agent(label: str):
        return SimpleNamespace(
            role="slide",
            skill_name="mural-presenter-v0-3",
            ownership_topology="single",
            assigned_slide_pages=(7,),
            slide_group_id="",
            trace_label=label,
            n_views=0,
            max_vision_edge=1600,
            model="main-multimodal-model",
            cfg={"model_base_url": "http://main-model"},
            deadline_monotonic=time.monotonic() + 60,
            ws=str(tmp_path),
            read_path=lambda path: str(tmp_path / path),
            log=lambda *_args, **_kwargs: None,
            rendered_output_hashes={7: label},
            viewed_output_hashes={},
            repair_required_reason="",
        )

    critic = json.dumps({
        "schema": "mural.vision-critic.v1",
        "verdict": "ready",
        "summary": "page is ready",
        "observations": [],
        "issues": [],
        "scan": {"visible_subjects": [], "text_regions": [], "regions": [],
                 "edges": {"top": "clear", "right": "clear", "bottom": "clear", "left": "clear"}},
    })
    with mock.patch(
        "core.model_call.call_vision_same_model_auxiliary",
        return_value=critic,
    ) as auxiliary:
        for index, color in enumerate(("red", "green", "blue"), start=1):
            Image.new("RGB", (160, 90), color).save(image_path)
            result = tools.vision_analyze(
                slide_agent("slide_07" if index == 1 else f"slide_07_r{index}"),
                "renders/slide_07.png",
                "inspect current pixels",
            )
            assert result["vision_verdict"] == "ready"
        Image.new("RGB", (160, 90), "yellow").save(image_path)
        final_agent = slide_agent("slide_07_r4")
        blocked = tools.vision_analyze(
            final_agent,
            "renders/slide_07.png",
            "inspect current pixels",
        )

    assert auxiliary.call_count == 3
    assert "repair_required" in blocked
    assert "重新委派不会重置" in blocked
    assert final_agent.repair_required_reason


def test_three_page_group_allows_nine_page_and_three_group_checks(
    tmp_path: Path,
) -> None:
    from PIL import Image

    renders = tmp_path / "renders"
    renders.mkdir()
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-3",
        ownership_topology="grouped",
        assigned_slide_pages=(1, 2, 3),
        slide_group_id="intro",
        trace_label="slide_group_intro",
        n_views=0,
        max_vision_edge=1600,
        model="main-multimodal-model",
        cfg={"model_base_url": "http://main-model"},
        deadline_monotonic=time.monotonic() + 60,
        ws=str(tmp_path),
        read_path=lambda path: str(tmp_path / path),
        log=lambda *_args, **_kwargs: None,
        rendered_output_hashes={},
        viewed_output_hashes={},
        repair_required_reason="",
    )
    critic = json.dumps({
        "schema": "mural.vision-critic.v1",
        "verdict": "ready",
        "summary": "ready",
        "observations": [],
        "issues": [],
        "scan": {"visible_subjects": [], "text_regions": [], "regions": [],
                 "edges": {"top": "clear", "right": "clear", "bottom": "clear", "left": "clear"}},
    })
    with mock.patch(
        "core.model_call.call_vision_same_model_auxiliary",
        return_value=critic,
    ) as auxiliary:
        for revision in range(1, 4):
            for page in (1, 2, 3):
                path = renders / f"slide_{page:02d}.png"
                Image.new("RGB", (160, 90), (revision * 40, page * 50, 80)).save(path)
                agent.rendered_output_hashes[page] = f"page-{page}-state-{revision}"
                result = tools.vision_analyze(
                    agent,
                    f"renders/slide_{page:02d}.png",
                    "inspect page",
                )
                assert result["vision_verdict"] == "ready"
                agent.viewed_output_hashes[page] = agent.rendered_output_hashes[page]
            sheet = renders / "contact-sheet-group-intro.png"
            Image.new("RGB", (300, 180), (revision * 50, 90, 120)).save(sheet)
            result = tools.vision_analyze(
                agent,
                "renders/contact-sheet-group-intro.png",
                "inspect group consistency",
            )
            assert result["vision_verdict"] == "ready"

        cached = tools.vision_analyze(
            agent,
            "renders/contact-sheet-group-intro.png",
            "inspect group consistency again",
        )
        assert cached["vision_cached"] is True
        Image.new("RGB", (300, 180), "black").save(
            renders / "contact-sheet-group-intro.png"
        )
        blocked = tools.vision_analyze(
            agent,
            "renders/contact-sheet-group-intro.png",
            "inspect fourth group state",
        )

    assert auxiliary.call_count == 12
    assert "repair_required" in blocked
    assert config.SLIDE_PAGE_MAX_INSPECTIONS == 3
    assert config.SLIDE_GROUP_MAX_INSPECTIONS == 3


def test_group_delivery_requires_current_pages_then_current_group_sheet(
    tmp_path: Path,
) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan").mkdir()
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    (tmp_path / "base.css").write_text(".slide{color:black}", encoding="utf-8")
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    rendered_hashes = {}
    paths = {}
    for page in (1, 2, 3):
        html = tmp_path / "slides" / f"slide_{page:02d}.html"
        html.write_text(
            f'<section class="slide"><main>page {page}</main></section>',
            encoding="utf-8",
        )
        png = tmp_path / "renders" / f"slide_{page:02d}.png"
        png.write_bytes(f"pixels-{page}".encode())
        render_digest = _render_input_digest(
            tmp_path,
            (tmp_path / "base.css").read_text(encoding="utf-8"),
            html.read_text(encoding="utf-8"),
        )
        file_digest = hashlib.sha256(html.read_bytes()).hexdigest()
        rendered_hashes[page] = file_digest
        paths[page] = str(html)
        (tmp_path / "_trace/slide-render-states" / f"page_{page:02d}.json").write_text(
            json.dumps({
                "hashes": [render_digest],
                "authoring_hashes": [render_digest],
                "attempt_hashes": {"slide_group_intro": [render_digest]},
                "published_hash": render_digest,
                "authoring_attempt_limit": 3,
                "quality_action_count": 0,
            }),
            encoding="utf-8",
        )
    sheet = tmp_path / "renders/contact-sheet-group-intro.png"
    sheet.write_bytes(b"group-pixels")
    sheet_digest = hashlib.sha256(sheet.read_bytes()).hexdigest()
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-3",
        ownership_topology="grouped",
        ws=str(tmp_path),
        assigned_slide_pages=(1, 2, 3),
        slide_group_id="intro",
        expected_output_paths=paths,
        expected_output_initial_hashes={page: "scaffold" for page in paths},
        expected_output_path="",
        rendered_output_hashes=rendered_hashes,
        viewed_output_hashes=dict(rendered_hashes),
        group_viewed_contact_hash="",
        group_viewed_page_hashes={},
        n_renders=3,
        n_views=3,
    )

    assert "未检查当前组联系表" in _slide_deliverable_gap(agent)
    agent.group_viewed_contact_hash = sheet_digest
    agent.group_viewed_page_hashes = dict(rendered_hashes)
    assert _slide_deliverable_gap(agent) == ""


def test_bounded_inspected_review_is_not_a_whole_deck_failure() -> None:
    class Parent:
        finalize_succeeded = True

    result = {
        "ok": False,
        "status": "needs_orchestrator",
        "blocking": "no",
        "final_pixels_inspected": True,
        "final_view_after_review": True,
        "review_changed": False,
        "attempt": 3,
    }
    assert _review_can_complete_needs_improvement(Parent(), result)
    result["blocking"] = "yes"
    assert not _review_can_complete_needs_improvement(Parent(), result)
    result["blocking"] = "no"
    result["final_pixels_inspected"] = False
    assert not _review_can_complete_needs_improvement(Parent(), result)


def test_review_attempt_count_survives_orchestrator_resume(tmp_path: Path) -> None:
    root = tmp_path / "_trace/subagents"
    for label in ("review", "review_retry2", "review_verify3"):
        (root / label).mkdir(parents=True)
    parent = SimpleNamespace(ws=str(tmp_path))
    assert _persisted_trace_attempts(parent, "review") == 3
    assert config.REVIEW_MAX_ATTEMPTS == 3


def test_review_cannot_repeat_until_delivery_surface_changes(tmp_path: Path) -> None:
    for relative in ("plan", "slides", "renders", "assets"):
        (tmp_path / relative).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section>before</section>\n", encoding="utf-8"
    )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet")
    parent = Agent(
        "review-fingerprint",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-3", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = True
    parent.child_outcomes = {"review": {
        "ok": False,
        "status": "needs_orchestrator",
        "exit_reason": "needs_orchestrator",
        "input_fingerprint": _review_delivery_fingerprint(str(tmp_path)),
    }}

    unchanged = _delegate(parent, {"tasks": [{"role": "review"}]})
    assert "均未发生变化" in unchanged
    assert "禁止在相同输入上重复验收" in unchanged

    # A render-capture repair may change only final pixels, not HTML/assets.
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet-after-repair")
    captured: list[dict] = []

    def fake_child(_parent, _index, spec):
        captured.append(dict(spec))
        return {
            "label": "review",
            "trace_label": "review_verify2",
            "role": "review",
            "ok": True,
            "status": "ready",
            "exit_reason": "text_response",
            "completed_pages": [],
            "incomplete_pages": [],
            "renders": 0,
            "views": 1,
            "final_render_after_review": True,
            "final_view_after_review": True,
            "finalize_attempted": True,
            "finalize_succeeded": True,
            "summary": "verified",
        }

    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        verified = json.loads(_delegate(parent, {"tasks": [{"role": "review"}]}))
    assert verified["status"] == "completed"
    assert captured[0]["_recovery_kind"] == "verification"


def test_review_round_must_start_from_current_contact_sheet(tmp_path: Path) -> None:
    review = SimpleNamespace(
        role="review",
        review_contact_sheet_inspected=False,
        ws=str(tmp_path),
        read_path=lambda path: str(tmp_path / path),
    )
    result = tools.vision_analyze(
        review,
        "renders/slide_01.png",
        "inspect one page",
    )
    assert "必须先打开当前 renders/contact-sheet.png" in result


def test_finish_gap_nudges_only_incomplete_single_pages(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (plan / "slide_01.md").write_text("# P1\n", encoding="utf-8")
    (plan / "slide_02.md").write_text("# P2\n", encoding="utf-8")

    class Agent:
        role = "orchestrator"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = True
        delegated_roles = {"slide"}
        child_outcomes = {
            "slide_01": {"ok": True},
            "slide_02": {"ok": False},
        }
        ws = str(tmp_path)

    gap = _finish_gap(Agent())
    assert "Slide 02:" in gap
    assert "Slide 01:" not in gap
    assert "不要创建新的 Slide Agent" in gap
    assert "operational retry" in gap


def test_finish_gap_retries_only_incomplete_complete_group(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    for page, group in ((1, "opening"), (2, "mechanism"), (3, "mechanism")):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n- production_group: {group}\n",
            encoding="utf-8",
        )

    class GroupedAgent:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-3"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = True
        revision_mode = False
        delegated_roles = {"slide"}
        child_outcomes = {
            "slide_group_opening": {
                "ok": True,
                "completed_pages": [1],
            },
            "slide_group_mechanism": {
                "ok": False,
                "completed_pages": [],
                "incomplete_pages": [2, 3],
            },
        }
        ws = str(tmp_path)

    gap = _finish_gap(GroupedAgent())
    assert "SlideGroup mechanism [02,03]:" in gap
    assert "SlideGroup opening" not in gap
    assert "不要创建新的 SlideGroup 实例" in gap
    assert "operational retry" in gap


def test_v03_has_one_skill_with_both_ownership_topologies() -> None:
    root = REPO / "skills/mural-presenter-v0.3"
    editions = sorted(
        path.name for path in root.iterdir()
        if path.is_dir() and (path / "SKILL.md").is_file()
    )
    assert editions == ["mural-presenter-v0-3"]
    content = (root / editions[0] / "SKILL.md").read_text(encoding="utf-8")
    assert "ownership_topology: single | grouped" in content
    assert "role=slide, pages=[NN]" in content
    assert "role=slide, group_id=GROUP" in content


def test_slide_stop_line_prevents_blind_edits_after_unresolved_third_view(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace" / "slide-render-states"
    trace.mkdir(parents=True)
    png = tmp_path / "renders" / "slide_01.png"
    png.parent.mkdir()
    png.write_bytes(b"current-pixels")
    attempt = "slide_01"
    (trace / "page_01.json").write_text(
        json.dumps({
            "limit": 8,
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "attempt_hashes": {attempt: ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )

    class SlideAgent:
        role = "slide"
        ws = str(tmp_path)
        trace_label = attempt
        assigned_slide_pages = (1,)
        repair_required_reason = "Vision Critic unresolved: text overlaps image"
        vision_critic_results = {}

    assert _slide_authoring_stop_line(SlideAgent()) is True
    SlideAgent.repair_required_reason = ""
    assert _slide_authoring_stop_line(SlideAgent()) is False


def test_research_brief_is_bounded_before_orchestrator_handoff(
    tmp_path: Path,
) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-3"
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.write_file(
        ResearchAgent(),
        "research/knowledge-brief.md",
        "x" * 10001,
    )
    assert "knowledge_brief_too_long" in result
    assert not (tmp_path / "research/knowledge-brief.md").exists()


def test_image_requires_contact_sheet_and_each_catalog_asset_at_full_resolution(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## hero\n- path: assets/hero.png\n",
        encoding="utf-8",
    )
    (assets / "hero.png").write_bytes(b"hero")
    (assets / "contact-sheet.png").write_bytes(b"sheet")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        vision_critic_results = {}

    agent = ImageAgent()
    gap = _image_required_fullres_gap(agent)
    assert "assets/contact-sheet.png" in gap
    assert "assets/hero.png" in gap
    agent.vision_critic_results = {
        "assets/contact-sheet.png": {
            "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
            "verdict": "ready",
        },
        "assets/hero.png": {
            "source_sha256": hashlib.sha256(b"hero").hexdigest(),
            "verdict": "ready",
        },
    }
    assert _image_required_fullres_gap(agent) == ""


def test_plan_validator_enforces_selected_ownership_topology(tmp_path: Path) -> None:
    script = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts/orchestrator.py"
    plan = tmp_path / "plan"
    plan.mkdir()

    def write_deck(topology: str) -> None:
        (plan / "deck.md").write_text(
            f"""# Deck Plan
- title: Ownership test

## Resolved deck brief
- language: en
- page_count: 3
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: deterministic topology fixture uses only typography
- ownership_topology: {topology}
- ownership_rationale: dependency structure
- rationale: deterministic test

## Audience and objective
Verify ownership.

## Narrative and page map
Three pages.

## Visual storyboard
Code visuals.

## Visual contract
Test system.

## Special pages
None.

## Theme Tokens
```css
:root {{
  --content-canvas: #eeeeee;
}}
```
""",
            encoding="utf-8",
        )

    def write_slides(groups: list[str]) -> None:
        for page_no, group in enumerate(groups, start=1):
            (plan / f"slide_{page_no:02d}.md").write_text(
                f"""# slide_{page_no:02d}
- role: test page
- page_type: content
- page_family: editorial-test
- production_group: {group}
- composition: editorial
- needs_bitmap: false

## Narrative
Advance the test.

## Evidence
- Deterministic fixture.

## On-screen copy (exact)
- title: Page {page_no}

## Composition blueprint
- focal: title

## Semantic visual need
Use a deliberate typographic hierarchy without raster assets.

## Speech beat
Explain page {page_no}.
""",
                encoding="utf-8",
            )

    write_deck("single")
    write_slides(["shared", "shared", "page-03"])
    invalid_single = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert invalid_single.returncode != 0
    assert "one unique production_group per page" in invalid_single.stderr

    write_slides(["page-01", "page-02", "page-03"])
    valid_single = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert valid_single.returncode == 0, valid_single.stderr
    assert "ownership:single" in valid_single.stdout

    trace = tmp_path / "_trace"
    trace.mkdir()
    (trace / "runtime-capabilities.json").write_text(
        json.dumps({"requirements": {"bitmap_required": True}}),
        encoding="utf-8",
    )
    required_bitmap = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert required_bitmap.returncode != 0
    assert "cannot be silently downgraded" in required_bitmap.stderr
    (trace / "runtime-capabilities.json").unlink()

    write_deck("grouped")
    write_slides(["mechanism", "isolated", "mechanism"])
    invalid_grouped = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert invalid_grouped.returncode != 0
    assert "contiguous pages" in invalid_grouped.stderr

    write_slides(["mechanism", "mechanism", "closing"])
    valid_grouped = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert valid_grouped.returncode == 0, valid_grouped.stderr
    assert "ownership:grouped" in valid_grouped.stdout


def test_plan_validator_warns_on_excessive_section_dividers(tmp_path: Path) -> None:
    script = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts/orchestrator.py"

    def make_fixture(name: str, page_types: list[str]) -> Path:
        root = tmp_path / name
        plan = root / "plan"
        plan.mkdir(parents=True)
        page_count = len(page_types)
        (plan / "deck.md").write_text(
            f"""# Deck Plan
- title: Divider quota test

## Resolved deck brief
- language: en
- page_count: {page_count}
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: deterministic structural fixture uses only typography
- ownership_topology: single
- ownership_rationale: every page is independently owned in this fixture
- rationale: validate special-page pacing without blocking delivery

## Audience and objective
Verify divider pacing.

## Narrative and page map
{page_count} deterministic pages.

## Visual storyboard
Code visuals.

## Visual contract
Test system.

## Special pages
Cover, closing, and only the listed dividers.

## Theme Tokens
```css
:root {{
  --content-canvas: #eeeeee;
}}
```
""",
            encoding="utf-8",
        )
        divider_index = 0
        for page_no, page_type in enumerate(page_types, start=1):
            special = page_type in {"cover", "section-divider", "closing"}
            if page_type == "section-divider":
                divider_index += 1
            special_meta = "- special_layout: centered\n" if special else "- composition: editorial\n"
            evidence = "" if special else "\n## Evidence\n- Deterministic fixture.\n"
            section_index = (
                f"\n- section_index: {divider_index:02d}"
                if page_type == "section-divider"
                else ""
            )
            (plan / f"slide_{page_no:02d}.md").write_text(
                f"""# slide_{page_no:02d}
- role: test page
- page_type: {page_type}
- page_family: editorial-test
- production_group: page-{page_no:02d}
{special_meta}- needs_bitmap: false

## Narrative
Advance the test.{evidence}
## On-screen copy (exact)
- title: Page {page_no}{section_index}

## Composition blueprint
- focal: title

## Semantic visual need
Use a deliberate typographic hierarchy without raster assets.

## Speech beat
Explain page {page_no}.
""",
                encoding="utf-8",
            )
        return root

    cases = {
        "short": ["cover", "content", "content", "content", "closing"],
        "balanced": [
            "cover",
            "section-divider",
            "content",
            "content",
            "content",
            "section-divider",
            "content",
            "content",
            "content",
            "closing",
        ],
        "over": [
            "cover",
            "section-divider",
            "content",
            "section-divider",
            "content",
            "section-divider",
            "content",
            "closing",
        ],
    }
    results = {}
    for name, page_types in cases.items():
        root = make_fixture(name, page_types)
        results[name] = subprocess.run(
            [
                sys.executable,
                str(script),
                "validate-plans",
                str(root),
                "--expected",
                str(len(page_types)),
            ],
            capture_output=True,
            text=True,
        )
        assert results[name].returncode == 0, results[name].stderr

    for name in ("short", "balanced"):
        assert "section-divider quota exceeded" not in results[name].stdout
        assert "special-page ratio exceeds 40%" not in results[name].stdout
    assert "[plan-warning] section-divider quota exceeded" in results["over"].stdout
    assert "[plan-warning] special-page ratio exceeds 40%" in results["over"].stdout


def test_revision_finish_gap_does_not_require_unchanged_pages(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    for page in range(1, 4):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n", encoding="utf-8"
        )

    class RevisionAgent:
        role = "orchestrator"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = False
        revision_mode = True
        revision_route = "complex_edit"
        delegated_roles = {"slide"}
        child_outcomes = {"slide_02": {"ok": True}}
        review_completed = False
        finalize_attempted = False
        finalize_succeeded = False
        ws = str(tmp_path)

    gap = _finish_gap(RevisionAgent())
    assert "Review 尚未完成" in gap
    assert "Slide 01:" not in gap
    assert "Slide 03:" not in gap


def test_pdf_ingestion_writes_complete_text_and_page_derivatives(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    source = inputs / "paper.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Principal finding: attachment evidence is preserved.")
    document.save(source)
    document.close()

    companion = inputs / "paper.pdf.md"
    assert _extract_material(source, companion)
    text = companion.read_text(encoding="utf-8")
    assert "Principal finding" in text
    assert "text_mode: native" in text
    assert (inputs / "paper.pdf.pages/page_001.png").is_file()
    layout = json.loads(
        (inputs / "paper.pdf.pages/page_001.json").read_text(encoding="utf-8")
    )
    assert layout["text_mode"] == "native"
    assert layout["image_pixels"][0] >= 1400


def test_office_ingestion_preserves_text_formulas_and_embedded_media(
    tmp_path: Path,
) -> None:
    from PIL import Image
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorksheetImage
    Presentation = pytest.importorskip("pptx").Presentation
    from pptx.util import Inches

    inputs = tmp_path / "inputs"
    inputs.mkdir()
    picture = inputs / "evidence.png"
    Image.new("RGB", (320, 180), "#3c78d8").save(picture)

    docx_path = inputs / "brief.docx"
    doc = Document()
    doc.add_heading("Policy brief", level=1)
    doc.add_paragraph("Target adoption is 42 percent.")
    doc.add_picture(str(picture))
    doc.save(docx_path)
    assert _extract_material(docx_path, inputs / "brief.docx.md")
    docx_text = (inputs / "brief.docx.md").read_text(encoding="utf-8")
    assert "Target adoption is 42 percent" in docx_text
    assert "embedded_media:" in docx_text
    assert list((inputs / "brief.docx.pages").glob("embedded_*.png"))

    pptx_path = inputs / "roadmap.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
    textbox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    textbox.text = "Launch timeline: pilot in Q3"
    deck.save(pptx_path)
    assert _extract_material(pptx_path, inputs / "roadmap.pptx.md")
    pptx_text = (inputs / "roadmap.pptx.md").read_text(encoding="utf-8")
    assert "Launch timeline: pilot in Q3" in pptx_text
    assert list((inputs / "roadmap.pptx.pages").glob("embedded_*.png"))

    xlsx_path = inputs / "portfolio.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(["Project", "Budget", "Double"])
    sheet.append(["Aurora", 125, "=B2*2"])
    sheet.add_image(WorksheetImage(str(picture)), "E2")
    workbook.save(xlsx_path)
    assert _extract_material(xlsx_path, inputs / "portfolio.xlsx.md")
    xlsx_text = (inputs / "portfolio.xlsx.md").read_text(encoding="utf-8")
    assert "Project | Budget | Double" in xlsx_text
    assert "Aurora | 125 | =B2*2" in xlsx_text
    assert list((inputs / "portfolio.xlsx.pages").glob("embedded_*.png"))


def test_material_staging_failure_is_structured_and_blocks_planning(
    tmp_path: Path,
) -> None:
    source = tmp_path / "unknown.bin"
    source.write_bytes(b"not a supported attachment")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)]}
    staged = stage_materials(workspace, seed)
    assert staged == ["inputs/01_unknown.bin"]
    failures = seed["_material_ingestion_failures"]
    assert failures[0]["error"] == "unsupported attachment type"
    _material_blocked(str(workspace), failures)
    blocked = (workspace / "research/material.md").read_text(encoding="utf-8")
    assert "status: material_blocked" in blocked
    assert "unknown.bin" in blocked


def test_standalone_image_attachment_routes_directly_to_image(
    tmp_path: Path,
) -> None:
    from PIL import Image

    source = tmp_path / "diagram.png"
    Image.new("RGB", (640, 360), "white").save(source)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)], "query": "用这张图制作一套演示"}
    assert stage_materials(workspace, seed) == ["inputs/01_diagram.png"]
    assert seed["_material_ingestion_failures"] == []
    profile = seed["_attachment_profile"]
    assert profile["material_stage"] == "omitted"
    assert profile["material_agent_required"] is False
    assert profile["visual_asset_paths"] == ["inputs/01_diagram.png"]

    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        runtime = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "_raw_user_query": seed["query"],
                "_attachment_profile": profile,
            },
            seed["_staged_materials"] if "_staged_materials" in seed else ["inputs/01_diagram.png"],
            plan_only=True,
        )
    assert "material" not in runtime["active_roles"]
    assert "image" in runtime["active_roles"]


def test_markdown_attachment_uses_direct_text_handoff_without_material(tmp_path: Path) -> None:
    source = tmp_path / "brief.md"
    source.write_text("# Launch brief\n\nRevenue target: 18%.\n", encoding="utf-8")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)], "query": "把附件做成演示"}
    staged = stage_materials(workspace, seed)
    profile = seed["_attachment_profile"]
    assert staged == ["inputs/01_brief.md"]
    assert profile["material_stage"] == "direct_text"
    assert profile["material_agent_required"] is False
    handoff = (workspace / "research/material.md").read_text(encoding="utf-8")
    assert "Harness generated this deterministic text handoff" in handoff
    assert "Revenue target: 18%." in handoff

    with mock.patch.dict(os.environ, {"SERPER_API_KEY": "search-key"}):
        runtime = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "_raw_user_query": seed["query"],
                "_attachment_profile": profile,
                "_evidence_scope_explicit": False,
            },
            staged,
            plan_only=True,
        )
    assert "research" not in runtime["active_roles"]
    assert runtime["workflow"]["research_reason"] == "direct_text_complete_no_unresolved"
    injected = _direct_text_injection(str(workspace), "zh")
    assert 'mode="verbatim"' in injected
    assert "Revenue target: 18%." in injected
    assert "无需分页读取" in injected


@pytest.mark.parametrize(
    ("query", "expected_stage", "visual", "style"),
    [
        ("读取这张截图里的数据并制作演示", "agent_required", False, False),
        ("根据这张图片的风格设计整套演示", "omitted", False, True),
        ("分析这张图的内容，并把原图放在封面", "mixed", True, False),
    ],
)
def test_image_attachment_intent_controls_material_route(
    tmp_path: Path,
    query: str,
    expected_stage: str,
    visual: bool,
    style: bool,
) -> None:
    from PIL import Image

    source = tmp_path / "input.png"
    Image.new("RGB", (640, 360), "#8899aa").save(source)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)], "query": query}
    stage_materials(workspace, seed)
    profile = seed["_attachment_profile"]
    assert profile["material_stage"] == expected_stage
    assert bool(profile["visual_asset_paths"]) is visual
    assert bool(profile["style_reference_paths"]) is style
    assert profile["material_agent_required"] is (expected_stage in {"agent_required", "mixed"})


def test_scanned_pdf_blocks_before_agent_when_main_model_is_text_only(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(source)
    document.close()
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"materials": [str(source)]}
    with mock.patch("core.run_batch.acfg.VISION_BACKEND", "disabled"):
        stage_materials(workspace, seed)
    assert seed["_material_ingestion_failures"]
    assert "text-only" in seed["_material_ingestion_failures"][0]["error"]


def test_attachment_only_is_a_tool_level_research_scope(tmp_path: Path) -> None:
    tool_names = {
        schema["name"]
        for schema in tools.agent_tools(
            "research", evidence_scope="attachment_only"
        )
    }
    assert "web_search" not in tool_names
    assert "web_extract" not in tool_names

    class Orchestrator:
        role = "orchestrator"
        material_required = True
        material_completed = False

    class Research:
        role = "research"

    assert "不得直接读取" in tools.read_file(Orchestrator(), "inputs/report.pdf")
    assert "不直接重读" in tools.read_file(Research(), "inputs/report.pdf.md")


def test_plan_only_acceptance_requires_material_research_and_plans(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "plan").mkdir()
    (tmp_path / "research/material.md").write_text("status: ready\n", encoding="utf-8")
    (tmp_path / "research/knowledge-brief.md").write_text("# Evidence\n", encoding="utf-8")
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    for page in (1, 2):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n", encoding="utf-8"
        )

    class PlanOnlyAgent:
        child_outcomes = {}
        tool_policy_violations = []
        workspace_policy_violations = []

    with mock.patch.dict(os.environ, {"CLEAN_PLAN_ONLY_EVAL": "1"}):
        ok, reason, detail = _accept(PlanOnlyAgent(), str(tmp_path))
    assert ok is True
    assert reason == "plan-only deliverables complete"
    assert detail["n_slide_plans"] == 2


def test_material_blocked_is_a_terminal_structured_acceptance_state(
    tmp_path: Path,
) -> None:
    class BlockedAgent:
        role = "orchestrator"
        skill_language = "zh"
        material_blocked = True
        material_required = True
        material_completed = False
        child_outcomes = {"material": {"ok": False, "status": "material_blocked"}}
        tool_policy_violations = []
        workspace_policy_violations = []

    assert _finish_gap(BlockedAgent()) == ""
    ok, reason, detail = _accept(BlockedAgent(), str(tmp_path))
    assert ok is False
    assert reason == "material_blocked"
    assert detail["material_blocked"] is True


def test_attachment_figure_crop_and_material_catalog_are_supported(tmp_path: Path) -> None:
    from PIL import Image

    page_dir = tmp_path / "inputs/paper.pdf.pages"
    page_dir.mkdir(parents=True)
    source = page_dir / "page_001.png"
    Image.new("RGB", (1600, 2000), "white").save(source)
    source.with_suffix(".json").write_text(
        json.dumps({"page_points": [800, 1000], "text_blocks": []}),
        encoding="utf-8",
    )
    script = (
        REPO
        / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts/image.py"
    )
    crop = subprocess.run(
        [
            sys.executable,
            str(script),
            "crop-material",
            str(tmp_path),
            "--source",
            "inputs/paper.pdf.pages/page_001.png",
            "--output",
            "assets/figure.png",
            "--box",
            "0.10,0.10,0.90,0.35",
        ],
        capture_output=True,
        text=True,
    )
    assert crop.returncode == 0, crop.stderr
    assert '"status": "PASS"' in crop.stdout
    assert (tmp_path / "assets/figure.png").is_file()

    full_page = subprocess.run(
        [
            sys.executable,
            str(script),
            "crop-material",
            str(tmp_path),
            "--source",
            "inputs/paper.pdf.pages/page_001.png",
            "--output",
            "assets/full-page.png",
            "--box",
            "0.00,0.00,1.00,1.00",
        ],
        capture_output=True,
        text=True,
    )
    assert full_page.returncode != 0
    assert "page facsimile" in full_page.stderr

    assets = tmp_path / "assets"
    assets.mkdir(exist_ok=True)
    (assets / "catalog.md").write_text(
        """# Asset catalog

## attachment-figure
- slides: 1
- kind: material
- path: assets/figure.png
- source: inputs/paper.pdf#page=1
- purpose: copied from attachment
""",
        encoding="utf-8",
    )
    accepted_catalog = subprocess.run(
        [sys.executable, str(script), "fetch", str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert accepted_catalog.returncode == 0, accepted_catalog.stderr
    assert "status:PASS" in accepted_catalog.stdout


def test_image_terminal_exposes_material_figure(tmp_path: Path) -> None:
    from PIL import Image

    script = (
        REPO
        / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts/image.py"
    )
    page_dir = tmp_path / "inputs/paper.pdf.pages"
    page_dir.mkdir(parents=True)
    Image.new("RGB", (1600, 2000), "white").save(page_dir / "page_001.png")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-3"
        ws = str(tmp_path)
        render_script = str(script)
        bash_relaxed = False
        bash_timeout = 30

    result = tools.terminal(
        ImageAgent(),
        f"python {script} crop-material . "
        "--source inputs/paper.pdf.pages/page_001.png "
        "--output assets/figure.png --box 0.1,0.1,0.9,0.35",
    )
    assert '"status": "PASS"' in result


def test_image_terminal_registers_user_image_and_catalog_accepts_it(tmp_path: Path) -> None:
    from PIL import Image

    script = (
        REPO
        / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts/image.py"
    )
    inputs = tmp_path / "inputs"
    inputs.mkdir(parents=True)
    Image.new("RGB", (800, 500), "#557799").save(inputs / "photo.jpg")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-3"
        ws = str(tmp_path)
        render_script = str(script)
        bash_relaxed = False
        bash_timeout = 30

    result = tools.terminal(
        ImageAgent(),
        f"python {script} register-user . "
        "--source inputs/photo.jpg --output assets/cover-photo.png",
    )
    assert '"status": "PASS"' in result
    assert (tmp_path / "assets/cover-photo.png").is_file()
    registry = json.loads(
        (tmp_path / "_trace/user-images.json").read_text(encoding="utf-8")
    )
    assert "assets/cover-photo.png" in registry["assets"]
    (tmp_path / "assets/catalog.md").write_text(
        """# Asset catalog

## cover-photo
- slides: 1
- kind: user
- path: assets/cover-photo.png
- source: inputs/photo.jpg
- purpose: user-supplied cover image
""",
        encoding="utf-8",
    )
    accepted = subprocess.run(
        [sys.executable, str(script), "fetch", str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert accepted.returncode == 0, accepted.stderr
    assert "status:PASS" in accepted.stdout


def test_v03_portable_font_bundle_is_self_contained(tmp_path: Path) -> None:
    edition = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3"
    shutil.copyfile(edition / "assets/base.css", tmp_path / "base.css")
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "slide_01.html").write_text(
        '<section class="slide"><h1 class="type-heavy">RAVE 渲染一致</h1></section>',
        encoding="utf-8",
    )

    result = subprocess.run(
        [sys.executable, str(edition / "scripts/_internal/font_bundle.py"), str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr
    css = (tmp_path / "base.css").read_text(encoding="utf-8")
    assert "DECK_FONT_BUNDLE_START" in css
    assert '@font-face {' in css
    manifest = json.loads(
        (tmp_path / "assets/fonts/manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["faces"]
    assert all((tmp_path / face["path"]).is_file() for face in manifest["faces"])


@pytest.mark.parametrize(
    ("script_name", "allowed", "forbidden"),
    [
        ("orchestrator.py", {"validate-plans", "scaffold", "sync-speech", "finalize", "audit"}, {"render", "fetch"}),
        ("image.py", {"register-user", "crop-material", "fetch", "finalize", "inspect", "remove-checkerboard"}, {"audit", "render"}),
        ("slide.py", {"render", "render-group"}, {"finalize", "audit", "fetch"}),
        ("review.py", {"sync-speech", "finalize"}, {"render", "audit", "fetch"}),
    ],
)
def test_role_entry_points_expose_only_their_public_actions(
    script_name: str,
    allowed: set[str],
    forbidden: set[str],
) -> None:
    scripts = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts"
    completed = subprocess.run(
        [sys.executable, str(scripts / script_name), "--help"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    for action in allowed:
        assert action in completed.stdout
    for action in forbidden:
        assert action not in completed.stdout


def test_v03_has_no_public_deck_entry_and_internal_is_not_model_readable(
    tmp_path: Path,
) -> None:
    scripts = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts"
    assert not (scripts / "deck.py").exists()

    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-3",
        ws=str(tmp_path),
    )
    blocked = tools.read_file(
        agent,
        "skills/mural-presenter-v0-3/scripts/_internal/deck_core.py",
    )
    assert "只执行能力面" in blocked


def test_v03_role_script_policy_blocks_cross_role_and_internal_execution(
    tmp_path: Path,
) -> None:
    scripts = REPO / "skills/mural-presenter-v0.3/mural-presenter-v0-3/scripts"
    image_script = scripts / "image.py"
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-3",
        ws=str(tmp_path),
        render_script=str(image_script),
        bash_relaxed=False,
        bash_timeout=30,
    )
    cross_role = tools.terminal(
        agent,
        f"python {scripts / 'orchestrator.py'} audit .",
    )
    assert "只执行能力面" in cross_role
    wrong_action = tools.terminal(agent, f"python {image_script} audit .")
    assert "不允许动作" in wrong_action
    internal = tools.terminal(
        agent,
        f"python {scripts / '_internal/bootstrap.py'} prepare .",
    )
    assert "只执行能力面" in internal
    internal_read = tools.terminal(
        agent,
        f"cat {scripts / '_internal/deck_core.py'}",
    )
    assert "只执行能力面" in internal_read


def test_material_and_research_receive_no_terminal_schema() -> None:
    for role in ("material", "research"):
        names = {
            schema["name"]
            for schema in tools.agent_tools(
                role,
                render_script="",
                skill_name="mural-presenter-v0-3",
            )
        }
        assert "terminal" not in names
