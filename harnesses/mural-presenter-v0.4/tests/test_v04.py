from __future__ import annotations

import json
import hashlib
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from pathlib import Path
from types import SimpleNamespace

import pymupdf
import pytest
from unittest import mock

HARNESS = Path(__file__).resolve().parents[1]
REPO = HARNESS.parents[1]
sys.path.insert(0, str(HARNESS))
V04_SCRIPTS = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts"
sys.path.insert(0, str(V04_SCRIPTS))

from core import tools  # noqa: E402
from core import config  # noqa: E402
from core import model_call  # noqa: E402
from core.agent_loop import (  # noqa: E402
    Agent,
    _ProgressGuard,
    _accept,
    _auto_retry_unstarted_slide_interruptions,
    _auto_continue_read_allowed,
    _call,
    _contract_fields,
    _delegate,
    _delegate_task,
    _direct_text_injection,
    _finish_gap,
    _image_handoff_declared_pages,
    _claim_image_repair_evidence,
    _image_catalog_has_current_fullres_coverage,
    _image_required_fullres_gap,
    _image_route_fingerprint,
    _claim_image_source_route_repair,
    _invalidate_review_after_image_repair,
    _recover_image_operational_handoff,
    _mandatory_fullres_review_pages,
    _material_blocked_decision,
    _material_handoff_gap,
    _material_has_ingestion_evidence,
    _open_vision_issues,
    _orchestrator_delivery_close_status,
    _persisted_trace_attempts,
    _reconcile_review_issues_ledger,
    _render_console_errors,
    _pages_with_active_media,
    _review_can_complete_needs_improvement,
    _review_budget_pending_page_views,
    _review_in_session_closure_nudge_limit,
    _review_has_current_contact_delivery,
    _review_has_current_inspected_delivery,
    _review_manifest_quality_findings,
    _review_preflight_repair_pages,
    _review_required_view_gap,
    _review_self_repair_gap,
    _review_typography_evidence,
    _render_input_digest,
    _review_delivery_fingerprint,
    _review_verification_allowed,
    _research_handoff_written,
    _revision_delegation_error,
    _revision_prompt,
    _slide_authoring_stop_line,
    _slide_assignment_envelope,
    _slide_group_active_page,
    _grouped_exhausted_page_tool_error,
    _slide_pending_group_view,
    _slide_pending_render_views,
    _slide_deliverable_gap,
    _slide_render_quality_actions,
    _slide_repair_issue,
    _style_lock_injection,
    _tool_results,
    _unresolved_vision_critic_issues,
    _update_vision_issue_ledger,
    _artifact_completion_retry_allowed,
    _is_durable_review_exit,
    _page_has_terminal_slide_failure,
    _outcome_for_page,
    _responsibility_pages_from_outcome,
    _review_should_enter_closure_only,
    DURABLE_REVIEW_EXIT_REASONS,
    REVIEW_CLOSURE_ONLY_TURN_FRACTION,
    REVIEW_CLOSURE_TAIL_BUDGET,
)
from core.run_batch import (  # noqa: E402
    _extract_material,
    _material_blocked,
    build_config,
    stage_materials,
)
from core.runtime_capabilities import (  # noqa: E402
    classify_research_mode,
    detect_runtime_capabilities,
    user_requires_bitmap,
)
from core.trace_mode import write_trace  # noqa: E402
from _internal import deck_core  # noqa: E402


def test_sync_speech_strips_outer_markdown_fence_and_redundant_label(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    deck = {"language": "zh"}
    plans = [
        {
            "number": 1,
            "title": "结论",
            "speech": (
                "```markdown\n"
                "讲稿内容\n\n"
                "请直接朗读这一句，并保留内联 `agent.run()`。\n"
                "```"
            ),
            "evidence": "",
        }
    ]
    monkeypatch.setattr(deck_core, "_load_plans", lambda _root, _expected: (deck, plans))

    deck_core.sync_speech(tmp_path, expected=1)

    speech = (tmp_path / "speech.md").read_text(encoding="utf-8")
    assert "```" not in speech
    assert "讲稿内容" not in speech
    assert "请直接朗读这一句" in speech
    assert "`agent.run()`" in speech


def test_initial_spoken_script_heading_is_canonical_with_legacy_compatibility() -> None:
    canonical = deck_core._sections(
        "## 初版口语讲稿\n这是可直接朗读的初版。\n"
    )
    english = deck_core._sections(
        "## Initial spoken script\nThis is directly speakable.\n"
    )
    legacy = deck_core._sections("## 讲稿节拍\n旧快照仍可恢复。\n")

    assert canonical["speech"] == "这是可直接朗读的初版。"
    assert english["speech"] == "This is directly speakable."
    assert legacy["speech"] == "旧快照仍可恢复。"


@pytest.mark.parametrize("role", ["orchestrator", "review"])
def test_agents_cannot_directly_edit_derived_speech_file(
    tmp_path: Path,
    role: str,
) -> None:
    class SpeechAgent:
        skill_name = "mural-presenter-v0.4"
        label = role
        ws = str(tmp_path)

        def __init__(self) -> None:
            self.role = role

        def safe(self, path: str) -> str:
            return str(tmp_path / path)

    agent = SpeechAgent()
    write_result = tools.write_file(agent, "speech.md", "manual copy\n")
    assert "确定性派生" in write_result
    assert not (tmp_path / "speech.md").exists()

    (tmp_path / "speech.md").write_text("derived\n", encoding="utf-8")
    patch_result = tools.patch(
        agent,
        path="speech.md",
        old_string="derived",
        new_string="manual",
    )
    assert "确定性派生" in patch_result
    assert (tmp_path / "speech.md").read_text(encoding="utf-8") == "derived\n"


def test_build_config_records_effective_openai_model_endpoint(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("MODEL_BACKEND", "openai")
    monkeypatch.setenv("STUDENT_BASE_URL", "http://student.example/v1/")
    monkeypatch.setenv("MODEL", "student-model")
    cfg = build_config(
        SimpleNamespace(batch="endpoint-test", mode="inference", workers=1, max_attempts=1)
    )
    assert cfg["model_base_url"] == "http://student.example/v1"
    assert cfg["model"] == "student-model"


def test_soft_asset_quality_marks_blurry_full_bleed_without_rejecting(
    tmp_path: Path,
) -> None:
    from PIL import Image

    path = tmp_path / "archive.png"
    Image.new("RGB", (1600, 900), (92, 92, 92)).save(path)
    report = deck_core._asset_visual_quality(path, "full-bleed")
    assert report["width"] == 1600
    assert report["height"] == 900
    assert not report["full_bleed_ready"]
    assert any("sharpness" in reason for reason in report["reasons"])


def test_high_confidence_layout_defects_are_structured_but_nonblocking(
    tmp_path: Path,
) -> None:
    geometry = {
        "page_type": "content",
        "frame": "content",
        "page_family": "editorial",
        "canvas_variant": "base",
        "text_boxes": [
            {
                "tag": "p",
                "class_name": "card-copy",
                "ancestor_chain": ["div.card", "main.page-body"],
                "text": "clipped body copy",
                "x": 100,
                "y": 100,
                "width": 300,
                "height": 80,
                "overflow_px_x": 0,
                "overflow_px_y": 12,
            },
            {
                "tag": "li",
                "class_name": "field-item",
                "ancestor_chain": ["ul.field-list", "main.page-body"],
                "text": "colliding list item",
                "x": 120,
                "y": 120,
                "width": 260,
                "height": 70,
                "overflow_px_x": 0,
                "overflow_px_y": 0,
            },
        ],
        "layout_blocks": [
            {
                "tag": "div",
                "class_name": "composition composition--sequence",
                "parent_key": "div.content-stage",
                "ancestor_chain": ["div.content-stage", "main.page-body"],
                "x": 80,
                "y": 180,
                "width": 1440,
                "height": 650,
                "overflow_px_x": 0,
                "overflow_px_y": 24,
                "overflow_mode_x": "visible",
                "overflow_mode_y": "visible",
            },
            {
                "tag": "article",
                "class_name": "card card-a",
                "parent_key": "div.grid-2",
                "ancestor_chain": ["div.grid-2", "div.content-stage"],
                "x": 100,
                "y": 300,
                "width": 500,
                "height": 300,
                "overflow_px_x": 0,
                "overflow_px_y": 0,
                "overflow_mode_x": "visible",
                "overflow_mode_y": "visible",
            },
            {
                "tag": "article",
                "class_name": "card card-b",
                "parent_key": "div.grid-2",
                "ancestor_chain": ["div.grid-2", "div.content-stage"],
                "x": 200,
                "y": 320,
                "width": 500,
                "height": 300,
                "overflow_px_x": 0,
                "overflow_px_y": 0,
                "overflow_mode_x": "visible",
                "overflow_mode_y": "visible",
            },
        ],
        "images": [],
        "boxes": {
            "header": {"x": 60, "y": 40, "width": 1480, "height": 100},
            "footer": {"x": 60, "y": 830, "width": 1480, "height": 30},
        },
    }
    defects = deck_core._page_layout_defects(1, geometry)
    assert {item["type"] for item in defects} == {
        "text_overflow",
        "text_collision",
        "block_overflow",
        "block_collision",
    }
    assert any("page-body" in item.get("ancestor_hint", "") for item in defects)
    assert all(item["id"].startswith("LAYOUT-") for item in defects)
    assert defects == deck_core._page_layout_defects(1, geometry)

    rendered = {"pages": [{"page": 1, "geometry": geometry}]}
    plans = [{
        "number": 1,
        "page_type": "content",
        "page_family": "editorial",
        "canvas_variant": "base",
        "show_footer": True,
        "needs_bitmap": False,
    }]
    audit = deck_core._geometry_audit(tmp_path, plans, rendered)
    assert audit["status"] == "PASS"
    assert audit["errors"] == []
    assert audit["layout_defects"]["01"]

    with mock.patch.dict(os.environ, {"MURAL_LAYOUT_DEFECT": "off"}):
        assert deck_core._page_layout_defects(1, geometry) == []


def test_layout_defect_action_is_shared_by_single_and_grouped_v03(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    state_path = tmp_path / "_trace/slide-render-states/page_01.json"
    state_path.write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "attempt_hashes": {"slide_group_intro": ["first"]},
            "authoring_hashes": ["first"],
            "layout_defects": [{"type": "text_overflow"}],
            "quality_action_count": 1,
        }),
        encoding="utf-8",
    )

    class Slide:
        role = "slide"
        skill_name = "mural-presenter-v0-4"
        trace_label = "slide_group_intro"
        ws = str(tmp_path)

    for topology in ("single", "grouped"):
        (tmp_path / "plan/deck.md").write_text(
            f"# Deck\n## Resolved deck brief\n- ownership_topology: {topology}\n",
            encoding="utf-8",
        )
        assert _slide_render_quality_actions(Slide(), 1) == 1

    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["attempt_hashes"]["slide_group_intro"].append("second")
    state["authoring_hashes"].append("second")
    state_path.write_text(json.dumps(state), encoding="utf-8")
    assert _slide_render_quality_actions(Slide(), 1) == 1
    state["attempt_hashes"]["slide_group_intro"].append("third")
    state["authoring_hashes"].append("third")
    state_path.write_text(json.dumps(state), encoding="utf-8")
    assert _slide_render_quality_actions(Slide(), 1) == 0


def test_content_header_variation_is_review_warning_not_delivery_failure(
    tmp_path: Path,
) -> None:
    plans = [
        {
            "number": number,
            "page_type": "content",
            "page_family": family,
            "canvas_variant": "base",
            "show_footer": True,
            "needs_bitmap": False,
        }
        for number, family in ((1, "editorial"), (2, "breathing-quote"))
    ]
    rendered = {
        "pages": [
            {
                "page": number,
                "geometry": {
                    "page_type": "content",
                    "page_family": family,
                    "canvas_variant": "base",
                    "frame": "content",
                    "boxes": {
                        "header": {
                            "x": 60,
                            "y": y,
                            "width": 1480,
                            "height": 100,
                        },
                        "footer": {
                            "x": 60,
                            "y": 830,
                            "width": 1480,
                            "height": 30,
                        },
                    },
                    "text_boxes": [],
                    "images": [],
                },
            }
            for number, family, y in (
                (1, "editorial", 40),
                (2, "breathing-quote", 430),
            )
        ]
    }

    audit = deck_core._geometry_audit(tmp_path, plans, rendered)

    assert audit["status"] == "PASS"
    assert audit["errors"] == []
    assert any("content header geometry differs" in item for item in audit["warnings"])


def test_layout_defects_do_not_turn_usable_delivery_into_rejection(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "assets").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- language: zh\n",
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_01.md").write_text("# Slide 01\n", encoding="utf-8")
    (tmp_path / "assets/catalog.md").write_text("# Catalog\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "present.html").write_text("<main>" + "x" * 600 + "</main>", encoding="utf-8")
    (tmp_path / "renders/render.json").write_text(
        json.dumps({
            "n_pages": 1,
            "blank_pages": [],
            "console_errors": [],
            "static_pages": [1],
            "special_page_geometry": {"status": "PASS"},
            "layout_defects": {
                "01": [{"id": "LAYOUT-TEST", "type": "text_collision"}],
            },
        }),
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        child_outcomes={},
        tool_policy_violations=[],
        workspace_policy_violations=[],
        material_blocked=False,
        n_renders=1,
        deck_language="zh",
        review_completed=True,
        final_render_after_review=True,
        final_view_after_review=True,
        finalize_attempted=True,
        finalize_succeeded=True,
        finalize_failure="",
        quality_status="needs_improvement",
        research_required=False,
        research_completed=False,
        material_required=False,
        material_completed=False,
        image_required=False,
        image_completed=False,
    )

    ok, reason, detail = _accept(agent, str(tmp_path))
    assert ok is True
    assert reason == "ok"
    assert detail["quality_status"] == "needs_improvement"


def test_visual_review_does_not_score_palette_gradient_or_card_style(
    tmp_path: Path,
) -> None:
    (tmp_path / "slides").mkdir()
    html = (
        '<section class="slide"><div class="card"></div><div class="card"></div>'
        '<div class="card"></div><style>.slide{background:radial-gradient(circle,#111,#000)}'
        '</style></section>'
    )
    plans = []
    for page in range(1, 5):
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(html, encoding="utf-8")
        plans.append({
            "number": page,
            "page_type": "content",
            "canvas_variant": "base",
            "needs_bitmap": False,
            "special_layout": "",
        })
    deck = {
        "sections": {"visual": "generic clean interface"},
        "tokens": {
            "--content-canvas": "#11143a",
            "--special-bg": "#11143a",
            "--accent": "#6366f1",
        },
    }
    assert deck_core._visual_review_notes(tmp_path, deck, plans) == []


def test_visual_review_leaves_theme_palette_judgment_to_model(
    tmp_path: Path,
) -> None:
    plans = [
        {
            "number": page,
            "page_type": "content",
            "canvas_variant": "base",
            "needs_bitmap": False,
            "special_layout": "",
        }
        for page in range(1, 5)
    ]
    deck = {
        "title": "多模态人工智能通识课",
        "resolved": "",
        "sections": {
            "视觉契约": "背景配方：基础画布使用深靛蓝纯色，全册统一背景。",
        },
        "tokens": {
            "--content-canvas": "#14182C",
            "--special-bg": "#0B0E1E",
        },
    }
    assert deck_core._visual_review_notes(tmp_path, deck, plans) == []


def test_body_mono_is_an_advisory_audit_warning(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "base.css").write_text(
        ".body-copy { font-family: var(--font-mono); }\n"
        ".tech-id { font-family: var(--font-mono); }\n",
        encoding="utf-8",
    )
    warnings = deck_core._typography_quality_warnings(tmp_path)
    assert len(warnings) == 1
    assert "body-copy" in warnings[0]
    assert "tech-id" not in warnings[0]


def test_v04_base_css_has_projectable_default_type_floor() -> None:
    source = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/assets/base.css"
    )
    css = source.read_text(encoding="utf-8")
    assert "--fs-body: 26px;" in css
    assert "--fs-min: 20px;" in css
    assert "--fs-micro: var(--fs-min);" in css
    match = re.search(r"(?ms)^body \{(.*?)^\}", css.split("html,\nbody", 1)[1])
    assert match is not None
    body = match.group(1)
    assert "font-size: var(--fs-body);" in body
    assert "line-height: var(--lh-body);" in body


def test_v04_rendered_typography_warns_below_twenty_pixel_floor() -> None:
    warnings = deck_core._rendered_typography_warnings(
        3,
        [
            {
                "text": "Source note that should remain projectable",
                "font_size": 19,
                "class_name": "source-note",
                "tag": "div",
            },
            {
                "text": "03 / METHOD",
                "font_size": 18,
                "class_name": "eyebrow",
                "tag": "span",
            },
        ],
    )
    joined = "\n".join(warnings)
    assert "below 20px" in joined
    assert "Source note" in joined
    assert "03 / METHOD" not in joined


def test_v04_page_render_surfaces_typography_feedback_to_slide_owner(
    tmp_path: Path, capsys,
) -> None:
    warnings = deck_core._page_typography_warnings(
        tmp_path,
        4,
        {
            "text_boxes": [
                {
                    "text": "Visible explanatory label",
                    "font_size": 17,
                    "font_weight": "400",
                    "font_family": "sans-serif",
                    "class_name": "diagram-label",
                    "tag": "div",
                }
            ]
        },
    )
    assert any("below 20px" in warning for warning in warnings)
    deck_core._print_page_render_feedback([], warnings)
    output = capsys.readouterr().out
    assert "layout-defects:none" in output
    assert "typography-warnings:" in output
    assert "ordinary visible labels at least 20px" in output
    assert "soft quality findings" in output


def test_v04_page_render_reports_clean_typography_feedback(capsys) -> None:
    deck_core._print_page_render_feedback([], [])
    output = capsys.readouterr().out
    assert "layout-defects:none" in output
    assert "typography-warnings:none" in output
    assert "typography-repair-order" not in output


def test_v04_attachment_attribution_and_derived_value_review_contract() -> None:
    edition = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4"
    checklist = (edition / "references/quality-checklist.md").read_text(
        encoding="utf-8"
    )
    review = (edition / "roles/review.md").read_text(encoding="utf-8")
    slide = (edition / "roles/slide.md").read_text(encoding="utf-8")
    assert "Derived from Table 3" in checklist
    assert "派生均值、差值、百分比和趋势" in review
    assert "拼写、语法与模板标记通扫" in review
    assert "Figure / Table / Equation" in slide
    assert "≥20px" in slide


def test_smiley_display_token_uses_its_real_metadata_weight(tmp_path: Path) -> None:
    source = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/assets/base.css"
    )
    (tmp_path / "base.css").write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
    deck_core._apply_theme_tokens(
        tmp_path,
        {"tokens": {"--font-display": "var(--font-heavy)"}},
    )
    css = (tmp_path / "base.css").read_text(encoding="utf-8")
    assert "--font-display-weight: 400;" in css
    assert "font-weight: var(--font-display-weight);" in css


def test_v04_visual_medium_contract_uses_role_specific_taxonomy(tmp_path: Path) -> None:
    plan = tmp_path / "slide_01.md"
    plan.write_text(
        """# slide_01
- role: evidence page
- page_type: content
- page_family: chart
- production_group: page-01
- composition: data-focus
- needs_bitmap: false
- primary_visual_medium: echarts

## Narrative
Compare the measured systems.

## Evidence
- A: 42; B: 31.

## On-screen copy (exact)
- title: Measured comparison

## Composition blueprint
- focal: chart

## Semantic visual need
Use a quantitative bar chart as the dominant visual.

## Speech beat
Explain the comparison.
""",
        encoding="utf-8",
    )
    parsed = deck_core._parse_slide(plan, 1)
    assert parsed["primary_visual_medium"] == "echarts"
    assert parsed["visual_medium_explicit"] is True

    plan.write_text(plan.read_text(encoding="utf-8").replace(
        "primary_visual_medium: echarts",
        "primary_visual_medium: bitmap-identity",
    ), encoding="utf-8")
    with pytest.raises(ValueError, match="requires needs_bitmap=true"):
        deck_core._parse_slide(plan, 1)

    plan.write_text(
        plan.read_text(encoding="utf-8").replace(
            "- primary_visual_medium: bitmap-identity\n", ""
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="primary_visual_medium"):
        deck_core._parse_slide(plan, 1)


def test_deck_routing_metadata_can_live_in_compact_global_sections(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        """# Deck Plan
- title: Test deck

## Resolved deck brief
- language: en
- page_count: 1
- audience: reviewers
- image_mode: mixed
- rationale: test

## Ownership
- ownership_topology: single
- ownership_rationale: pages are self-contained

## Bitmap strategy
- bitmap_strategy: active
- bitmap_rationale: the cover benefits from a hero

## Theme Tokens
```css
:root { --content-canvas: #f5f5f5; }
```
""",
        encoding="utf-8",
    )
    parsed = deck_core._parse_deck(tmp_path, expected=1)
    assert parsed["ownership_topology"] == "single"
    assert parsed["bitmap_strategy"] == "active"
    assert any("accepted outside Resolved deck brief" in item for item in parsed["warnings"])


def test_special_page_ignores_harmless_content_composition_hint(tmp_path: Path) -> None:
    plan = tmp_path / "slide_01.md"
    plan.write_text(
        """# slide_01
- role: cover
- page_type: cover
- page_family: special-hero
- production_group: bookends
- special_layout: split
- composition: visual-split
- needs_bitmap: false
- primary_visual_medium: editorial-typography

## Narrative
Open the deck.

## On-screen copy (exact)
- title: A clear title

## Semantic visual need
Use a strong typographic focal point.

## Speech beat
Introduce the topic.
""",
        encoding="utf-8",
    )
    parsed = deck_core._parse_slide(plan, 1)
    assert parsed["special_layout"] == "split"
    assert parsed["composition"] == ""
    assert parsed["composition_explicit"] is False


def test_orchestrator_can_reclassify_only_unstarted_image_failed_page(
    tmp_path: Path,
) -> None:
    (tmp_path / "renders").mkdir()
    agent = SimpleNamespace(
        role="orchestrator",
        label="orchestrator",
        skill_name="mural-presenter-v0-4",
        revision_mode=False,
        delegated_roles=["image", "slide"],
        image_failed_pages=(3,),
        ws=str(tmp_path),
    )
    assert tools._model_write_error(agent, "plan/slide_03.md") is None

    (tmp_path / "renders/slide_03.png").write_bytes(b"pixels")
    error = tools._model_write_error(agent, "plan/slide_03.md")
    assert error is not None
    assert "primary_visual_medium 已冻结" in error

    assert tools._model_write_error(agent, "plan/slide_04.md") is not None


def test_v04_legacy_bitmap_medium_is_normalized_to_source_route(tmp_path: Path) -> None:
    plan = tmp_path / "slide_01.md"
    plan.write_text(
        """# slide_01
- role: cover
- page_type: cover
- page_family: archival
- production_group: bookends
- special_layout: split
- needs_bitmap: true
- primary_visual_medium: bitmap-atmosphere

## Narrative
Open the story.

## On-screen copy (exact)
- title: The story

## Composition blueprint
- focal: generated hero

## Semantic visual need
Use an original visual series.

## Speech beat
Introduce the story.
""",
        encoding="utf-8",
    )
    parsed = deck_core._parse_slide(plan, 1)
    assert parsed["primary_visual_medium"] == "bitmap-generated"


def test_v04_asset_catalog_must_match_orchestrator_source_route(
    tmp_path: Path,
) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/hero.png").write_bytes(b"generated-pixels")
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## hero\n"
        "- slides: 1\n- kind: generated\n- path: assets/hero.png\n"
        "- source: generated\n",
        encoding="utf-8",
    )
    fragment = '<section id="slide-01"><img src="assets/hero.png"></section>'
    generated = [{
        "number": 1,
        "needs_bitmap": True,
        "primary_visual_medium": "bitmap-generated",
    }]
    assert deck_core._validate_asset_requirements(
        tmp_path, generated, [(1, fragment)]
    )

    real = [{
        "number": 1,
        "needs_bitmap": True,
        "primary_visual_medium": "bitmap-real",
    }]
    with pytest.raises(ValueError, match="not silently switch acquisition methods"):
        deck_core._validate_asset_requirements(tmp_path, real, [(1, fragment)])

    (tmp_path / "assets/real.png").write_bytes(b"real-pixels")
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## real\n"
        "- slides: 1\n- kind: real\n- path: assets/real.png\n"
        "- source: https://example.test/real.png\n\n## generated\n"
        "- slides: 1\n- kind: generated\n- path: assets/hero.png\n"
        "- source: generated\n",
        encoding="utf-8",
    )
    both = (
        '<section id="slide-01"><img src="assets/real.png">'
        '<img src="assets/hero.png"></section>'
    )
    # `mixed-*` means bitmap plus code-native visual, not silent mixing of
    # acquisition sources. Every raster must follow the frozen route.
    with pytest.raises(ValueError, match="not silently switch acquisition methods"):
        deck_core._validate_asset_requirements(tmp_path, real, [(1, both)])


def test_v04_prepare_bundles_portable_echarts_vendor(tmp_path: Path) -> None:
    deck_core.prepare(tmp_path)
    vendor = tmp_path / "assets/vendor/echarts.min.js"
    assert vendor.is_file()
    assert vendor.stat().st_size > 500_000


def test_v04_echarts_rejects_remote_or_missing_vendor(tmp_path: Path) -> None:
    plan = {"primary_visual_medium": "echarts"}
    deck = {}
    remote = (
        '<script src="https://cdn.example/echarts.min.js"></script>'
        '<script>echarts.init(document.body)</script>'
    )
    with pytest.raises(ValueError, match="portable bundled"):
        deck_core._validate_fragment(tmp_path / "slide_01.html", 1, remote, plan, deck)
    missing = '<script>echarts.init(document.body)</script>'
    with pytest.raises(ValueError, match="does not load the bundled vendor"):
        deck_core._validate_fragment(tmp_path / "slide_01.html", 1, missing, plan, deck)


def test_v04_echarts_allows_callbacks_but_names_forbidden_dynamic_api(
    tmp_path: Path,
) -> None:
    plan = {"primary_visual_medium": "echarts"}
    deck = {}
    valid = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        '<script>var chart=echarts.init(document.body);'
        'chart.on("click", function (params) { return params.name; });</script>'
    )
    # Reaching the ordinary fragment-structure check proves the callback was not
    # mistaken for the capital-F Function constructor.
    with pytest.raises(ValueError, match="exactly one section fragment"):
        deck_core._validate_fragment(tmp_path / "slide_01.html", 1, valid, plan, deck)

    forbidden = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        '<script>var chart=echarts.init(document.body); fetch("/remote")</script>'
    )
    with pytest.raises(ValueError, match=r"forbidden dynamic API `fetch`"):
        deck_core._validate_fragment(
            tmp_path / "slide_01.html", 1, forbidden, plan, deck
        )

    constructor = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        '<script>var chart=echarts.init(document.body); new Function("return 1")()</script>'
    )
    with pytest.raises(ValueError, match=r"Function\("):
        deck_core._validate_fragment(
            tmp_path / "slide_01.html", 1, constructor, plan, deck
        )


def test_v04_echarts_canvas_typography_is_a_review_signal(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "slides/slide_04.html").write_text(
        "<script>var chart=echarts.init(document.body);"
        "chart.setOption({axisLabel:{fontSize:12},legend:{fontSize:18}});</script>",
        encoding="utf-8",
    )
    warnings = deck_core._echarts_typography_warnings(tmp_path, 4)
    assert len(warnings) == 1
    assert "12px" in warnings[0]
    assert "16px" in warnings[0]


def test_v04_substantive_media_accepts_useful_seven_percent_visual(
    tmp_path: Path,
) -> None:
    plan = {
        "number": 7,
        "page_type": "content",
        "page_family": "editorial",
        "canvas_variant": "default",
        "show_footer": False,
        "primary_visual_medium": "bitmap-identity",
        "needs_bitmap": True,
        "visual_evidence_role": "supporting",
    }
    geometry = {
        "page_type": "content",
        "page_family": "editorial",
        "canvas_variant": "default",
        "frame": "content",
        "boxes": {},
        "text_boxes": [],
        "images": [{"width": 435, "height": 263, "opacity": 1}],
        "media": [{
            "kind": "bitmap",
            "width": 435,
            "height": 263,
            "area_ratio": 0.0795,
        }],
    }
    result = deck_core._geometry_audit(
        tmp_path,
        [plan],
        {"pages": [{"page": 7, "geometry": geometry}]},
    )
    assert result["media_mismatches"] == {}
    assert not any("substantive visual threshold" in item for item in result["warnings"])


def test_v04_collective_bitmap_atlas_is_substantive_but_not_a_small_singleton(
    tmp_path: Path,
) -> None:
    plan = {
        "number": 6,
        "page_type": "content",
        "page_family": "species-atlas",
        "canvas_variant": "default",
        "show_footer": False,
        "primary_visual_medium": "mixed",
        "needs_bitmap": True,
        "visual_evidence_role": "identity",
    }
    thumbnails = [
        {"kind": "bitmap", "width": 200, "height": 160, "area_ratio": 0.0222}
        for _ in range(8)
    ]
    geometry = {
        "page_type": "content",
        "page_family": "species-atlas",
        "canvas_variant": "default",
        "frame": "content",
        "boxes": {},
        "text_boxes": [],
        "images": [{"width": 200, "height": 160, "opacity": 1}] * 8,
        "media": [
            *thumbnails,
            {"kind": "svg", "width": 500, "height": 300, "area_ratio": 0.1042},
        ],
    }
    result = deck_core._geometry_audit(
        tmp_path,
        [plan],
        {"pages": [{"page": 6, "geometry": geometry}]},
    )
    assert result["media_mismatches"] == {}
    assert geometry["media_inventory"]["collective_bitmap_count"] == 8
    assert "bitmap" in geometry["media_inventory"]["substantive_kinds"]


def test_failed_image_page_requires_explicit_plan_reclassification(
    tmp_path: Path,
) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "_trace").mkdir()
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## unavailable-photo\n"
        "- slides: 2\n- kind: real\n- path: assets/missing.jpg\n"
        "- source: https://example.test/missing.jpg\n",
        encoding="utf-8",
    )
    plans = [{"number": 2, "needs_bitmap": True}]
    fragments = [(2, '<section id="slide-02">SVG fallback</section>')]
    with pytest.raises(FileNotFoundError, match="missing or empty"):
        deck_core._validate_asset_requirements(tmp_path, plans, fragments)

    (tmp_path / "_trace/image-handoff.json").write_text(
        json.dumps({
            "schema": "mural.image-handoff.v1",
            "ready_pages": [],
            "failed_pages": [2],
        }),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="requires a resolved raster asset"):
        deck_core._validate_asset_requirements(tmp_path, plans, fragments)

    plans[0]["needs_bitmap"] = False
    plans[0]["primary_visual_medium"] = "svg-diagram"
    entries = deck_core._validate_asset_requirements(tmp_path, plans, fragments)
    assert [entry["id"] for entry in entries] == ["unavailable-photo"]


def test_catalog_failed_asset_is_a_valid_partial_image_handoff(
    tmp_path: Path,
) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## unavailable-photo\n"
        "- slides: 2\n- kind: real\n- path: assets/missing.jpg\n"
        "- source: https://example.test/missing.jpg\n- status: failed\n",
        encoding="utf-8",
    )
    plans = [{
        "number": 2,
        "needs_bitmap": True,
        "primary_visual_medium": "bitmap-real",
    }]
    assert deck_core._validate_asset_requirements(tmp_path, plans) == []


def test_failed_material_candidate_needs_no_crop_registry_and_finalizes(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## unavailable-material-crop\n"
        "- slides: 2\n- kind: material\n"
        "- path: assets/unavailable-material-crop.png\n"
        "- source: inputs/report.pdf.pages/page_018.png\n"
        "- purpose: planned documentary image\n"
        "- crop: bounded source figure\n- display: inline\n"
        "- status: failed\n"
        "- failure_reason: source figure is too small for a faithful crop\n",
        encoding="utf-8",
    )
    plans = [{
        "number": 2,
        "needs_bitmap": True,
        "primary_visual_medium": "bitmap-material",
    }]

    entries = deck_core._parse_catalog(tmp_path)
    assert entries[0]["status"] == "failed"
    assert not (tmp_path / "_trace/material-figures.json").exists()
    with mock.patch.object(deck_core, "_load_plans", return_value=({}, plans)):
        deck_core.finalize_assets(tmp_path)

    output = capsys.readouterr().out
    assert "status:PASS" in output
    assert "bitmap_ready:none" in output
    assert "failed:02" in output
    final_catalog = (assets / "catalog.md").read_text(encoding="utf-8")
    assert "status: failed" in final_catalog
    assert "source figure is too small" in final_catalog


def test_failed_user_candidate_needs_no_user_registry_and_finalizes(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## unavailable-user-image\n"
        "- slides: 4\n- kind: user\n"
        "- path: assets/unavailable-user-image.png\n"
        "- source: inputs/portrait.png\n"
        "- purpose: user-provided portrait\n"
        "- crop: preserve subject\n- display: inline\n"
        "- status: failed\n"
        "- failure_reason: source image is unreadable\n",
        encoding="utf-8",
    )
    plans = [{
        "number": 4,
        "needs_bitmap": True,
        "primary_visual_medium": "bitmap-material",
    }]

    entries = deck_core._parse_catalog(tmp_path)
    assert entries[0]["status"] == "failed"
    assert not (tmp_path / "_trace/user-images.json").exists()
    with mock.patch.object(deck_core, "_load_plans", return_value=({}, plans)):
        deck_core.finalize_assets(tmp_path)

    output = capsys.readouterr().out
    assert "status:PASS" in output
    assert "bitmap_ready:none" in output
    assert "failed:04" in output
    final_catalog = (assets / "catalog.md").read_text(encoding="utf-8")
    assert "status: failed" in final_catalog
    assert "source image is unreadable" in final_catalog


def test_failed_catalog_candidate_does_not_poison_page_with_usable_asset() -> None:
    entries = [
        {"id": "good", "slides": [3, 4], "status": "ready"},
        {"id": "bad-p3", "slides": [3], "status": "failed"},
        {"id": "bad-p5", "slides": [5], "status": "failed"},
    ]
    assert deck_core._catalog_failed_only_pages(entries) == {5}


def test_catalog_rejects_ambiguous_asset_status(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## maybe\n- slides: 1\n- kind: real\n"
        "- path: assets/maybe.jpg\n- source: https://example.test/maybe.jpg\n"
        "- status: partial\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="status must be ready or failed"):
        deck_core._parse_catalog(tmp_path)


def test_finalize_keeps_page_ready_when_one_of_several_assets_failed(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    from PIL import Image

    assets = tmp_path / "assets"
    assets.mkdir()
    Image.new("RGB", (640, 360), "#27758c").save(assets / "usable.jpg")
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## usable\n- slides: 3\n- kind: real\n"
        "- path: assets/usable.jpg\n- source: https://example.test/usable\n"
        "- purpose: usable place photo\n- crop: centered\n- display: inline\n"
        "- status: ready\n\n## unavailable\n- slides: 3\n- kind: real\n"
        "- path: assets/unavailable.jpg\n- source: https://example.test/unavailable\n"
        "- purpose: optional second photo\n- crop: centered\n- display: inline\n"
        "- status: failed\n- failure_reason: provider returned 429\n",
        encoding="utf-8",
    )
    plans = [{
        "number": 3,
        "needs_bitmap": True,
        "primary_visual_medium": "bitmap-real",
    }]
    with mock.patch.object(deck_core, "_load_plans", return_value=({}, plans)):
        deck_core.finalize_assets(tmp_path)
    output = capsys.readouterr().out
    assert "bitmap_ready:03" in output
    assert "failed:none" in output
    final_catalog = (assets / "catalog.md").read_text(encoding="utf-8")
    assert "## usable" in final_catalog
    assert "## unavailable" in final_catalog
    assert "status: failed" in final_catalog
    assert "failure_reason: provider returned 429" in final_catalog


def test_v04_placeholder_and_synthetic_weight_are_review_signals(tmp_path: Path) -> None:
    assert deck_core._PLACEHOLDER_COPY.search("汇报人：XXX")
    support = {"deck-test-smiley": [(400, 400)]}
    warnings = deck_core._unsupported_rendered_font_weights(
        1,
        [{
            "text": "A strong title",
            "font_weight": "700",
            "font_family": "Deck-Test-Smiley, sans-serif",
        }],
        support,
    )
    assert len(warnings) == 1
    assert "no real 700 weight" in warnings[0]


def test_divider_with_one_follower_emits_soft_actionable_warning(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    deck = {
        "ownership_topology": "single",
        "bitmap_strategy": "none",
        "language": "en",
        "warnings": [],
    }
    plans = [
        {"number": 1, "production_group": "bookends", "needs_bitmap": False,
         "page_type": "cover", "composition_explicit": True,
         "composition_blueprint": "hero", "composition": "freeform"},
        {"number": 2, "production_group": "p2", "needs_bitmap": False,
         "page_type": "section-divider", "composition_explicit": True,
         "composition_blueprint": "pause", "composition": "freeform"},
        {"number": 3, "production_group": "p3", "needs_bitmap": False,
         "page_type": "content", "composition_explicit": True,
         "composition_blueprint": "evidence", "composition": "editorial"},
        {"number": 4, "production_group": "bookends", "needs_bitmap": False,
         "page_type": "closing", "composition_explicit": True,
         "composition_blueprint": "close", "composition": "freeform"},
    ]
    with mock.patch.object(
        deck_core, "_load_plans_for_validation", return_value=(deck, plans)
    ):
        assert deck_core.validate_plans(tmp_path, 4) == plans
    output = capsys.readouterr().out
    assert "status:PASS" in output
    assert "leads only 1 content page" in output
    assert "remove the standalone divider" in output


def test_solo_cover_with_bookends_passes_validation(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Solo cover with production_group=bookends passes; no warning emitted."""
    deck = {
        "ownership_topology": "single",
        "bitmap_strategy": "not-beneficial",
        "bitmap_rationale": "all pages are code visual",
        "language": "en",
        "warnings": [],
    }
    plans = [
        {"number": 1, "production_group": "bookends", "needs_bitmap": False,
         "page_type": "cover", "composition_explicit": True,
         "composition_blueprint": "hero", "composition": "freeform"},
        {"number": 2, "production_group": "p2", "needs_bitmap": False,
         "page_type": "content", "composition_explicit": True,
         "composition_blueprint": "evidence", "composition": "data-focus"},
    ]
    with mock.patch.object(
        deck_core, "_load_plans_for_validation", return_value=(deck, plans)
    ):
        assert deck_core.validate_plans(tmp_path, 2) == plans
    output = capsys.readouterr().out
    assert "status:PASS" in output
    assert "solo cover" not in output


def test_solo_cover_with_unique_group_emits_soft_warning(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Solo cover with non-bookends group passes but emits a soft warning."""
    deck = {
        "ownership_topology": "single",
        "bitmap_strategy": "not-beneficial",
        "bitmap_rationale": "all pages are code visual",
        "language": "en",
        "warnings": [],
    }
    plans = [
        {"number": 1, "production_group": "cover-hero", "needs_bitmap": False,
         "page_type": "cover", "composition_explicit": True,
         "composition_blueprint": "hero", "composition": "freeform"},
        {"number": 2, "production_group": "p2", "needs_bitmap": False,
         "page_type": "content", "composition_explicit": True,
         "composition_blueprint": "evidence", "composition": "data-focus"},
    ]
    with mock.patch.object(
        deck_core, "_load_plans_for_validation", return_value=(deck, plans)
    ):
        assert deck_core.validate_plans(tmp_path, 2) == plans
    output = capsys.readouterr().out
    assert "status:PASS" in output
    assert "solo cover" in output
    assert "bookends is recommended" in output


def test_plan_validation_aggregates_deck_and_page_parse_errors(
    tmp_path: Path,
) -> None:
    plan_dir = tmp_path / "plan"
    plan_dir.mkdir()
    (plan_dir / "slide_01.md").write_text("broken one", encoding="utf-8")
    (plan_dir / "slide_02.md").write_text("broken two", encoding="utf-8")

    with mock.patch.object(
        deck_core, "_parse_deck", side_effect=ValueError("deck.md missing Theme Tokens")
    ):
        with pytest.raises(ValueError) as error:
            deck_core._load_plans_for_validation(tmp_path, expected=2)

    message = str(error.value)
    assert "plan validation found 3 issue(s)" in message
    assert "deck.md missing Theme Tokens" in message
    assert "slide_01.md first line must be exactly `# slide_01`" in message
    assert "slide_02.md first line must be exactly `# slide_02`" in message


def test_orchestrator_validate_plans_prints_clean_failure_without_traceback(
    tmp_path: Path,
) -> None:
    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
    )
    result = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "2"],
        capture_output=True,
        text=True,
    )

    assert result.returncode == 2
    assert result.stdout.startswith("status:FAIL\n")
    assert "plan validation found" in result.stdout
    assert "Traceback" not in result.stdout
    assert "Traceback" not in result.stderr


def test_review_finalize_prints_clean_contract_failure_without_traceback(
    tmp_path: Path,
) -> None:
    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/review.py"
    )
    result = subprocess.run(
        [sys.executable, str(script), "finalize", str(tmp_path), "--expected", "2"],
        capture_output=True,
        text=True,
    )

    assert result.returncode == 2
    assert result.stdout.startswith("status:FAIL\n")
    assert "Traceback" not in result.stdout
    assert "Traceback" not in result.stderr


def test_theme_tokens_reject_css_self_reference() -> None:
    with pytest.raises(ValueError, match="invalid self-references: --font-mono"):
        deck_core._theme_tokens(
            ":root { --content-canvas: #f2eee4; --font-mono: var(--font-mono); }"
        )


def test_slide_write_guard_preserves_scaffold_owned_root_metadata(
    tmp_path: Path,
) -> None:
    page = tmp_path / "slide_03.html"
    opening = (
        '<section class="slide" id="slide-03" data-slide="03" '
        'data-page-type="content" data-page-family="editorial" '
        'data-frame="content" data-canvas-variant="base">'
    )
    page.write_text(opening + "<main>scaffold</main></section>", encoding="utf-8")
    agent = SimpleNamespace(role="slide")
    broken = '<section class="slide" id="slide-03"><main>done</main></section>'
    with mock.patch.object(tools, "_is_current_variant_skill", return_value=True):
        error = tools._slide_root_contract_error(
            agent, "slides/slide_03.html", str(page), broken
        )
        assert error and "data-slide='03'" in error
        assert tools._slide_root_contract_error(
            agent,
            "slides/slide_03.html",
            str(page),
            opening + "<main>done</main></section>",
        ) is None
        assert tools._slide_root_contract_error(
            agent,
            "slides/slide_03.html",
            str(page),
            opening
            + "<main><section class=\"finding\"><h2>Result</h2></section>"
            + "<section class=\"evidence\"><p>Evidence</p></section></main></section>",
        ) is None
        sibling = (
            opening
            + "<main>done</main></section>"
            + '<section class="slide"><main>extra root</main></section>'
        )
        error = tools._slide_root_contract_error(
            agent, "slides/slide_03.html", str(page), sibling
        )
        assert error and "顶层兄弟节点" in error
        unbalanced = opening + "<main><section>broken</main></section>"
        assert tools._slide_root_contract_error(
            agent, "slides/slide_03.html", str(page), unbalanced
        )


def test_renderer_contract_allows_nested_semantic_sections() -> None:
    plan = {
        "number": 5,
        "page_type": "content",
        "page_family": "mechanism",
        "composition": "visual-split",
        "canvas_variant": "base",
        "title": "Two techniques",
        "subtitle": "Evidence and search",
        "eyebrow": "METHOD",
        "show_footer": False,
        "primary_visual_medium": "code-visual",
    }
    deck = {"pages": 10, "footer": ""}
    skeleton = deck_core._slide_skeleton(plan, deck)
    nested = skeleton.replace(
        "<!-- SLIDE_AGENT_FILL_END -->",
        '<section class="tech"><h2>Privileged information</h2></section>\n'
        '<section class="tech"><h2>Adaptive search</h2></section>\n'
        "        <!-- SLIDE_AGENT_FILL_END -->",
    )
    deck_core._validate_fragment(Path("slide_05.html"), 5, nested, plan, deck)

    sibling = nested + '<section class="slide">extra root</section>'
    with pytest.raises(ValueError, match="exactly one section fragment"):
        deck_core._validate_fragment(Path("slide_05.html"), 5, sibling, plan, deck)


def test_scaffold_force_cannot_destroy_authored_or_rendered_pages(
    tmp_path: Path,
) -> None:
    slides = tmp_path / "slides"
    renders = tmp_path / "renders"
    slides.mkdir()
    renders.mkdir()
    page = slides / "slide_01.html"
    original = '<section class="slide"><main>completed page</main></section>\n'
    page.write_text(original, encoding="utf-8")
    (renders / "slide_01.png").write_bytes(b"stable-render")

    with mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ):
        with pytest.raises(ValueError, match="refusing scaffold --force") as error:
            deck_core.scaffold_from_plans(tmp_path, 1, force=True)

    assert "authored HTML: slide_01.html" in str(error.value)
    assert "render checkpoints: slide_01.png" in str(error.value)
    assert page.read_text(encoding="utf-8") == original


def test_trace_sidecar_distinguishes_real_user_from_tool_result_wrappers(
    tmp_path: Path,
) -> None:
    messages = [
        {"role": "user", "content": "real query"},
        {"role": "assistant", "content": [{"type": "tool_use", "id": "t1"}]},
        {
            "role": "user",
            "content": [
                {"type": "tool_result", "tool_use_id": "t1", "content": "ok"}
            ],
        },
        {"role": "user", "content": "[系统] continue"},
    ]
    write_trace(tmp_path, messages, [], {}, "inference")
    origins = json.loads(
        (tmp_path / "message-origins.json").read_text(encoding="utf-8")
    )["messages"]
    assert [row["semantic_origin"] for row in origins] == [
        "end_user",
        "model_response",
        "harness_tool_result",
        "harness_control",
    ]
    assert json.loads((tmp_path / "messages.json").read_text(encoding="utf-8")) == messages


def test_workspace_audit_allows_empty_harness_tmp_and_result_json(
    tmp_path: Path,
) -> None:
    (tmp_path / "tmp").mkdir()
    (tmp_path / "result.json").write_text('{"status":"completed"}', encoding="utf-8")
    assert deck_core.audit_workspace(tmp_path) == []
    (tmp_path / "tmp" / "residue.tmp").write_text("x", encoding="utf-8")
    violations = deck_core.audit_workspace(tmp_path)
    assert any("retained files" in item for item in violations)


def test_render_state_budget_is_scoped_to_page_lifecycle(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "_trace" / "slide-render-states").mkdir(parents=True)
    (tmp_path / "base.css").write_text("body { color: black; }", encoding="utf-8")
    (tmp_path / "slides" / "slide_01.html").write_text(
        '<section class="slide"><main>completed page</main></section>',
        encoding="utf-8",
    )
    state = tmp_path / "_trace" / "slide-render-states" / "page_01.json"
    state.write_text(json.dumps({
        "limit": 8,
        "authoring_attempt_limit": 3,
        "hashes": ["first", "second"],
        "authoring_hashes": ["first", "second"],
        "attempt_hashes": {"slide_01": ["first", "second"]},
        "published_hash": "second",
    }), encoding="utf-8")

    def publish(_root, _source, target):
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b"png")

    with mock.patch.dict(os.environ, {"MURAL_RENDER_ATTEMPT_ID": "slide_01_r2"}), mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ), mock.patch.object(
        deck_core, "_skeleton_is_unfilled", return_value=False
    ), mock.patch.object(
        deck_core, "_validate_fragment", return_value=None
    ), mock.patch.object(
        deck_core, "_run_renderer", return_value=None
    ), mock.patch.object(
        deck_core, "atomic_copy", side_effect=publish
    ):
        deck_core.render(tmp_path, 1, 1)

    payload = json.loads(state.read_text(encoding="utf-8"))
    assert payload["limit_scope"] == "page_lifecycle"
    assert payload["authoring_attempt_limit"] == 3
    assert len(payload["authoring_hashes"]) == 3
    assert len(payload["attempt_hashes"]["slide_01_r2"]) == 1
    assert (tmp_path / "renders" / "slide_01.png").is_file()

    (tmp_path / "slides" / "slide_01.html").write_text(
        '<section class="slide"><main>fourth state</main></section>',
        encoding="utf-8",
    )
    with mock.patch.dict(os.environ, {"MURAL_RENDER_ATTEMPT_ID": "slide_01_r3"}), mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ), mock.patch.object(
        deck_core, "_skeleton_is_unfilled", return_value=False
    ), mock.patch.object(
        deck_core, "_validate_fragment", return_value=None
    ):
        with pytest.raises(ValueError, match="page lifecycle"):
            deck_core.render(tmp_path, 1, 1)


def test_invalid_fragment_does_not_consume_render_state_budget(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "base.css").write_text("body { color: black; }", encoding="utf-8")
    (tmp_path / "slides" / "slide_01.html").write_text(
        '<section class="slide"><main>completed page</main></section>',
        encoding="utf-8",
    )
    with mock.patch.object(
        deck_core,
        "_load_plans",
        return_value=({"language": "en"}, [{"number": 1}]),
    ), mock.patch.object(
        deck_core, "_skeleton_is_unfilled", return_value=False
    ), mock.patch.object(
        deck_core, "_validate_fragment", side_effect=ValueError("missing data-slide")
    ):
        with pytest.raises(ValueError, match="missing data-slide"):
            deck_core.render(tmp_path, 1, 1)

    assert not (tmp_path / "_trace" / "slide-render-states" / "page_01.json").exists()


def test_incremental_finalize_checkpoint_tracks_published_pixels(tmp_path: Path) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders/.page_01").mkdir(parents=True)
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    (tmp_path / "base.css").write_text(".slide{color:#111}\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide" data-slide="01">A</section>\n', encoding="utf-8"
    )
    (tmp_path / "renders/slide_01.png").write_bytes(b"png")
    digest = deck_core._page_digest(tmp_path, 1)
    deck_core._publish_render_state(
        tmp_path,
        1,
        tmp_path / "_trace/slide-render-states/page_01.json",
        [digest],
        {"slide_01": [digest]},
        digest,
    )
    isolated = {
        "pages": [{"page": 1, "png": "old", "blank": False, "geometry": {}}],
        "console_errors": [],
    }
    (tmp_path / "renders/.page_01/render.json").write_text(
        json.dumps(isolated), encoding="utf-8"
    )
    manifest = deck_core._incremental_manifest_seed(tmp_path, [1])
    assert manifest["pages"][0]["page"] == 1
    assert deck_core._render_checkpoint_current(tmp_path, 1, {1})

    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide" data-slide="01">B</section>\n', encoding="utf-8"
    )
    assert not deck_core._render_checkpoint_current(tmp_path, 1, {1})


def test_page_render_fingerprint_changes_when_referenced_asset_changes(
    tmp_path: Path,
) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "assets").mkdir()
    (tmp_path / "base.css").write_text(".slide{color:#111}\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide"><img src="assets/hero.png"></section>',
        encoding="utf-8",
    )
    asset = tmp_path / "assets/hero.png"
    asset.write_bytes(b"before")
    before = deck_core._page_digest(tmp_path, 1)
    asset.write_bytes(b"after")
    after = deck_core._page_digest(tmp_path, 1)
    assert before != after


def test_runtime_capabilities_omit_optional_stages_without_inputs_or_keys(
    tmp_path: Path,
) -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {"model": "test", "model_base_url": "http://model", "enable_image_gen": True},
            [],
            plan_only=True,
        )
    assert profile["active_roles"] == ["slide", "review"]
    assert profile["omitted_roles"] == ["material", "research", "image"]
    assert profile["workflow"]["image"] == "bitmap_unavailable"
    assert profile["workflow"]["research_mode"] == "open_research"
    assert profile["tools"]["web_search"] is False
    assert profile["tools"]["image_generate"] is False
    assert "api_key" not in json.dumps(profile).lower()


def test_runtime_capabilities_route_material_directly_without_search() -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "image-secret", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {"model": "test", "model_base_url": "http://model", "enable_image_gen": True},
            ["inputs/report.pdf.md"],
            plan_only=True,
        )
    assert profile["active_roles"] == ["material", "image", "slide", "review"]
    assert profile["workflow"]["material_handoff"] == "orchestrator"
    assert profile["workflow"]["research"] == "omitted"


def test_attachment_can_route_material_figures_without_search_or_generation() -> None:
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "enable_image_gen": True,
                "_raw_user_query": "仅使用附件制作演示",
            },
            ["inputs/report.pdf.md"],
            plan_only=True,
        )
    assert profile["active_roles"] == ["material", "image", "slide", "review"]
    assert profile["workflow"]["research_mode"] == "attachment_only"
    assert profile["workflow"]["image"] == "user_or_material_bitmap_only"
    assert profile["tools"]["web_search"] is False
    assert profile["tools"]["image_generate"] is False


@pytest.mark.parametrize(
    ("query", "has_material", "expected"),
    [
        ("把以下文字润色成三页演示", False, "off"),
        ("仅使用附件内容，不要联网", True, "attachment_only"),
        ("基于附件制作并核验最新政策", True, "verify_external"),
        ("介绍台风白海豚的形成和影响", False, "open_research"),
    ],
)
def test_research_mode_is_classified_before_model_call(
    query: str,
    has_material: bool,
    expected: str,
) -> None:
    mode, _reason = classify_research_mode(query, has_material=has_material)
    assert mode == expected


def test_web_search_sets_locale_and_retries_one_simplified_query() -> None:
    calls = []

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def raise_for_status(self):
            return None

        def json(self):
            return self.payload

    def fake_post(_url, **kwargs):
        calls.append(kwargs["json"])
        return Response({"organic": []} if len(calls) == 1 else {
            "organic": [{"title": "官方资料", "link": "https://example.com", "snippet": "证据"}]
        })

    class Research:
        search_enabled = True
        evidence_scope = "open_research"

    with mock.patch.dict(os.environ, {"SERPER_API_KEY": "test-key"}), mock.patch(
        "core.tools.requests.post", side_effect=fake_post
    ):
        payload = json.loads(
            tools.web_search(Research(), '"台风白海豚" 形成 路径 site:example.invalid')
        )

    assert len(calls) == 2
    assert calls[0]["hl"] == "zh-cn" and calls[0]["gl"] == "cn"
    assert "site:" not in calls[1]["q"] and '"' not in calls[1]["q"]
    assert payload["retried_with_simplified_query"] is True
    assert payload["results"][0]["url"] == "https://example.com"


def test_verify_external_keeps_factual_gate_but_allows_image_asset_search(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "research" / "material.md").write_text(
        "# Material\n- unresolved_items: []\n", encoding="utf-8"
    )

    class AgentStub:
        search_enabled = True
        evidence_scope = "verify_external"
        ws = str(tmp_path)

        def __init__(self, role: str):
            self.role = role

    assert "当前没有可核验项" in tools._external_evidence_gate(AgentStub("research"))
    assert tools._external_evidence_gate(
        AgentStub("image"), purpose="visual_asset_acquisition"
    ) == ""


def test_verify_external_research_search_budget_stops_synonym_loops() -> None:
    class Research:
        role = "research"
        evidence_scope = "verify_external"

    agent = Research()
    assert all(tools._consume_research_search_budget(agent) == "" for _ in range(6))
    message = tools._consume_research_search_budget(agent)
    assert "budget_exhausted" in message
    assert "unresolved/boundary" in message


def test_short_open_research_budget_is_page_aware_and_deduplicates() -> None:
    class Research:
        role = "research"
        evidence_scope = "open_research"
        requested_slide_count = 8

    agent = Research()
    assert tools._consume_research_search_budget(agent, "成都 熊猫") == ""
    assert "duplicate_query" in tools._consume_research_search_budget(
        agent, "成都   熊猫"
    )
    for index in range(5):
        assert tools._consume_research_search_budget(agent, f"query {index}") == ""
    assert "budget_exhausted" in tools._consume_research_search_budget(
        agent, "ninth unique query"
    )


def test_web_extract_budget_is_page_aware_and_deduplicates() -> None:
    class Research:
        role = "research"
        evidence_scope = "open_research"
        requested_slide_count = 8

    agent = Research()
    assert tools._consume_web_extract_budget(agent, "https://example.com/a#top") == ""
    assert "duplicate_url" in tools._consume_web_extract_budget(
        agent, "https://example.com/a#details"
    )
    for index in range(5):
        assert tools._consume_web_extract_budget(
            agent, f"https://example.com/{index}"
        ) == ""
    assert "budget_exhausted" in tools._consume_web_extract_budget(
        agent, "https://example.com/seventh"
    )


def test_image_search_has_separate_bounded_acquisition_budget() -> None:
    class Image:
        role = "image"
        evidence_scope = "verify_external"
        requested_slide_count = 10

    agent = Image()
    assert all(tools._consume_research_search_budget(agent) == "" for _ in range(12))
    assert "budget_exhausted" in tools._consume_research_search_budget(agent)


def test_image_allows_only_one_merged_replacement_pass(tmp_path: Path) -> None:
    class Image:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        role_script = "skills/mural-presenter-v0-4/scripts/image.py"
        bash_relaxed = False
        bash_timeout = 30

    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/a.jpg").write_bytes(b"current-image")
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## a\n- path: assets/a.jpg\n",
        encoding="utf-8",
    )
    digest = hashlib.sha256(b"current-image").hexdigest()
    agent = Image()
    agent.vision_critic_results = {
        "assets/a.jpg": {"source_sha256": digest, "verdict": "repair_required"}
    }
    command = (
        "python skills/mural-presenter-v0-4/scripts/image.py "
        "fetch . --replace"
    )
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        run.return_value.returncode = 0
        run.return_value.stdout = "ok"
        run.return_value.stderr = ""
        assert "ok" in tools.terminal(agent, command)
        blocked = tools.terminal(agent, command)
    assert "image_replacement_budget" in blocked
    assert run.call_count == 1


def test_failed_image_replacement_attempt_also_consumes_the_single_pass(
    tmp_path: Path,
) -> None:
    class Image:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        role_script = "skills/mural-presenter-v0-4/scripts/image.py"
        bash_relaxed = False
        bash_timeout = 30

    (tmp_path / "assets").mkdir()
    agent = Image()
    command = "python skills/mural-presenter-v0-4/scripts/image.py fetch . --replace"
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        run.return_value.returncode = 2
        run.return_value.stdout = "status:FAIL\nasset download failed: provider 429"
        run.return_value.stderr = ""
        first = tools.terminal(agent, command)
        second = tools.terminal(agent, command)
    assert "provider 429" in first
    assert "image_replacement_budget" in second
    assert run.call_count == 1


def test_image_finalize_records_current_catalog_digest(tmp_path: Path) -> None:
    class Image:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        role_script = "skills/mural-presenter-v0-4/scripts/image.py"
        bash_relaxed = False
        bash_timeout = 30

    assets = tmp_path / "assets"
    assets.mkdir()
    catalog = assets / "catalog.md"
    catalog.write_text("# Asset catalog\n", encoding="utf-8")
    agent = Image()
    command = "python skills/mural-presenter-v0-4/scripts/image.py finalize ."
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        run.return_value.returncode = 0
        run.return_value.stdout = (
            "status:PASS\nassets:2\nbitmap_ready:01,06\nfailed:09"
        )
        run.return_value.stderr = ""
        result = tools.terminal(agent, command)
    assert "status:PASS" in result
    assert agent.image_finalized_catalog_digest == hashlib.sha256(
        catalog.read_bytes()
    ).hexdigest()
    assert agent.image_ready_pages == (1, 6)
    assert agent.image_failed_pages == (9,)
    assert "bitmap_ready:01,06" in agent.image_handoff_summary


def test_stale_patch_forces_current_file_read_before_one_retry(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    catalog = assets / "catalog.md"
    catalog.write_text("# Asset catalog\n- source: current\n", encoding="utf-8")
    trace_dir = tmp_path / "trace"
    trace_dir.mkdir()

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        started = time.time()
        remote_tool_concurrency = 1
        n_views = 0
        n_renders = 0
        image_by_tool = {}
        vision_critic_results = {}

        def safe(self, path):
            return str(tmp_path / path)

        def log(self, _message):
            return None

    agent = ImageAgent()
    agent.trace_dir = trace_dir
    first = tools.patch(
        agent,
        path="assets/catalog.md",
        old_string="- source: stale",
        new_string="- source: replacement",
    )
    assert "下一轮只开放" in first
    assert agent._patch_recovery_path == "assets/catalog.md"

    use = SimpleNamespace(
        id="tool-read-stale-patch",
        name="read_file",
        input={"path": "assets/catalog.md"},
    )
    with mock.patch("core.agent_loop.tools.dispatch", return_value="1|# Asset catalog"):
        _tool_results(agent, [use], 1, [])
    assert agent._patch_recovery_path == ""

    repeated = tools.patch(
        agent,
        path="assets/catalog.md",
        old_string="- source: stale",
        new_string="- source: replacement",
    )
    assert "stale_patch_repeated" in repeated


def test_identical_patch_is_rejected_as_actionable_noop(tmp_path: Path) -> None:
    target = tmp_path / "plan" / "slide_03.md"
    target.parent.mkdir()
    target.write_text("- eyebrow: KEY FEATURE · LENOV0 X POWER\n", encoding="utf-8")
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        revision_mode=False,
    )
    result = tools.patch(
        agent,
        mode="replace",
        path="plan/slide_03.md",
        old_string="LENOV0 X POWER",
        new_string="LENOV0 X POWER",
    )
    assert "noop_identical" in result
    assert "O/0" in result
    assert "不要原样重试" in result
    assert target.read_text(encoding="utf-8") == "- eyebrow: KEY FEATURE · LENOV0 X POWER\n"


def test_image_replacement_does_not_require_every_asset_fullres(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/a.jpg").write_bytes(b"a")
    (tmp_path / "assets/b.jpg").write_bytes(b"b")
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n\n## a\n- path: assets/a.jpg\n\n"
        "## b\n- path: assets/b.jpg\n",
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        bash_relaxed=False,
        bash_timeout=30,
        vision_critic_results={
            "assets/a.jpg": {
                "source_sha256": hashlib.sha256(b"a").hexdigest(),
                "verdict": "ready",
            }
        },
    )
    command = "python skills/mural-presenter-v0-4/scripts/image.py fetch . --replace"
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        run.return_value.returncode = 0
        run.return_value.stdout = "ok"
        run.return_value.stderr = ""
        result = tools.terminal(agent, command)
    assert "ok" in result
    assert run.call_count == 1


def test_image_replacement_rejects_sleep_wrapper(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/catalog.md").write_text(
        "# Asset catalog\n", encoding="utf-8"
    )
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        bash_relaxed=False,
        bash_timeout=30,
        vision_critic_results={},
    )
    command = (
        "sleep 300 && python skills/mural-presenter-v0-4/scripts/image.py "
        "fetch . --replace"
    )
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        result = tools.terminal(agent, command)
    assert "sleep" in result or "未批准" in result or "不允许" in result
    assert run.call_count == 0


def test_review_page_scoped_css_stays_in_slide_fragment(tmp_path: Path) -> None:
    (tmp_path / "base.css").write_text(
        "#slide-03 .hero { width: 40%; }\n", encoding="utf-8"
    )
    agent = SimpleNamespace(
        role="review",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        safe=lambda path: str(tmp_path / path),
    )
    blocked = tools.patch(
        agent,
        path="base.css",
        old_string="#slide-03 .hero { width: 40%; }",
        new_string="#slide-03 .hero { width: 55%; }",
    )
    assert "page_scoped_css_in_base" in blocked


def test_image_plain_fetch_is_blocked_after_contact_sheet_diagnosis(
    tmp_path: Path,
) -> None:
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        bash_relaxed=False,
        bash_timeout=30,
        vision_critic_results={
            "assets/contact-sheet.png": {"verdict": "repair_required"}
        },
    )
    command = "python skills/mural-presenter-v0-4/scripts/image.py fetch ."
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        blocked = tools.terminal(agent, command)
    assert "image_post_critic_replacement" in blocked
    assert run.call_count == 0


def test_image_role_output_filter_is_canonicalized_and_finalize_receipt_saved(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    catalog = assets / "catalog.md"
    catalog.write_text("# Asset catalog\n", encoding="utf-8")
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        role_script="skills/mural-presenter-v0-4/scripts/image.py",
        bash_relaxed=False,
        bash_timeout=30,
        vision_critic_results={},
        log=lambda _message: None,
    )
    command = (
        "python skills/mural-presenter-v0-4/scripts/image.py finalize . "
        "2>&1 | tail -10"
    )
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        run.return_value.returncode = 0
        run.return_value.stdout = "status:PASS\nassets:1\nbitmap_ready:01\nfailed:none"
        run.return_value.stderr = ""
        result = tools.terminal(agent, command)
    assert "status:PASS" in result
    assert run.call_args.args[0] == (
        "python skills/mural-presenter-v0-4/scripts/image.py finalize ."
    )
    assert agent.image_finalized_catalog_digest == hashlib.sha256(
        catalog.read_bytes()
    ).hexdigest()


def test_image_plain_fetch_runs_once_then_requires_replacement(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        role_script="skills/mural-presenter-v0-4/scripts/image.py",
        bash_relaxed=False,
        bash_timeout=30,
        vision_critic_results={},
        log=lambda _message: None,
    )
    command = "python skills/mural-presenter-v0-4/scripts/image.py fetch ."
    with mock.patch.object(tools, "_bash_mutation_error", return_value=None), mock.patch(
        "core.tools.subprocess.run"
    ) as run:
        run.return_value.returncode = 1
        run.return_value.stdout = "download failed"
        run.return_value.stderr = ""
        first = tools.terminal(agent, command)
        second = tools.terminal(agent, command)
    assert "exit_code=1" in first
    assert "image_initial_fetch_budget" in second
    assert run.call_count == 1


def test_explicit_bitmap_requirement_is_persisted_in_runtime_contract() -> None:
    assert user_requires_bitmap("须含高质量相关图片") is True
    assert user_requires_bitmap("能配图就多配图，封面可以生成") is True
    assert user_requires_bitmap("尽量多使用真实照片") is True
    assert user_requires_bitmap(
        "不要米黄纸张模板，也不要电蓝科技风。能配图就多配图：地点用真图。"
    ) is True
    assert user_requires_bitmap("Images: Include relevant images. Images must be high quality.") is True
    assert user_requires_bitmap("不要使用图片，只做数据图") is False
    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "search", "IMAGE_API_KEY": "image", "OPENAI_API_KEY": ""},
    ):
        profile = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "enable_image_gen": True,
                "_raw_user_query": "制作演示，须含高质量相关图片",
            },
            [],
            plan_only=True,
        )
    assert profile["requirements"] == {
        "bitmap_required": True,
        "source": "raw_user_query",
        "failure_policy": "keep_needs_bitmap_and_return_image_blocked",
    }


def test_image_generation_provider_failure_preserves_requested_model(
    tmp_path: Path,
) -> None:
    class Response:
        ok = False
        status_code = 429
        text = ""

        @staticmethod
        def json():
            return {
                "error": {
                    "message": "All channels for model gpt-image-2 are unavailable"
                }
            }

    class ImageAgent:
        enable_image_gen = True
        img_base = "http://tokenhub"
        img_key = "secret"
        image_model = "gpt-image-2-adobe-2"
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    with mock.patch("core.tools.requests.post", return_value=Response()):
        result = tools.image_generate(ImageAgent(), "editorial portrait")
    assert "[provider_unavailable]" in result
    assert "requested_model=gpt-image-2-adobe-2" in result
    assert "http_status=429" in result
    assert "不得把 needs_bitmap:true 静默改为 false" in result


def test_runtime_capabilities_inject_only_available_roles_and_tools(tmp_path: Path) -> None:
    profile = {
        "active_roles": ["slide", "review"],
        "omitted_roles": ["material", "research", "image"],
        "workflow": {
            "material": "omitted",
            "research": "omitted",
            "material_handoff": "none",
            "image": "bitmap_unavailable",
            "research_mode": "off",
            "research_reason": "test",
            "material_stage": "omitted",
        },
        "tools": {
            "web_search": False,
            "web_extract": False,
            "image_generate": False,
            "vision_analyze": True,
        },
    }
    agent = Agent(
        "runtime-capability-test",
        str(tmp_path),
        "制作一套演示",
        {
            "skill_name": "mural-presenter-v0-4",
            "skill_language": "zh",
            "_runtime_capabilities": profile,
        },
    )
    tool_names = {schema["name"] for schema in agent.tool_schemas}
    assert agent.research_required is False
    assert agent.role_card_injected is True
    assert '<injected_role_card path="roles/orchestrator.md">' in agent.system
    assert "Orchestrator 角色卡已由 Harness 完整注入" in agent.system
    assert "Research：omitted" in agent.system
    assert "可委派角色：slide, review" in agent.system
    assert "web_search" not in tool_names
    assert "image_generate" not in tool_names
    rejected = _delegate_task(
        agent,
        {"tasks": [{"goal": "Research: collect sources"}]},
    )
    assert "运行时能力检查已省略" in rejected


def test_agent_tools_remove_search_generation_and_vision() -> None:
    names = {
        schema["name"]
        for schema in tools.agent_tools(
            "image",
            enable_web_search=False,
            enable_image_gen=False,
            enable_vision=False,
        )
    }
    assert "web_search" not in names
    assert "web_extract" not in names
    assert "image_generate" not in names
    assert "vision_analyze" not in names


def test_research_tool_surface_supports_focused_brief_patch() -> None:
    names = {schema["name"] for schema in tools.agent_tools("research")}
    assert {"read_file", "write_file", "patch", "web_search", "web_extract"} <= names


def test_image_tool_surface_supports_catalog_patch() -> None:
    names = {schema["name"] for schema in tools.agent_tools("image")}
    assert "patch" in names
    assert "web_search" in names
    assert "web_extract" not in names


def test_pre_review_partial_image_failure_cannot_reopen_image(tmp_path: Path) -> None:
    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)
        child_outcomes = {"image": {"ok": True, "exit_reason": "partial_ready"}}
        children = []
        material_completed = False
        research_completed = False
        research_required = False
        image_completed = True
        revision_mode = False

    result = _delegate_task(Parent(), {"tasks": [{
        "role": "image",
        "repair": True,
        "goal": "再次搜索首轮 failed 的 P05 图片",
    }]})
    assert "真实 material/generated 来源改路" in result
    assert "当前不满足来源改路条件" in result


def test_redundant_completed_image_request_routes_unrendered_ready_slide(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_03.md").write_text(
        "# slide_03\n- needs_bitmap: true\n"
        "- primary_visual_medium: bitmap-real\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)
        child_outcomes = {"image": {"ok": True, "exit_reason": "text_response"}}
        children = []
        image_completed = True
        image_ready_pages = (3,)

    parent = Parent()
    result = _delegate_task(parent, {"tasks": [{
        "role": "image",
        "goal": "继续为 P03 补更多候选图",
    }]})
    assert "P03" in result
    assert "不要再次委派 Image" in result
    assert parent.pending_slide_delegation_pages == (3,)


def test_review_driven_image_repair_cannot_overwrite_catalog(tmp_path: Path) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    catalog = assets / "catalog.md"
    original = "# Asset catalog\n\n## existing\n- path: assets/existing.png\n"
    catalog.write_text(original, encoding="utf-8")
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        image_repair_mode=True,
        ws=str(tmp_path),
    )
    result = tools.write_file(agent, "assets/catalog.md", "# replacement\n")
    assert "image_repair_catalog_preservation" in result
    assert catalog.read_text(encoding="utf-8") == original


def test_review_asset_quality_can_claim_only_one_deck_level_image_repair(
    tmp_path: Path,
) -> None:
    review = SimpleNamespace(
        role="review",
        final_text=(
            "status: needs_orchestrator\n"
            "issue_type: asset_quality\n"
            "pages: 03\n"
            "evidence: source image is visibly watermarked\n"
        ),
    )
    parent = SimpleNamespace(ws=str(tmp_path), children=[review])
    claimed, reason = _claim_image_repair_evidence(parent)
    assert claimed is True
    assert reason == ""
    claimed_again, repeated = _claim_image_repair_evidence(parent)
    assert claimed_again is False
    assert "已经消费" in repeated
    state = json.loads(
        (tmp_path / "_trace/image-repair-state.json").read_text(encoding="utf-8")
    )
    assert state["status"] == "claimed"
    assert len(state["review_evidence_sha256"]) == 64


def test_image_repair_reopens_same_review_and_invalidates_stale_delivery(
    tmp_path: Path,
) -> None:
    (tmp_path / "assets").mkdir()
    (tmp_path / "assets/replacement.png").write_bytes(b"new pixels")
    parent = SimpleNamespace(
        ws=str(tmp_path),
        review_completed=True,
        final_render_after_review=True,
        final_view_after_review=True,
        finalize_attempted=True,
        finalize_succeeded=True,
        finalize_failure="",
        quality_status="needs_improvement",
        child_outcomes={
            "review": {
                "role": "review",
                "ok": True,
                "status": "needs_improvement",
                "exit_reason": "needs_improvement",
                "issue_type": "asset_quality",
                "input_fingerprint": "before-image-repair",
                "attempt": 1,
            }
        },
    )
    _invalidate_review_after_image_repair(parent)
    reopened = parent.child_outcomes["review"]
    assert reopened["ok"] is False
    assert reopened["status"] == "needs_orchestrator"
    assert reopened["exit_reason"] == "external_repair_completed"
    assert parent.review_completed is False
    assert parent.final_render_after_review is False
    assert parent.final_view_after_review is False
    assert parent.finalize_attempted is False
    assert parent.finalize_succeeded is False
    allowed, reason = _review_verification_allowed(parent, reopened)
    assert allowed is True
    assert reason == ""


def test_reviewed_needs_improvement_forces_orchestrator_delivery(
    tmp_path: Path,
) -> None:
    captured = {}
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="orchestrator",
        model="pptagent",
        system="system",
        tool_schemas=[
            {"name": "read_file", "input_schema": {"type": "object"}},
            {"name": "delegate_task", "input_schema": {"type": "object"}},
            {"name": "terminal", "input_schema": {"type": "object"}},
        ],
        child_outcomes={"review": {"issue_type": "page_authoring"}},
        review_completed=True,
        finalize_succeeded=True,
        final_view_after_review=True,
        quality_status="needs_improvement",
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )
    assert _orchestrator_delivery_close_status(agent) == "needs_improvement"

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent, [{"role": "user", "content": "continue"}])
    assert captured["tools"] == []
    assert "needs_improvement" in captured["messages"][-1]["content"]


def test_unconsumed_asset_quality_keeps_one_image_repair_route(
    tmp_path: Path,
) -> None:
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="orchestrator",
        child_outcomes={"review": {"issue_type": "asset_quality"}},
        review_completed=True,
        finalize_succeeded=True,
        final_view_after_review=True,
        quality_status="needs_improvement",
    )
    assert _orchestrator_delivery_close_status(agent) == ""
    state = tmp_path / "_trace/image-repair-state.json"
    state.parent.mkdir()
    state.write_text("{}", encoding="utf-8")
    assert _orchestrator_delivery_close_status(agent) == "needs_improvement"


def test_orchestrator_finalize_forces_missing_slide_delegation(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_01.md").write_text("# slide_01\n", encoding="utf-8")
    (tmp_path / "plan/slide_02.md").write_text("# slide_02\n", encoding="utf-8")
    (tmp_path / "renders/slide_01.png").write_bytes(b"pixels")
    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
    )
    result = tools.terminal(
        agent,
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py finalize . --expected 2",
    )
    assert "unstarted_slide_responsibility" in result
    assert "P02" in result
    assert agent.pending_slide_delegation_pages == (2,)


def test_finalize_responsibility_gate_defers_failed_bitmap_page_for_reclassification(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_02.md").write_text(
        "# slide_02\n- needs_bitmap: true\n"
        "- primary_visual_medium: bitmap-real\n",
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_03.md").write_text(
        "# slide_03\n- needs_bitmap: true\n"
        "- primary_visual_medium: bitmap-real\n",
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        image_failed_pages=(3,),
    )
    result = tools.terminal(
        agent,
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py audit . --expected 3",
    )
    assert "P02" in result
    assert "暂不进入责任门：P03" in result
    assert agent.pending_slide_delegation_pages == (2,)


def test_tool_executor_rejects_hidden_mutation_while_slide_responsibility_pending(
    tmp_path: Path,
) -> None:
    target = tmp_path / "plan/slide_03.md"
    target.parent.mkdir()
    target.write_text("before", encoding="utf-8")
    agent = SimpleNamespace(
        role="orchestrator",
        pending_slide_delegation_pages=(3,),
        started=time.time(),
        log=lambda *_args, **_kwargs: None,
        remote_tool_concurrency=1,
    )
    use = SimpleNamespace(
        id="hidden-patch",
        name="write_file",
        input={"path": "plan/slide_03.md", "content": "after"},
    )
    result = _tool_results(agent, [use], 1, [])
    assert "unstarted_slide_responsibility" in str(result)
    assert target.read_text(encoding="utf-8") == "before"


def test_pending_missing_slide_requires_and_clears_complete_delegation(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    slides = tmp_path / "slides"
    plan.mkdir()
    slides.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (plan / "slide_03.md").write_text(
        "# slide_03\n- page_type: content\n- production_group: places\n"
        "- needs_bitmap: false\n- primary_visual_medium: svg-diagram\n",
        encoding="utf-8",
    )
    (slides / "slide_03.html").write_text("<section></section>", encoding="utf-8")

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)
        pending_slide_delegation_pages = (3,)

    parent = Parent()
    missing = _delegate_task(parent, {"tasks": [{
        "role": "image",
        "goal": "不要再次委派 Image",
    }]})
    assert "缺页责任门只接受 Slide 责任单元" in missing
    assert parent.pending_slide_delegation_pages == (3,)

    accepted_payload = json.dumps({"status": "completed", "executed": 1})
    with mock.patch("core.agent_loop._delegate", return_value=accepted_payload):
        accepted = _delegate_task(parent, {"tasks": [{
            "role": "slide",
            "pages": [3],
            "goal": "按重分类后的 SVG 合同制作 P03",
        }]})
    assert accepted == accepted_payload
    assert parent.pending_slide_delegation_pages == ()


def test_pending_missing_slide_survives_downstream_delegation_rejection(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    slides = tmp_path / "slides"
    plan.mkdir()
    slides.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (plan / "slide_03.md").write_text(
        "# slide_03\n- page_type: content\n- production_group: places\n"
        "- needs_bitmap: false\n- primary_visual_medium: svg-diagram\n",
        encoding="utf-8",
    )
    (slides / "slide_03.html").write_text("<section></section>", encoding="utf-8")
    parent = SimpleNamespace(
        skill_language="zh",
        skill_name="mural-presenter-v0-4",
        response_language="zh",
        query_language_hint="zh",
        ws=str(tmp_path),
        pending_slide_delegation_pages=(3,),
    )
    with mock.patch(
        "core.agent_loop._delegate",
        return_value="delegate_task 错误：synthetic downstream rejection",
    ):
        result = _delegate_task(parent, {"tasks": [{
            "role": "slide",
            "pages": [3],
            "goal": "制作 P03",
        }]})
    assert "synthetic downstream rejection" in result
    assert parent.pending_slide_delegation_pages == (3,)


def test_pending_slide_tool_surface_exposes_only_delegate_without_read_state() -> None:
    captured = {}
    agent = SimpleNamespace(
        ws=".",
        role="orchestrator",
        model="pptagent",
        system="system",
        tool_schemas=[
            {"name": "read_file", "input_schema": {"type": "object"}},
            {"name": "delegate_task", "input_schema": {"type": "object"}},
        ],
        pending_slide_delegation_pages=(3, 4),
        model_base_url="http://unused",
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent, [{"role": "user", "content": "continue"}])
    assert [schema["name"] for schema in captured["tools"]] == ["delegate_task"]


def test_checked_single_slide_hides_vision_until_new_render(tmp_path: Path) -> None:
    captured = {}
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="slide",
        model="pptagent",
        system="system",
        tool_schemas=[
            {"name": "vision_analyze", "input_schema": {"type": "object"}},
            {"name": "patch", "input_schema": {"type": "object"}},
        ],
        assigned_slide_pages=(4,),
        slide_group_id="",
        rendered_output_hashes={4: "current-pixels"},
        viewed_output_hashes={4: "current-pixels"},
        trace_label="slide_04",
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent, [{"role": "user", "content": "continue"}])
    assert [schema["name"] for schema in captured["tools"]] == ["patch"]


def test_group_slide_requires_one_current_consistency_sheet_view(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    sheet = renders / "contact-sheet-group-bookends.png"
    sheet.write_bytes(b"current group pixels")
    captured = {}
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="slide",
        model="pptagent",
        system="system",
        tool_schemas=[
            {
                "name": "vision_analyze",
                "input_schema": {
                    "type": "object",
                    "properties": {"image": {"type": "string"}},
                },
            },
            {"name": "patch", "input_schema": {"type": "object"}},
        ],
        assigned_slide_pages=(1, 5),
        slide_group_id="bookends",
        rendered_output_hashes={1: "p1", 5: "p5"},
        viewed_output_hashes={1: "p1", 5: "p5"},
        group_viewed_contact_hash="",
        group_viewed_page_hashes={},
        trace_label="slide_group_bookends",
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )

    assert _slide_pending_group_view(agent) == (
        "renders/contact-sheet-group-bookends.png"
    )

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent, [{"role": "user", "content": "continue"}])
    assert [schema["name"] for schema in captured["tools"]] == ["vision_analyze"]
    assert captured["tools"][0]["input_schema"]["properties"]["image"]["enum"] == [
        "renders/contact-sheet-group-bookends.png"
    ]


def test_stale_group_sheet_exposes_only_render_group_terminal(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    sheet = renders / "contact-sheet-group-bookends.png"
    sheet.write_bytes(b"old group pixels")
    captured = {}
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="slide",
        model="pptagent",
        system="system",
        tool_schemas=[
            {"name": "terminal", "input_schema": {"type": "object"}},
            {"name": "vision_analyze", "input_schema": {"type": "object"}},
            {"name": "patch", "input_schema": {"type": "object"}},
        ],
        assigned_slide_pages=(1, 5),
        slide_group_id="bookends",
        rendered_output_hashes={1: "p1-new", 5: "p5-same"},
        viewed_output_hashes={1: "p1-new", 5: "p5-same"},
        group_rendered_contact_hash=hashlib.sha256(sheet.read_bytes()).hexdigest(),
        group_rendered_page_hashes={1: "p1-old", 5: "p5-same"},
        group_viewed_contact_hash="",
        group_viewed_page_hashes={},
        trace_label="slide_group_bookends",
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent, [{"role": "user", "content": "continue"}])
    assert [schema["name"] for schema in captured["tools"]] == ["terminal"]
    assert "render-group" in captured["messages"][-1]["content"]


def test_duplicate_image_query_closes_only_search_tool() -> None:
    search_agent = SimpleNamespace(
        role="image",
        evidence_scope="open_research",
        requested_slide_count=5,
    )
    assert tools._consume_research_search_budget(search_agent, "same query") == ""
    duplicate = tools._consume_research_search_budget(search_agent, "same query")
    assert "Image 搜索阶段现已关闭" in duplicate
    assert search_agent.image_search_closed is True

    captured = {}
    agent = SimpleNamespace(
        ws=".",
        role="image",
        image_search_closed=True,
        model="pptagent",
        system="system",
        tool_schemas=[
            {"name": "web_search", "input_schema": {"type": "object"}},
            {"name": "image_generate", "input_schema": {"type": "object"}},
            {"name": "write_file", "input_schema": {"type": "object"}},
        ],
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent, [{"role": "user", "content": "continue"}])
    assert [schema["name"] for schema in captured["tools"]] == [
        "image_generate",
        "write_file",
    ]


def test_orchestrator_cannot_mutate_image_owned_assets(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()
    catalog = tmp_path / "assets/catalog.md"
    catalog.write_text("# Asset catalog\n", encoding="utf-8")
    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
    )
    result = tools.patch(
        agent,
        path="assets/catalog.md",
        old_string="# Asset catalog",
        new_string="# Rewritten catalog",
    )
    assert "Image 独占" in result
    assert catalog.read_text(encoding="utf-8") == "# Asset catalog\n"


def test_single_topology_rejects_slide_group(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        ws = str(tmp_path)

    result = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "SlideGroup hero [01,02]: build both pages"}]},
    )
    assert "ownership_topology: grouped" in result


def test_single_topology_accepts_only_complete_special_memory_groups(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    fixtures = {
        1: ("cover", "bookends"),
        4: ("section-divider", "dividers"),
        8: ("section-divider", "dividers"),
        12: ("closing", "bookends"),
    }
    for number, (page_type, group) in fixtures.items():
        (plan / f"slide_{number:02d}.md").write_text(
            f"# P{number}\n- page_type: {page_type}\n- production_group: {group}\n",
            encoding="utf-8",
        )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [
            {"role": "slide", "group_id": "bookends", "pages": [1, 12]},
            {"role": "slide", "group_id": "dividers", "pages": [4, 8]},
        ]})
    assert result == "delegated"
    assert [item["label"] for item in delegate.call_args.args[1]["tasks"]] == [
        "slide_group_bookends",
        "slide_group_dividers",
    ]
    assert tools._is_grouped_skill(SimpleNamespace(
        role="slide",
        assigned_slide_pages=(1, 12),
        slide_group_id="bookends",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
    ))

    incomplete = _delegate_task(Parent(), {"tasks": [{
        "role": "slide",
        "group_id": "bookends",
        "pages": [1],
    }]})
    assert "完整的 `bookends`" in incomplete

    split_member = _delegate_task(Parent(), {"tasks": [{
        "role": "slide",
        "pages": [12],
    }]})
    assert "不能拆成 Single" in split_member

    envelope = _slide_assignment_envelope(
        "只制作第 1 页封面。",
        (1, 12),
        "bookends",
        "zh",
    )
    assert "P01, P12" in envelope
    assert "必须完成全部分配页" in envelope
    assert "不得把任一成员称为另一个责任单元" in envelope
    assert envelope.rfind("[Harness 页面所有权合同]") > envelope.rfind(
        "只制作第 1 页封面。"
    )


def test_solo_bookends_delegation_accepted_at_runtime(
    tmp_path: Path,
) -> None:
    """Solo cover with bookends can be dispatched as group_id=bookends pages=[1]."""
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (plan / "slide_01.md").write_text(
        "# P1\n- page_type: cover\n- production_group: bookends\n",
        encoding="utf-8",
    )
    (plan / "slide_02.md").write_text(
        "# P2\n- page_type: content\n- production_group: p2\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [
            {"role": "slide", "group_id": "bookends", "pages": [1]},
        ]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"][0]["label"] == "slide_group_bookends"


def test_group_exhausted_page_is_frozen_but_unfinished_sibling_remains_available(
    tmp_path: Path,
) -> None:
    """Regression for deck #165: P01 must not consume or strand P10."""
    states = tmp_path / "_trace" / "slide-render-states"
    slides = tmp_path / "slides"
    renders = tmp_path / "renders"
    states.mkdir(parents=True)
    slides.mkdir()
    renders.mkdir()
    p1 = slides / "slide_01.html"
    p10 = slides / "slide_10.html"
    p1.write_text("<section>verified P1</section>", encoding="utf-8")
    p10.write_text("<section>scaffold P10</section>", encoding="utf-8")
    p1_hash = hashlib.sha256(p1.read_bytes()).hexdigest()
    p1_png = renders / "slide_01.png"
    p1_png.write_bytes(b"verified-p1-pixels")
    p1_png_hash = hashlib.sha256(p1_png.read_bytes()).hexdigest()
    (states / "page_01.json").write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "attempt_hashes": {"slide_group_bookends": ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        role="slide",
        ws=str(tmp_path),
        trace_label="slide_group_bookends",
        assigned_slide_pages=(1, 10),
        slide_group_id="bookends",
        expected_output_paths={1: str(p1), 10: str(p10)},
        rendered_output_hashes={1: p1_hash},
        viewed_output_hashes={1: p1_hash},
        vision_critic_results={
            "renders/slide_01.png": {
                "source_sha256": p1_png_hash,
                "verdict": "repair_required",
                "issues": [{"type": "missing_chart", "severity": "major"}],
            }
        },
    )

    assert _slide_group_active_page(agent) == 10
    blocked_patch = _grouped_exhausted_page_tool_error(
        agent, "patch", {"path": "slides/slide_01.html"}
    )
    blocked_render = _grouped_exhausted_page_tool_error(
        agent,
        "terminal",
        {"command": "python skills/mural-presenter-v0-4/scripts/slide.py render . --page 1"},
    )
    assert "group_page_frozen" in blocked_patch
    assert "P10" in blocked_patch
    assert "group_page_frozen" in blocked_render
    assert _grouped_exhausted_page_tool_error(
        agent, "patch", {"path": "slides/slide_10.html"}
    ) == ""
    assert _grouped_exhausted_page_tool_error(
        agent,
        "terminal",
        {
            "command": (
                "python skills/mural-presenter-v0-4/scripts/slide.py render-group . "
                "--group bookends --pages 01,10"
            )
        },
    ) == ""


def test_single_topology_accepts_one_planned_divider_without_group_id(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (plan / "slide_08.md").write_text(
        "# P8\n- page_type: section-divider\n- production_group: dividers\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [{
            "role": "slide",
            "pages": [8],
        }]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"][0]["label"] == "slide_08"

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [{
            "role": "slide",
            "group_id": "dividers",
            "pages": [8],
        }]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"][0]["label"] == "slide_group_dividers"


def test_grouped_topology_accepts_one_planned_special_page(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    (plan / "slide_08.md").write_text(
        "# P8\n- page_type: section-divider\n- production_group: dividers\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [{
            "role": "slide",
            "group_id": "dividers",
            "pages": [8],
        }]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"][0]["label"] == "slide_group_dividers"


def test_single_topology_full_wave_accepts_bookends_and_one_divider(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    for page in range(1, 14):
        if page == 1:
            page_type, group = "cover", "bookends"
        elif page == 13:
            page_type, group = "closing", "bookends"
        elif page == 8:
            page_type, group = "section-divider", "dividers"
        else:
            page_type, group = "content", f"page-{page:02d}"
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n- page_type: {page_type}\n- production_group: {group}\n",
            encoding="utf-8",
        )

    class Parent:
        skill_language = "en"
        skill_name = "mural-presenter-v0-4"
        response_language = "en"
        query_language_hint = "en"
        ws = str(tmp_path)

    tasks = [{
        "role": "slide",
        "group_id": "bookends",
        "pages": [1, 13],
    }]
    tasks.extend(
        {"role": "slide", "pages": [page]}
        for page in range(2, 13)
    )
    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": tasks})
    assert result == "delegated"
    labels = [item["label"] for item in delegate.call_args.args[1]["tasks"]]
    assert labels == ["slide_group_bookends"] + [
        f"slide_{page:02d}" for page in range(2, 13)
    ]


def test_grouped_topology_accepts_only_complete_slide_groups(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(
            Parent(),
            {"tasks": [{"goal": "SlideGroup mechanism [03,04,05]: build group"}]},
        )
    assert result == "delegated"
    specs = delegate.call_args.args[1]["tasks"]
    assert specs == [{
        "role": "slide",
        "label": "slide_group_mechanism",
        "group_id": "mechanism",
        "pages": [3, 4, 5],
        "task": "SlideGroup mechanism [03,04,05]: build group",
    }]

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        isolated = _delegate_task(
            Parent(),
            {"tasks": [{"goal": "SlideGroup closing [13]: build isolated closing"}]},
        )
    assert isolated == "delegated"
    assert delegate.call_args.args[1]["tasks"][0]["pages"] == [13]

    duplicate = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "SlideGroup closing [13,13]: build closing"}]},
    )
    assert "use [13], never [13,13]" in duplicate

    rejected = _delegate_task(
        Parent(),
        {"tasks": [{"goal": "Slide 03: build one page"}]},
    )
    assert "ownership_topology: single" in rejected


def test_structured_delegation_derives_labels_without_goal_syntax(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [{
            "role": "slide",
            "group_id": "mechanism",
            "pages": [3, 4, 5],
            "goal": "兑现跨页机制编码",
        }]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"] == [{
        "role": "slide",
        "label": "slide_group_mechanism",
        "group_id": "mechanism",
        "pages": [3, 4, 5],
        "task": "兑现跨页机制编码",
    }]


def test_non_slide_delegation_ignores_provider_copied_page_scope(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)

    with mock.patch("core.agent_loop._delegate", return_value="delegated") as delegate:
        result = _delegate_task(Parent(), {"tasks": [{
            "role": "image",
            "pages": [1, 14],
            "group_id": "visuals",
            "goal": "准备 P01 与 P14 的位图",
        }]})
    assert result == "delegated"
    assert delegate.call_args.args[1]["tasks"] == [{
        "role": "image",
        "label": "image",
        "task": "准备 P01 与 P14 的位图",
    }]


def test_style_lock_is_extracted_once_for_visual_workers(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n\n## 视觉契约\n档案纸纹；禁通用深蓝金。\n\n"
        "## 特殊页\n封面使用真实档案主体。\n\n"
        "## 主题变量\n```css\n:root { --content-canvas: #eee; }\n```\n",
        encoding="utf-8",
    )
    lock = _style_lock_injection(str(tmp_path), "zh")
    assert "<style_lock" in lock
    assert "档案纸纹" in lock
    assert "--content-canvas" in lock
    assert "通用暗底模板" in lock


def test_vision_issue_ledger_requires_changed_pixel_closure(tmp_path: Path) -> None:
    agent = SimpleNamespace(ws=str(tmp_path), trace_label="slide_03", role="slide")
    scan = {
        "visible_subjects": ["one portrait"],
        "text_regions": ["title readable"],
        "regions": ["left image, right copy"],
        "edges": {edge: "clear" for edge in ("top", "right", "bottom", "left")},
    }
    issue = {
        "severity": "major",
        "type": "text_overlap",
        "location": "lower right",
        "evidence": "two labels visibly overlap",
        "suggested_fix": "separate the labels",
    }
    _update_vision_issue_ledger(agent, "renders/slide_03.png", "hash-a", {
        "verdict": "repair_required", "summary": "overlap", "scan": scan,
        "issues": [issue],
    })
    assert len(_open_vision_issues(agent)) == 1

    _update_vision_issue_ledger(agent, "renders/slide_03.png", "hash-a", {
        "verdict": "ready", "summary": "generic ready", "scan": scan, "issues": [],
    })
    assert len(_open_vision_issues(agent)) == 1

    _update_vision_issue_ledger(agent, "renders/slide_03.png", "hash-b", {
        "verdict": "ready", "summary": "overlap is visibly gone", "scan": scan,
        "issues": [],
    })
    assert _open_vision_issues(agent) == []
    ledger = json.loads((tmp_path / "_trace/vision-issues.json").read_text())
    assert ledger["issues"][0]["opened_pixel_sha256"] == "hash-a"
    assert ledger["issues"][0]["closed_pixel_sha256"] == "hash-b"


def test_ready_minor_vision_observation_does_not_open_blocking_ledger(
    tmp_path: Path,
) -> None:
    agent = SimpleNamespace(ws=str(tmp_path), trace_label="image")
    scan = {
        "visible_subjects": [], "text_regions": [], "regions": [],
        "edges": {"top": "ok", "right": "ok", "bottom": "ok", "left": "ok"},
    }
    _update_vision_issue_ledger(agent, "assets/contact-sheet.png", "hash-a", {
        "verdict": "ready",
        "summary": "usable set",
        "scan": scan,
        "issues": [{
            "severity": "minor",
            "type": "contact_sheet_balance",
            "location": "row 2",
            "evidence": "preview grid has an empty cell",
        }],
    })
    assert not (tmp_path / "_trace/vision-issues.json").exists()


def test_source_image_findings_never_enter_final_slide_ledger(tmp_path: Path) -> None:
    agent = SimpleNamespace(
        ws=str(tmp_path),
        trace_label="image",
        role="image",
        vision_critic_results={
            "inputs/paper.pdf.pages/page_029.png": {
                "verdict": "repair_required",
                "summary": "immutable source label is clipped",
                "source_sha256": "source-hash",
            },
        },
    )
    result = {
        "verdict": "repair_required",
        "summary": "the immutable source has a clipped paper label",
        "scan": {
            "visible_subjects": ["paper chart"],
            "text_regions": ["axis label"],
            "regions": ["plot"],
            "edges": {edge: "clear" for edge in ("top", "right", "bottom", "left")},
        },
        "issues": [{
            "severity": "major",
            "type": "source_label_clipping",
            "location": "right edge",
            "evidence": "label is clipped in the source paper page",
        }],
    }
    _update_vision_issue_ledger(
        agent,
        "inputs/paper.pdf.pages/page_029.png",
        "source-hash",
        result,
    )
    _update_vision_issue_ledger(
        agent,
        "assets/figure9.png",
        "asset-hash",
        result,
    )
    assert not (tmp_path / "_trace/vision-issues.json").exists()
    assert _unresolved_vision_critic_issues(agent) == []

    trace = tmp_path / "_trace"
    trace.mkdir()
    (trace / "vision-issues.json").write_text(json.dumps({
        "schema": "mural.vision-issues.v1",
        "issues": [{
            "id": "VIS-LEGACY",
            "status": "open",
            "page": None,
            "source": "inputs/paper.pdf.pages/page_029.png",
        }],
    }))
    assert _open_vision_issues(agent) == []


def test_image_contact_sheet_is_triage_only_but_fullres_issue_blocks(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "contact-sheet.png").write_bytes(b"sheet")
    (assets / "hero.png").write_bytes(b"hero")
    scan = {
        "visible_subjects": ["candidate"],
        "text_regions": [],
        "regions": ["thumbnail"],
        "edges": {edge: "clear" for edge in ("top", "right", "bottom", "left")},
    }
    agent = SimpleNamespace(
        ws=str(tmp_path),
        trace_label="image",
        role="image",
        vision_critic_results={
            "assets/contact-sheet.png": {
                "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
                "verdict": "repair_required",
                "summary": "thumbnail looks suspicious",
            },
            "assets/hero.png": {
                "source_sha256": hashlib.sha256(b"hero").hexdigest(),
                "verdict": "repair_required",
                "summary": "visible watermark on original",
            },
        },
    )
    unresolved = _unresolved_vision_critic_issues(agent)
    assert [source for source, _ in unresolved] == ["assets/hero.png"]

    _update_vision_issue_ledger(agent, "assets/contact-sheet.png", "sheet-hash", {
        "verdict": "repair_required",
        "summary": "thumbnail candidate",
        "scan": scan,
        "issues": [{
            "severity": "major",
            "type": "possible_watermark",
            "location": "thumbnail 2",
            "evidence": "small mark",
        }],
    })
    ledger = tmp_path / "_trace/vision-issues.json"
    assert not ledger.exists()


def test_review_contact_sheets_are_triage_only_but_page_issue_blocks(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "contact-sheet-special.png").write_bytes(b"sheet")
    (renders / "slide_02.png").write_bytes(b"page")
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="review",
        vision_critic_results={
            "renders/contact-sheet-special.png": {
                "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
                "verdict": "repair_required",
                "summary": "thumbnail candidate",
            },
            "renders/slide_02.png": {
                "source_sha256": hashlib.sha256(b"page").hexdigest(),
                "verdict": "repair_required",
                "summary": "current page defect",
            },
        },
    )
    unresolved = _unresolved_vision_critic_issues(agent)
    assert [source for source, _ in unresolved] == ["renders/slide_02.png"]
    scan = {
        "visible_subjects": ["candidate"],
        "text_regions": [],
        "regions": ["thumbnail"],
        "edges": {edge: "clear" for edge in ("top", "right", "bottom", "left")},
    }
    _update_vision_issue_ledger(
        agent,
        "renders/contact-sheet-special.png",
        "sheet-hash",
        {
            "verdict": "repair_required",
            "summary": "thumbnail candidate",
            "scan": scan,
            "issues": [{
                "severity": "major",
                "type": "candidate_layout",
                "location": "P02 thumbnail",
                "evidence": "possible overlap",
            }],
        },
    )
    assert not (tmp_path / "_trace/vision-issues.json").exists()


def test_delegation_reloads_canonical_topology_instead_of_cached_hint(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    deck = tmp_path / "plan/deck.md"

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        query_language_hint = "zh"
        ws = str(tmp_path)
        ownership_topology = "grouped"

    parent = Parent()
    deck.write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    with mock.patch("core.agent_loop._delegate", return_value="delegated"):
        result = _delegate_task(
            parent,
            {"tasks": [{"goal": "Slide 01: build cover"}]},
        )
    assert result == "delegated"
    assert parent.ownership_topology == "single"
    assert tools._is_grouped_skill(parent) is False

    deck.write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    with mock.patch("core.agent_loop._delegate", return_value="delegated"):
        result = _delegate_task(
            parent,
            {"tasks": [{"goal": "SlideGroup cover [01]: build cover"}]},
        )
    assert result == "delegated"
    assert parent.ownership_topology == "grouped"
    assert tools._is_grouped_skill(parent) is True


def test_v03_keeps_research_and_revision_contracts(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    agent = Agent(
        "grouped-contract-test",
        str(tmp_path),
        "制作八页演示",
        {
            "skill_name": "mural-presenter-v0-4",
            "skill_language": "zh",
        },
    )
    assert agent.research_required is True

    prompt = _revision_prompt(
        {"query": "介绍台风白海豚", "lang": "zh"},
        {"instruction": "把主题纠正为台风白海豚"},
        grouped=True,
    )
    assert "Review: mode=simple_edit" in prompt
    assert "plan/revision-impact.md" in prompt
    assert "group_id: GROUP" in prompt
    assert "不得把受影响页组拆成单页" in prompt

    class RevisionOrchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = True

    error = tools._model_write_error(
        RevisionOrchestrator(), "slides/slide_02.html"
    )
    assert error is not None
    assert "不允许 Orchestrator 直接修改" in error


def test_image_and_slide_share_dependency_aware_wave(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )

    class Parent:
        skill_language = "zh"
        skill_name = "mural-presenter-v0-4"
        response_language = "zh"
        child_outcomes = {}
        ws = str(tmp_path)

    parent = Parent()
    with mock.patch("core.agent_loop._delegate", return_value="dependency-aware") as delegate:
        result = _delegate_task(
            parent,
            {
                "tasks": [
                    {"goal": "Image: resolve all raster assets"},
                    {"goal": "Slide 03: complete the bitmap-independent page"},
                ]
            },
        )
    assert result == "dependency-aware"
    assert [spec["role"] for spec in delegate.call_args.args[1]["tasks"]] == ["image", "slide"]


def test_bitmap_slide_releases_when_image_finishes_without_waiting_for_other_slides(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("speech\n", encoding="utf-8")
    for page, needs_bitmap in ((1, False), (2, True)):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- production_group: page-{page:02d}\n"
            f"- needs_bitmap: {str(needs_bitmap).lower()}\n",
            encoding="utf-8",
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            "<section></section>\n", encoding="utf-8"
        )

    profile = {
        "active_roles": ["image", "slide", "review"],
        "workflow": {
            "research": "omitted",
            "research_mode": "off",
            "material_stage": "omitted",
            "image": "bitmap_available",
        },
        "tools": {
            "web_search": False,
            "web_extract": False,
            "image_generate": True,
            "vision_analyze": True,
        },
    }
    parent = Agent(
        "dependency-scheduler",
        str(tmp_path),
        "build",
        {
            "skill_name": "mural-presenter-v0-4",
            "skill_language": "zh",
            "_runtime_capabilities": profile,
        },
    )
    parent.research_required = False
    events: list[str] = []
    lock = threading.Lock()

    def fake_child(_parent, _index, spec):
        role = spec["role"]
        label = spec["label"]
        if role == "image":
            time.sleep(0.03)
            with lock:
                events.append("image_done")
        elif label == "slide_01":
            time.sleep(0.15)
            with lock:
                events.append("plain_done")
        else:
            with lock:
                events.append("bitmap_started")
        return {
            "label": label,
            "role": role,
            "ok": True,
            "exit_reason": "completed",
            "completed_pages": list(spec.get("pages") or []),
            "incomplete_pages": [],
            "renders": 0,
            "views": 0,
            "summary": "done",
        }

    specs = [
        {"role": "image", "label": "image", "task": "resolve"},
        {"role": "slide", "label": "slide_01", "pages": [1], "task": "plain"},
        {"role": "slide", "label": "slide_02", "pages": [2], "task": "bitmap"},
    ]
    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        result = json.loads(_delegate(parent, {"tasks": specs}))

    assert result["failed"] == 0
    assert events.index("image_done") < events.index("bitmap_started")
    assert events.index("bitmap_started") < events.index("plain_done")


def test_completed_slide_cannot_be_reopened_in_normal_production(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_01.md").write_text(
        "# P1\n- needs_bitmap: false\n", encoding="utf-8"
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    parent = Agent(
        "repair-reopen",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.child_outcomes = {"slide_01": {"ok": True}}

    ordinary = json.loads(_delegate(parent, {"tasks": [{
        "role": "slide", "label": "slide_01", "pages": [1], "task": "retry"
    }]}))
    assert ordinary["status"] == "already_completed"

    with mock.patch("core.agent_loop._run_child") as run_child:
        repaired = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide",
            "label": "slide_01",
            "pages": [1],
            "task": "fix confirmed duplicate credit",
            "repair": True,
        }]}))
    assert run_child.call_count == 0
    assert repaired["status"] == "already_completed"
    assert repaired["executed"] == 0


def test_repair_required_slide_is_handed_to_review_without_r2(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_01.md").write_text(
        "# P1\n- needs_bitmap: false\n", encoding="utf-8"
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section>authored</section>\n", encoding="utf-8"
    )
    parent = Agent(
        "repair-handoff",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.child_outcomes = {"slide_01": {
        "ok": True,
        "status": "repair_required",
        "repair_issue": {
            "status": "repair_required",
            "pages": [1],
            "issue_type": "page_authoring",
        },
    }}
    with mock.patch("core.agent_loop._run_child") as run_child:
        result = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide", "label": "slide_01", "pages": [1], "repair": True,
        }]}))
    assert result["status"] == "already_completed"
    assert result["executed"] == 0
    assert run_child.call_count == 0


def test_scheduler_only_asset_pending_slide_runs_once_after_image_recovery(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_01.md").write_text(
        "# P1\n- needs_bitmap: true\n- primary_visual_medium: bitmap-material\n",
        encoding="utf-8",
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    parent = Agent(
        "asset-release",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.image_completed = True
    parent.image_ready_pages = (1,)
    parent.child_outcomes = {"slide_01": {
        "role": "slide",
        "ok": False,
        "status": "asset_pending",
        "exit_reason": "asset_pending",
        "trace_label": None,
        "completed_pages": [],
        "attempt": 1,
    }}
    completed = {
        "label": "slide_01",
        "trace_label": "slide_01",
        "role": "slide",
        "ok": True,
        "status": "ready",
        "exit_reason": "text_response",
        "completed_pages": [1],
        "incomplete_pages": [],
        "renders": 1,
        "views": 1,
        "summary": "ready",
    }
    with mock.patch("core.agent_loop._run_child", return_value=completed) as run_child:
        result = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide", "label": "slide_01", "pages": [1], "task": "build"
        }]}))
    assert result["executed"] == 1
    assert result["failed"] == 0
    assert run_child.call_count == 1
    assert run_child.call_args.args[2].get("_recovery_kind") is None


def test_one_shot_shared_env_maps_to_isolated_same_model_aux() -> None:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(HARNESS)
    env["VISION_BACKEND"] = "one_shot"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "from core import config; "
            "print(config.VISION_BACKEND_REQUESTED, config.VISION_BACKEND)",
        ],
        capture_output=True,
        text=True,
        env=env,
        check=True,
    )
    assert result.stdout.strip() == "one_shot same_model_aux"


def test_gemini_backend_alias_maps_to_external_model() -> None:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(HARNESS)
    env["VISION_BACKEND"] = "gemini"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            "from core import config; print(config.VISION_BACKEND)",
        ],
        capture_output=True,
        text=True,
        env=env,
        check=True,
    )
    assert result.stdout.strip() == "external_model"


def test_same_model_vision_critic_uses_fresh_history_free_request() -> None:
    captured = {}

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[SimpleNamespace(
            type="text",
            text='{"schema":"mural.vision-critic.v1","verdict":"ready",'
                 '"summary":"clear","observations":[],"issues":[]}',
        )])

    agent = SimpleNamespace(
        model="main-multimodal-model",
        cfg={"model_base_url": "http://main-model"},
        log=lambda *_args, **_kwargs: None,
        deadline_monotonic=time.monotonic() + 60,
    )
    with mock.patch("core.model_call.call_with_tools", side_effect=fake_call):
        result = model_call.call_vision_same_model_auxiliary(
            agent=agent,
            image_bytes=b"pixels",
            media_type="image/png",
            query="check clipping",
            system_prompt="isolated critic system",
        )

    assert '"verdict":"ready"' in result
    assert captured["model"] == agent.model
    assert captured["system"] == "isolated critic system"
    assert captured["tools"] is None
    assert captured["thinking"] is False
    assert len(captured["messages"]) == 1
    assert captured["messages"][0]["role"] == "user"
    assert captured["messages"][0]["content"][1]["text"] == "check clipping"


def test_vision_tool_returns_structured_text_from_same_model_aux(tmp_path: Path) -> None:
    from PIL import Image

    renders = tmp_path / "renders"
    renders.mkdir()
    Image.new("RGB", (1600, 900), "white").save(renders / "slide_01.png")
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ownership_topology="single",
        slide_pixel_inspections=0,
        assigned_slide_pages=(1,),
        slide_group_id="",
        n_views=0,
        max_vision_edge=1600,
        render_script="skills/mural-presenter-v0-4/scripts/slide.py",
        model="main-multimodal-model",
        cfg={"model_base_url": "http://main-model"},
        deadline_monotonic=time.monotonic() + 60,
        ws=str(tmp_path),
        read_path=lambda path: str(tmp_path / path),
        log=lambda *_args, **_kwargs: None,
    )
    critic = json.dumps({
        "schema": "mural.vision-critic.v1",
        "verdict": "repair_required",
        "summary": "右下图片孤立且过小",
        "observations": ["主标题位于左上"],
        "issues": [{
            "severity": "major",
            "type": "isolated_small_image",
            "location": "right-bottom",
            "evidence": "图片仅占画布约 8%",
            "suggested_fix": "扩大为右侧主视觉区",
        }],
    }, ensure_ascii=False)
    with mock.patch("core.tools.config.VISION_BACKEND", "same_model_aux"), mock.patch(
        "core.tools.model_call.call_vision_same_model_auxiliary",
        return_value=critic,
    ) as auxiliary:
        result = tools.vision_analyze(
            agent,
            image="renders/slide_01.png",
            query="检查孤立小图与四边安全区",
        )

    assert result["vision_verdict"] == "repair_required"
    assert result["vision_result"]["issues"][0]["type"] == "isolated_small_image"
    assert result["vision_backend"] == "same_model_aux:main-multimodal-model"
    assert result["source_sha256"]
    assert auxiliary.call_args.kwargs["system_prompt"].startswith(
        "You are an isolated visual agent."
    )
    assert auxiliary.call_args.kwargs["query"].startswith("FIXED OPEN SCAN")
    assert "检查孤立小图与四边安全区" in auxiliary.call_args.kwargs["query"]


def test_vision_tool_result_keeps_pixels_out_of_main_agent_history(tmp_path: Path) -> None:
    trace_dir = tmp_path / "trace"
    (trace_dir / "images").mkdir(parents=True)
    agent = SimpleNamespace(
        role="image",
        started=time.time(),
        log=lambda *_args, **_kwargs: None,
        remote_tool_concurrency=1,
        ws=str(tmp_path),
        trace_dir=trace_dir,
        n_views=0,
        image_by_tool={},
        vision_critic_results={},
    )
    value = {
        "vision_analysis": '{"verdict":"ready","issues":[]}',
        "vision_result": {"verdict": "ready", "issues": []},
        "vision_verdict": "ready",
        "vision_summary": "asset is usable",
        "image_b64": __import__("base64").b64encode(b"pixels").decode(),
        "media_type": "image/png",
        "path": "assets/item.png",
        "source_sha256": __import__("hashlib").sha256(b"pixels").hexdigest(),
        "vision_backend": "same_model_aux:test",
    }
    use = SimpleNamespace(
        id="tool-vision-1",
        name="vision_analyze",
        input={"image": "assets/item.png", "query": "inspect asset"},
    )
    with mock.patch("core.agent_loop.tools.dispatch", return_value=value):
        results = _tool_results(agent, [use], 1, [])

    assert isinstance(results[0]["content"], str)
    assert '"verdict":"ready"' in results[0]["content"]
    assert "image" not in results[0]["content"]
    assert (trace_dir / "images/view_001.png").read_bytes() == b"pixels"


def test_review_edit_invalidates_finalize_pixels_and_contact_sheet(
    tmp_path: Path,
) -> None:
    trace_dir = tmp_path / "trace"
    (trace_dir / "images").mkdir(parents=True)
    agent = SimpleNamespace(
        role="review",
        started=time.time(),
        log=lambda *_args, **_kwargs: None,
        remote_tool_concurrency=1,
        ws=str(tmp_path),
        trace_dir=trace_dir,
        n_views=0,
        image_by_tool={},
        vision_critic_results={},
        review_changed=False,
        final_render_after_review=True,
        final_view_after_review=True,
        finalize_succeeded=True,
        review_contact_sheet_inspected=True,
        review_revision_rounds=0,
        review_mutation_since_contact_scan=False,
    )
    use = SimpleNamespace(
        id="tool-patch-1",
        name="patch",
        input={"path": "slides/slide_01.html", "old": "before", "new": "after"},
    )
    with mock.patch("core.agent_loop.tools.dispatch", return_value="已编辑 slides/slide_01.html"):
        _tool_results(agent, [use], 1, [])

    assert agent.review_changed
    assert not agent.final_render_after_review
    assert not agent.final_view_after_review
    assert not agent.finalize_succeeded
    assert not agent.review_contact_sheet_inspected
    assert agent.review_mutation_since_contact_scan

    agent.final_render_after_review = True
    vision_use = SimpleNamespace(
        id="tool-vision-review-1",
        name="vision_analyze",
        input={"image": "renders/contact-sheet.png", "query": "whole-deck scan"},
    )
    vision_value = {
        "vision_analysis": '{"verdict":"ready","issues":[]}',
        "vision_result": {"verdict": "ready", "issues": []},
        "vision_verdict": "ready",
        "vision_summary": "whole deck is inspectable",
        "image_b64": __import__("base64").b64encode(b"sheet-pixels").decode(),
        "media_type": "image/png",
        "path": "renders/contact-sheet.png",
        "source_sha256": hashlib.sha256(b"sheet-pixels").hexdigest(),
        "vision_backend": "same_model_aux:test",
    }
    with mock.patch("core.agent_loop.tools.dispatch", return_value=vision_value):
        _tool_results(agent, [vision_use], 2, [])

    assert agent.review_revision_rounds == 1
    assert not agent.review_mutation_since_contact_scan

    # Re-reading the same finalized contact sheet without a new mutation is
    # inspection, not another effective repair round.
    with mock.patch("core.agent_loop.tools.dispatch", return_value=vision_value):
        _tool_results(agent, [vision_use], 3, [])
    assert agent.review_revision_rounds == 1


def test_review_three_round_stop_line_accepts_current_needs_improvement(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    page = renders / "slide_08.png"
    page.write_bytes(b"current-page")
    digest = hashlib.sha256(b"current-page").hexdigest()
    agent = SimpleNamespace(
        role="review",
        ws=str(tmp_path),
        review_revision_rounds=config.REVIEW_MAX_ATTEMPTS,
        required_review_pages=(8,),
        review_viewed_page_hashes={8: digest},
        review_changed=True,
        review_contact_sheet_inspected=True,
        finalize_succeeded=True,
        final_view_after_review=True,
        final_text=(
            "status: needs_orchestrator\n"
            "issue_type: page_authoring\n"
            "pages: 08\n"
            "evidence: minor hierarchy issue remains after bounded repair\n"
            "blocking: no\n"
            "final_pixels_inspected: yes"
        ),
    )

    assert _review_budget_pending_page_views(agent) == []
    assert _finish_gap(agent) == ""

    page.write_bytes(b"newer-page")
    assert _review_budget_pending_page_views(agent) == [8]


def test_review_closure_budget_stays_in_same_session() -> None:
    agent = SimpleNamespace(role="review", review_revision_rounds=0)
    assert _review_in_session_closure_nudge_limit(agent) == 3
    agent.review_revision_rounds = 1
    assert _review_in_session_closure_nudge_limit(agent) == 2
    agent.review_revision_rounds = 2
    assert _review_in_session_closure_nudge_limit(agent) == 1
    agent.review_revision_rounds = 3
    assert _review_in_session_closure_nudge_limit(agent) == 0

    agent.role = "slide"
    agent.review_revision_rounds = 0
    assert _review_in_session_closure_nudge_limit(agent) == 0


def test_review_locally_repairable_no_edit_handoff_is_nudged_in_session() -> None:
    agent = SimpleNamespace(
        role="review",
        review_changed=False,
        final_text=(
            "status: needs_orchestrator\n"
            "issue_type: page_authoring\n"
            "pages: 07, 08\n"
            "evidence: current pixels show overlap\n"
            "blocking: no\n"
            "final_pixels_inspected: yes"
        ),
    )
    gap = _review_self_repair_gap(agent)
    assert "patch" in gap
    assert "不要把可本地修复的问题退回 Orchestrator" in gap

    agent.review_changed = True
    assert _review_self_repair_gap(agent) == ""
    agent.review_changed = False
    agent.final_text = agent.final_text.replace(
        "issue_type: page_authoring", "issue_type: render_capture"
    )
    assert _review_self_repair_gap(agent) == ""


def test_deepseek_no_thinking_is_explicitly_disabled() -> None:
    captured = {}

    class Response:
        usage = None
        content = []

    class Messages:
        def create(self, **kwargs):
            captured.update(kwargs)
            return Response()

    class Client:
        messages = Messages()

    with mock.patch.object(model_call, "_client", return_value=Client()):
        result = model_call.call_with_tools(
            model="bailian/deepseek-v4-flash-0731",
            system="test",
            messages=[{"role": "user", "content": "test"}],
            tools=None,
            max_tokens=16,
            thinking=False,
            effort="high",
            max_retries=1,
        )

    assert result is not None
    assert captured["thinking"] == {"type": "disabled"}
    assert "output_config" not in captured


def test_research_receives_verbatim_raw_user_query_from_harness(tmp_path: Path) -> None:
    raw_query = '做个 8 页 PPT，介绍台风“白海豚”'
    child = Agent(
        "raw-query-test",
        str(tmp_path),
        "Research: investigate the Chinese white dolphin species",
        {
            "skill_name": "mural-presenter-v0-4",
            "skill_language": "zh",
            "_raw_user_query": raw_query,
        },
        role="research",
        label="research",
    )

    assert json.dumps(raw_query, ensure_ascii=False) in child.system
    assert "委派 goal 只是 Orchestrator 的研究假设" in child.system


def test_research_receives_latest_revision_as_separate_user_context(
    tmp_path: Path,
) -> None:
    child = Agent(
        "revision-query-test",
        str(tmp_path),
        "Research: verify the proposed topic",
        {
            "skill_name": "mural-presenter-v0-4",
            "skill_language": "zh",
            "_raw_user_query": "介绍台风白海豚",
            "_revision_instruction": "不是动物，改回台风主题",
        },
        role="research",
        label="research",
    )

    assert "介绍台风白海豚" in child.system
    assert "不是动物，改回台风主题" in child.system
    assert "明确的最新修改优先" in child.system


def test_revision_prompt_routes_simple_and_complex_edits() -> None:
    prompt = _revision_prompt(
        {"query": "介绍台风白海豚", "lang": "zh"},
        {"instruction": "把主题纠正为台风白海豚"},
    )
    assert "Review: mode=simple_edit" in prompt
    assert "plan/revision-impact.md" in prompt
    assert "主题/实体纠正" in prompt
    assert "不能当成局部换字" in prompt


def test_revision_route_is_locked_and_complex_requires_impact_map(
    tmp_path: Path,
) -> None:
    class Parent:
        revision_mode = True
        revision_route = ""
        ws = str(tmp_path)

    parent = Parent()
    error = _revision_delegation_error(
        parent,
        [{"role": "review", "task": "Review: mode=simple_edit; page 2 title"}],
    )
    assert error == ""
    assert parent.revision_route == "simple_edit"
    assert "只能使用唯一" in _revision_delegation_error(
        parent,
        [{"role": "slide", "task": "Slide 02: rewrite page"}],
    )

    complex_parent = Parent()
    assert "revision-impact.md" in _revision_delegation_error(
        complex_parent,
        [{"role": "research", "task": "Research: resolve the named storm"}],
    )
    impact = tmp_path / "plan/revision-impact.md"
    impact.parent.mkdir(parents=True, exist_ok=True)
    impact.write_text("# Impact\n- facts: all pages\n", encoding="utf-8")
    assert _revision_delegation_error(
        complex_parent,
        [{"role": "research", "task": "Research: resolve the named storm"}],
    ) == ""
    assert complex_parent.revision_route == "complex_edit"


def test_revision_orchestrator_cannot_patch_slide_html_directly() -> None:
    class RevisionOrchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = True

    error = tools._model_write_error(
        RevisionOrchestrator(), "slides/slide_02.html"
    )
    assert error is not None
    assert "不允许 Orchestrator 直接修改" in error


def test_initial_orchestrator_cannot_rewrite_plan_after_slide_production() -> None:
    agent = SimpleNamespace(
        role="orchestrator",
        label="orchestrator",
        skill_name="mural-presenter-v0-4",
        revision_mode=False,
        delegated_roles={"research", "slide"},
    )
    error = tools._model_write_error(agent, "plan/slide_04.md")
    assert error is not None
    assert "已冻结" in error
    assert "不得" in error


def test_initial_page_plans_use_structured_batch_tool(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- page_count: 8\n",
        encoding="utf-8",
    )

    class Orchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        ws = str(tmp_path)
        render_script = str(
            REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
        )

        def safe(self, path):
            return str(tmp_path / path)

    agent = Orchestrator()
    single = tools.write_file(agent, "plan/slide_01.md", "# slide 1\n")
    assert "不得逐页串行" in single
    direct_manifest = tools.write_file(agent, "plan/plan-batch.json", "{}")
    assert "write_plan_batch" in direct_manifest

    files = [
        {"path": f"plan/slide_{page:02d}.md", "content": f"# slide {page}\n"}
        for page in range(1, 5)
    ]
    applied = tools.write_plan_batch(agent, files)
    assert "pages=01,02,03,04" in applied
    assert all((plan / f"slide_{page:02d}.md").is_file() for page in range(1, 5))
    assert not (plan / "plan-batch.json").exists()

    stringified = [
        {"path": "plan/slide_05.md", "content": "# slide 5\n"},
        {"path": "plan/slide_06.md", "content": "# slide 6\n"},
    ]
    decoded = tools.write_plan_batch(
        agent,
        json.dumps(stringified, ensure_ascii=False),
    )
    assert "pages=05,06" in decoded

    corrected = [
        {
            "path": f"plan/slide_{page:02d}.md",
            "content": f"# corrected slide {page}\n",
        }
        for page in range(1, 5)
    ]
    recovered_overwrite = tools.write_plan_batch(agent, corrected)
    assert "auto_replace_existing=true" in recovered_overwrite
    assert (plan / "slide_01.md").read_text(encoding="utf-8") == "# corrected slide 1\n"
    assert not (plan / "plan-batch.json").exists()

    corrected_again = [
        {
            "path": f"plan/slide_{page:02d}.md",
            "content": f"# corrected again slide {page}\n",
        }
        for page in range(1, 5)
    ]
    exhausted = tools.write_plan_batch(agent, corrected_again)
    assert "[auto_recovery_exhausted]" in exhausted
    assert (plan / "slide_01.md").read_text(encoding="utf-8") == "# corrected slide 1\n"
    recovered = tools.write_plan_batch(
        agent,
        corrected_again,
        replace_existing=True,
    )
    assert "replace_existing=true" in recovered
    assert (plan / "slide_01.md").read_text(encoding="utf-8") == "# corrected again slide 1\n"

    bad_files = [
        {"path": "plan/slide_07.md", "content": "# slide 7\n"},
        {"path": "plan/slide_09.md", "content": "# slide 9\n"},
    ]
    nonconsecutive = tools.write_plan_batch(agent, bad_files)
    assert "[non_contiguous]" in nonconsecutive
    assert not (plan / "slide_07.md").exists()
    assert not (plan / "slide_09.md").exists()

    tool_names = {
        schema["name"]
        for schema in tools.agent_tools(
            "orchestrator", skill_name="mural-presenter-v0-4"
        )
    }
    assert "write_file" in tool_names
    assert "write_plan_batch" in tool_names
    assert "write_files" not in tool_names

    batch_schema = next(
        schema
        for schema in tools.agent_tools(
            "orchestrator", skill_name="mural-presenter-v0-4"
        )
        if schema["name"] == "write_plan_batch"
    )
    assert batch_schema["input_schema"]["properties"]["replace_existing"]["type"] == "boolean"
    assert batch_schema["input_schema"]["properties"]["files"]["maxItems"] == 32


def test_plan_batch_overwrite_is_blocked_after_page_production_starts(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "slides/slide_01.html").write_text(
        '<section class="slide"></section>',
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_01.md").write_text("# original 1", encoding="utf-8")

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        delegated_roles = set()
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    files = [
        {"path": "plan/slide_01.md", "content": "# corrected 1"},
        {"path": "plan/slide_02.md", "content": "# corrected 2"},
    ]
    result = tools.write_plan_batch(
        Orchestrator(),
        files,
        replace_existing=True,
    )
    assert "[production_started]" in result
    omitted_flag = tools.write_plan_batch(Orchestrator(), files)
    assert "[production_started]" in omitted_flag
    assert (tmp_path / "plan/slide_01.md").read_text(encoding="utf-8") == "# original 1"


def test_dispatch_normalizes_plan_batch_alias_to_files() -> None:
    captured = {}

    def fake_batch(agent, files, replace_existing=False):
        captured["files"] = files
        captured["replace_existing"] = replace_existing
        return "status:PASS"

    payload = [{"path": "plan/slide_01.md", "content": "# slide_01"}]
    with mock.patch.dict(tools.BUILTINS, {"write_plan_batch": fake_batch}):
        result = tools.dispatch(
            SimpleNamespace(),
            "write_plan_batch",
            {"plan_batch": payload, "replace_existing": True},
        )
    assert result == "status:PASS"
    assert captured == {"files": payload, "replace_existing": True}


def test_write_file_dispatch_recovers_only_unambiguous_single_slide_path() -> None:
    received: list[tuple[str, str]] = []

    def fake_write_file(_agent, path, content):
        received.append((path, content))
        return "ok"

    class SingleSlide:
        role = "slide"
        assigned_slide_pages = (6,)

    class GroupedSlide:
        role = "slide"
        assigned_slide_pages = (6, 7)

    class Orchestrator:
        role = "orchestrator"
        assigned_slide_pages = (6,)

    with mock.patch.dict(tools.BUILTINS, {"write_file": fake_write_file}):
        assert tools.dispatch(
            SingleSlide(), "write_file", {"content": "<section>page</section>"}
        ) == "ok"
        grouped = tools.dispatch(
            GroupedSlide(), "write_file", {"content": "<section>group</section>"}
        )
        orchestrator = tools.dispatch(
            Orchestrator(), "write_file", {"content": "brief"}
        )

    assert received == [("slides/slide_06.html", "<section>page</section>")]
    assert "缺少必填参数 ['path']" in grouped
    assert "缺少必填参数 ['path']" in orchestrator


def test_write_plan_batch_recovers_singletons_and_large_transport_batches(tmp_path: Path) -> None:
    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        ws = str(tmp_path)
        role_script = str(
            REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
        )

        def safe(self, path):
            candidate = Path(path)
            if candidate.is_absolute():
                return str(candidate)
            return str(tmp_path / candidate)

    agent = Orchestrator()
    one = [{"path": "plan/slide_01.md", "content": "# one"}]
    singleton_result = tools.write_plan_batch(agent, one)
    assert "status:PASS" in singleton_result
    assert "did not overwrite any other plan" in singleton_result
    assert "do not switch to serial write_file" in singleton_result
    assert (tmp_path / "plan" / "slide_01.md").read_text() == "# one"
    seven = [
        {"path": f"plan/slide_{page:02d}.md", "content": f"# {page}"}
        for page in range(2, 9)
    ]
    assert "status:PASS" in tools.write_plan_batch(agent, seven)
    assert (tmp_path / "plan" / "slide_08.md").read_text() == "# 8"
    inferred = [{"content": "# slide_09\n- role: inferred path"}]
    assert "status:PASS" in tools.write_plan_batch(agent, inferred)
    assert (tmp_path / "plan" / "slide_09.md").exists()
    malformed_transport = (
        '[{"content":"# slide_10\\n- title: "Quoted title"",'
        '"path":"plan/slide_10.md"}]'
    )
    assert "status:PASS" in tools.write_plan_batch(agent, malformed_transport)
    assert "Quoted title" in (tmp_path / "plan" / "slide_10.md").read_text()
    collapsed_content_only = (
        '[{"content":'
        + json.dumps(
            '# slide_11\n- role: first "quoted" page with enough contract detail',
            ensure_ascii=False,
        )
        + ', "content":'
        + json.dumps(
            "# slide_12\n- role: second page with enough contract detail",
            ensure_ascii=False,
        )
        + "]"
    )
    assert "status:PASS" in tools.write_plan_batch(agent, collapsed_content_only)
    assert "first \"quoted\" page" in (
        tmp_path / "plan" / "slide_11.md"
    ).read_text()
    assert (tmp_path / "plan" / "slide_12.md").exists()
    long_compatibility_batch = [
        {"path": f"plan/slide_{page:02d}.md", "content": f"# slide_{page:02d}\n"}
        for page in range(13, 26)
    ]
    long_result = tools.write_plan_batch(agent, long_compatibility_batch)
    assert "status:PASS" in long_result
    assert (tmp_path / "plan" / "slide_25.md").is_file()
    valid_duplicate_keys = (
        '[{"content":'
        + json.dumps(
            "# slide_11\n- role: updated first with enough contract detail",
            ensure_ascii=False,
        )
        + ', "content":'
        + json.dumps(
            "# slide_12\n- role: updated second with enough contract detail",
            ensure_ascii=False,
        )
        + "}]"
    )
    assert "status:PASS" in tools.write_plan_batch(
        agent,
        valid_duplicate_keys,
        replace_existing=True,
    )
    assert "updated first" in (tmp_path / "plan" / "slide_11.md").read_text()
    duplicate = [
        {"path": "plan/slide_01.md", "content": "# one"},
        {"path": "plan/slide_01.md", "content": "# duplicate"},
    ]
    assert "[duplicate]" in tools.write_plan_batch(agent, duplicate)
    reversed_pages = [
        {"path": "plan/slide_02.md", "content": "# two"},
        {"path": "plan/slide_01.md", "content": "# one"},
    ]
    assert "[out_of_order]" in tools.write_plan_batch(agent, reversed_pages)


def test_pending_read_only_accepts_exact_continuation_and_clears_at_eof(
    tmp_path: Path,
) -> None:
    path = tmp_path / "research" / "material.md"
    path.parent.mkdir()
    path.write_text(
        "\n".join(f"line {index:04d} " + ("x" * 90) for index in range(300)),
        encoding="utf-8",
    )

    class Research:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Research()
    first = tools.read_file(agent, "research/material.md")
    assert "续读 offset=" in first
    required = tools.pending_read_requirement(agent)
    assert required is not None
    wrong = tools.read_file(agent, "research/material.md", offset=required[1] + 1)
    assert "pending_read_only" in wrong
    while (required := tools.pending_read_requirement(agent)) is not None:
        result = tools.read_file(agent, required[0], offset=required[1])
    assert "pending_read 已解除" in result
    assert tools.pending_read_error(agent) is None


def test_pending_read_advances_when_model_requests_a_short_chunk(tmp_path: Path) -> None:
    path = tmp_path / "base.css"
    path.write_text("\n".join(f"line {index:04d}" for index in range(810)), encoding="utf-8")

    class Review:
        role = "review"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Review()
    agent._pending_read_continuations = {"base.css": 710}
    chunk = tools.read_file(agent, "base.css", offset=710, limit=100)
    assert "pending_read 已解除" in chunk
    assert tools.pending_read_requirement(agent) is None


def test_pending_read_does_not_accept_one_line_model_ping_pong(tmp_path: Path) -> None:
    path = tmp_path / "plan" / "deck.md"
    path.parent.mkdir()
    path.write_text(
        "\n".join(f"deck line {index:02d}" for index in range(80)),
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        _pending_read_continuations={"plan/deck.md": 38},
        read_path=lambda value: str(tmp_path / value),
    )
    result = tools.read_file(agent, "plan/deck.md", offset=38, limit=1)
    assert "38|deck line 37" in result
    assert "80|deck line 79" in result
    assert "pending_read 已解除" in result
    assert tools.pending_read_requirement(agent) is None


def test_slide_closes_read_surface_after_repeating_completed_immutable_input(
    tmp_path: Path,
) -> None:
    for relative, text in {
        "plan/deck.md": "deck\n",
        "plan/slide_02.md": "plan\n",
        "slides/slide_02.html": "<section></section>\n",
        "assets/catalog.md": "# catalog\n",
    }.items():
        target = tmp_path / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")
    agent = SimpleNamespace(
        ws=str(tmp_path),
        role="slide",
        skill_name="mural-presenter-v0-4",
        assigned_slide_pages=(2,),
        _pending_read_continuations={},
        read_path=lambda value: str(tmp_path / value),
    )
    for relative in (
        "plan/deck.md",
        "plan/slide_02.md",
        "slides/slide_02.html",
        "assets/catalog.md",
    ):
        assert "EOF" in tools.read_file(agent, relative)
    result = tools.read_file(agent, "assets/catalog.md")
    assert "无需重复" in result
    assert agent._slide_inputs_read_closed is True


def test_orchestrator_cannot_read_material_when_research_is_enabled(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "research" / "material.md").write_text("raw", encoding="utf-8")

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        research_required = True
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    result = tools.read_file(Orchestrator(), "research/material.md")
    assert "research_route_violation" in result
    assert "knowledge-brief.md" in result


def test_orchestrator_does_not_page_slide_layout_library_or_base_css(tmp_path: Path) -> None:
    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Orchestrator()
    assert "Slide 的实现词汇库" in tools.read_file(
        agent, "skills/mural-presenter-v0-4/references/layout-patterns.md"
    )
    assert "不要分页读取 base.css" in tools.read_file(agent, "base.css")
    assert tools.pending_read_requirement(agent) is None


def test_final_single_missing_plan_can_resume_with_write_file(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- page_count: 3\n",
        encoding="utf-8",
    )
    for page in (1, 2):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# slide {page}\n", encoding="utf-8"
        )

    class Orchestrator:
        role = "orchestrator"
        label = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.write_file(
        Orchestrator(), "plan/slide_03.md", "# recovered slide 3\n"
    )
    assert result.startswith("已写入")
    assert (plan / "slide_03.md").is_file()


def test_slide_repair_required_is_a_structured_nonfatal_handoff(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace/slide-render-states"
    trace.mkdir(parents=True)
    (trace / "page_07.json").write_text(
        json.dumps({
            "limit": 8,
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "hashes": ["a", "b", "c"],
            "attempt_hashes": {"slide_07": ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )

    class Slide:
        role = "slide"
        ws = str(tmp_path)
        assigned_slide_pages = (7,)
        repair_required_reason = ""
        final_text = """status: repair_required
pages: 07
issue_type: shared_system
evidence: inherited two-column grid squeezes the primary content
proposed_fix: override the page grid to one column
final_pixels_inspected: yes
"""

    issue = _slide_repair_issue(Slide())
    assert issue is not None
    assert issue["status"] == "repair_required"
    assert issue["pages"] == [7]
    assert issue["issue_type"] == "shared_system"
    assert issue["render_budget_exhausted_pages"] == [7]


def test_review_must_open_flagged_page_current_pixels(tmp_path: Path) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    png = renders / "slide_07.png"
    png.write_bytes(b"current-pixels")

    class Review:
        role = "review"
        ws = str(tmp_path)
        review_contact_sheet_inspected = True
        required_review_pages = (7,)
        review_viewed_page_hashes = {}

    assert "P07" in _review_required_view_gap(Review())
    Review.review_viewed_page_hashes = {
        7: __import__("hashlib").sha256(png.read_bytes()).hexdigest()
    }
    assert _review_required_view_gap(Review()) == ""


def test_review_gets_compact_render_summary_without_pending_pagination(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(
        json.dumps(
            {
                "mode": "incremental-finalize",
                "n_pages": 1,
                "console_errors": [],
                "visual_review_notes": ["all bitmap pages use a heavy mask"],
                "special_page_geometry": {"warnings": []},
                "pages": [
                    {
                        "page": 1,
                        "png": str(renders / "slide_01.png"),
                        "blank": False,
                        "static": False,
                        "geometry": {
                            "page_type": "content",
                            "text_boxes": [
                                {
                                    "text": "too small",
                                    "font_size": 14,
                                    "overflow_x": False,
                                    "overflow_y": False,
                                }
                            ],
                            "images": [],
                        },
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    class Review:
        role = "review"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Review()
    result = tools.read_file(agent, "renders/render.json")
    assert "Review compact render summary" in result
    assert "all bitmap pages use a heavy mask" in result
    assert "too small" in result
    assert len(result) < 7800
    assert tools.pending_read_requirement(agent) is None
    agent._pending_read_continuations = {"renders/render.json": 200}
    offset_result = tools.read_file(agent, "renders/render.json", offset=500, limit=400)
    assert "Review compact render summary" in offset_result
    assert "所有 offset" in offset_result
    assert tools.pending_read_requirement(agent) is None


def test_orchestrator_gets_compact_render_summary_without_pending_pagination(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(
        json.dumps({"n_pages": 1, "pages": [{"page": 1, "geometry": {}}]}),
        encoding="utf-8",
    )

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def read_path(self, name):
            return str(tmp_path / name)

    agent = Orchestrator()
    result = tools.read_file(agent, "renders/render.json")
    assert "Orchestrator compact render summary" in result
    assert "不重新诊断或重派已完成 Slide" in result
    assert tools.pending_read_requirement(agent) is None
    agent._pending_read_continuations = {"renders/render.json": 200}
    offset_result = tools.read_file(agent, "renders/render.json", offset=500, limit=400)
    assert "Orchestrator compact render summary" in offset_result
    assert "所有 offset" in offset_result
    assert tools.pending_read_requirement(agent) is None


def test_review_terminal_cannot_page_raw_render_manifest(tmp_path: Path) -> None:
    (tmp_path / "renders").mkdir()
    (tmp_path / "renders/render.json").write_text(
        json.dumps({"n_pages": 1, "pages": []}), encoding="utf-8"
    )
    agent = SimpleNamespace(
        role="review",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
    )
    result = tools.terminal(agent, 'grep -n "slide_01" renders/render.json')
    assert "review_compact_render_summary" in result
    assert '"n_pages":1' in result


@pytest.mark.parametrize(
    ("role", "path", "expected"),
    [
        ("material", "inputs/paper.pdf.md", True),
        ("material", "inputs/notes.txt", True),
        ("material", "research/material.md", True),
        ("research", "research/material.md", True),
        ("research", "research/knowledge-brief.md", True),
        ("orchestrator", "research/material.md", True),
        ("orchestrator", "research/knowledge-brief.md", True),
        ("orchestrator", "plan/deck.md", True),
        ("orchestrator", "plan/slide_15.md", True),
        ("image", "research/knowledge-brief.md", True),
        ("image", "plan/deck.md", True),
        ("image", "assets/catalog.md", True),
        ("slide", "plan/deck.md", True),
        ("slide", "plan/slide_01.md", True),
        ("slide", "slides/slide_01.html", True),
        ("review", "plan/deck.md", True),
        ("review", "plan/slide_07.md", True),
        ("review", "slides/slide_07.html", True),
        ("review", "slides/not-a-slide.html", False),
        ("review", "renders/render.json", False),
    ],
)
def test_canonical_evidence_pending_reads_auto_continue_without_model_turn(
    role: str,
    path: str,
    expected: bool,
) -> None:
    agent = SimpleNamespace(role=role, skill_name="mural-presenter-v0-4")
    assert _auto_continue_read_allowed(agent, path) is expected


def test_bitmap_and_high_confidence_findings_select_full_resolution_review(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_01.md").write_text(
        "# slide_01\n- needs_bitmap: true\n- composition: visual-split\n",
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_02.md").write_text(
        "# slide_02\n- needs_bitmap: false\n- composition: matrix\n",
        encoding="utf-8",
    )
    (tmp_path / "renders/render.json").write_text(
        json.dumps({
            "pages": [],
            "special_page_geometry": {"warnings": []},
            "layout_defects": {"03": [{"type": "text_collision"}]},
        }),
        encoding="utf-8",
    )
    parent = SimpleNamespace(ws=str(tmp_path))
    assert _mandatory_fullres_review_pages(parent) == {1, 3}


def test_metadata_heavy_cover_is_not_forced_to_fullres_without_warning(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_01.md").write_text(
        "# slide_01\n- page_type: cover\n- needs_bitmap: false\n"
        "- composition: editorial\n",
        encoding="utf-8",
    )
    (tmp_path / "renders/render.json").write_text(
        json.dumps({
            "pages": [{
                "page": 1,
                "geometry": {
                    "text_boxes": [{"text": "metadata" * 100} for _ in range(20)]
                },
            }],
            "special_page_geometry": {"warnings": []},
            "layout_defects": {},
        }),
        encoding="utf-8",
    )
    parent = SimpleNamespace(ws=str(tmp_path))
    assert _mandatory_fullres_review_pages(parent) == set()


def test_generic_geometry_warnings_and_dense_composition_are_not_mandatory(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/slide_02.md").write_text(
        "- page_type: content\n- needs_bitmap: false\n- composition: matrix\n",
        encoding="utf-8",
    )
    (tmp_path / "renders/render.json").write_text(json.dumps({
        "pages": [{"page": 2, "geometry": {
            "text_boxes": [{"text": "dense" * 200} for _ in range(20)]
        }}],
        "special_page_geometry": {"warnings": [
            "page 2: 20 body leaves below 24px",
            "page 2: possible DOM text overflow",
        ]},
        "layout_defects": {},
        "media_mismatches": {},
        "placeholder_flags": {},
    }), encoding="utf-8")
    assert _mandatory_fullres_review_pages(
        SimpleNamespace(ws=str(tmp_path))
    ) == set()


def test_layout_and_placeholder_findings_remain_mandatory(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "renders/render.json").write_text(json.dumps({
        "pages": [],
        "special_page_geometry": {"warnings": []},
        "layout_defects": {"03": [{"type": "block_collision"}]},
        "media_mismatches": {},
        "placeholder_flags": {"06": ["TODO"]},
    }), encoding="utf-8")
    assert _mandatory_fullres_review_pages(
        SimpleNamespace(ws=str(tmp_path))
    ) == {3, 6}


def test_review_cannot_override_current_vision_critic_repair_required(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    png = renders / "slide_07.png"
    png.write_bytes(b"defective-current-pixels")
    digest = __import__("hashlib").sha256(png.read_bytes()).hexdigest()

    class Review:
        role = "review"
        ws = str(tmp_path)
        review_contact_sheet_inspected = True
        required_review_pages = ()
        review_viewed_page_hashes = {}
        vision_critic_results = {
            "renders/slide_07.png": {
                "source_sha256": digest,
                "verdict": "repair_required",
                "summary": "isolated small image remains",
                "issues": [],
            }
        }

    assert "Vision Critic" in _review_required_view_gap(Review())
    png.write_bytes(b"new-pixels-not-yet-criticized")
    assert _review_required_view_gap(Review()) == ""


def test_contract_parser_ignores_free_prose() -> None:
    fields = _contract_fields(
        "The page may be ready later.\nstatus: repair_required\npages: 07\n"
    )
    assert fields == {"status": "repair_required", "pages": "07"}


def test_page_vision_budget_is_three_states_across_redispatch(tmp_path: Path) -> None:
    from PIL import Image

    renders = tmp_path / "renders"
    renders.mkdir()
    image_path = renders / "slide_07.png"

    def slide_agent(label: str):
        return SimpleNamespace(
            role="slide",
            skill_name="mural-presenter-v0-4",
            ownership_topology="single",
            assigned_slide_pages=(7,),
            slide_group_id="",
            trace_label=label,
            n_views=0,
            max_vision_edge=1600,
            model="main-multimodal-model",
            cfg={"model_base_url": "http://main-model"},
            deadline_monotonic=time.monotonic() + 60,
            ws=str(tmp_path),
            read_path=lambda path: str(tmp_path / path),
            log=lambda *_args, **_kwargs: None,
            rendered_output_hashes={7: label},
            viewed_output_hashes={},
            repair_required_reason="",
        )

    critic = json.dumps({
        "schema": "mural.vision-critic.v1",
        "verdict": "ready",
        "summary": "page is ready",
        "observations": [],
        "issues": [],
        "scan": {"visible_subjects": [], "text_regions": [], "regions": [],
                 "edges": {"top": "clear", "right": "clear", "bottom": "clear", "left": "clear"}},
    })
    with mock.patch(
        "core.model_call.call_vision_same_model_auxiliary",
        return_value=critic,
    ) as auxiliary:
        for index, color in enumerate(("red", "green", "blue"), start=1):
            Image.new("RGB", (160, 90), color).save(image_path)
            result = tools.vision_analyze(
                slide_agent("slide_07" if index == 1 else f"slide_07_r{index}"),
                "renders/slide_07.png",
                "inspect current pixels",
            )
            assert result["vision_verdict"] == "ready"
        Image.new("RGB", (160, 90), "yellow").save(image_path)
        final_agent = slide_agent("slide_07_r4")
        blocked = tools.vision_analyze(
            final_agent,
            "renders/slide_07.png",
            "inspect current pixels",
        )

    assert auxiliary.call_count == 3
    assert "repair_required" in blocked
    assert "重新委派不会重置" in blocked
    assert final_agent.repair_required_reason


def test_three_page_group_allows_nine_page_and_three_group_checks(
    tmp_path: Path,
) -> None:
    from PIL import Image

    renders = tmp_path / "renders"
    renders.mkdir()
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ownership_topology="grouped",
        assigned_slide_pages=(1, 2, 3),
        slide_group_id="intro",
        trace_label="slide_group_intro",
        n_views=0,
        max_vision_edge=1600,
        model="main-multimodal-model",
        cfg={"model_base_url": "http://main-model"},
        deadline_monotonic=time.monotonic() + 60,
        ws=str(tmp_path),
        read_path=lambda path: str(tmp_path / path),
        log=lambda *_args, **_kwargs: None,
        rendered_output_hashes={},
        viewed_output_hashes={},
        repair_required_reason="",
    )
    critic = json.dumps({
        "schema": "mural.vision-critic.v1",
        "verdict": "ready",
        "summary": "ready",
        "observations": [],
        "issues": [],
        "scan": {"visible_subjects": [], "text_regions": [], "regions": [],
                 "edges": {"top": "clear", "right": "clear", "bottom": "clear", "left": "clear"}},
    })
    with mock.patch(
        "core.model_call.call_vision_same_model_auxiliary",
        return_value=critic,
    ) as auxiliary:
        for revision in range(1, 4):
            for page in (1, 2, 3):
                path = renders / f"slide_{page:02d}.png"
                Image.new("RGB", (160, 90), (revision * 40, page * 50, 80)).save(path)
                agent.rendered_output_hashes[page] = f"page-{page}-state-{revision}"
                result = tools.vision_analyze(
                    agent,
                    f"renders/slide_{page:02d}.png",
                    "inspect page",
                )
                assert result["vision_verdict"] == "ready"
                agent.viewed_output_hashes[page] = agent.rendered_output_hashes[page]
            sheet = renders / "contact-sheet-group-intro.png"
            Image.new("RGB", (300, 180), (revision * 50, 90, 120)).save(sheet)
            result = tools.vision_analyze(
                agent,
                "renders/contact-sheet-group-intro.png",
                "inspect group consistency",
            )
            assert result["vision_verdict"] == "ready"

        cached = tools.vision_analyze(
            agent,
            "renders/contact-sheet-group-intro.png",
            "inspect group consistency again",
        )
        assert cached["vision_cached"] is True
        Image.new("RGB", (300, 180), "black").save(
            renders / "contact-sheet-group-intro.png"
        )
        blocked = tools.vision_analyze(
            agent,
            "renders/contact-sheet-group-intro.png",
            "inspect fourth group state",
        )

    assert auxiliary.call_count == 12
    assert "repair_required" in blocked
    assert config.SLIDE_PAGE_MAX_INSPECTIONS == 3
    assert config.SLIDE_GROUP_MAX_INSPECTIONS == 3


def test_group_delivery_requires_current_pages_then_current_group_sheet(
    tmp_path: Path,
) -> None:
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan").mkdir()
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    (tmp_path / "base.css").write_text(".slide{color:black}", encoding="utf-8")
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    rendered_hashes = {}
    paths = {}
    for page in (1, 2, 3):
        html = tmp_path / "slides" / f"slide_{page:02d}.html"
        html.write_text(
            f'<section class="slide"><main>page {page}</main></section>',
            encoding="utf-8",
        )
        png = tmp_path / "renders" / f"slide_{page:02d}.png"
        png.write_bytes(f"pixels-{page}".encode())
        render_digest = _render_input_digest(
            tmp_path,
            (tmp_path / "base.css").read_text(encoding="utf-8"),
            html.read_text(encoding="utf-8"),
        )
        file_digest = hashlib.sha256(html.read_bytes()).hexdigest()
        rendered_hashes[page] = file_digest
        paths[page] = str(html)
        (tmp_path / "_trace/slide-render-states" / f"page_{page:02d}.json").write_text(
            json.dumps({
                "hashes": [render_digest],
                "authoring_hashes": [render_digest],
                "attempt_hashes": {"slide_group_intro": [render_digest]},
                "published_hash": render_digest,
                "authoring_attempt_limit": 3,
                "quality_action_count": 0,
            }),
            encoding="utf-8",
        )
    sheet = tmp_path / "renders/contact-sheet-group-intro.png"
    sheet.write_bytes(b"group-pixels")
    sheet_digest = hashlib.sha256(sheet.read_bytes()).hexdigest()
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ownership_topology="grouped",
        ws=str(tmp_path),
        assigned_slide_pages=(1, 2, 3),
        slide_group_id="intro",
        expected_output_paths=paths,
        expected_output_initial_hashes={page: "scaffold" for page in paths},
        expected_output_path="",
        rendered_output_hashes=rendered_hashes,
        viewed_output_hashes=dict(rendered_hashes),
        group_viewed_contact_hash="",
        group_viewed_page_hashes={},
        n_renders=3,
        n_views=3,
    )

    assert "未检查当前组联系表" in _slide_deliverable_gap(agent)
    agent.group_viewed_contact_hash = sheet_digest
    agent.group_viewed_page_hashes = dict(rendered_hashes)
    assert _slide_deliverable_gap(agent) == ""


def test_group_contact_sheet_view_is_not_shadowed_by_page_same_pixel_guard(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    trace_dir = tmp_path / "trace"
    (trace_dir / "images").mkdir(parents=True)
    renders.mkdir()
    sheet = renders / "contact-sheet-group-bookends.png"
    sheet.write_bytes(b"current-group-pixels")
    sheet_hash = hashlib.sha256(sheet.read_bytes()).hexdigest()
    rendered = {1: "page-1-current", 24: "page-24-current"}
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ownership_topology="grouped",
        ws=str(tmp_path),
        trace_dir=trace_dir,
        assigned_slide_pages=(1, 24),
        slide_group_id="bookends",
        rendered_output_hashes=dict(rendered),
        viewed_output_hashes=dict(rendered),
        group_rendered_contact_hash=sheet_hash,
        group_rendered_page_hashes=dict(rendered),
        group_viewed_contact_hash="",
        group_viewed_page_hashes={},
        started=time.time(),
        log=lambda *_args, **_kwargs: None,
        remote_tool_concurrency=1,
        n_views=2,
        n_renders=2,
        image_by_tool={},
        vision_critic_results={},
        repair_required_reason="",
    )
    assert _slide_pending_render_views(agent) == ()
    assert _slide_pending_group_view(agent) == (
        "renders/contact-sheet-group-bookends.png"
    )
    use = SimpleNamespace(
        id="tool-group-sheet",
        name="vision_analyze",
        input={
            "image": "renders/contact-sheet-group-bookends.png",
            "query": "inspect bookend coherence",
        },
    )
    vision_value = {
        "vision_analysis": '{"verdict":"ready","issues":[]}',
        "vision_result": {"verdict": "ready", "issues": []},
        "vision_verdict": "ready",
        "vision_summary": "bookends are coherent",
        "image_b64": __import__("base64").b64encode(b"sheet-pixels").decode(),
        "media_type": "image/png",
        "path": "renders/contact-sheet-group-bookends.png",
        "source_sha256": sheet_hash,
        "vision_backend": "same_model_aux:test",
    }
    with mock.patch("core.agent_loop.tools.dispatch", return_value=vision_value):
        result = _tool_results(agent, [use], 1, [])

    assert "current_pixels_already_inspected" not in str(result)
    assert agent.group_viewed_contact_hash == sheet_hash
    assert agent.group_viewed_page_hashes == rendered
    assert _slide_pending_group_view(agent) == ""


def test_group_contact_ready_cannot_override_open_fullres_page_issue(
    tmp_path: Path,
) -> None:
    """Regression for deck #154: page-local pixels outrank the montage."""
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "_trace/slide-render-states").mkdir(parents=True)
    for page in (1, 3):
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            f"<section>P{page}</section>", encoding="utf-8"
        )
    p1_html = hashlib.sha256(
        (tmp_path / "slides/slide_01.html").read_bytes()
    ).hexdigest()
    p1_png = tmp_path / "renders/slide_01.png"
    p1_png.write_bytes(b"p1-overlap-pixels")
    p1_png_hash = hashlib.sha256(p1_png.read_bytes()).hexdigest()
    (tmp_path / "_trace/slide-render-states/page_01.json").write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "authoring_hashes": [p1_html],
            "quality_action_count": 0,
        }),
        encoding="utf-8",
    )
    group_source = "renders/contact-sheet-group-bookends.png"
    sheet = tmp_path / group_source
    sheet.write_bytes(b"small-clean-looking-montage")
    sheet_hash = hashlib.sha256(sheet.read_bytes()).hexdigest()
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        trace_label="slide_group_bookends",
        assigned_slide_pages=(1, 3),
        slide_group_id="bookends",
        bash_timeout=30,
        expected_output_paths={
            1: str(tmp_path / "slides/slide_01.html"),
            3: str(tmp_path / "slides/slide_03.html"),
        },
        rendered_output_hashes={1: p1_html},
        viewed_output_hashes={1: p1_html},
        group_rendered_contact_hash=sheet_hash,
        group_rendered_page_hashes={1: p1_html, 3: ""},
        group_viewed_contact_hash=sheet_hash,
        group_viewed_page_hashes={1: p1_html, 3: ""},
        vision_critic_results={
            "renders/slide_01.png": {
                "source_sha256": p1_png_hash,
                "verdict": "repair_required",
                "summary": "标题与代码块重叠",
                "issues": [{"type": "text_overlap", "severity": "major"}],
            },
            group_source: {
                "source_sha256": sheet_hash,
                "verdict": "ready",
                "summary": "组级节奏整齐，缩略图未见异常",
                "issues": [],
            },
        },
        repair_required_reason="Vision Critic unresolved: 标题与代码块重叠",
        active_slide_page=1,
        role_script="skills/mural-presenter-v0-4/scripts/slide.py",
        bash_relaxed=False,
        safe=lambda path: str(tmp_path / path),
    )

    assert _slide_group_active_page(agent) == 1
    assert _slide_pending_group_view(agent) == ""
    assert _unresolved_vision_critic_issues(agent)[0][0] == "renders/slide_01.png"
    # After #154: no hard group_page_order rejection — write and render
    # succeed; Skill guides sequencing, Harness only provides soft feedback.
    result = tools.write_file(
        agent, "slides/slide_03.html", "<section>premature P3</section>"
    )
    assert "group_page_order" not in result
    render_result = tools.terminal(
        agent,
        "python skills/mural-presenter-v0-4/scripts/slide.py render-group . "
        "--group bookends --pages 01,03",
    )
    assert "group_page_order" not in render_result


def test_group_contact_sheet_view_is_not_blocked_by_authoring_stop_line(
    tmp_path: Path,
) -> None:
    """When all pages are at render budget but a fresh group sheet needs
    inspection, vision_analyze must succeed — the stop-line guard must not
    shadow the group-view responsibility (mirrors L4716 same-pixel exemption).
    """
    renders = tmp_path / "renders"
    trace_dir = tmp_path / "trace"
    (trace_dir / "images").mkdir(parents=True)
    renders.mkdir()
    # Set up render budget = exhausted (3/3) for both pages
    states_dir = tmp_path / "_trace" / "slide-render-states"
    states_dir.mkdir(parents=True)
    for page in (1, 24):
        (states_dir / f"page_{page:02d}.json").write_text(
            json.dumps({
                "authoring_attempt_limit": 3,
                "authoring_hashes": ["h1", "h2", "h3"],
            }),
            encoding="utf-8",
        )
    # Fresh group sheet needing inspection
    sheet = renders / "contact-sheet-group-bookends.png"
    sheet.write_bytes(b"fresh-group-pixels-after-budget")
    sheet_hash = hashlib.sha256(sheet.read_bytes()).hexdigest()
    rendered = {1: "page-1-final", 24: "page-24-final"}
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ownership_topology="grouped",
        ws=str(tmp_path),
        trace_dir=trace_dir,
        trace_label="slide_bookends",
        assigned_slide_pages=(1, 24),
        slide_group_id="bookends",
        rendered_output_hashes=dict(rendered),
        viewed_output_hashes=dict(rendered),
        group_rendered_contact_hash=sheet_hash,
        group_rendered_page_hashes=dict(rendered),
        group_viewed_contact_hash="",
        group_viewed_page_hashes={},
        started=time.time(),
        log=lambda *_args, **_kwargs: None,
        remote_tool_concurrency=1,
        n_views=6,
        n_renders=6,
        image_by_tool={},
        vision_critic_results={},
        repair_required_reason="",
    )
    # Preconditions: budget exhausted, no pending page views, group view pending
    assert _slide_authoring_stop_line(agent) is True
    assert _slide_pending_render_views(agent) == ()
    assert _slide_pending_group_view(agent) == (
        "renders/contact-sheet-group-bookends.png"
    )
    use = SimpleNamespace(
        id="tool-group-budget",
        name="vision_analyze",
        input={
            "image": "renders/contact-sheet-group-bookends.png",
            "query": "inspect bookend coherence after final renders",
        },
    )
    vision_value = {
        "vision_analysis": '{"verdict":"ready","issues":[]}',
        "vision_result": {"verdict": "ready", "issues": []},
        "vision_verdict": "ready",
        "vision_summary": "bookends coherent",
        "image_b64": __import__("base64").b64encode(b"sheet-pixels").decode(),
        "media_type": "image/png",
        "path": "renders/contact-sheet-group-bookends.png",
        "source_sha256": sheet_hash,
        "vision_backend": "same_model_aux:test",
    }
    with mock.patch("core.agent_loop.tools.dispatch", return_value=vision_value):
        result = _tool_results(agent, [use], 1, [])

    assert "slide_authoring_budget" not in str(result)
    assert "current_pixels_already_inspected" not in str(result)
    assert agent.group_viewed_contact_hash == sheet_hash
    assert agent.group_viewed_page_hashes == rendered
    assert _slide_pending_group_view(agent) == ""


def test_missing_group_sheet_requires_render_group_before_global_stop_line(
    tmp_path: Path,
) -> None:
    """All terminal members still retain the final group consistency pass."""
    slides = tmp_path / "slides"
    renders = tmp_path / "renders"
    states = tmp_path / "_trace" / "slide-render-states"
    slides.mkdir()
    renders.mkdir()
    states.mkdir(parents=True)
    rendered: dict[int, str] = {}
    viewed: dict[int, str] = {}
    critic: dict[str, dict] = {}
    for page in (1, 10):
        html = slides / f"slide_{page:02d}.html"
        html.write_text(f"<section>finished P{page}</section>", encoding="utf-8")
        digest = hashlib.sha256(html.read_bytes()).hexdigest()
        rendered[page] = digest
        viewed[page] = digest
        png = renders / f"slide_{page:02d}.png"
        png.write_bytes(f"pixels-{page}".encode())
        critic[f"renders/slide_{page:02d}.png"] = {
            "source_sha256": hashlib.sha256(png.read_bytes()).hexdigest(),
            "verdict": "repair_required",
            "issues": [{"type": "bounded_issue", "severity": "major"}],
        }
        (states / f"page_{page:02d}.json").write_text(
            json.dumps({
                "authoring_attempt_limit": 3,
                "authoring_hashes": ["a", "b", "c"],
            }),
            encoding="utf-8",
        )
    agent = SimpleNamespace(
        role="slide",
        ws=str(tmp_path),
        trace_label="slide_group_bookends",
        assigned_slide_pages=(1, 10),
        slide_group_id="bookends",
        expected_output_paths={
            1: str(slides / "slide_01.html"),
            10: str(slides / "slide_10.html"),
        },
        rendered_output_hashes=rendered,
        viewed_output_hashes=viewed,
        vision_critic_results=critic,
    )

    assert _slide_group_active_page(agent) == 0
    assert _slide_authoring_stop_line(agent) is True
    assert _slide_pending_group_view(agent) == "__render_group_required__"


def test_group_sheet_provenance_blocks_old_cache_after_member_changes(
    tmp_path: Path,
) -> None:
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text(
        "# Deck\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    sheet = tmp_path / "renders/contact-sheet-group-intro.png"
    sheet.write_bytes(b"old group sheet")
    sheet_hash = hashlib.sha256(sheet.read_bytes()).hexdigest()
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        assigned_slide_pages=(1, 2),
        slide_group_id="intro",
        rendered_output_hashes={1: "page-1-new", 2: "page-2-same"},
        viewed_output_hashes={1: "page-1-new", 2: "page-2-same"},
        group_rendered_contact_hash=sheet_hash,
        group_rendered_page_hashes={1: "page-1-old", 2: "page-2-same"},
        group_viewed_contact_hash=sheet_hash,
        group_viewed_page_hashes={1: "page-1-old", 2: "page-2-same"},
    )
    assert _slide_pending_group_view(agent) == "__render_group_required__"
    blocked = tools.vision_analyze(
        agent,
        "renders/contact-sheet-group-intro.png",
        "inspect group consistency",
    )
    assert "不是由当前全部页面像素生成" in blocked


def test_review_does_not_block_on_source_asset_issue_hidden_by_final_crop(
    tmp_path: Path,
) -> None:
    (tmp_path / "_trace").mkdir()
    (tmp_path / "_trace/vision-issues.json").write_text(
        json.dumps({
            "schema": "mural.vision-issues.v1",
            "issues": [{
                "id": "VIS-ASSET",
                "status": "open",
                "page": None,
                "source": "assets/source-with-watermark.jpg",
                "type": "watermark_present",
            }],
        }),
        encoding="utf-8",
    )
    agent = SimpleNamespace(
        role="review",
        ws=str(tmp_path),
        required_review_pages=(),
        review_contact_sheet_inspected=True,
        review_viewed_page_hashes={},
        vision_critic_results={},
    )
    assert _review_required_view_gap(agent) == ""


def test_bounded_inspected_review_is_not_a_whole_deck_failure() -> None:
    class Parent:
        finalize_succeeded = True

    result = {
        "ok": False,
        "status": "needs_orchestrator",
        "blocking": "no",
        "final_pixels_inspected": True,
        "final_view_after_review": True,
        "finalize_succeeded": True,
        "review_changed": False,
        "attempt": 3,
    }
    assert _review_can_complete_needs_improvement(Parent(), result)
    result["blocking"] = "yes"
    assert not _review_can_complete_needs_improvement(Parent(), result)
    result["blocking"] = "no"
    result["final_pixels_inspected"] = False
    assert not _review_can_complete_needs_improvement(Parent(), result)
    result["final_pixels_inspected"] = True
    result["finalize_succeeded"] = False
    assert not _review_can_complete_needs_improvement(Parent(), result)


def test_bounded_review_recovery_requires_current_mandatory_pixels(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "contact-sheet.png").write_bytes(b"sheet")
    page = renders / "slide_02.png"
    page.write_bytes(b"current pixels")
    digest = hashlib.sha256(page.read_bytes()).hexdigest()
    agent = SimpleNamespace(
        role="review",
        ws=str(tmp_path),
        finalize_succeeded=True,
        final_view_after_review=True,
        review_contact_sheet_inspected=True,
        required_review_pages=(2,),
        review_viewed_page_hashes={2: digest},
    )

    assert _review_has_current_inspected_delivery(agent)
    page.write_bytes(b"changed after view")
    assert not _review_has_current_inspected_delivery(agent)


def test_current_manifest_quality_findings_forbid_plain_ready(
    tmp_path: Path,
) -> None:
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(
        json.dumps(
            {
                "layout_defects": {"06": [{"type": "block_overflow"}]},
                "typography_flags": {},
                "media_mismatches": {"01": {"reason": "low_texture"}},
                "placeholder_flags": {"09": ["XXX"]},
            }
        ),
        encoding="utf-8",
    )
    agent = SimpleNamespace(ws=str(tmp_path))

    assert _review_manifest_quality_findings(agent) == [
        "layout_defects:06",
        "media_mismatches:01",
        "placeholder_flags:09",
    ]


def test_review_attempt_count_survives_orchestrator_resume(tmp_path: Path) -> None:
    root = tmp_path / "_trace/subagents"
    for label in ("review", "review_retry2", "review_verify3"):
        (root / label).mkdir(parents=True)
    parent = SimpleNamespace(ws=str(tmp_path))
    assert _persisted_trace_attempts(parent, "review") == 3
    assert config.REVIEW_MAX_ATTEMPTS == 3


def test_review_cannot_repeat_until_delivery_surface_changes(tmp_path: Path) -> None:
    for relative in ("plan", "slides", "renders", "assets"):
        (tmp_path / relative).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section>before</section>\n", encoding="utf-8"
    )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet")
    parent = Agent(
        "review-fingerprint",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = True
    parent.child_outcomes = {"review": {
        "ok": False,
        "status": "needs_orchestrator",
        "exit_reason": "needs_orchestrator",
        "input_fingerprint": _review_delivery_fingerprint(str(tmp_path)),
    }}

    unchanged = _delegate(parent, {"tasks": [{"role": "review"}]})
    assert "均未发生变化" in unchanged
    assert "禁止在相同输入上重复验收" in unchanged

    # A render-capture repair may change only final pixels, not HTML/assets.
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet-after-repair")
    captured: list[dict] = []

    def fake_child(_parent, _index, spec):
        captured.append(dict(spec))
        return {
            "label": "review",
            "trace_label": "review_verify2",
            "role": "review",
            "ok": True,
            "status": "ready",
            "exit_reason": "text_response",
            "completed_pages": [],
            "incomplete_pages": [],
            "renders": 0,
            "views": 1,
            "final_render_after_review": True,
            "final_view_after_review": True,
            "finalize_attempted": True,
            "finalize_succeeded": True,
            "summary": "verified",
        }

    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        verified = json.loads(_delegate(parent, {"tasks": [{"role": "review"}]}))
    assert verified["status"] == "completed"
    assert captured[0]["_recovery_kind"] == "verification"


def test_page_local_finalize_blocker_routes_to_preflight_review(
    tmp_path: Path,
) -> None:
    for relative in ("plan", "slides", "renders", "assets"):
        (tmp_path / relative).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    for page in (1, 8):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- needs_bitmap: false\n", encoding="utf-8"
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            f"<section data-slide='{page:02d}'>page</section>\n", encoding="utf-8"
        )
        (tmp_path / f"renders/slide_{page:02d}.png").write_bytes(b"pixels")

    parent = Agent(
        "preflight-review",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = False
    parent.finalize_failure = (
        "ValueError: slide_08.html must not move shared "
        "`#slide-08 .page-footer`; blocked properties: z-index"
    )
    assert _review_preflight_repair_pages(parent) == (8,)

    captured: list[dict] = []

    def fake_child(_parent, _index, spec):
        captured.append(dict(spec))
        return {
            "label": "review",
            "trace_label": "review",
            "role": "review",
            "ok": True,
            "status": "ready",
            "exit_reason": "text_response",
            "completed_pages": [],
            "incomplete_pages": [],
            "renders": 2,
            "views": 2,
            "final_render_after_review": True,
            "final_view_after_review": True,
            "finalize_attempted": True,
            "finalize_succeeded": True,
            "finalize_failure": "",
            "summary": "repaired and verified",
        }

    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        result = json.loads(_delegate(parent, {"tasks": [{"role": "review"}]}))

    assert result["status"] == "completed"
    assert captured[0]["_review_preflight_repair_pages"] == [8]
    assert captured[0]["required_review_pages"] == [8]
    assert "mode=preflight_repair" in captured[0]["task"]
    assert "不要等待不存在的联系表" in captured[0]["task"]
    assert "不要读取 `_trace/**`" in captured[0]["task"]
    assert "<harness_review_issues>" in captured[0]["task"]
    assert "先读取 `_trace/review-issues.json`" not in captured[0]["task"]
    assert parent.finalize_succeeded is True


def test_review_round_must_start_from_current_contact_sheet(tmp_path: Path) -> None:
    review = SimpleNamespace(
        role="review",
        review_contact_sheet_inspected=False,
        ws=str(tmp_path),
        read_path=lambda path: str(tmp_path / path),
    )
    result = tools.vision_analyze(
        review,
        "renders/slide_01.png",
        "inspect one page",
    )
    assert "必须先打开当前 renders/contact-sheet.png" in result


def test_finish_gap_nudges_only_incomplete_single_pages(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (plan / "slide_01.md").write_text("# P1\n", encoding="utf-8")
    (plan / "slide_02.md").write_text("# P2\n", encoding="utf-8")

    class Agent:
        role = "orchestrator"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = True
        delegated_roles = {"slide"}
        child_outcomes = {
            "slide_01": {"ok": True},
            "slide_02": {"ok": False},
        }
        ws = str(tmp_path)

    gap = _finish_gap(Agent())
    assert "Slide 02:" in gap
    assert "Slide 01:" not in gap
    assert "不要创建新的 Slide Agent" in gap
    assert "operational retry" in gap


def test_finish_gap_retries_only_incomplete_complete_group(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: grouped\n",
        encoding="utf-8",
    )
    for page, group in ((1, "opening"), (2, "mechanism"), (3, "mechanism")):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n- production_group: {group}\n",
            encoding="utf-8",
        )

    class GroupedAgent:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = True
        revision_mode = False
        delegated_roles = {"slide"}
        child_outcomes = {
            "slide_group_opening": {
                "ok": True,
                "completed_pages": [1],
            },
            "slide_group_mechanism": {
                "ok": False,
                "completed_pages": [],
                "incomplete_pages": [2, 3],
            },
        }
        ws = str(tmp_path)

    gap = _finish_gap(GroupedAgent())
    assert "SlideGroup mechanism [02,03]:" in gap
    assert "SlideGroup opening" not in gap
    assert "不要创建新的 SlideGroup 实例" in gap
    assert "operational retry" in gap


def test_v03_has_one_skill_with_both_ownership_topologies() -> None:
    root = REPO / "skills/mural-presenter-v0.4"
    editions = sorted(
        path.name for path in root.iterdir()
        if path.is_dir() and (path / "SKILL.md").is_file()
    )
    assert editions == ["mural-presenter-v0-4"]
    content = (root / editions[0] / "SKILL.md").read_text(encoding="utf-8")
    assert "ownership_topology: single | grouped" in content
    assert "role=slide, pages=[NN]" in content
    assert "role=slide, group_id=GROUP" in content


def test_slide_stop_line_prevents_blind_edits_after_unresolved_third_view(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace" / "slide-render-states"
    trace.mkdir(parents=True)
    png = tmp_path / "renders" / "slide_01.png"
    png.parent.mkdir()
    png.write_bytes(b"current-pixels")
    attempt = "slide_01"
    (trace / "page_01.json").write_text(
        json.dumps({
            "limit": 8,
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "attempt_hashes": {attempt: ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )

    class SlideAgent:
        role = "slide"
        ws = str(tmp_path)
        trace_label = attempt
        assigned_slide_pages = (1,)
        repair_required_reason = "Vision Critic unresolved: text overlaps image"
        vision_critic_results = {}
        rendered_output_hashes = {1: "third-render"}
        viewed_output_hashes = {1: "third-render"}

    assert _slide_authoring_stop_line(SlideAgent()) is True
    SlideAgent.repair_required_reason = ""
    assert _slide_authoring_stop_line(SlideAgent()) is True


def test_clean_third_pixel_state_stops_edits_without_false_repair_handoff(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace" / "slide-render-states"
    trace.mkdir(parents=True)
    (trace / "page_01.json").write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "attempt_hashes": {"slide_01": ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )

    class SlideAgent:
        role = "slide"
        ws = str(tmp_path)
        trace_label = "slide_01"
        assigned_slide_pages = (1,)
        repair_required_reason = ""
        final_text = "status: ready\npages: 01\nfinal_pixels_inspected: yes"
        vision_critic_results = {}
        rendered_output_hashes = {1: "third-render"}
        viewed_output_hashes = {1: "third-render"}

    agent = SlideAgent()
    assert _slide_authoring_stop_line(agent) is True
    assert _slide_repair_issue(agent) is None


def test_third_render_gets_its_final_vision_check_before_stop_line(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace" / "slide-render-states"
    trace.mkdir(parents=True)
    (trace / "page_01.json").write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "attempt_hashes": {"slide_01": ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )

    class SlideAgent:
        role = "slide"
        ws = str(tmp_path)
        trace_label = "slide_01"
        assigned_slide_pages = (1,)
        rendered_output_hashes = {1: "third-render"}
        viewed_output_hashes = {1: "second-render"}

    agent = SlideAgent()
    assert _slide_pending_render_views(agent) == (1,)
    assert _slide_authoring_stop_line(agent) is False
    agent.viewed_output_hashes[1] = "third-render"
    assert _slide_pending_render_views(agent) == ()
    assert _slide_authoring_stop_line(agent) is True


def test_same_turn_tool_batch_cannot_edit_after_third_pixel_view(
    tmp_path: Path,
) -> None:
    trace = tmp_path / "_trace" / "slide-render-states"
    trace.mkdir(parents=True)
    (trace / "page_01.json").write_text(
        json.dumps({
            "authoring_attempt_limit": 3,
            "authoring_hashes": ["a", "b", "c"],
            "attempt_hashes": {"slide_01": ["a", "b", "c"]},
        }),
        encoding="utf-8",
    )
    slide = tmp_path / "slides" / "slide_01.html"
    slide.parent.mkdir()
    slide.write_text("before", encoding="utf-8")

    class SlideAgent:
        role = "slide"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        trace_label = "slide_01"
        assigned_slide_pages = (1,)
        rendered_output_hashes = {1: "third-render"}
        viewed_output_hashes = {1: "second-render"}
        vision_critic_results = {}
        repair_required_reason = ""
        started = time.time()
        remote_tool_concurrency = 1
        n_views = 0
        n_renders = 0
        image_by_tool = {}
        trace_dir = tmp_path / "_trace" / "subagents" / "slide_01"

        def safe(self, path):
            return str(tmp_path / path)

        def log(self, _message):
            return None

    agent = SlideAgent()
    (agent.trace_dir / "images").mkdir(parents=True)
    calls = [
        SimpleNamespace(
            id="vision",
            name="vision_analyze",
            input={"image": "renders/slide_01.png", "query": "check"},
        ),
        SimpleNamespace(
            id="patch",
            name="patch",
            input={
                "mode": "replace",
                "path": "slides/slide_01.html",
                "old_string": "before",
                "new_string": "after",
            },
        ),
        SimpleNamespace(
            id="render",
            name="terminal",
            input={
                "command": (
                    "python skills/mural-presenter-v0-4/scripts/slide.py "
                    "render . --page 1"
                )
            },
        ),
    ]
    vision_value = {
        "path": "renders/slide_01.png",
        "image_b64": "eA==",
        "source_sha256": "third-render",
        "vision_verdict": "ready",
        "vision_summary": "clean",
        "vision_result": {"issues": []},
    }
    with mock.patch("core.agent_loop.tools.dispatch", side_effect=[vision_value]):
        results = _tool_results(agent, calls, 1, [])
    assert slide.read_text(encoding="utf-8") == "before"
    assert "slide_authoring_budget" in str(results[-1])
    assert agent.n_renders == 0


def test_slide_first_draft_is_rendered_before_a_patch_storm() -> None:
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        prerender_slide_mutations={2: tools._SLIDE_PRE_RENDER_MUTATION_LIMIT},
        render_attempted_pages=set(),
        rendered_output_hashes={},
        viewed_output_hashes={},
    )
    message = tools._slide_prerender_mutation_error(
        agent, "slides/slide_02.html"
    )
    assert message is not None
    assert "先运行" in message
    agent.render_attempted_pages = {2}
    # A render *attempt* alone no longer disables the cadence guard forever;
    # the dispatcher resets the interval counter when that attempt occurs.
    assert tools._slide_prerender_mutation_error(agent, "slides/slide_02.html")
    agent.prerender_slide_mutations[2] = 0
    assert tools._slide_prerender_mutation_error(agent, "slides/slide_02.html") is None


def test_slide_must_inspect_successful_render_before_more_patches() -> None:
    agent = SimpleNamespace(
        role="slide",
        skill_name="mural-presenter-v0-4",
        prerender_slide_mutations={2: 0},
        rendered_output_hashes={2: "current-pixels"},
        viewed_output_hashes={2: "older-pixels"},
    )
    message = tools._slide_prerender_mutation_error(agent, "slides/slide_02.html")
    assert message is not None
    assert "vision_analyze" in message
    agent.viewed_output_hashes[2] = "current-pixels"
    assert tools._slide_prerender_mutation_error(agent, "slides/slide_02.html") is None


def test_research_brief_is_bounded_before_orchestrator_handoff(
    tmp_path: Path,
) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请制作 8 页演示文稿"}

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    ceiling = tools.research_brief_hard_ceiling(agent)
    result = tools.write_file(
        agent,
        "research/knowledge-brief.md",
        "x" * (ceiling + 1),
    )
    assert "knowledge_brief_too_long" in result
    assert not (tmp_path / "research/knowledge-brief.md").exists()


def test_research_brief_limit_scales_with_requested_page_count(tmp_path: Path) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "Please prepare a 20-slide deck."}

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    assert tools.research_brief_char_limit(agent) == 19000
    result = tools.write_file(
        agent,
        "research/knowledge-brief.md",
        "x" * 15000,
    )
    assert "已写入" in result


def test_research_brief_limit_understands_slide_count_ranges(tmp_path: Path) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        requested_slide_count = 0
        cfg = {"_raw_user_query": "请做一份 25到35页 的电影美学分享"}

    # Use the upper end of the requested range and retain the global cap.
    assert tools.research_brief_char_limit(ResearchAgent()) == 24000


def test_research_brief_range_does_not_parse_date_suffix(tmp_path: Path) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        requested_slide_count = 0
        cfg = {"_raw_user_query": "2025-01-15页后交付，具体页数你决定"}

    # Falls back to the neutral 15-page allowance instead of treating 01-15
    # from an ISO date as a requested slide-count range.
    assert tools.research_brief_char_limit(ResearchAgent()) == 15750


def test_research_patch_cannot_bypass_brief_ceiling(tmp_path: Path) -> None:
    path = tmp_path / "research/knowledge-brief.md"
    path.parent.mkdir()
    original = "x" * 19199 + "Z"
    path.write_text(original, encoding="utf-8")

    class ResearchAgent:
        role = "research"
        label = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        requested_slide_count = 0
        cfg = {"_raw_user_query": "请做 8 页演示"}

        def safe(self, relative):
            return str(tmp_path / relative)

    result = tools.patch(
        ResearchAgent(),
        path="research/knowledge-brief.md",
        old_string="Z",
        new_string="Y" * 200,
    )
    assert "knowledge_brief_too_long" in result
    assert path.read_text(encoding="utf-8") == original


def test_attachment_verification_accepts_first_complete_brief(tmp_path: Path) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    assert tools.research_brief_char_limit(agent) == 20000
    result = tools.write_file(
        agent,
        "research/knowledge-brief.md",
        "x" * 18103,
    )
    assert "已写入" in result


def test_research_handoff_receipt_recovers_only_exact_accepted_brief(
    tmp_path: Path,
) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        requested_slide_count = 8
        cfg = {"_raw_user_query": "请制作 8 页演示文稿"}

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "# Knowledge Brief\n\n## Facts\n" + "evidence\n" * 80
    result = tools.write_file(agent, "research/knowledge-brief.md", content)
    assert "已写入" in result
    assert tools.research_handoff_is_valid(agent) is True

    (tmp_path / "research/knowledge-brief.md").write_text(
        content + "tampered\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(agent) is False


def test_research_receipt_locks_brief_against_patch(tmp_path: Path) -> None:
    """After write_file creates a receipt, subsequent patch is rejected."""

    class ResearchAgent:
        role = "research"
        label = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        requested_slide_count = 8
        cfg = {"_raw_user_query": "请制作 8 页演示文稿"}

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "# Knowledge Brief\n\n## Facts\nold fact\n"
    assert "已写入" in tools.write_file(
        agent, "research/knowledge-brief.md", content
    )
    assert tools.research_handoff_is_valid(agent) is True
    result = tools.patch(
        agent,
        path="research/knowledge-brief.md",
        old_string="old fact",
        new_string="new fact",
    )
    assert "knowledge_brief_locked" in result
    assert (tmp_path / "research/knowledge-brief.md").read_text("utf-8") == content


def test_research_blocked_parent_rejects_new_delegation(tmp_path: Path) -> None:
    parent = SimpleNamespace(
        revision_mode=False,
        child_outcomes={},
        available_roles={"material", "research", "image", "slide", "review"},
        research_required=True,
        research_completed=False,
        research_blocked=True,
        ws=str(tmp_path),
    )
    result = _delegate(parent, {"tasks": [{
        "role": "research",
        "label": "research",
        "task": "再次研究",
    }]})
    assert "Research 已进入 research_blocked" in result


def test_unstarted_operational_slide_retries_inside_same_wave() -> None:
    messages = []
    parent = SimpleNamespace(log=messages.append)
    specs = [
        {"role": "slide", "label": "slide_04", "pages": [4], "task": "P04"},
        {"role": "slide", "label": "slide_05", "pages": [5], "task": "P05"},
    ]
    results = [
        {
            "role": "slide",
            "label": "slide_04",
            "ok": False,
            "exit_reason": "api_failed",
            "attempt": 1,
            "completed_pages": [],
            "repair_issue": None,
        },
        {
            "role": "slide",
            "label": "slide_05",
            "ok": True,
            "exit_reason": "repair_required_handoff",
            "attempt": 1,
            "completed_pages": [5],
            "repair_issue": {"status": "repair_required", "pages": [5]},
        },
    ]
    recovered = {
        **results[0],
        "ok": True,
        "exit_reason": "text_response",
        "attempt": 2,
        "completed_pages": [4],
    }
    with mock.patch("core.agent_loop._run_child", return_value=recovered) as run:
        updated = _auto_retry_unstarted_slide_interruptions(
            parent, specs, results, limit=12
        )
    assert updated[0]["ok"] is True
    assert updated[1]["exit_reason"] == "repair_required_handoff"
    assert run.call_count == 1
    retry_spec = run.call_args.args[2]
    assert retry_spec["label"] == "slide_04"
    assert retry_spec["_recovery_kind"] == "operational"
    assert any("slide_04" in message for message in messages)


def test_attachment_reproduction_brief_cannot_point_back_to_material(
    tmp_path: Path,
) -> None:
    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请忠实重绘附件中的两个 Figure/表格"}

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.write_file(
        ResearchAgent(),
        "research/knowledge-brief.md",
        "## F2\n- 数据：见 material.md §5.4 完整表\n",
    )
    assert "knowledge_brief_indirect_artifact" in result


def test_image_requires_contact_sheet_but_not_every_catalog_asset_at_full_resolution(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## hero\n- path: assets/hero.png\n",
        encoding="utf-8",
    )
    (assets / "hero.png").write_bytes(b"hero")
    (assets / "contact-sheet.png").write_bytes(b"sheet")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        vision_critic_results = {}

    agent = ImageAgent()
    gap = _image_required_fullres_gap(agent)
    assert "assets/contact-sheet.png" in gap
    agent.vision_critic_results = {
        "assets/contact-sheet.png": {
            "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
            "verdict": "ready",
        },
    }
    agent.image_finalized_catalog_digest = hashlib.sha256(
        (assets / "catalog.md").read_bytes()
    ).hexdigest()
    assert _image_required_fullres_gap(agent) == ""


def test_image_all_failed_handoff_closes_without_impossible_contact_sheet(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    plan = tmp_path / "plan"
    assets.mkdir()
    plan.mkdir()
    (plan / "slide_01.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (plan / "slide_04.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n"
        "## cover\n- slides: 1\n- kind: real\n- path: assets/cover.png\n"
        "- status: failed\n- failure_reason: search unavailable\n\n"
        "## detail\n- slides: 4\n- kind: real\n- path: assets/detail.png\n"
        "- status: failed\n- failure_reason: no faithful candidate\n",
        encoding="utf-8",
    )

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        final_text = "bitmap_ready: none\nfailed: 01, 04 (source unavailable)"
        vision_critic_results = {}

        def log(self, _message):
            return None

    agent = ImageAgent()
    assert _image_required_fullres_gap(agent) == ""
    assert agent.image_ready_pages == ()
    assert agent.image_failed_pages == (1, 4)


def test_one_preproduction_image_source_route_repair_requires_real_change(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    trace = tmp_path / "_trace"
    renders = tmp_path / "renders"
    plan.mkdir()
    trace.mkdir()
    renders.mkdir()
    slide = plan / "slide_01.md"
    slide.write_text(
        "- needs_bitmap: true\n- primary_visual_medium: bitmap-real\n",
        encoding="utf-8",
    )
    (trace / "attachment-manifest.json").write_text(
        json.dumps({"visual_asset_paths": ["inputs/report.pdf.pages/page_001.png"]}),
        encoding="utf-8",
    )
    before = _image_route_fingerprint(str(tmp_path))
    slide.write_text(
        "- needs_bitmap: true\n- primary_visual_medium: bitmap-material\n",
        encoding="utf-8",
    )
    parent = SimpleNamespace(
        ws=str(tmp_path),
        image_failed_pages=(1,),
        runtime_capabilities={
            "inputs": {"visual_asset_paths": ["inputs/report.pdf.pages/page_001.png"]},
            "tools": {"image_generate": True},
        },
    )
    previous = {
        "ok": True,
        "status": "partial_ready",
        "input_fingerprint": before,
    }
    allowed, reason = _claim_image_source_route_repair(parent, previous)
    assert allowed is True
    assert reason == ""
    repeated, repeated_reason = _claim_image_source_route_repair(parent, previous)
    assert repeated is False
    assert "already consumed" in repeated_reason


def test_image_explicit_partial_handoff_closes_after_replacement_budget(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_01.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (plan / "slide_04.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## cover\n- slides: 1\n- path: assets/cover.jpg\n\n"
        "## food\n- slides: 4\n- path: assets/food.jpg\n",
        encoding="utf-8",
    )
    (assets / "cover.jpg").write_bytes(b"cover")
    (assets / "food.jpg").write_bytes(b"food")
    (assets / "contact-sheet.png").write_bytes(b"sheet")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        final_text = "bitmap_ready: 01\nfailed: 04 (subject mismatch)"
        vision_critic_results = {
            "assets/contact-sheet.png": {
                "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
                "verdict": "ready",
            },
            "assets/cover.jpg": {
                "source_sha256": hashlib.sha256(b"cover").hexdigest(),
                "verdict": "ready",
            },
            "assets/food.jpg": {
                "source_sha256": hashlib.sha256(b"food").hexdigest(),
                "verdict": "repair_required",
                "summary": "real Sichuan dish, not the requested hotpot",
            },
        }

        def log(self, _message):
            return None

    agent = ImageAgent()
    assert _image_required_fullres_gap(agent) == ""
    assert agent.image_ready_pages == (1,)
    assert agent.image_failed_pages == (4,)


def test_image_contact_sheet_regeneration_does_not_stale_current_originals(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## hero\n- slides: 1\n- path: assets/hero.jpg\n",
        encoding="utf-8",
    )
    (assets / "hero.jpg").write_bytes(b"current hero")
    (assets / "contact-sheet.png").write_bytes(b"regenerated sheet")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        vision_critic_results = {
            "assets/contact-sheet.png": {
                "source_sha256": hashlib.sha256(b"older sheet").hexdigest(),
                "verdict": "ready",
            },
            "assets/hero.jpg": {
                "source_sha256": hashlib.sha256(b"current hero").hexdigest(),
                "verdict": "ready",
            },
        }

    assert _image_required_fullres_gap(ImageAgent()) == ""


def test_image_second_infra_failure_recovers_only_current_verified_pages(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_01.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (plan / "slide_09.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## cover\n- slides: 1\n- path: assets/cover.png\n",
        encoding="utf-8",
    )
    (assets / "cover.png").write_bytes(b"verified cover")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        exit_reason = "api_failed"
        vision_critic_results = {
            "assets/cover.png": {
                "source_sha256": hashlib.sha256(b"verified cover").hexdigest(),
                "verdict": "ready",
            }
        }

    agent = ImageAgent()
    assert not _recover_image_operational_handoff(agent, 1)
    assert _recover_image_operational_handoff(agent, 2)
    assert agent.image_ready_pages == (1,)
    assert agent.image_failed_pages == (9,)
    assert "bitmap_ready: 01" in agent.image_handoff_summary
    assert "failed: 09" in agent.image_handoff_summary


def test_image_second_infra_failure_can_release_all_pages_as_failed(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_04.md").write_text("- needs_bitmap: true\n", encoding="utf-8")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        exit_reason = "runtime_timeout"
        vision_critic_results = {}

    agent = ImageAgent()
    assert _recover_image_operational_handoff(agent, 2)
    assert agent.image_ready_pages == ()
    assert agent.image_failed_pages == (4,)


@pytest.mark.parametrize(
    "exit_reason", ["incomplete_closure", "stalled_repetition", "max_turns"]
)
def test_image_bounded_failure_preserves_current_finalize_page_partition(
    tmp_path: Path,
    exit_reason: str,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_01.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (plan / "slide_04.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n"
        "## cover-main\n- slides: 1\n- path: assets/cover.jpg\n\n"
        "## cover-support\n- slides: 1\n- path: assets/support.jpg\n\n"
        "## food\n- slides: 4\n- path: assets/food.jpg\n",
        encoding="utf-8",
    )
    (assets / "cover.jpg").write_bytes(b"cover")
    (assets / "support.jpg").write_bytes(b"support")
    (assets / "food.jpg").write_bytes(b"food")
    (assets / "contact-sheet.png").write_bytes(b"sheet")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        vision_critic_results = {
            "assets/contact-sheet.png": {
                "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
                "verdict": "ready",
            },
            "assets/cover.jpg": {
                "source_sha256": hashlib.sha256(b"cover").hexdigest(),
                "verdict": "ready",
            },
            "assets/support.jpg": {
                "source_sha256": hashlib.sha256(b"support").hexdigest(),
                "verdict": "repair_required",
            },
            "assets/food.jpg": {
                "source_sha256": hashlib.sha256(b"food").hexdigest(),
                "verdict": "ready",
            },
        }

    agent = ImageAgent()
    agent.exit_reason = exit_reason
    agent.image_ready_pages = (4,)
    agent.image_failed_pages = (1,)
    agent.image_handoff_summary = "status:PASS\nbitmap_ready:04\nfailed:01"
    agent.image_finalized_catalog_digest = hashlib.sha256(
        (assets / "catalog.md").read_bytes()
    ).hexdigest()
    assert _image_catalog_has_current_fullres_coverage(agent)
    assert _recover_image_operational_handoff(agent, 1)
    assert agent.image_ready_pages == (4,)
    assert agent.image_failed_pages == (1,)


def test_image_incomplete_closure_does_not_recover_without_fullres_coverage(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_01.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## cover\n- slides: 1\n- path: assets/cover.jpg\n",
        encoding="utf-8",
    )
    (assets / "cover.jpg").write_bytes(b"cover")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        exit_reason = "incomplete_closure"
        # The original itself is current, but the required locator-sheet pass
        # never happened; bounded semantic recovery must therefore stay shut.
        vision_critic_results = {
            "assets/cover.jpg": {
                "source_sha256": hashlib.sha256(b"cover").hexdigest(),
                "verdict": "ready",
            }
        }

    agent = ImageAgent()
    assert not _image_catalog_has_current_fullres_coverage(agent)
    assert not _recover_image_operational_handoff(agent, 1)


def test_image_catalog_change_after_finalize_blocks_semantic_closure(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "slide_01.md").write_text("- needs_bitmap: true\n", encoding="utf-8")
    catalog = assets / "catalog.md"
    catalog.write_text(
        "# Asset catalog\n\n## cover\n- slides: 1\n- path: assets/cover.jpg\n",
        encoding="utf-8",
    )
    (assets / "cover.jpg").write_bytes(b"cover")
    (assets / "contact-sheet.png").write_bytes(b"sheet")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        final_text = "bitmap_ready: 01\nfailed: none"
        vision_critic_results = {
            "assets/contact-sheet.png": {
                "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
                "verdict": "ready",
            },
            "assets/cover.jpg": {
                "source_sha256": hashlib.sha256(b"cover").hexdigest(),
                "verdict": "ready",
            },
        }

        def log(self, _message):
            return None

    agent = ImageAgent()
    agent.image_finalized_catalog_digest = hashlib.sha256(
        catalog.read_bytes()
    ).hexdigest()
    assert _image_required_fullres_gap(agent) == ""
    catalog.write_text(catalog.read_text(encoding="utf-8") + "\n- source: changed\n")
    assert "catalog 尚未在当前内容上完成最终确认" in _image_required_fullres_gap(agent)
    agent.exit_reason = "incomplete_closure"
    assert not _recover_image_operational_handoff(agent, 1)


def test_echarts_hyphenated_rich_key_is_reported(tmp_path: Path) -> None:
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "slide_10.html").write_text(
        """
        <script>
        var chart = echarts.init(document.getElementById('chart'));
        chart.setOption({
          yAxis: {axisLabel: {
            formatter: '{tag-open|OPEN}',
            rich: {'tag-open': {fontSize: 18}}
          }}
        });
        </script>
        """,
        encoding="utf-8",
    )
    warnings = deck_core._echarts_typography_warnings(tmp_path, 10)
    assert any("tag-open" in warning and "literal template" in warning for warning in warnings)


def test_echarts_dynamic_hyphenated_rich_key_is_reported(tmp_path: Path) -> None:
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "slide_10.html").write_text(
        """
        <script>
        var rows = [{name: 'Model', cls: 'open-weight'}];
        var labels = rows.map(function (row) {
          return '{name|' + row.name + '}\\n{tag-' + row.cls + '|' + row.cls + '}';
        });
        var chart = echarts.init(document.getElementById('chart'));
        chart.setOption({yAxis: {data: labels, axisLabel: {rich: {
          name: {fontSize: 20},
          'tag-open-weight': {fontSize: 18},
          'tag-agent-tuned': {fontSize: 18}
        }}}});
        </script>
        """,
        encoding="utf-8",
    )
    warnings = deck_core._echarts_typography_warnings(tmp_path, 10)
    combined = "\n".join(warnings)
    assert "tag-<dynamic>" in combined
    assert "tag-open-weight" in combined
    assert "tag-agent-tuned" in combined
    assert "literal template" in combined


def test_finalize_recovery_quarantines_only_renderer_owned_preview_files(
    tmp_path: Path,
) -> None:
    temporary = tmp_path / "tmp"
    temporary.mkdir()
    stale = temporary / "preview_slide_04.5c4olkgf.html"
    unrelated = temporary / "user-note.txt"
    stale.write_text("<html>stale preview</html>", encoding="utf-8")
    unrelated.write_text("keep visible as an audit violation", encoding="utf-8")

    moved = deck_core.recover_stale_render_previews(tmp_path)

    assert moved == ["tmp/preview_slide_04.5c4olkgf.html"]
    assert not stale.exists()
    assert unrelated.read_text(encoding="utf-8") == "keep visible as an audit violation"
    quarantine = tmp_path / "_trace" / "noncanonical-artifacts"
    assert any(path.read_text(encoding="utf-8") == "<html>stale preview</html>" for path in quarantine.iterdir())


def test_progress_guard_stops_varied_goals_bypassing_same_closed_unit_policy(
    tmp_path: Path,
) -> None:
    guard = _ProgressGuard(str(tmp_path))
    action = "ok"
    reason = ""
    for page in (1, 2, 3, 4):
        use = SimpleNamespace(
            name="delegate_task",
            input={"tasks": [{"role": "slide", "pages": [page], "goal": f"unique-{page}"}]},
        )
        result = {
            "type": "tool_result",
            "content": (
                f"责任单元 slide_{page:02d} 已结束，exit_reason=text_response；"
                "这不是允许创建新 Slide Agent 的基础设施中断。"
            ),
        }
        action, reason = guard.observe([use], [result])
    assert action == "stop"
    assert "权限或阶段错误" in reason


def test_progress_guard_allows_three_review_cycles_with_real_pixel_progress(
    tmp_path: Path,
) -> None:
    slide = tmp_path / "slides/slide_01.html"
    slide.parent.mkdir()
    slide.write_text("<section>v0</section>", encoding="utf-8")
    guard = _ProgressGuard(str(tmp_path))
    use = SimpleNamespace(
        name="terminal",
        input={"command": "python scripts/review.py finalize . --expected 1"},
    )
    result = {"type": "tool_result", "content": "status:PASS"}

    for cycle in range(1, 4):
        slide.write_text(f"<section>v{cycle}</section>", encoding="utf-8")
        action, reason = guard.observe([use], [result])
        assert action != "stop", reason


def test_image_handoff_accepts_markdown_emphasis_and_failed_bullets() -> None:
    text = """
## 最终状态

**bitmap_ready: 01, 05, 07**
- `assets/cover.jpg`（slide 01）— ready.

**failed:**
- `04 (subject_mismatch)` — wrong landmark.
- `06 (no_negative_space)` — unsuitable crop.
- `08 (subject_mismatch)` — wrong subject.

## 遗留问题
- source image is 5472×3648 and was not downloaded.
"""
    assert _image_handoff_declared_pages(text, "bitmap_ready") == {1, 5, 7}
    assert _image_handoff_declared_pages(text, "failed") == {4, 6, 8}


def test_image_partial_handoff_rejects_omitted_planned_bitmap_page(
    tmp_path: Path,
) -> None:
    assets = tmp_path / "assets"
    assets.mkdir()
    plan = tmp_path / "plan"
    plan.mkdir()
    for page in (1, 4, 8):
        (plan / f"slide_{page:02d}.md").write_text(
            "- needs_bitmap: true\n", encoding="utf-8"
        )
    (assets / "catalog.md").write_text(
        "# Asset catalog\n\n## cover\n- slides: 1\n- path: assets/cover.jpg\n\n"
        "## food\n- slides: 4\n- path: assets/food.jpg\n",
        encoding="utf-8",
    )
    (assets / "cover.jpg").write_bytes(b"cover")
    (assets / "food.jpg").write_bytes(b"food")
    (assets / "contact-sheet.png").write_bytes(b"sheet")

    class ImageAgent:
        role = "image"
        ws = str(tmp_path)
        final_text = "bitmap_ready: 01\nfailed: 04 (subject mismatch)"
        _image_replacement_passes = 1
        vision_critic_results = {
            "assets/contact-sheet.png": {
                "source_sha256": hashlib.sha256(b"sheet").hexdigest(),
                "verdict": "ready",
            },
            "assets/cover.jpg": {
                "source_sha256": hashlib.sha256(b"cover").hexdigest(),
                "verdict": "ready",
            },
            "assets/food.jpg": {
                "source_sha256": hashlib.sha256(b"food").hexdigest(),
                "verdict": "repair_required",
                "summary": "subject mismatch",
            },
        }

        def log(self, _message):
            return None

    gap = _image_required_fullres_gap(ImageAgent())
    assert "finalize receipt 未覆盖全部 needs_bitmap 页面" in gap


def test_current_variant_terminal_blocks_trace_introspection(tmp_path: Path) -> None:
    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
    )
    assert "trace_internal" in tools.terminal(
        agent, "tail -20 _trace/subagents/image/tool_log.json"
    )


def test_image_catalog_noop_and_rewrite_budget_force_local_patch(tmp_path: Path) -> None:
    (tmp_path / "assets").mkdir()

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    agent = ImageAgent()
    first = "# Asset catalog\n\n## hero\n- path: assets/a.jpg\n"
    assert "已写入" in tools.write_file(agent, "assets/catalog.md", first)
    assert "image_catalog_noop" in tools.write_file(agent, "assets/catalog.md", first)
    second = first.replace("a.jpg", "b.jpg")
    assert "已写入" in tools.write_file(agent, "assets/catalog.md", second)
    third = second.replace("b.jpg", "c.jpg")
    assert "image_catalog_rewrite_budget" in tools.write_file(
        agent, "assets/catalog.md", third
    )


def test_requested_page_count_blocks_added_closing_page(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        requested_slide_count = 8
        delegated_roles = set()
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    agent = Orchestrator()
    wrong_deck = "# Deck\n## Resolved deck brief\n- page_count: 9\n"
    assert "page_count_contract" in tools.write_file(
        agent, "plan/deck.md", wrong_deck
    )
    assert not (tmp_path / "plan/deck.md").exists()

    too_far = [{"path": "plan/slide_09.md", "content": "# slide_09\n"}]
    assert "page_count_contract" in tools.write_plan_batch(agent, too_far)

    command = (
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py "
        "validate-plans . --expected 9"
    )
    assert "page_count_contract" in tools.terminal(agent, command)
    no_expected = (
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py "
        "validate-plans ."
    )
    assert "page_count_contract" in tools.terminal(agent, no_expected)

    correct_deck = "# Deck\n## Resolved deck brief\n- page_count: 8\n"
    assert "已写入" in tools.write_file(agent, "plan/deck.md", correct_deck)


def test_structured_slide_count_contract_rejects_wrong_page_count(tmp_path: Path) -> None:
    (tmp_path / "plan").mkdir()

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        requested_slide_count = 1
        delegated_roles = set()
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    agent = Orchestrator()
    three_pages = "# Deck\n## Resolved deck brief\n- page_count: 3\n"
    assert "page_count_contract" in tools.write_file(
        agent, "plan/deck.md", three_pages
    )
    wrong_validate = (
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py "
        "validate-plans . --expected 3"
    )
    assert "page_count_contract" in tools.terminal(agent, wrong_validate)
    one_page = "# Deck\n## Resolved deck brief\n- page_count: 1\n"
    assert "已写入" in tools.write_file(agent, "plan/deck.md", one_page)


def test_explicit_color_exclusion_rejects_renamed_large_surface_only(
    tmp_path: Path,
) -> None:
    (tmp_path / "plan").mkdir()

    class Orchestrator:
        role = "orchestrator"
        skill_name = "mural-presenter-v0-4"
        revision_mode = False
        requested_slide_count = 0
        delegated_roles = set()
        ws = str(tmp_path)
        cfg = {
            "_raw_user_query": (
                "不要米黄纸张模板，也不要电蓝科技风；配色从海岸和红砖取色。"
            )
        }

        def safe(self, path):
            return str(tmp_path / path)

    agent = Orchestrator()
    conflict = (
        "# Deck\n## 视觉契约\n"
        "- 画布 = 米白墙 #F2E7D3（不是米黄纸张）\n"
        "```css\n:root { --content-canvas: #F5EBD6; /* 浅沙米白 */ }\n```\n"
    )
    result = tools.write_file(agent, "plan/deck.md", conflict)
    assert "explicit_surface_color_conflict" in result
    assert not (tmp_path / "plan/deck.md").exists()

    compliant = (
        "# Deck\n## 视觉契约\n"
        "- 画布 = 海雾青灰 #D9EFF1\n"
        "- 小面积 accent = 沙金 #E8C88A\n"
        "```css\n:root { --content-canvas: #D9EFF1; --surface: #C7D9D8; }\n```\n"
    )
    assert "已写入" in tools.write_file(agent, "plan/deck.md", compliant)

    explanatory = (
        "# Deck\n## 视觉契约\n"
        "- 配色命题：用户不用米黄纸张，也不用电蓝科技风。\n"
        "- `--content-canvas`：海雾白 `#EAF1F4`（与米黄纸张明显不同）。\n"
        "- `--surface`：冷灰青 `#C7D9D8`（不是电蓝）。\n"
    )
    assert "已写入" in tools.write_file(agent, "plan/deck.md", explanatory)


def test_patch_accepts_legacy_value_aliases_at_dispatch_edge(tmp_path: Path) -> None:
    target = tmp_path / "assets/catalog.md"
    target.parent.mkdir()
    target.write_text("- download: old-url\n", encoding="utf-8")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.dispatch(
        ImageAgent(),
        "patch",
        {
            "path": "assets/catalog.md",
            "old_value": "old-url",
            "new_value": "new-url",
        },
    )
    assert result == "已编辑 assets/catalog.md"
    assert "new-url" in target.read_text(encoding="utf-8")


def test_plan_validator_enforces_selected_ownership_topology(tmp_path: Path) -> None:
    script = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
    plan = tmp_path / "plan"
    plan.mkdir()

    def write_deck(topology: str) -> None:
        (plan / "deck.md").write_text(
            f"""# Deck Plan
- title: Ownership test

## Resolved deck brief
- language: en
- page_count: 3
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: deterministic topology fixture uses only typography
- ownership_topology: {topology}
- ownership_rationale: dependency structure
- rationale: deterministic test

## Audience and objective
Verify ownership.

## Narrative and page map
Three pages.

## Visual storyboard
Code visuals.

## Visual contract
Test system.

## Special pages
None.

## Theme Tokens
```css
:root {{
  --content-canvas: #eeeeee;
}}
```
""",
            encoding="utf-8",
        )

    def write_slides(groups: list[str]) -> None:
        for page_no, group in enumerate(groups, start=1):
            (plan / f"slide_{page_no:02d}.md").write_text(
                f"""# slide_{page_no:02d}
- role: test page
- page_type: content
- page_family: editorial-test
- production_group: {group}
- composition: editorial
- needs_bitmap: false
- primary_visual_medium: editorial-typography

## Narrative
Advance the test.

## Evidence
- Deterministic fixture.

## On-screen copy (exact)
- title: Page {page_no}

## Composition blueprint
- focal: title

## Semantic visual need
Use a deliberate typographic hierarchy without raster assets.

## Speech beat
Explain page {page_no}.
""",
                encoding="utf-8",
            )

    write_deck("single")
    write_slides(["shared", "shared", "page-03"])
    invalid_single = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert invalid_single.returncode != 0
    assert "one unique production_group per content page" in invalid_single.stdout

    write_slides(["page-01", "page-02", "page-03"])
    valid_single = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert valid_single.returncode == 0, valid_single.stderr
    assert "ownership:single" in valid_single.stdout

    trace = tmp_path / "_trace"
    trace.mkdir()
    (trace / "runtime-capabilities.json").write_text(
        json.dumps({"requirements": {"bitmap_required": True}}),
        encoding="utf-8",
    )
    required_bitmap = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert required_bitmap.returncode != 0
    assert "cannot be silently downgraded" in required_bitmap.stdout
    (trace / "runtime-capabilities.json").unlink()

    write_deck("grouped")
    write_slides(["mechanism", "isolated", "mechanism"])
    invalid_grouped = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert invalid_grouped.returncode != 0
    assert "contiguous pages" in invalid_grouped.stdout

    write_slides(["mechanism", "mechanism", "closing"])
    valid_grouped = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "3"],
        capture_output=True,
        text=True,
    )
    assert valid_grouped.returncode == 0, valid_grouped.stderr
    assert "ownership:grouped" in valid_grouped.stdout


def test_plan_validator_enforces_mandatory_special_page_groups(tmp_path: Path) -> None:
    script = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
    plan = tmp_path / "plan"
    plan.mkdir()

    def write_deck(page_count: int) -> None:
        for old in plan.glob("slide_*.md"):
            old.unlink()
        (plan / "deck.md").write_text(
            f"""# Deck Plan
- title: Special group test

## Resolved deck brief
- language: en
- page_count: {page_count}
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: typography-only fixture
- ownership_topology: single
- ownership_rationale: independent pages
- rationale: deterministic test

## Audience and objective
Verify special-page group enforcement.

## Narrative and page map
Test pages.

## Visual storyboard
Code visuals.

## Visual contract
Test system.

## Special pages
Bookends and dividers.

## Theme Tokens
```css
:root {{
  --content-canvas: #eeeeee;
}}
```
""",
            encoding="utf-8",
        )

    def write_page(number: int, page_type: str, group: str) -> None:
        layout = ""
        if page_type == "cover":
            layout = "- special_layout: centered\n"
        elif page_type == "closing":
            layout = "- special_layout: centered\n"
        elif page_type == "section-divider":
            layout = "- special_layout: number-copy\n- section_index: 1\n"
        (plan / f"slide_{number:02d}.md").write_text(
            f"""# slide_{number:02d}
- role: test
- page_type: {page_type}
- page_family: test-family
- production_group: {group}
{layout}- needs_bitmap: false
- primary_visual_medium: editorial-typography

## Narrative
Test.

## Evidence
- Fixture.

## On-screen copy (exact)
- title: Page {number}

## Composition blueprint
- focal: title

## Semantic visual need
Typographic hierarchy only.

## Speech beat
Explain page {number}.
""",
            encoding="utf-8",
        )

    def run_validate(expected: int):
        return subprocess.run(
            [sys.executable, str(script), "validate-plans", str(tmp_path),
             "--expected", str(expected)],
            capture_output=True, text=True,
        )

    # cover with wrong production_group -> must fail
    write_deck(3)
    write_page(1, "cover", "wrong-cover")
    write_page(2, "content", "page-02")
    write_page(3, "closing", "bookends")
    result = run_validate(3)
    assert result.returncode != 0
    assert "bookends" in result.stdout

    # closing with wrong production_group -> must fail
    write_page(1, "cover", "bookends")
    write_page(3, "closing", "wrong-closing")
    result = run_validate(3)
    assert result.returncode != 0
    assert "bookends" in result.stdout

    # 2 section-dividers not using "dividers" -> must fail
    write_deck(6)
    write_page(1, "cover", "bookends")
    write_page(2, "content", "page-02")
    write_page(3, "section-divider", "chapter-a")
    write_page(4, "content", "page-04")
    write_page(5, "section-divider", "chapter-b")
    write_page(6, "closing", "bookends")
    result = run_validate(6)
    assert result.returncode != 0
    assert "dividers" in result.stdout

    # lone section-divider with its own group -> PASS (exception preserved)
    write_deck(4)
    write_page(1, "cover", "bookends")
    write_page(2, "section-divider", "chapter-intro")
    write_page(3, "content", "page-03")
    write_page(4, "closing", "bookends")
    result = run_validate(4)
    assert result.returncode == 0, result.stdout

    # correct fixture: bookends + 2 dividers in "dividers" -> PASS
    write_deck(7)
    write_page(1, "cover", "bookends")
    write_page(2, "content", "page-02")
    write_page(3, "section-divider", "dividers")
    write_page(4, "content", "page-04")
    write_page(5, "section-divider", "dividers")
    write_page(6, "content", "page-06")
    write_page(7, "closing", "bookends")
    result = run_validate(7)
    assert result.returncode == 0, result.stdout
    assert "ownership:single" in result.stdout


def test_plan_validator_rejects_reserved_group_names_on_wrong_page_types(
    tmp_path: Path,
) -> None:
    script = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
    plan = tmp_path / "plan"
    plan.mkdir()

    def write_deck(page_count: int, topology: str = "single") -> None:
        for old in plan.glob("slide_*.md"):
            old.unlink()
        (plan / "deck.md").write_text(
            f"""# Deck Plan
- title: Reserved group test

## Resolved deck brief
- language: en
- page_count: {page_count}
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: typography-only fixture
- ownership_topology: {topology}
- ownership_rationale: test
- rationale: deterministic test

## Audience and objective
Verify reserved group enforcement.

## Narrative and page map
Test pages.

## Visual storyboard
Code visuals.

## Visual contract
Test system.

## Special pages
As needed.

## Theme Tokens
```css
:root {{
  --content-canvas: #eeeeee;
}}
```
""",
            encoding="utf-8",
        )

    def write_page(number: int, page_type: str, group: str) -> None:
        layout = ""
        if page_type == "cover":
            layout = "- special_layout: centered\n"
        elif page_type == "closing":
            layout = "- special_layout: centered\n"
        elif page_type == "section-divider":
            layout = "- special_layout: number-copy\n- section_index: 1\n"
        (plan / f"slide_{number:02d}.md").write_text(
            f"""# slide_{number:02d}
- role: test
- page_type: {page_type}
- page_family: test-family
- production_group: {group}
{layout}- composition: editorial
- needs_bitmap: false
- primary_visual_medium: editorial-typography

## Narrative
Test.

## Evidence
- Fixture.

## On-screen copy (exact)
- title: Page {number}

## Composition blueprint
- focal: title

## Semantic visual need
Typographic hierarchy only.

## Speech beat
Explain page {number}.
""",
            encoding="utf-8",
        )

    def run_validate(expected: int):
        return subprocess.run(
            [sys.executable, str(script), "validate-plans", str(tmp_path),
             "--expected", str(expected)],
            capture_output=True, text=True,
        )

    # content page using reserved "bookends" in single topology -> must fail
    write_deck(3)
    write_page(1, "content", "bookends")
    write_page(2, "content", "page-02")
    write_page(3, "content", "page-03")
    result = run_validate(3)
    assert result.returncode != 0
    assert "reserved for cover and closing" in result.stdout

    # content page using reserved "dividers" in single topology -> must fail
    write_deck(3)
    write_page(1, "content", "page-01")
    write_page(2, "content", "dividers")
    write_page(3, "content", "page-03")
    result = run_validate(3)
    assert result.returncode != 0
    assert "reserved for section-divider" in result.stdout

    # content page using "bookends" in grouped topology (bypasses contiguity) -> must fail
    write_deck(3, topology="grouped")
    write_page(1, "content", "bookends")
    write_page(2, "content", "mechanism")
    write_page(3, "content", "mechanism")
    result = run_validate(3)
    assert result.returncode != 0
    assert "reserved for cover and closing" in result.stdout

    # content page using "dividers" in grouped topology -> must fail
    write_deck(3, topology="grouped")
    write_page(1, "content", "dividers")
    write_page(2, "content", "mechanism")
    write_page(3, "content", "mechanism")
    result = run_validate(3)
    assert result.returncode != 0
    assert "reserved for section-divider" in result.stdout


def test_plan_validator_warns_on_excessive_section_dividers(tmp_path: Path) -> None:
    script = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"

    def make_fixture(name: str, page_types: list[str]) -> Path:
        root = tmp_path / name
        plan = root / "plan"
        plan.mkdir(parents=True)
        page_count = len(page_types)
        (plan / "deck.md").write_text(
            f"""# Deck Plan
- title: Divider quota test

## Resolved deck brief
- language: en
- page_count: {page_count}
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: deterministic structural fixture uses only typography
- ownership_topology: single
- ownership_rationale: every page is independently owned in this fixture
- rationale: validate special-page pacing without blocking delivery

## Audience and objective
Verify divider pacing.

## Narrative and page map
{page_count} deterministic pages.

## Visual storyboard
Code visuals.

## Visual contract
Test system.

## Special pages
Cover, closing, and only the listed dividers.

## Theme Tokens
```css
:root {{
  --content-canvas: #eeeeee;
}}
```
""",
            encoding="utf-8",
        )
        divider_index = 0
        for page_no, page_type in enumerate(page_types, start=1):
            special = page_type in {"cover", "section-divider", "closing"}
            if page_type == "section-divider":
                divider_index += 1
            special_meta = "- special_layout: centered\n" if special else "- composition: editorial\n"
            evidence = "" if special else "\n## Evidence\n- Deterministic fixture.\n"
            section_index = (
                f"\n- section_index: {divider_index:02d}"
                if page_type == "section-divider"
                else ""
            )
            group = (
                "bookends" if page_type in {"cover", "closing"}
                else "dividers" if page_type == "section-divider"
                else f"page-{page_no:02d}"
            )
            (plan / f"slide_{page_no:02d}.md").write_text(
                f"""# slide_{page_no:02d}
- role: test page
- page_type: {page_type}
- page_family: editorial-test
- production_group: {group}
{special_meta}- needs_bitmap: false
- primary_visual_medium: editorial-typography

## Narrative
Advance the test.{evidence}
## On-screen copy (exact)
- title: Page {page_no}{section_index}

## Composition blueprint
- focal: title

## Semantic visual need
Use a deliberate typographic hierarchy without raster assets.

## Speech beat
Explain page {page_no}.
""",
                encoding="utf-8",
            )
        return root

    cases = {
        "short": ["cover", "content", "content", "content", "closing"],
        "balanced": [
            "cover",
            "section-divider",
            "content",
            "content",
            "content",
            "section-divider",
            "content",
            "content",
            "content",
            "closing",
        ],
        "over": [
            "cover",
            "section-divider",
            "content",
            "section-divider",
            "content",
            "section-divider",
            "content",
            "closing",
        ],
    }
    results = {}
    for name, page_types in cases.items():
        root = make_fixture(name, page_types)
        results[name] = subprocess.run(
            [
                sys.executable,
                str(script),
                "validate-plans",
                str(root),
                "--expected",
                str(len(page_types)),
            ],
            capture_output=True,
            text=True,
        )
        assert results[name].returncode == 0, results[name].stderr

    for name in ("short", "balanced"):
        assert "section-divider quota exceeded" not in results[name].stdout
        assert "special-page ratio exceeds 40%" not in results[name].stdout
    assert "[plan-warning] section-divider quota exceeded" in results["over"].stdout
    assert "[plan-warning] special-page ratio exceeds 40%" in results["over"].stdout


def test_revision_finish_gap_does_not_require_unchanged_pages(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    for page in range(1, 4):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page}\n", encoding="utf-8"
        )

    class RevisionAgent:
        role = "orchestrator"
        skill_language = "zh"
        material_required = False
        research_required = False
        research_completed = False
        revision_mode = True
        revision_route = "complex_edit"
        delegated_roles = {"slide"}
        child_outcomes = {"slide_02": {"ok": True}}
        review_completed = False
        finalize_attempted = False
        finalize_succeeded = False
        ws = str(tmp_path)

    gap = _finish_gap(RevisionAgent())
    assert "Review 尚未完成" in gap
    assert "Slide 01:" not in gap
    assert "Slide 03:" not in gap


def test_pdf_ingestion_writes_complete_text_and_page_derivatives(tmp_path: Path) -> None:
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    source = inputs / "paper.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Principal finding: attachment evidence is preserved.")
    document.save(source)
    document.close()

    companion = inputs / "paper.pdf.md"
    assert _extract_material(source, companion)
    text = companion.read_text(encoding="utf-8")
    assert "Principal finding" in text
    assert "text_mode: native" in text
    assert (inputs / "paper.pdf.pages/page_001.png").is_file()
    layout = json.loads(
        (inputs / "paper.pdf.pages/page_001.json").read_text(encoding="utf-8")
    )
    assert layout["text_mode"] == "native"
    assert layout["image_pixels"][0] >= 1400


def test_office_ingestion_preserves_text_formulas_and_embedded_media(
    tmp_path: Path,
) -> None:
    from PIL import Image
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorksheetImage
    Presentation = pytest.importorskip("pptx").Presentation
    from pptx.util import Inches

    inputs = tmp_path / "inputs"
    inputs.mkdir()
    picture = inputs / "evidence.png"
    Image.new("RGB", (320, 180), "#3c78d8").save(picture)

    docx_path = inputs / "brief.docx"
    doc = Document()
    doc.add_heading("Policy brief", level=1)
    doc.add_paragraph("Target adoption is 42 percent.")
    doc.add_picture(str(picture))
    doc.save(docx_path)
    assert _extract_material(docx_path, inputs / "brief.docx.md")
    docx_text = (inputs / "brief.docx.md").read_text(encoding="utf-8")
    assert "Target adoption is 42 percent" in docx_text
    assert "embedded_media:" in docx_text
    assert list((inputs / "brief.docx.pages").glob("embedded_*.png"))

    pptx_path = inputs / "roadmap.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
    textbox = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(6), Inches(1))
    textbox.text = "Launch timeline: pilot in Q3"
    deck.save(pptx_path)
    assert _extract_material(pptx_path, inputs / "roadmap.pptx.md")
    pptx_text = (inputs / "roadmap.pptx.md").read_text(encoding="utf-8")
    assert "Launch timeline: pilot in Q3" in pptx_text
    assert list((inputs / "roadmap.pptx.pages").glob("embedded_*.png"))

    xlsx_path = inputs / "portfolio.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portfolio"
    sheet.append(["Project", "Budget", "Double"])
    sheet.append(["Aurora", 125, "=B2*2"])
    sheet.add_image(WorksheetImage(str(picture)), "E2")
    workbook.save(xlsx_path)
    assert _extract_material(xlsx_path, inputs / "portfolio.xlsx.md")
    xlsx_text = (inputs / "portfolio.xlsx.md").read_text(encoding="utf-8")
    assert "Project | Budget | Double" in xlsx_text
    assert "Aurora | 125 | =B2*2" in xlsx_text
    assert list((inputs / "portfolio.xlsx.pages").glob("embedded_*.png"))


def test_material_staging_failure_is_structured_and_blocks_planning(
    tmp_path: Path,
) -> None:
    source = tmp_path / "unknown.bin"
    source.write_bytes(b"not a supported attachment")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)]}
    staged = stage_materials(workspace, seed)
    assert staged == ["inputs/01_unknown.bin"]
    failures = seed["_material_ingestion_failures"]
    assert failures[0]["error"] == "unsupported attachment type"
    _material_blocked(str(workspace), failures)
    blocked = (workspace / "research/material.md").read_text(encoding="utf-8")
    assert "status: material_blocked" in blocked
    assert "unknown.bin" in blocked


def test_standalone_image_attachment_routes_directly_to_image(
    tmp_path: Path,
) -> None:
    from PIL import Image

    source = tmp_path / "diagram.png"
    Image.new("RGB", (640, 360), "white").save(source)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)], "query": "用这张图制作一套演示"}
    assert stage_materials(workspace, seed) == ["inputs/01_diagram.png"]
    assert seed["_material_ingestion_failures"] == []
    profile = seed["_attachment_profile"]
    assert profile["material_stage"] == "omitted"
    assert profile["material_agent_required"] is False
    assert profile["visual_asset_paths"] == ["inputs/01_diagram.png"]

    with mock.patch.dict(
        os.environ,
        {"SERPER_API_KEY": "", "IMAGE_API_KEY": "", "OPENAI_API_KEY": ""},
    ):
        runtime = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "_raw_user_query": seed["query"],
                "_attachment_profile": profile,
            },
            seed["_staged_materials"] if "_staged_materials" in seed else ["inputs/01_diagram.png"],
            plan_only=True,
        )
    assert "material" not in runtime["active_roles"]
    assert "image" in runtime["active_roles"]


def test_pdf_requested_visuals_expose_rendered_pages_as_material_assets(
    tmp_path: Path,
) -> None:
    source = tmp_path / "advertisement.pdf"
    document = pymupdf.open()
    page = document.new_page(width=640, height=360)
    page.insert_text((72, 100), "Yoga product hero")
    document.save(source)
    document.close()
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {
        "attachments": [str(source)],
        "query": "Use high-quality images from the advertisement to showcase the product.",
    }
    staged = stage_materials(workspace, seed)
    profile = seed["_attachment_profile"]
    assert "inputs/01_advertisement.pdf" in staged
    assert profile["material_stage"] == "mixed"
    assert profile["material_agent_required"] is True
    assert profile["records"][0]["needs_visual_reuse"] is True
    assert profile["visual_asset_paths"] == [
        "inputs/01_advertisement.pdf.pages/page_001.png"
    ]
    assert (workspace / profile["visual_asset_paths"][0]).is_file()


def test_markdown_attachment_uses_direct_text_handoff_without_material(tmp_path: Path) -> None:
    source = tmp_path / "brief.md"
    source.write_text("# Launch brief\n\nRevenue target: 18%.\n", encoding="utf-8")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)], "query": "把附件做成演示"}
    staged = stage_materials(workspace, seed)
    profile = seed["_attachment_profile"]
    assert staged == ["inputs/01_brief.md"]
    assert profile["material_stage"] == "direct_text"
    assert profile["material_agent_required"] is False
    handoff = (workspace / "research/material.md").read_text(encoding="utf-8")
    assert "Harness generated this deterministic text handoff" in handoff
    assert "Revenue target: 18%." in handoff

    with mock.patch.dict(os.environ, {"SERPER_API_KEY": "search-key"}):
        runtime = detect_runtime_capabilities(
            {
                "model": "test",
                "model_base_url": "http://model",
                "_raw_user_query": seed["query"],
                "_attachment_profile": profile,
                "_evidence_scope_explicit": False,
            },
            staged,
            plan_only=True,
        )
    assert "research" not in runtime["active_roles"]
    assert runtime["workflow"]["research_reason"] == "direct_text_complete_no_unresolved"
    injected = _direct_text_injection(str(workspace), "zh")
    assert 'mode="verbatim"' in injected
    assert "Revenue target: 18%." in injected
    assert "无需分页读取" in injected


@pytest.mark.parametrize(
    ("query", "expected_stage", "visual", "style"),
    [
        ("读取这张截图里的数据并制作演示", "agent_required", False, False),
        ("根据这张图片的风格设计整套演示", "omitted", False, True),
        ("分析这张图的内容，并把原图放在封面", "mixed", True, False),
    ],
)
def test_image_attachment_intent_controls_material_route(
    tmp_path: Path,
    query: str,
    expected_stage: str,
    visual: bool,
    style: bool,
) -> None:
    from PIL import Image

    source = tmp_path / "input.png"
    Image.new("RGB", (640, 360), "#8899aa").save(source)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"attachments": [str(source)], "query": query}
    stage_materials(workspace, seed)
    profile = seed["_attachment_profile"]
    assert profile["material_stage"] == expected_stage
    assert bool(profile["visual_asset_paths"]) is visual
    assert bool(profile["style_reference_paths"]) is style
    assert profile["material_agent_required"] is (expected_stage in {"agent_required", "mixed"})


def test_scanned_pdf_blocks_before_agent_when_main_model_is_text_only(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scan.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(source)
    document.close()
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    seed = {"materials": [str(source)]}
    with mock.patch("core.run_batch.acfg.VISION_BACKEND", "disabled"):
        stage_materials(workspace, seed)
    assert seed["_material_ingestion_failures"]
    assert "text-only" in seed["_material_ingestion_failures"][0]["error"]


def test_attachment_only_is_a_tool_level_research_scope(tmp_path: Path) -> None:
    tool_names = {
        schema["name"]
        for schema in tools.agent_tools(
            "research", evidence_scope="attachment_only"
        )
    }
    assert "web_search" not in tool_names
    assert "web_extract" not in tool_names

    class Orchestrator:
        role = "orchestrator"
        material_required = True
        material_completed = False

    class Research:
        role = "research"

    assert "不得直接读取" in tools.read_file(Orchestrator(), "inputs/report.pdf")
    assert "不直接重读" in tools.read_file(Research(), "inputs/report.pdf.md")


def test_plan_only_acceptance_requires_material_research_and_plans(
    tmp_path: Path,
) -> None:
    (tmp_path / "research").mkdir()
    (tmp_path / "plan").mkdir()
    (tmp_path / "research/material.md").write_text("status: ready\n", encoding="utf-8")
    (tmp_path / "research/knowledge-brief.md").write_text("# Evidence\n", encoding="utf-8")
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    for page in (1, 2):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n", encoding="utf-8"
        )

    class PlanOnlyAgent:
        child_outcomes = {}
        tool_policy_violations = []
        workspace_policy_violations = []

    with mock.patch.dict(os.environ, {"CLEAN_PLAN_ONLY_EVAL": "1"}):
        ok, reason, detail = _accept(PlanOnlyAgent(), str(tmp_path))
    assert ok is True
    assert reason == "plan-only deliverables complete"
    assert detail["n_slide_plans"] == 2


def test_material_blocked_is_a_terminal_structured_acceptance_state(
    tmp_path: Path,
) -> None:
    class BlockedAgent:
        role = "orchestrator"
        skill_language = "zh"
        material_blocked = True
        material_required = True
        material_completed = False
        child_outcomes = {"material": {"ok": False, "status": "material_blocked"}}
        tool_policy_violations = []
        workspace_policy_violations = []

    assert _finish_gap(BlockedAgent()) == ""
    ok, reason, detail = _accept(BlockedAgent(), str(tmp_path))
    assert ok is False
    assert reason == "material_blocked"
    assert detail["material_blocked"] is True


def test_attachment_figure_crop_and_material_catalog_are_supported(tmp_path: Path) -> None:
    from PIL import Image

    page_dir = tmp_path / "inputs/paper.pdf.pages"
    page_dir.mkdir(parents=True)
    source = page_dir / "page_001.png"
    Image.new("RGB", (1600, 2000), "white").save(source)
    source.with_suffix(".json").write_text(
        json.dumps({"page_points": [800, 1000], "text_blocks": []}),
        encoding="utf-8",
    )
    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/image.py"
    )
    crop = subprocess.run(
        [
            sys.executable,
            str(script),
            "crop-material",
            str(tmp_path),
            "--source",
            "inputs/paper.pdf.pages/page_001.png",
            "--output",
            "assets/figure.png",
            "--box",
            "0.10,0.10,0.90,0.35",
        ],
        capture_output=True,
        text=True,
    )
    assert crop.returncode == 0, crop.stderr
    assert '"status": "PASS"' in crop.stdout
    assert (tmp_path / "assets/figure.png").is_file()

    full_page = subprocess.run(
        [
            sys.executable,
            str(script),
            "crop-material",
            str(tmp_path),
            "--source",
            "inputs/paper.pdf.pages/page_001.png",
            "--output",
            "assets/full-page.png",
            "--box",
            "0.00,0.00,1.00,1.00",
        ],
        capture_output=True,
        text=True,
    )
    assert full_page.returncode != 0
    assert full_page.stdout.startswith("status:FAIL\n")
    assert "page facsimile" in full_page.stdout
    assert "Traceback" not in full_page.stdout
    assert "Traceback" not in full_page.stderr

    justified_page = subprocess.run(
        [
            sys.executable,
            str(script),
            "crop-material",
            str(tmp_path),
            "--source",
            "inputs/paper.pdf.pages/page_001.png",
            "--output",
            "assets/justified-page.png",
            "--box",
            "0.00,0.00,1.00,1.00",
            "--facsimile-justification",
            "The original full-page arrangement is the historical evidence being discussed.",
        ],
        capture_output=True,
        text=True,
    )
    assert justified_page.returncode == 0, justified_page.stderr
    registry = json.loads(
        (tmp_path / "_trace/material-figures.json").read_text(encoding="utf-8")
    )
    assert registry["assets"]["assets/justified-page.png"]["facsimile"] is True

    assets = tmp_path / "assets"
    assets.mkdir(exist_ok=True)
    (assets / "catalog.md").write_text(
        """# Asset catalog

## attachment-figure
- slides: 1
- kind: material
- path: assets/figure.png
- source: inputs/paper.pdf#page=1
- purpose: copied from attachment
""",
        encoding="utf-8",
    )
    accepted_catalog = subprocess.run(
        [sys.executable, str(script), "fetch", str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert accepted_catalog.returncode == 0, accepted_catalog.stderr
    assert "status:PASS" in accepted_catalog.stdout


def test_image_terminal_exposes_material_figure(tmp_path: Path) -> None:
    from PIL import Image

    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/image.py"
    )
    page_dir = tmp_path / "inputs/paper.pdf.pages"
    page_dir.mkdir(parents=True)
    Image.new("RGB", (1600, 2000), "white").save(page_dir / "page_001.png")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        render_script = str(script)
        bash_relaxed = False
        bash_timeout = 30

    result = tools.terminal(
        ImageAgent(),
        f"python {script} crop-material . "
        "--source inputs/paper.pdf.pages/page_001.png "
        "--output assets/figure.png --box 0.1,0.1,0.9,0.35",
    )
    assert '"status": "PASS"' in result


def test_image_material_crop_failure_budget_routes_to_code_redraw(tmp_path: Path) -> None:
    from PIL import Image

    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/image.py"
    )
    page_dir = tmp_path / "inputs/paper.pdf.pages"
    page_dir.mkdir(parents=True)
    source = page_dir / "page_005.png"
    Image.new("RGB", (1600, 2000), "white").save(source)
    source.with_suffix(".json").write_text(json.dumps({
        "page_points": [800, 1000],
        "text_blocks": [{
            "bbox_pdf": [80, 100, 720, 700],
            "chars": 1800,
        }],
    }))

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        render_script = str(script)
        bash_relaxed = False
        bash_timeout = 30

    agent = ImageAgent()
    command = (
        f"python {script} crop-material . "
        "--source inputs/paper.pdf.pages/page_005.png "
        "--output assets/figure2.png --box 0.1,0.1,0.9,0.7"
    )
    first = tools.terminal(agent, command)
    second = tools.terminal(agent, command)
    third = tools.terminal(agent, command)
    assert "crop is text-heavy" in first
    assert "material_crop_route_required" in second
    assert "material_crop_budget" in third
    assert "HTML/CSS/SVG/ECharts" in third

    assert "crop_escalation_required" in second

    # Renaming the output must not reset the bounded attachment-page failure pass.
    renamed = command.replace("assets/figure2.png", "assets/figure2-b.png")
    assert "crop is text-heavy" in tools.terminal(agent, renamed)
    renamed_again = command.replace("assets/figure2.png", "assets/figure2-c.png")
    assert "material_source_crop_budget" in tools.terminal(agent, renamed_again)


def test_image_terminal_registers_user_image_and_catalog_accepts_it(tmp_path: Path) -> None:
    from PIL import Image

    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/image.py"
    )
    inputs = tmp_path / "inputs"
    inputs.mkdir(parents=True)
    Image.new("RGB", (800, 500), "#557799").save(inputs / "photo.jpg")

    class ImageAgent:
        role = "image"
        skill_name = "mural-presenter-v0-4"
        ws = str(tmp_path)
        render_script = str(script)
        bash_relaxed = False
        bash_timeout = 30

    result = tools.terminal(
        ImageAgent(),
        f"python {script} register-user . "
        "--source inputs/photo.jpg --output assets/cover-photo.png",
    )
    assert '"status": "PASS"' in result
    assert (tmp_path / "assets/cover-photo.png").is_file()
    registry = json.loads(
        (tmp_path / "_trace/user-images.json").read_text(encoding="utf-8")
    )
    assert "assets/cover-photo.png" in registry["assets"]
    (tmp_path / "assets/catalog.md").write_text(
        """# Asset catalog

## cover-photo
- slides: 1
- kind: user
- path: assets/cover-photo.png
- source: inputs/photo.jpg
- purpose: user-supplied cover image
""",
        encoding="utf-8",
    )
    accepted = subprocess.run(
        [sys.executable, str(script), "fetch", str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert accepted.returncode == 0, accepted.stderr
    assert "status:PASS" in accepted.stdout


def test_v03_portable_font_bundle_is_self_contained(tmp_path: Path) -> None:
    edition = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4"
    shutil.copyfile(edition / "assets/base.css", tmp_path / "base.css")
    slides = tmp_path / "slides"
    slides.mkdir()
    (slides / "slide_01.html").write_text(
        '<section class="slide"><h1 class="type-heavy">RAVE 渲染一致</h1></section>',
        encoding="utf-8",
    )

    result = subprocess.run(
        [sys.executable, str(edition / "scripts/_internal/font_bundle.py"), str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stderr
    css = (tmp_path / "base.css").read_text(encoding="utf-8")
    assert "DECK_FONT_BUNDLE_START" in css
    assert '@font-face {' in css
    manifest = json.loads(
        (tmp_path / "assets/fonts/manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["faces"]
    assert all((tmp_path / face["path"]).is_file() for face in manifest["faces"])
    css_before = (tmp_path / "base.css").read_bytes()
    manifest_before = (tmp_path / "assets/fonts/manifest.json").read_bytes()
    rerun = subprocess.run(
        [sys.executable, str(edition / "scripts/_internal/font_bundle.py"), str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert rerun.returncode == 0, rerun.stderr
    assert (tmp_path / "base.css").read_bytes() == css_before
    assert (tmp_path / "assets/fonts/manifest.json").read_bytes() == manifest_before
    (slides / "slide_01.html").write_text(
        '<section class="slide"><h1 class="type-heavy">RAVE 渲染一致，新增字形</h1></section>',
        encoding="utf-8",
    )
    changed = subprocess.run(
        [sys.executable, str(edition / "scripts/_internal/font_bundle.py"), str(tmp_path)],
        capture_output=True,
        text=True,
    )
    assert changed.returncode == 0, changed.stderr
    changed_manifest = json.loads(
        (tmp_path / "assets/fonts/manifest.json").read_text(encoding="utf-8")
    )
    assert changed_manifest["deck_id"] != manifest["deck_id"]


@pytest.mark.parametrize(
    ("script_name", "allowed", "forbidden"),
    [
        ("orchestrator.py", {"validate-plans", "scaffold", "sync-speech", "finalize", "audit"}, {"render", "fetch"}),
        ("image.py", {"register-user", "crop-material", "fetch", "finalize", "inspect", "remove-checkerboard"}, {"audit", "render"}),
        ("slide.py", {"render", "render-group"}, {"finalize", "audit", "fetch"}),
        ("review.py", {"sync-speech", "finalize"}, {"render", "audit", "fetch"}),
    ],
)
def test_role_entry_points_expose_only_their_public_actions(
    script_name: str,
    allowed: set[str],
    forbidden: set[str],
) -> None:
    scripts = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts"
    completed = subprocess.run(
        [sys.executable, str(scripts / script_name), "--help"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    for action in allowed:
        assert action in completed.stdout
    for action in forbidden:
        assert action not in completed.stdout


def test_v03_has_no_public_deck_entry_and_internal_is_not_model_readable(
    tmp_path: Path,
) -> None:
    scripts = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts"
    assert not (scripts / "deck.py").exists()

    agent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
    )
    blocked = tools.read_file(
        agent,
        "skills/mural-presenter-v0-4/scripts/_internal/deck_core.py",
    )
    assert "只执行能力面" in blocked
    assert "直接按该错误回执修改当前 slide HTML" in blocked


def test_canvas_missing_script_error_is_directly_actionable(tmp_path: Path) -> None:
    path = tmp_path / "slide_03.html"
    text = '<section id="slide-03" data-slide="03"></section>'
    plan = {"primary_visual_medium": "canvas-diagram"}
    with pytest.raises(ValueError) as caught:
        deck_core._validate_fragment(path, 3, text, plan, {})
    message = str(caught.value)
    assert "contains no script" in message
    assert "<canvas></canvas>" in message
    assert "edit the current slide HTML directly" in message
    assert "do not inspect Skill scripts" in message


def test_scaffold_root_class_error_names_exact_missing_tokens(tmp_path: Path) -> None:
    path = tmp_path / "slide_11.html"
    plan = {
        "page_type": "comparison",
        "page_family": "comparison",
        "canvas_variant": "base",
        "composition": "comparison",
        "primary_visual_medium": "editorial-typography",
    }
    text = (
        '<section class="slide" id="slide-11" data-slide="11" '
        'data-page-type="comparison" data-page-family="comparison" '
        'data-frame="content" data-canvas-variant="base" '
        'data-composition="comparison"></section>'
    )
    with pytest.raises(ValueError) as caught:
        deck_core._validate_fragment(path, 11, text, plan, {})
    message = str(caught.value)
    assert "restore the missing class token(s)" in message
    assert "frame-content" in message
    assert "page-comparison" in message
    assert "family-comparison" in message
    assert "composition-comparison" in message
    assert "canvas-variant--base" in message


def test_slide_plan_heading_error_requires_exact_canonical_line(tmp_path: Path) -> None:
    path = tmp_path / "slide_01.md"
    path.write_text("# slide_01 — Cover\n", encoding="utf-8")
    with pytest.raises(ValueError) as caught:
        deck_core._parse_slide(path, 1)
    message = str(caught.value)
    assert "exactly `# slide_01`" in message
    assert "nothing after the page number" in message
    assert "remove any dash/title suffix" in message


def test_deck_title_metadata_error_names_exact_bullet_syntax(tmp_path: Path) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        "# Deck Plan\n\ntitle: A bare field\n\n## Resolved deck brief\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError) as caught:
        deck_core._parse_deck(tmp_path)
    message = str(caught.value)
    assert "exact bullet `- title: <deck title>`" in message
    assert "bare `title:` line" in message


def test_v03_role_script_policy_blocks_cross_role_and_internal_execution(
    tmp_path: Path,
) -> None:
    scripts = REPO / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts"
    image_script = scripts / "image.py"
    agent = SimpleNamespace(
        role="image",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        render_script=str(image_script),
        bash_relaxed=False,
        bash_timeout=30,
    )
    cross_role = tools.terminal(
        agent,
        f"python {scripts / 'orchestrator.py'} audit .",
    )
    assert "只执行能力面" in cross_role
    wrong_action = tools.terminal(agent, f"python {image_script} audit .")
    assert "不允许动作" in wrong_action
    internal = tools.terminal(
        agent,
        f"python {scripts / '_internal/bootstrap.py'} prepare .",
    )
    assert "只执行能力面" in internal
    internal_read = tools.terminal(
        agent,
        f"cat {scripts / '_internal/deck_core.py'}",
    )
    assert "只执行能力面" in internal_read


def test_material_and_research_receive_no_terminal_schema() -> None:
    for role in ("material", "research"):
        names = {
            schema["name"]
            for schema in tools.agent_tools(
                role,
                render_script="",
                skill_name="mural-presenter-v0-4",
            )
        }
        assert "terminal" not in names


def test_plan_validator_warns_when_closing_is_a_multi_item_content_page(
    tmp_path: Path,
) -> None:
    plan = tmp_path / "plan"
    plan.mkdir()
    (plan / "deck.md").write_text(
        """# Deck Plan
- title: Closing warning fixture

## Resolved deck brief
- language: en
- page_count: 2
- audience: reviewers
- image_mode: code only
- bitmap_strategy: not-beneficial
- bitmap_rationale: deterministic typography fixture
- ownership_topology: single
- ownership_rationale: independent content
- rationale: test warning

## Audience and objective
Verify closing semantics.

## Narrative and page map
One content page and one closing.

## Visual storyboard
Typography.

## Visual contract
Test.

## Special pages
One closing.

## Theme Tokens
```css
:root { --content-canvas: #eeeeee; }
```
""",
        encoding="utf-8",
    )
    common = """
## Narrative
Advance the fixture.

## On-screen copy (exact)
- title: {title}

## Composition blueprint
- focal: title

## Semantic visual need
Use strong typography.

## Speech beat
Explain the page.
"""
    (plan / "slide_01.md").write_text(
        "# slide_01\n- role: content\n- page_type: content\n"
        "- page_family: editorial\n- production_group: page-01\n"
        "- composition: editorial\n- needs_bitmap: false\n"
        "- primary_visual_medium: editorial-typography\n"
        "\n## Evidence\n- Established conclusion.\n"
        + common.format(title="Established conclusion"),
        encoding="utf-8",
    )
    (plan / "slide_02.md").write_text(
        "# slide_02\n- role: closing\n- page_type: closing\n"
        "- page_family: editorial\n- production_group: page-02\n"
        "- special_layout: centered\n- needs_bitmap: false\n"
        "- primary_visual_medium: editorial-typography\n"
        "\n## Evidence\n1. First action\n2. Second action\n3. Third action\n"
        + common.format(title="Three actions"),
        encoding="utf-8",
    )
    script = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/scripts/orchestrator.py"
    )
    result = subprocess.run(
        [sys.executable, str(script), "validate-plans", str(tmp_path), "--expected", "2"],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout
    assert "closing slide_02 carries 3 separate evidence/list items" in result.stdout


# ---------------------------------------------------------------------------
# Artifact-completion recovery: bounded retry for max_turns slide failures
# ---------------------------------------------------------------------------


def test_artifact_completion_retry_allowed_first_max_turns_failure(
    tmp_path: Path,
) -> None:
    """First max_turns failure with no artifacts allows exactly one retry."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_11.md").write_text(
        "# P11\n- needs_bitmap: false\n- primary_visual_medium: code-visual\n",
        encoding="utf-8",
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_11.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    parent = Agent(
        "retry-test",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.image_completed = False
    parent.child_outcomes = {"slide_11": {
        "role": "slide",
        "ok": False,
        "status": "max_turns",
        "exit_reason": "max_turns",
        "trace_label": "slide_11",
        "completed_pages": [],
        "repair_issue": None,
        "attempt": 1,
    }}
    # No PNG exists → artifact_completion retry is allowed
    assert not (tmp_path / "renders" / "slide_11.png").exists()
    completed = {
        "label": "slide_11",
        "trace_label": "slide_11_retry2",
        "role": "slide",
        "ok": True,
        "status": "ready",
        "exit_reason": "text_response",
        "completed_pages": [11],
        "incomplete_pages": [],
        "renders": 1,
        "views": 1,
        "summary": "ready",
    }
    with mock.patch("core.agent_loop._run_child", return_value=completed) as run_child:
        result = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide", "label": "slide_11", "pages": [11], "task": "build"
        }]}))
    assert result["executed"] == 1, f"Expected 1 execution, got: {result}"
    assert result["failed"] == 0
    assert run_child.call_count == 1
    spec_passed = run_child.call_args.args[2]
    assert spec_passed.get("_recovery_kind") == "artifact_completion"


def test_artifact_completion_second_failure_is_terminal(
    tmp_path: Path,
) -> None:
    """After attempt>=2 (second failure), the page becomes terminal and cannot be retried."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_11.md").write_text(
        "# P11\n- needs_bitmap: false\n- primary_visual_medium: code-visual\n",
        encoding="utf-8",
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_11.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    parent = Agent(
        "terminal-test",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.image_completed = False
    parent.child_outcomes = {"slide_11": {
        "role": "slide",
        "ok": False,
        "status": "max_turns",
        "exit_reason": "max_turns",
        "trace_label": "slide_11_retry2",
        "completed_pages": [],
        "repair_issue": None,
        "attempt": 2,
    }}
    # attempt=2 → terminal, delegation must be rejected
    assert _page_has_terminal_slide_failure(parent, 11)
    result = _delegate(parent, {"tasks": [{
        "role": "slide", "label": "slide_11", "pages": [11], "task": "build"
    }]})
    assert "已结束" in result
    assert "max_turns" in result


def test_finalize_does_not_reactivate_pending_for_terminal_slide(
    tmp_path: Path,
) -> None:
    """End-to-end: tools.terminal(finalize) returns missing_slide_artifact for
    terminal pages, does NOT set pending_slide_delegation_pages, sets
    terminal_slide_failures; subsequent _call hides all tools and injects
    the terminal exit instruction."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_11.md").write_text(
        "# P11\n- needs_bitmap: false\n- primary_visual_medium: code-visual\n",
        encoding="utf-8",
    )
    (tmp_path / "plan/slide_01.md").write_text(
        "# P1\n- needs_bitmap: false\n- primary_visual_medium: editorial-typography\n",
        encoding="utf-8",
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "slides/slide_01.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    (tmp_path / "slides/slide_11.html").write_text(
        "<section></section>\n", encoding="utf-8"
    )
    # P01 has a PNG (completed), P11 does not (terminal failure)
    (tmp_path / "renders/slide_01.png").write_bytes(b"\x89PNG\r\n")
    parent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        child_outcomes={"slide_11": {
            "role": "slide",
            "ok": False,
            "status": "max_turns",
            "exit_reason": "max_turns",
            "trace_label": "slide_11_retry2",
            "completed_pages": [],
            "repair_issue": None,
            "attempt": 2,
        }},
        pending_slide_delegation_pages=(),
        terminal_slide_failures=(),
        requested_slide_count=2,
        image_completed=False,
        image_failed_pages=(),
        bash_relaxed=False,
        render_script="",
        role_script="skills/mural-presenter-v0-4/scripts/orchestrator.py",
        log=lambda *_args, **_kwargs: None,
    )
    # Step 1: call tools.terminal with a real finalize command
    result = tools.terminal(
        parent,
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py finalize . --expected 2",
    )
    # Assertions: terminal path fires, pending NOT set, terminal_slide_failures set
    assert "missing_slide_artifact" in result, f"Expected missing_slide_artifact, got: {result}"
    assert getattr(parent, "pending_slide_delegation_pages", ()) == ()
    assert getattr(parent, "terminal_slide_failures", ()) == (11,)

    # Step 2: verify the next model call hides tools and injects terminal message
    captured = {}
    agent_for_call = SimpleNamespace(
        ws=str(tmp_path),
        role="orchestrator",
        model="pptagent",
        system="system",
        tool_schemas=[
            {"name": "read_file", "input_schema": {"type": "object"}},
            {"name": "delegate_task", "input_schema": {"type": "object"}},
            {"name": "terminal", "input_schema": {"type": "object"}},
        ],
        child_outcomes=parent.child_outcomes,
        pending_slide_delegation_pages=(),
        terminal_slide_failures=(11,),
        _patch_recovery_path="",
        image_search_closed=False,
        _slide_inputs_read_closed=False,
        review_revision_rounds=0,
        review_completed=False,
        finalize_succeeded=False,
        final_view_after_review=False,
        quality_status="",
        thinking=False,
        effort="high",
        first_response_timeout_s=30,
        active_response_timeout_s=30,
        deadline_monotonic=time.monotonic() + 60,
        nova_raw=None,
        log=lambda *_args, **_kwargs: None,
        token_limit=lambda: 1000,
    )

    def fake_call(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(content=[])

    with mock.patch("core.agent_loop.model_call.call_with_tools", side_effect=fake_call):
        _call(agent_for_call, [{"role": "user", "content": "continue"}])
    assert captured["tools"] == [], "Tools must be empty when terminal_slide_failures is set"
    assert "missing_slide_artifact" in captured["messages"][-1]["content"]
    assert "P11" in captured["messages"][-1]["content"]


# ---------------------------------------------------------------------------
# Grouped slide: artifact-completion retry and terminal for slide_group_*
# ---------------------------------------------------------------------------


def test_grouped_artifact_completion_retry_allowed(tmp_path: Path) -> None:
    """Grouped slide (bookends [1,12]): first max_turns with missing PNGs allows retry."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    for page in (1, 12):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# P{page}\n- needs_bitmap: false\n- primary_visual_medium: editorial-typography\n"
            f"- production_group: bookends\n",
            encoding="utf-8",
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            "<section></section>\n", encoding="utf-8"
        )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    # P1 got a render (completed), P12 did not
    (tmp_path / "renders/slide_01.png").write_bytes(b"\x89PNG\r\n")
    parent = Agent(
        "grouped-retry-test",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.image_completed = False
    parent.child_outcomes = {"slide_group_bookends": {
        "role": "slide",
        "ok": False,
        "status": "max_turns",
        "exit_reason": "max_turns",
        "trace_label": "slide_group_bookends",
        "completed_pages": [1],
        "incomplete_pages": [12],
        "repair_issue": None,
        "attempt": 1,
    }}
    # P12 has no PNG → retry is allowed despite P1 being completed
    assert not (tmp_path / "renders" / "slide_12.png").exists()
    assert not _page_has_terminal_slide_failure(parent, 12)
    assert not _page_has_terminal_slide_failure(parent, 1)
    completed = {
        "label": "slide_group_bookends",
        "trace_label": "slide_group_bookends_retry2",
        "role": "slide",
        "ok": True,
        "status": "ready",
        "exit_reason": "text_response",
        "completed_pages": [1, 12],
        "incomplete_pages": [],
        "renders": 2,
        "views": 2,
        "summary": "ready",
    }
    with mock.patch("core.agent_loop._run_child", return_value=completed) as run_child:
        result = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide", "label": "slide_group_bookends",
            "group_id": "bookends", "pages": [1, 12], "task": "build bookends"
        }]}))
    assert result["executed"] == 1, f"Expected 1 execution, got: {result}"
    assert result["failed"] == 0
    assert run_child.call_count == 1
    spec_passed = run_child.call_args.args[2]
    assert spec_passed.get("_recovery_kind") == "artifact_completion"


def test_grouped_second_failure_is_terminal(tmp_path: Path) -> None:
    """Grouped slide (bookends [1,12]): attempt=2 with missing P12 PNG → terminal."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    for page in (1, 12):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# P{page}\n- needs_bitmap: false\n- primary_visual_medium: editorial-typography\n"
            f"- production_group: bookends\n",
            encoding="utf-8",
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            "<section></section>\n", encoding="utf-8"
        )
    # P1 has PNG (from first attempt), P12 still missing
    (tmp_path / "renders/slide_01.png").write_bytes(b"\x89PNG\r\n")
    parent = Agent(
        "grouped-terminal-test",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.image_completed = False
    parent.child_outcomes = {"slide_group_bookends": {
        "role": "slide",
        "ok": False,
        "status": "max_turns",
        "exit_reason": "max_turns",
        "trace_label": "slide_group_bookends_retry2",
        "completed_pages": [1],
        "incomplete_pages": [12],
        "repair_issue": None,
        "attempt": 2,
    }}
    # Both pages covered by the group are terminal
    assert _page_has_terminal_slide_failure(parent, 12)
    assert _page_has_terminal_slide_failure(parent, 1)
    # Delegation must be rejected
    result = _delegate(parent, {"tasks": [{
        "role": "slide", "label": "slide_group_bookends",
        "group_id": "bookends", "pages": [1, 12], "task": "build bookends"
    }]})
    assert "已结束" in result
    assert "max_turns" in result


def test_grouped_finalize_terminal_via_tools(tmp_path: Path) -> None:
    """End-to-end: finalize recognizes grouped terminal failure, excludes from pending."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "slides").mkdir()
    (tmp_path / "renders").mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    for page in (1, 5, 12):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# P{page}\n- needs_bitmap: false\n- primary_visual_medium: editorial-typography\n",
            encoding="utf-8",
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            "<section></section>\n", encoding="utf-8"
        )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    # P5 completed (has PNG), P1 completed (from bookends first attempt), P12 missing
    (tmp_path / "renders/slide_01.png").write_bytes(b"\x89PNG\r\n")
    (tmp_path / "renders/slide_05.png").write_bytes(b"\x89PNG\r\n")
    parent = SimpleNamespace(
        role="orchestrator",
        skill_name="mural-presenter-v0-4",
        ws=str(tmp_path),
        child_outcomes={"slide_group_bookends": {
            "role": "slide",
            "ok": False,
            "status": "max_turns",
            "exit_reason": "max_turns",
            "trace_label": "slide_group_bookends_retry2",
            "completed_pages": [1],
            "incomplete_pages": [12],
            "repair_issue": None,
            "attempt": 2,
        }},
        pending_slide_delegation_pages=(),
        terminal_slide_failures=(),
        requested_slide_count=3,
        image_completed=False,
        image_failed_pages=(),
        bash_relaxed=False,
        render_script="",
        role_script="skills/mural-presenter-v0-4/scripts/orchestrator.py",
        log=lambda *_args, **_kwargs: None,
    )
    result = tools.terminal(
        parent,
        "python skills/mural-presenter-v0-4/scripts/orchestrator.py finalize . --expected 3",
    )
    # P12 is terminal (grouped), only P12 is missing among non-terminal pages = none delegatable
    assert "missing_slide_artifact" in result, f"Expected missing_slide_artifact, got: {result}"
    assert getattr(parent, "pending_slide_delegation_pages", ()) == ()
    assert 12 in getattr(parent, "terminal_slide_failures", ())


def test_grouped_no_retry_when_all_pngs_exist(tmp_path: Path) -> None:
    """If all responsibility pages already have PNGs, no artifact-completion retry."""
    (tmp_path / "renders").mkdir()
    (tmp_path / "renders/slide_01.png").write_bytes(b"\x89PNG\r\n")
    (tmp_path / "renders/slide_12.png").write_bytes(b"\x89PNG\r\n")
    previous = {
        "role": "slide",
        "ok": False,
        "status": "max_turns",
        "exit_reason": "max_turns",
        "trace_label": "slide_group_bookends",
        "completed_pages": [1, 12],
        "incomplete_pages": [],
        "repair_issue": None,
        "attempt": 1,
    }
    assert not _artifact_completion_retry_allowed(
        previous, ws=str(tmp_path), pages=(1, 12)
    ), "Must NOT allow retry when all PNGs already exist"


# ---------------------------------------------------------------------------
# Regression: deck #149 — material_blocked status classification
# ---------------------------------------------------------------------------

def _make_material_text(
    *,
    status: str = "material_blocked",
    blocking: str | None = None,
    blocking_reason: str | None = None,
    evidence_items: int = 0,
    has_source_refs: bool = False,
    extra_bulk: int = 0,
) -> str:
    """Build a synthetic material.md for testing status classification."""
    lines = [f"status: {status}", "evidence_scope: partial"]
    if blocking is not None:
        lines.append(f"blocking: {blocking}")
    if blocking_reason is not None:
        lines.append(f"blocking_reason: {blocking_reason}")
    lines.append("")
    lines.append("## 1. 附件清单")
    if has_source_refs:
        lines.append("- 附件 A: paper.pdf (page 1-20)")
    for i in range(evidence_items):
        lines.append(f"- Evidence #{i}: data point from page {i + 1}")
    if extra_bulk:
        lines.append("\n" + "x" * extra_bulk)
    return "\n".join(lines)


class TestMaterialBlockedDecision:
    """Unit tests for _material_blocked_decision state classification."""

    def test_not_blocked_status_returns_empty(self) -> None:
        text = "status: ready\nevidence_scope: full\n\n- item 1\n"
        assert _material_blocked_decision(text) == ""

    def test_explicit_blocking_yes_with_valid_reason_returns_blocked(self) -> None:
        text = _make_material_text(
            blocking="yes",
            blocking_reason="unreadable_required_page",
        )
        assert _material_blocked_decision(text) == "blocked"

    def test_explicit_blocking_yes_no_usable_evidence_returns_blocked(self) -> None:
        text = _make_material_text(
            blocking="yes",
            blocking_reason="no_usable_evidence",
        )
        assert _material_blocked_decision(text) == "blocked"

    def test_explicit_blocking_yes_unsupported_format_returns_blocked(self) -> None:
        text = _make_material_text(
            blocking="yes",
            blocking_reason="unsupported_format",
        )
        assert _material_blocked_decision(text) == "blocked"

    def test_blocking_yes_with_invalid_reason_and_evidence_returns_ready(self) -> None:
        """blocking:yes but with a non-standard reason and real evidence = ready."""
        text = _make_material_text(
            blocking="yes",
            blocking_reason="some_concepts_missing",
            evidence_items=20,
            has_source_refs=True,
            extra_bulk=2000,
        )
        assert _material_blocked_decision(text) == "ready_with_gaps"

    def test_no_blocking_field_with_substantial_evidence_returns_ready(self) -> None:
        """Legacy agent: wrote material_blocked without blocking: yes, but has content."""
        text = _make_material_text(
            evidence_items=25,
            has_source_refs=True,
            extra_bulk=1500,
        )
        assert _material_blocked_decision(text) == "ready_with_gaps"

    def test_no_blocking_field_without_evidence_returns_blocked(self) -> None:
        """Legacy agent: wrote material_blocked, file is a stub = truly blocked."""
        text = _make_material_text(evidence_items=2, extra_bulk=0)
        assert _material_blocked_decision(text) == "blocked"

    def test_short_file_always_blocked(self) -> None:
        """A file under 2 KB cannot be a usable handoff."""
        text = "status: material_blocked\n\n- one item\n- two items\n"
        assert len(text) < 2048
        assert _material_blocked_decision(text) == "blocked"

    def test_large_unstructured_file_returns_blocked(self) -> None:
        """Large file without structured items or source refs = blocked."""
        text = (
            "status: material_blocked\n\n"
            + ("This is plain prose without structure. " * 200)
        )
        assert len(text) >= 2048
        assert _material_blocked_decision(text) == "blocked"


class TestMaterialHasIngestionEvidence:
    """Unit tests for _material_has_ingestion_evidence heuristic."""

    def test_short_returns_false(self) -> None:
        assert not _material_has_ingestion_evidence("short text\n" * 10)

    def test_structured_with_refs_returns_true(self) -> None:
        text = _make_material_text(
            evidence_items=20,
            has_source_refs=True,
            extra_bulk=1500,
        )
        assert _material_has_ingestion_evidence(text)

    def test_structured_without_refs_returns_false(self) -> None:
        """Evidence items without any source references = suspicious."""
        header = "status: material_blocked\n\n"
        items = "\n".join(f"- Generic point number {i}" for i in range(30))
        text = header + items + "\n" + "x" * 2000
        assert len(text) >= 2048
        assert not _material_has_ingestion_evidence(text)

    def test_real_world_deck149_pattern(self) -> None:
        """Simulate deck #149: 22 KB file with PDF evidence and page refs."""
        header = (
            "status: material_blocked\n"
            "evidence_scope: partial\n"
            "unresolved_items:\n"
            "  - concept: adversarial training details\n"
            "    reason: not found in attachments\n\n"
        )
        sections = (
            "## 1. 附件清单\n"
            "- 附件 1: paper.pdf (page 1-12)\n"
            "- 附件 2: supplement.pdf (page 1-8)\n\n"
            "## 2. 主要发现\n"
        )
        evidence = "\n".join(
            f"- Finding {i}: metric value from Table {i % 4 + 1} (page {i + 2})"
            for i in range(40)
        )
        figures = "\n\n## 3. 视觉交接\n" + "\n".join(
            f"- Figure {i}: inputs/paper.pdf.pages/page_{i + 1:03d}.png 候选框 0.1,0.2,0.8,0.6"
            for i in range(5)
        )
        text = header + sections + evidence + figures
        assert len(text) > 2048
        assert _material_has_ingestion_evidence(text)
        assert _material_blocked_decision(text) == "ready_with_gaps"


# ---------------------------------------------------------------------------
# Regression: deck #149 — parent status propagation for material partial
# ---------------------------------------------------------------------------

class TestMaterialParentPropagation:
    """Verify that material partial/blocked states propagate correctly to parent."""

    def test_blocked_result_sets_parent_material_blocked(self) -> None:
        """When child result has status=material_blocked, parent.material_blocked=True."""

        class BlockedAgent:
            role = "orchestrator"
            skill_language = "zh"
            material_blocked = True
            material_required = True
            material_completed = False
            child_outcomes = {"material": {"ok": False, "status": "material_blocked"}}
            tool_policy_violations = []
            workspace_policy_violations = []

        assert _finish_gap(BlockedAgent()) == ""
        ok, reason, detail = _accept(BlockedAgent(), "/tmp/nonexist")
        assert ok is False
        assert reason == "material_blocked"
        assert detail["material_blocked"] is True

    def test_partial_result_does_not_set_parent_material_blocked(self) -> None:
        """When material is partial (ok=True, status=ready), orchestrator proceeds."""

        class PartialAgent:
            role = "orchestrator"
            skill_language = "zh"
            material_blocked = False
            material_required = True
            material_completed = True
            research_required = True
            research_completed = False
            child_outcomes = {"material": {"ok": True, "status": "ready"}}
            tool_policy_violations = []
            workspace_policy_violations = []

        gap = _finish_gap(PartialAgent())
        assert "Material" not in gap
        assert not PartialAgent.material_blocked

    def test_decision_classifies_evidence_file_as_ready_with_gaps(
        self, tmp_path: Path
    ) -> None:
        """A blocked-labelled evidence file is classified as ready_with_gaps."""
        (tmp_path / "research").mkdir()
        material_text = _make_material_text(
            evidence_items=30,
            has_source_refs=True,
            extra_bulk=2000,
        )
        (tmp_path / "research/material.md").write_text(
            material_text, encoding="utf-8"
        )

        decision = _material_blocked_decision(material_text)
        assert decision == "ready_with_gaps"

    def test_decision_classifies_explicit_blocking_file_as_blocked(
        self, tmp_path: Path
    ) -> None:
        """A valid blocking contract is classified as a true terminal block."""
        (tmp_path / "research").mkdir()
        material_text = _make_material_text(
            blocking="yes",
            blocking_reason="unreadable_required_page",
            evidence_items=5,
            has_source_refs=True,
            extra_bulk=2000,
        )
        (tmp_path / "research/material.md").write_text(
            material_text, encoding="utf-8"
        )

        decision = _material_blocked_decision(material_text)
        assert decision == "blocked"


class TestMaterialInSessionCorrection:
    """Recoverable status mistakes are first returned to the same Material Agent."""

    @staticmethod
    def _agent(tmp_path: Path) -> SimpleNamespace:
        return SimpleNamespace(role="material", ws=str(tmp_path))

    def test_readable_evidence_requests_ready_with_gaps_rewrite(
        self, tmp_path: Path
    ) -> None:
        (tmp_path / "research").mkdir()
        text = _make_material_text(
            evidence_items=30,
            has_source_refs=True,
            extra_bulk=2000,
        )
        (tmp_path / "research/material.md").write_text(text, encoding="utf-8")
        gap = _material_handoff_gap(self._agent(tmp_path))
        assert "status 改为 ready" in gap
        assert "unresolved_items" in gap

    def test_explicit_true_block_does_not_request_rewrite(
        self, tmp_path: Path
    ) -> None:
        (tmp_path / "research").mkdir()
        text = _make_material_text(
            blocking="yes",
            blocking_reason="unreadable_required_page",
            evidence_items=30,
            has_source_refs=True,
            extra_bulk=2000,
        )
        (tmp_path / "research/material.md").write_text(text, encoding="utf-8")
        assert _material_handoff_gap(self._agent(tmp_path)) == ""

    def test_ready_material_does_not_request_rewrite(self, tmp_path: Path) -> None:
        (tmp_path / "research").mkdir()
        text = _make_material_text(
            evidence_items=30,
            has_source_refs=True,
            extra_bulk=2000,
        ).replace("status: material_blocked", "status: ready", 1)
        (tmp_path / "research/material.md").write_text(text, encoding="utf-8")
        assert _material_handoff_gap(self._agent(tmp_path)) == ""


# ---------------------------------------------------------------------------
# Regression: deck #149 — typography_flags must not force required_review_pages
# ---------------------------------------------------------------------------

def test_typography_flags_do_not_force_required_review_pages(tmp_path: Path) -> None:
    """Pages with only typography_flags should not be in mandatory full-res set."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/slide_07.md").write_text(
        "- page_type: content\n- needs_bitmap: false\n",
        encoding="utf-8",
    )
    (tmp_path / "renders").mkdir()
    manifest = {
        "pages": [],
        "typography_flags": {"7": "Font X not available, fallback used"},
        "media_mismatches": {},
        "placeholder_flags": {},
        "layout_defects": {},
        "special_page_geometry": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )

    parent = SimpleNamespace(ws=str(tmp_path))
    pages = _mandatory_fullres_review_pages(parent)
    assert 7 not in pages, (
        "typography_flags alone must not force page into required_review_pages"
    )


def test_media_mismatches_still_force_required_review_pages(tmp_path: Path) -> None:
    """Pages with media_mismatches should still be in mandatory full-res set."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/slide_05.md").write_text(
        "- page_type: content\n- needs_bitmap: false\n",
        encoding="utf-8",
    )
    (tmp_path / "renders").mkdir()
    manifest = {
        "pages": [],
        "typography_flags": {},
        "media_mismatches": {"5": "Image placeholder still visible"},
        "placeholder_flags": {},
        "layout_defects": {},
        "special_page_geometry": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )

    parent = SimpleNamespace(ws=str(tmp_path))
    pages = _mandatory_fullres_review_pages(parent)
    assert 5 in pages, "media_mismatches must still force page into required set"


# ---------------------------------------------------------------------------
# Regression: Figure fidelity — markArea syntax validation
# ---------------------------------------------------------------------------

def test_echarts_markarea_invalid_array_syntax_rejected(tmp_path: Path) -> None:
    """markArea with xAxis:[start,end] array form is detected and rejected."""
    plan = {"primary_visual_medium": "echarts"}
    deck = {}
    invalid_html = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        "<script>"
        "var chart = echarts.init(document.getElementById('chart'));"
        "chart.setOption({"
        "  xAxis: {type:'category', data:['V1','V2','V3','V4','V5']},"
        "  yAxis: {type:'value'},"
        "  series: [{type:'line', data:[1,2,3,4,5],"
        "    markArea: {data: [{xAxis: ['V1','V4']}]}"
        "  }]"
        "});"
        "</script>"
    )
    with pytest.raises(ValueError, match="markArea.*endpoint-pair"):
        deck_core._validate_fragment(
            tmp_path / "slide_16.html", 16, invalid_html, plan, deck
        )


def test_echarts_markarea_valid_endpoint_pair_accepted(tmp_path: Path) -> None:
    """markArea with correct [[{xAxis:'start'},{xAxis:'end'}]] passes validation."""
    plan = {"primary_visual_medium": "echarts"}
    deck = {}
    valid_html = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        "<script>"
        "var chart = echarts.init(document.getElementById('chart'));"
        "chart.setOption({"
        "  xAxis: {type:'category', data:['V1','V2','V3','V4','V5']},"
        "  yAxis: {type:'value'},"
        "  series: [{type:'line', data:[1,2,3,4,5],"
        "    markArea: {data: [[{xAxis:'V1'},{xAxis:'V4'}]]}"
        "  }]"
        "});"
        "</script>"
    )
    # Should not raise markArea error; will raise "exactly one section fragment"
    # because there's no data-slide wrapper, which proves markArea check passed.
    with pytest.raises(ValueError, match="exactly one section fragment"):
        deck_core._validate_fragment(
            tmp_path / "slide_16.html", 16, valid_html, plan, deck
        )


def test_echarts_markarea_yaxis_array_also_rejected(tmp_path: Path) -> None:
    """markArea with yAxis:[...] array form is also invalid."""
    plan = {"primary_visual_medium": "echarts"}
    deck = {}
    invalid_html = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        "<script>"
        "var chart = echarts.init(document.getElementById('chart'));"
        "chart.setOption({"
        "  xAxis: {type:'value'},"
        "  yAxis: {type:'value'},"
        "  series: [{type:'scatter', data:[[1,2]],"
        "    markArea: {data: [{yAxis: [10, 50]}]}"
        "  }]"
        "});"
        "</script>"
    )
    with pytest.raises(ValueError, match="markArea.*endpoint-pair"):
        deck_core._validate_fragment(
            tmp_path / "slide_16.html", 16, invalid_html, plan, deck
        )


def test_echarts_markarea_allows_top_level_multi_axis_array(tmp_path: Path) -> None:
    """A legal xAxis:[{...}] declaration must not resemble markArea value arrays."""
    plan = {"primary_visual_medium": "echarts"}
    deck = {}
    valid_html = (
        '<script src="assets/vendor/echarts.min.js"></script>'
        "<script>"
        "var chart = echarts.init(document.getElementById('chart'));"
        "chart.setOption({"
        "xAxis:[{type:'category',data:['A','B']}],"
        "yAxis:{type:'value'},"
        "series:[{type:'line',data:[1,2],"
        "markArea:{data:[[{xAxis:'A'},{xAxis:'B'}]]}}]"
        "});"
        "</script>"
    )
    with pytest.raises(ValueError, match="exactly one section fragment"):
        deck_core._validate_fragment(
            tmp_path / "slide_16.html", 16, valid_html, plan, deck
        )


def test_deck151_partial_image_handoff_releases_reclassified_failed_page(
    tmp_path: Path,
) -> None:
    """A bounded Image partial handoff must not deadlock unstarted Slides.

    Deck #151 durably finalized four assets and marked two source routes failed,
    but the historical runtime put every missing page behind a delegate-only
    gate.  The Orchestrator could therefore neither reclassify the failed pages
    nor run them, and repeatedly redispatched the singleton Image role.  Model
    the same control-plane sequence with one ready bitmap page, one failed
    bitmap page, and one code-only page.
    """
    plan = tmp_path / "plan"
    slides = tmp_path / "slides"
    renders = tmp_path / "renders"
    plan.mkdir()
    slides.mkdir()
    renders.mkdir()
    (plan / "deck.md").write_text(
        "# Deck\n## Resolved deck brief\n- ownership_topology: single\n",
        encoding="utf-8",
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    for page, needs_bitmap, medium in (
        (1, True, "bitmap-generated"),
        (2, True, "bitmap-real"),
        (3, False, "echarts"),
    ):
        (plan / f"slide_{page:02d}.md").write_text(
            f"# P{page:02d}\n- page_type: content\n"
            f"- production_group: p{page:02d}\n"
            f"- needs_bitmap: {'true' if needs_bitmap else 'false'}\n"
            f"- primary_visual_medium: {medium}\n",
            encoding="utf-8",
        )
        (slides / f"slide_{page:02d}.html").write_text(
            f'<section id="slide-{page:02d}"></section>\n', encoding="utf-8"
        )

    parent = Agent(
        "deck151-regression",
        str(tmp_path),
        "build the deck",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.available_roles = {"image", "slide", "review"}

    def first_wave(_parent, _index, spec):
        role = str(spec.get("role") or "")
        label = str(spec.get("label") or role)
        if role == "image":
            return {
                "label": label,
                "trace_label": "image",
                "role": "image",
                "ok": True,
                "status": "partial_ready",
                "exit_reason": "partial_ready_handoff",
                "attempt": 1,
                "completed_pages": [],
                "incomplete_pages": [],
                "image_ready_pages": [1],
                "image_failed_pages": [2],
                "image_handoff_summary": (
                    "bitmap_ready: 01\nfailed: 02 (search unavailable)"
                ),
                "summary": "bounded partial Image handoff",
                "renders": 0,
                "views": 1,
            }
        pages = list(spec.get("pages") or [])
        return {
            "label": label,
            "trace_label": label,
            "role": "slide",
            "ok": True,
            "status": "ready",
            "exit_reason": "text_response",
            "attempt": 1,
            "completed_pages": pages,
            "incomplete_pages": [],
            "summary": "ready",
            "renders": len(pages),
            "views": len(pages),
        }

    with mock.patch("core.agent_loop._run_child", side_effect=first_wave):
        first = json.loads(_delegate(parent, {"tasks": [
            {"role": "image", "label": "image", "task": "resolve assets"},
            {"role": "slide", "label": "slide_01", "pages": [1], "task": "P01"},
            {"role": "slide", "label": "slide_02", "pages": [2], "task": "P02"},
            {"role": "slide", "label": "slide_03", "pages": [3], "task": "P03"},
        ]}))

    assert first["failed"] == 1
    assert parent.image_completed is True
    assert parent.image_ready_pages == (1,)
    assert parent.image_failed_pages == (2,)
    assert parent.child_outcomes["slide_02"]["status"] == "asset_pending"
    assert parent.child_outcomes["slide_02"]["trace_label"] is None
    assert getattr(parent, "pending_slide_delegation_pages", ()) == ()

    # The Orchestrator remains able to make the bounded pre-production plan
    # correction instead of being trapped behind a delegate-only tool surface.
    (plan / "slide_02.md").write_text(
        "# P02\n- page_type: content\n- production_group: p02\n"
        "- needs_bitmap: false\n- primary_visual_medium: canvas\n",
        encoding="utf-8",
    )
    with mock.patch("core.agent_loop._run_child", side_effect=first_wave):
        second = json.loads(_delegate(parent, {"tasks": [{
            "role": "slide",
            "label": "slide_02",
            "pages": [2],
            "task": "Build the reclassified Canvas fallback without another Image run",
        }]}))

    assert second["failed"] == 0
    assert second["succeeded_labels"] == ["slide_02"]
    assert parent.child_outcomes["slide_02"]["ok"] is True


def test_review_typography_evidence_blocks_generic_dismissal(tmp_path: Path) -> None:
    """Typography quality findings prevent Review from returning plain ready."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text("# deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_05.md").write_text(
        "- page_type: content\n- needs_bitmap: false\n", encoding="utf-8"
    )
    (tmp_path / "renders").mkdir()
    manifest = {
        "pages": [],
        "typography_flags": {
            "5": ["ECharts axis labels 12px, minimum 16px required"],
            "7": ["Legend text 14px on dense chart"],
        },
        "media_mismatches": {},
        "placeholder_flags": {},
        "layout_defects": {},
        "special_page_geometry": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )

    parent = SimpleNamespace(ws=str(tmp_path))
    findings = _review_manifest_quality_findings(parent)
    assert any("typography_flags" in f for f in findings), (
        "typography_flags must appear in quality findings that prevent plain ready"
    )


def test_review_typography_evidence_zh_content(tmp_path: Path) -> None:
    """_review_typography_evidence produces zh tag with page numbers and items."""
    (tmp_path / "renders").mkdir()
    manifest = {
        "pages": [],
        "typography_flags": {
            "3": ["Axis labels 12px", "Legend 14px", "Footnote 11px"],
            "7": ["Canvas text 13px"],
        },
        "media_mismatches": {},
        "placeholder_flags": {},
        "layout_defects": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    result = _review_typography_evidence(str(tmp_path), "zh")
    assert "<harness_typography_evidence>" in result
    assert "</harness_typography_evidence>" in result
    assert "P3:" in result
    assert "P7:" in result
    assert "投影字号" in result
    assert "有意为之" in result


def test_review_typography_evidence_en_content(tmp_path: Path) -> None:
    """_review_typography_evidence produces en tag with page numbers."""
    (tmp_path / "renders").mkdir()
    manifest = {
        "pages": [],
        "typography_flags": {
            "2": ["Body text 14px below 20px minimum"],
        },
        "media_mismatches": {},
        "placeholder_flags": {},
        "layout_defects": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    result = _review_typography_evidence(str(tmp_path), "en")
    assert "<harness_typography_evidence>" in result
    assert "P2:" in result
    assert "Generic dismissal" in result
    assert "pixel reason or repair" in result


def test_review_typography_evidence_caps_8_pages_3_items(tmp_path: Path) -> None:
    """At most 8 pages and 3 items per page in the evidence."""
    (tmp_path / "renders").mkdir()
    flags = {
        str(i): [f"item_{j}" for j in range(5)]
        for i in range(1, 12)
    }
    manifest = {"pages": [], "typography_flags": flags}
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    result = _review_typography_evidence(str(tmp_path), "zh")
    assert result.count("P") <= 8 + 1  # 8 page entries + possible partial
    for page_key in list(flags.keys())[:8]:
        entry_start = f"P{page_key}:"
        if entry_start in result:
            section = result.split(entry_start)[1].split(";")[0]
            assert section.count(",") <= 2  # at most 3 items (2 commas)


def test_review_typography_evidence_empty_manifest(tmp_path: Path) -> None:
    """Empty typography_flags returns empty string."""
    (tmp_path / "renders").mkdir()
    manifest = {"pages": [], "typography_flags": {}}
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    assert _review_typography_evidence(str(tmp_path), "zh") == ""


def test_review_typography_evidence_missing_render_json(tmp_path: Path) -> None:
    """Missing render.json returns empty string."""
    (tmp_path / "renders").mkdir()
    assert _review_typography_evidence(str(tmp_path), "en") == ""


def test_preflight_repair_allows_changed_scaffold_owned(tmp_path: Path) -> None:
    """Deck 153: finalize failure with 'changed scaffold-owned' triggers preflight repair."""
    for d in ("plan", "slides", "renders", "assets"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    for page in (1, 5):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- needs_bitmap: false\n", encoding="utf-8"
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            f"<section data-slide='{page:02d}'>content</section>\n", encoding="utf-8"
        )
        (tmp_path / f"renders/slide_{page:02d}.png").write_bytes(b"pixels")

    parent = Agent(
        "preflight-scaffold",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = False
    parent.finalize_failure = (
        "ValueError: slide_05.html changed scaffold-owned data-page-type "
        "from 'content' to 'special'"
    )
    assert _review_preflight_repair_pages(parent) == (5,)


def test_preflight_repair_allows_root_classes_violation(tmp_path: Path) -> None:
    """Deck 153 variant: 'root classes' marker also triggers preflight repair."""
    for d in ("plan", "slides", "renders", "assets"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    for page in (1, 3):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- needs_bitmap: false\n", encoding="utf-8"
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            f"<section data-slide='{page:02d}'>content</section>\n", encoding="utf-8"
        )
        (tmp_path / f"renders/slide_{page:02d}.png").write_bytes(b"pixels")

    parent = Agent(
        "preflight-root",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = False
    parent.finalize_failure = (
        "ValueError: slide_03.html: root classes must not include "
        "'slide-special'; reserved for scaffold"
    )
    assert _review_preflight_repair_pages(parent) == (3,)


def test_is_durable_review_exit_known_reasons(tmp_path: Path) -> None:
    """All known bounded Review exit reasons qualify for durable handoff."""
    assert _is_durable_review_exit("incomplete_closure")
    assert _is_durable_review_exit("stalled_repetition")
    assert _is_durable_review_exit("max_turns")
    assert _is_durable_review_exit("required_review_page_not_inspected")


def test_is_durable_review_exit_rejects_unknown() -> None:
    """Unknown/infrastructure exit reasons do not qualify."""
    assert not _is_durable_review_exit("api_failed")
    assert not _is_durable_review_exit("runtime_timeout")
    assert not _is_durable_review_exit("text_response")
    assert not _is_durable_review_exit("")
    assert not _is_durable_review_exit(None)


def test_durable_review_exit_reasons_constant() -> None:
    """DURABLE_REVIEW_EXIT_REASONS is a frozenset with exactly 4 entries."""
    assert isinstance(DURABLE_REVIEW_EXIT_REASONS, frozenset)
    assert len(DURABLE_REVIEW_EXIT_REASONS) == 4
    assert "required_review_page_not_inspected" in DURABLE_REVIEW_EXIT_REASONS


def test_durable_review_handoff_required_page_not_inspected(tmp_path: Path) -> None:
    """Deck 154: required_review_page_not_inspected produces durable handoff via _run_child."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    for page in (1, 2, 3):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- needs_bitmap: false\n", encoding="utf-8"
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            f"<section data-slide='{page:02d}'>content</section>\n", encoding="utf-8"
        )
        (tmp_path / f"renders/slide_{page:02d}.png").write_bytes(
            f"pixels-page{page}".encode()
        )
    manifest = {
        "pages": [{"page": i, "status": "ok"} for i in range(1, 4)],
        "typography_flags": {},
        "media_mismatches": {},
        "placeholder_flags": {},
        "layout_defects": {},
        "special_page_geometry": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet")

    child = Agent(
        "review-bounded",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    child.role = "review"
    child.exit_reason = "required_review_page_not_inspected"
    child.finalize_succeeded = True
    child.final_view_after_review = True
    child.review_contact_sheet_inspected = True
    child.required_review_pages = (2,)
    child.review_viewed_page_hashes = {
        2: hashlib.sha256(b"pixels-page2").hexdigest()
    }

    assert _review_has_current_inspected_delivery(child)
    assert _is_durable_review_exit(child.exit_reason)


def test_durable_review_missing_fullres_is_honest_needs_improvement_floor(
    tmp_path: Path,
) -> None:
    """Deck 176: current contact pixels survive bounded incomplete page review."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"current-sheet")
    (tmp_path / "renders/slide_01.png").write_bytes(b"current-page")

    child = Agent(
        "review-partial-current",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    child.role = "review"
    child.finalize_succeeded = True
    child.final_render_after_review = True
    child.final_view_after_review = True
    child.review_contact_sheet_inspected = True
    child.required_review_pages = (1,)
    child.review_viewed_page_hashes = {}

    assert _review_has_current_contact_delivery(child)
    child.finalize_succeeded = False
    assert not _review_has_current_contact_delivery(child)
    child.finalize_succeeded = True
    child.review_contact_sheet_inspected = False
    assert not _review_has_current_contact_delivery(child)
    child.review_contact_sheet_inspected = True
    assert not _review_has_current_inspected_delivery(child)
    assert "P01" in _review_required_view_gap(child)

    result = {
        "ok": False,
        "status": "needs_orchestrator",
        "blocking": "no",
        "exit_reason": "needs_orchestrator",
        "final_pixels_inspected": True,
        "finalize_succeeded": True,
        "final_view_after_review": True,
        "review_changed": True,
    }
    assert _review_can_complete_needs_improvement(child, result)


def test_durable_review_handoff_no_review_r2_label(tmp_path: Path) -> None:
    """Durable handoff produces needs_orchestrator, not a review_r2 delegation."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    for page in (1, 2):
        (tmp_path / f"plan/slide_{page:02d}.md").write_text(
            f"# slide_{page:02d}\n- needs_bitmap: false\n", encoding="utf-8"
        )
        (tmp_path / f"slides/slide_{page:02d}.html").write_text(
            f"<section data-slide='{page:02d}'>content</section>\n", encoding="utf-8"
        )
        (tmp_path / f"renders/slide_{page:02d}.png").write_bytes(
            f"pixels-{page}".encode()
        )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet")
    manifest = {
        "pages": [{"page": i, "status": "ok"} for i in range(1, 3)],
        "typography_flags": {"1": ["axis 14px"]},
        "media_mismatches": {},
        "placeholder_flags": {},
        "layout_defects": {},
        "special_page_geometry": {},
    }
    (tmp_path / "renders/render.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )

    parent = Agent(
        "orch-bounded",
        str(tmp_path),
        "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = True

    def fake_child(_parent, _index, spec):
        return {
            "label": "review",
            "trace_label": "review",
            "role": "review",
            "ok": False,
            "exit_reason": "needs_orchestrator",
            "status": "needs_orchestrator",
            "renders": 2,
            "views": 2,
            "completed_pages": [],
            "incomplete_pages": [],
            "final_render_after_review": True,
            "final_view_after_review": True,
            "finalize_attempted": True,
            "finalize_succeeded": True,
            "finalize_failure": "",
            "review_changed": False,
            "blocking": "no",
            "issue_type": "bounded_review_incomplete",
            "pages": "inspected delivery surface",
            "evidence": "bounded Review ended with max_turns",
            "final_pixels_inspected": True,
            "input_fingerprint": "",
            "required_review_pages": [1],
            "repair_issue": None,
            "attempt": 1,
            "trace_mode": {},
            "summary": (
                "status: needs_orchestrator\n"
                "pages: inspected delivery surface\n"
                "issue_type: bounded_review_incomplete\n"
                "evidence: bounded Review ended with max_turns\n"
                "blocking: no\n"
                "final_pixels_inspected: yes"
            ),
        }

    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        result = json.loads(_delegate(parent, {"tasks": [{"role": "review"}]}))

    assert result["status"] == "completed"
    child_outcome = parent.child_outcomes.get("review", {})
    assert child_outcome.get("status") == "needs_improvement"
    assert "review_r2" not in json.dumps(result)


# ─── Review closure-only mode (Deck 157 fix) ───────────────────────────


def test_review_closure_only_triggers_on_high_turn_fraction() -> None:
    """Review enters closure-only when >70% turns used with unfinalized patches."""
    agent = SimpleNamespace(
        role="review",
        max_turns=100,
        review_patches_since_finalize=2,
        review_closure_only=False,
    )
    assert not _review_should_enter_closure_only(agent, 69)
    assert _review_should_enter_closure_only(agent, 70)
    assert _review_should_enter_closure_only(agent, 95)


def test_review_batch_may_use_many_patch_calls_before_first_finalize() -> None:
    """Patch calls in one coordinated batch are not pixel revision rounds."""
    agent = SimpleNamespace(
        role="review",
        max_turns=100,
        review_patches_since_finalize=8,
        review_closure_only=False,
    )
    assert not _review_should_enter_closure_only(agent, 10)


def test_review_closure_only_does_not_trigger_without_patches() -> None:
    """No unfinalized patches means no closure-only, even at high turns."""
    agent = SimpleNamespace(
        role="review",
        max_turns=100,
        review_patches_since_finalize=0,
        review_closure_only=False,
    )
    assert not _review_should_enter_closure_only(agent, 95)


def test_review_closure_only_not_for_non_review() -> None:
    """Closure-only check returns False for non-Review roles."""
    agent = SimpleNamespace(
        role="slide",
        max_turns=100,
        review_patches_since_finalize=10,
        review_closure_only=False,
    )
    assert not _review_should_enter_closure_only(agent, 95)


def test_review_closure_only_sticky() -> None:
    """Once review_closure_only is True, it stays True regardless of counters."""
    agent = SimpleNamespace(
        role="review",
        max_turns=100,
        review_patches_since_finalize=0,
        review_closure_only=True,
    )
    assert _review_should_enter_closure_only(agent, 5)


def test_review_closure_only_tool_restriction(tmp_path: Path) -> None:
    """In closure-only mode, _call restricts tools to terminal and vision_analyze."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "_trace_label": "review",
    }
    agent = Agent("test-sid", str(tmp_path), "test task", cfg, role="review")
    agent.review_patches_since_finalize = 10
    agent.review_closure_only = False
    agent.turn = 80
    agent.max_turns = 100
    agent.tool_schemas = [
        {"name": "terminal"},
        {"name": "vision_analyze"},
        {"name": "patch"},
        {"name": "read_file"},
    ]
    fake_response = SimpleNamespace(
        content=[SimpleNamespace(type="text", text="done")],
        stop_reason="end_turn",
        usage=SimpleNamespace(input_tokens=10, output_tokens=10),
    )
    with mock.patch("core.model_call.call_with_tools", return_value=fake_response):
        _call(agent, [{"role": "user", "content": "test"}], with_tools=True)
    assert agent.review_closure_only is True


def test_three_completed_rounds_hide_mutation_tools(tmp_path: Path) -> None:
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    agent = Agent(
        "three-round-stop", str(tmp_path), "review",
        {"_raw_user_query": "test"}, role="review",
    )
    agent.review_revision_rounds = config.REVIEW_MAX_ATTEMPTS
    agent.required_review_pages = ()
    agent.review_contact_sheet_inspected = True
    agent.tool_schemas = [
        {"name": "terminal"}, {"name": "vision_analyze"},
        {"name": "patch"}, {"name": "read_file"},
    ]
    fake_response = SimpleNamespace(
        content=[SimpleNamespace(type="text", text="done")],
        stop_reason="end_turn",
        usage=SimpleNamespace(input_tokens=10, output_tokens=10),
    )
    with mock.patch(
        "core.model_call.call_with_tools", return_value=fake_response
    ) as model_call:
        _call(agent, [{"role": "user", "content": "test"}], with_tools=True)
    assert model_call.call_args.kwargs["tools"] == []


def test_open_vision_issue_is_mandatory_without_plan_flag(tmp_path: Path) -> None:
    for relative in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / relative).mkdir()
    (tmp_path / "plan/deck.md").write_text("# Deck\n", encoding="utf-8")
    (tmp_path / "plan/slide_07.md").write_text(
        "# slide_07\n- needs_bitmap: false\n- composition: editorial\n",
        encoding="utf-8",
    )
    (tmp_path / "slides/slide_07.html").write_text(
        "<section data-slide='07'>page</section>\n", encoding="utf-8",
    )
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "renders/slide_07.png").write_bytes(b"pixels")
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"sheet")
    (tmp_path / "renders/render.json").write_text(
        json.dumps({"pages": [{"page": 7}]}), encoding="utf-8",
    )
    (tmp_path / "_trace/vision-issues.json").write_text(json.dumps({
        "schema": "mural.vision-issues.v1",
        "issues": [{
            "id": "VIS-P07", "status": "open", "page": 7,
            "type": "text_overlap", "severity": "major",
            "summary": "labels overlap",
        }],
    }), encoding="utf-8")
    parent = Agent(
        "vision-required", str(tmp_path), "build",
        {"skill_name": "mural-presenter-v0-4", "skill_language": "zh"},
    )
    parent.research_required = False
    parent.finalize_attempted = True
    parent.finalize_succeeded = True
    captured = []

    def fake_child(_parent, _index, spec):
        captured.append(dict(spec))
        return {
            "label": "review", "role": "review", "ok": True,
            "status": "ready", "exit_reason": "text_response",
            "completed_pages": [], "incomplete_pages": [],
            "renders": 0, "views": 2,
            "final_render_after_review": True,
            "final_view_after_review": True,
            "finalize_attempted": True, "finalize_succeeded": True,
            "summary": "verified",
        }

    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        result = json.loads(_delegate(parent, {"tasks": [{"role": "review"}]}))
    assert result["status"] == "completed"
    assert captured[0]["required_review_pages"] == [7]


def test_stale_review_not_promoted_to_ok(tmp_path: Path) -> None:
    """review_incomplete_current_pixels stays ok=False — stale pixels never deliverable."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("slide_count: 2\n", encoding="utf-8")
    for p in (1, 2):
        (tmp_path / f"plan/slide_{p:02d}.md").write_text(f"# Slide {p}\n", encoding="utf-8")
        (tmp_path / f"slides/slide_{p:02d}.html").write_text(f"<div>S{p}</div>", encoding="utf-8")
        (tmp_path / f"renders/slide_{p:02d}.png").write_bytes(b"PNG" + bytes([p]) * 100)
    (tmp_path / "speech.md").write_text("# Speech\n", encoding="utf-8")
    (tmp_path / "renders/render.json").write_text(
        json.dumps({"pages": [{"page": 1}, {"page": 2}]}), encoding="utf-8"
    )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"PNGCS" + b"\x00" * 100)
    parent_cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "child_wall_timeout_s": 120,
        "_trace_label": "orchestrator",
    }
    parent = Agent("sid", str(tmp_path), "orchestrate", parent_cfg, role="orchestrator")
    parent.finalize_attempted = True
    parent.finalize_succeeded = True

    def fake_child(p, idx, spec):
        return {
            "label": "review",
            "role": "review",
            "ok": False,
            "status": "needs_orchestrator",
            "blocking": "no",
            "exit_reason": "review_incomplete_current_pixels",
            "final_pixels_inspected": False,
            "finalize_succeeded": False,
            "finalize_attempted": True,
            "finalize_failure": "",
            "final_render_after_review": False,
            "final_view_after_review": False,
            "review_changed": True,
            "review_patches_since_finalize": 5,
            "input_fingerprint": "fp1",
            "attempt": 1,
            "trace_mode": {},
            "summary": "stale",
        }

    with mock.patch("core.agent_loop._run_child", side_effect=fake_child):
        _delegate(parent, {"tasks": [{"role": "review"}]})

    child_outcome = parent.child_outcomes.get("review", {})
    assert child_outcome.get("ok") is False
    assert child_outcome.get("exit_reason") == "review_incomplete_current_pixels"
    assert parent.review_completed is False


def test_delivery_close_status_not_for_stale_pixels(tmp_path: Path) -> None:
    """_orchestrator_delivery_close_status returns '' when pixels are stale."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "_trace_label": "orchestrator",
    }
    agent = Agent("sid", str(tmp_path), "orchestrate", cfg, role="orchestrator")
    agent.review_completed = False
    agent.finalize_succeeded = False
    agent.final_view_after_review = False
    agent.child_outcomes = {
        "review": {
            "role": "review",
            "ok": False,
            "exit_reason": "review_incomplete_current_pixels",
            "status": "needs_orchestrator",
        }
    }
    status = _orchestrator_delivery_close_status(agent)
    assert status == ""


def test_closure_tail_same_identity_bounded() -> None:
    """Closure tail uses same Agent identity and is bounded to REVIEW_CLOSURE_TAIL_BUDGET."""
    assert isinstance(REVIEW_CLOSURE_TAIL_BUDGET, int)
    assert 4 <= REVIEW_CLOSURE_TAIL_BUDGET <= 10


def test_closure_tail_success_enables_needs_improvement(tmp_path: Path) -> None:
    """After closure tail finalizes + inspects, _review_can_complete_needs_improvement works."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("slide_count: 1\n", encoding="utf-8")
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "_trace_label": "orchestrator",
    }
    parent = Agent("sid", str(tmp_path), "orchestrate", cfg, role="orchestrator")
    result = {
        "ok": False,
        "status": "needs_orchestrator",
        "blocking": "no",
        "exit_reason": "text_response",
        "final_pixels_inspected": True,
        "finalize_succeeded": True,
        "final_view_after_review": True,
        "review_changed": True,
    }
    assert _review_can_complete_needs_improvement(parent, result) is True


def test_closure_tail_still_stale_not_deliverable(tmp_path: Path) -> None:
    """If closure tail exhausts without finalize, result stays ok=False."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "_trace_label": "orchestrator",
    }
    parent = Agent("sid", str(tmp_path), "orchestrate", cfg, role="orchestrator")
    result = {
        "ok": False,
        "status": "needs_orchestrator",
        "blocking": "no",
        "exit_reason": "review_incomplete_current_pixels",
        "final_pixels_inspected": False,
        "finalize_succeeded": False,
        "final_view_after_review": False,
        "review_changed": True,
    }
    assert _review_can_complete_needs_improvement(parent, result) is False


def test_orchestrator_blocks_second_review_after_stale_exit(tmp_path: Path) -> None:
    """After review_incomplete_current_pixels, Orchestrator gets terminal failure instruction."""
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("slide_count: 2\n", encoding="utf-8")
    for p in (1, 2):
        (tmp_path / f"plan/slide_{p:02d}.md").write_text(f"# S{p}\n", encoding="utf-8")
        (tmp_path / f"slides/slide_{p:02d}.html").write_text(f"<div>{p}</div>", encoding="utf-8")
        (tmp_path / f"renders/slide_{p:02d}.png").write_bytes(b"PNG" + bytes([p]) * 50)
    (tmp_path / "speech.md").write_text("# S\n", encoding="utf-8")
    (tmp_path / "renders/render.json").write_text(
        json.dumps({"pages": [{"page": 1}, {"page": 2}]}), encoding="utf-8"
    )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"PNGCS" + b"\x00" * 50)
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "child_wall_timeout_s": 120,
        "_trace_label": "orchestrator",
    }
    agent = Agent("sid", str(tmp_path), "orchestrate", cfg, role="orchestrator")
    agent.finalize_attempted = True
    agent.finalize_succeeded = True
    agent.review_completed = False
    agent.child_outcomes = {
        "review": {
            "role": "review",
            "ok": False,
            "exit_reason": "review_incomplete_current_pixels",
            "status": "needs_orchestrator",
            "input_fingerprint": "fp1",
            "attempt": 1,
        }
    }
    agent.delegated_roles = ["review"]
    with mock.patch(
        "core.agent_loop._review_delivery_fingerprint", return_value="fp1"
    ):
        result = _delegate(agent, {"tasks": [{"role": "review", "label": "review"}]})
    assert "重复委派" in result or "均未发生变化" in result or "禁止" in result
    assert "review_r2" not in result


def test_review_terminal_failure_deterministic_stop(tmp_path: Path) -> None:
    """review_incomplete_current_pixels triggers deterministic terminal stop.

    Requirements:
    - spec has no explicit label (uses {"role": "review"})
    - _run_child is mocked to return review_incomplete_current_pixels
    - Only one child runs
    - parent.review_terminal_failure is set
    - Second _delegate_task call is blocked without spawning a child
    - No review_r2
    """
    for d in ("plan", "slides", "renders", "assets", "_trace"):
        (tmp_path / d).mkdir()
    (tmp_path / "plan/deck.md").write_text("slide_count: 2\n", encoding="utf-8")
    for p in (1, 2):
        (tmp_path / f"plan/slide_{p:02d}.md").write_text(f"# S{p}\n", encoding="utf-8")
        (tmp_path / f"slides/slide_{p:02d}.html").write_text(f"<div>{p}</div>", encoding="utf-8")
        (tmp_path / f"renders/slide_{p:02d}.png").write_bytes(b"PNG" + bytes([p]) * 50)
    (tmp_path / "speech.md").write_text("# S\n", encoding="utf-8")
    (tmp_path / "renders/render.json").write_text(
        json.dumps({"pages": [{"page": 1}, {"page": 2}]}), encoding="utf-8"
    )
    (tmp_path / "renders/contact-sheet.png").write_bytes(b"PNGCS" + b"\x00" * 50)
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "child_wall_timeout_s": 120,
        "_trace_label": "orchestrator",
    }
    agent = Agent("sid", str(tmp_path), "orchestrate", cfg, role="orchestrator")
    agent.finalize_attempted = True
    agent.finalize_succeeded = True
    agent.review_completed = False
    agent.child_outcomes = {}
    agent.delegated_roles = []

    run_child_calls = []

    def mock_run_child(parent, index, spec):
        run_child_calls.append(spec)
        return {
            "label": "review",
            "trace_label": "review_t1",
            "role": "review",
            "ok": False,
            "exit_reason": "review_incomplete_current_pixels",
            "renders": 0,
            "views": 0,
            "completed_pages": [],
            "incomplete_pages": [],
            "final_render_after_review": False,
            "final_view_after_review": False,
            "finalize_attempted": True,
            "finalize_succeeded": False,
            "finalize_failure": "",
            "trace_mode": {},
            "status": "needs_orchestrator",
            "repair_issue": None,
            "image_ready_pages": [],
            "image_failed_pages": [],
            "image_handoff_summary": "",
            "required_review_pages": [],
            "attempt": 1,
            "blocking": "no",
            "issue_type": "review_incomplete_current_pixels",
            "pages": "stale",
            "evidence": "closure tail exhausted",
            "input_fingerprint": "fp_stale",
            "final_pixels_inspected": False,
            "review_changed": False,
            "summary": "status: needs_orchestrator\nissue_type: review_incomplete_current_pixels",
        }

    with mock.patch("core.agent_loop._run_child", side_effect=mock_run_child):
        result1 = _delegate(agent, {"tasks": [{"role": "review"}]})

    # Only one child was spawned
    assert len(run_child_calls) == 1
    assert run_child_calls[0].get("role") == "review"

    # Terminal failure flag is set
    assert bool(getattr(agent, "review_terminal_failure", False))

    # Outcome is stored as ok=False
    review_outcome = agent.child_outcomes.get("review")
    assert review_outcome is not None
    assert review_outcome["ok"] is False
    assert review_outcome["exit_reason"] == "review_incomplete_current_pixels"

    # review_completed is NOT set (stale pixels are not deliverable)
    assert not bool(getattr(agent, "review_completed", False))

    # Second delegate_task call is blocked without spawning a child
    run_child_calls.clear()
    result2 = _delegate_task(agent, {"tasks": [{"role": "review"}]})
    assert len(run_child_calls) == 0
    assert "review_closure_failed" in result2
    assert "review_r2" not in result2

    # No review_r2 in any result
    assert "review_r2" not in result1
    assert "review_r2" not in (json.dumps(agent.child_outcomes, ensure_ascii=False))


def test_review_terminal_failure_blocks_all_tools(tmp_path: Path) -> None:
    """After review_terminal_failure is set, _tool_results blocks all tool calls."""
    (tmp_path / "plan").mkdir()
    (tmp_path / "plan/deck.md").write_text("slide_count: 1\n", encoding="utf-8")
    cfg = {
        "model": "test-model",
        "api_base": "http://test",
        "response_language": "en",
        "deck_timeout_s": 300,
        "child_wall_timeout_s": 120,
        "_trace_label": "orchestrator",
    }
    agent = Agent("sid", str(tmp_path), "orchestrate", cfg, role="orchestrator")
    agent.review_terminal_failure = True

    fake_call = SimpleNamespace(
        id="call_1", name="read_file", type="tool_use",
        input={"path": "plan/deck.md"},
    )
    tool_log: list = []
    results = _tool_results(agent, [fake_call], 5, tool_log)
    assert len(results) == 1
    result_text = results[0].get("content", "") if isinstance(results[0], dict) else ""
    if isinstance(results[0], dict) and isinstance(results[0].get("content"), list):
        result_text = results[0]["content"][0].get("text", "")
    elif isinstance(results[0], dict) and isinstance(results[0].get("content"), str):
        result_text = results[0]["content"]
    assert "review_closure_failed" in result_text


def test_wide_table_skill_contract() -> None:
    """Skill role card forbids shrinking projection text below 20px for wide tables."""
    role_card = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/roles/review.md"
    )
    text = role_card.read_text(encoding="utf-8")
    assert "< 20px" in text or "<20px" in text or "< 20 px" in text
    assert "减维" in text or "fewer columns" in text
    assert "聚合" in text or "aggregate" in text or "summarize" in text
    assert "message-led" in text
    assert "notes" in text


# --- #158 knowledge-brief budget fix tests ---


def test_brief_hard_ceiling_formula() -> None:
    """Hard ceiling = min(40000, int(recommended * 1.6))."""

    class Agent10:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = "/tmp/fakews"
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

    agent = Agent10()
    rec = tools.research_brief_char_limit(agent)
    assert rec == 20000
    assert tools.research_brief_hard_ceiling(agent) == 32000

    class Agent18:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = "/tmp/fakews"
        cfg = {"_raw_user_query": "请基于附件制作 18 页组会汇报"}
        requested_slide_count = 0

    agent18 = Agent18()
    assert tools.research_brief_char_limit(agent18) == 28000
    assert tools.research_brief_hard_ceiling(agent18) == 40000

    class DefaultAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = "/tmp/fakews"
        cfg = {"_raw_user_query": "请做 8 页演示"}
        requested_slide_count = 0

    d = DefaultAgent()
    rec_d = tools.research_brief_char_limit(d)
    assert tools.research_brief_hard_ceiling(d) == min(40000, int(rec_d * 1.6))


def test_over_budget_brief_accepted_on_first_write(tmp_path: Path) -> None:
    """Brief at 25927 chars with recommended=20000 is accepted once with warning."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    assert tools.research_brief_char_limit(agent) == 20000
    content = "x" * 25927
    result = tools.write_file(agent, "research/knowledge-brief.md", content)
    assert "已写入" in result
    assert "超出推荐预算" in result or "超推荐" in result or "over" in result.lower()
    assert "锁定" in result
    assert (tmp_path / "research/knowledge-brief.md").read_text("utf-8") == content
    assert tools.research_handoff_is_valid(agent) is True


def test_over_budget_receipt_is_v2_with_correct_fields(tmp_path: Path) -> None:
    """Receipt for over-budget brief is version 2 with over_budget=True."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "x" * 25927
    tools.write_file(agent, "research/knowledge-brief.md", content)
    import json as _json
    receipt = _json.loads(
        (tmp_path / "_trace/research-handoff.json").read_text("utf-8")
    )
    assert receipt["version"] == 2
    assert receipt["over_budget"] is True
    assert receipt["characters"] == 25927
    assert receipt["recommended_characters"] == 20000
    assert receipt["hard_ceiling_characters"] == 32000


def test_above_ceiling_brief_rejected(tmp_path: Path) -> None:
    """Brief above hard ceiling (>32000) is hard-rejected."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    result = tools.write_file(
        ResearchAgent(),
        "research/knowledge-brief.md",
        "x" * 32001,
    )
    assert "knowledge_brief_too_long" in result
    assert not (tmp_path / "research/knowledge-brief.md").exists()


def test_normal_brief_under_recommended_accepted(tmp_path: Path) -> None:
    """Brief under recommended budget is accepted normally without warning."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "x" * 15000
    result = tools.write_file(agent, "research/knowledge-brief.md", content)
    assert "已写入" in result
    assert "超出推荐预算" not in result
    assert "锁定" not in result
    assert tools.research_handoff_is_valid(agent) is True


def test_second_write_after_receipt_locked(tmp_path: Path) -> None:
    """Second write_file to knowledge-brief is blocked after receipt exists."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    first_content = "x" * 25927
    result1 = tools.write_file(agent, "research/knowledge-brief.md", first_content)
    assert "已写入" in result1
    import hashlib as _hl
    first_hash = _hl.sha256(first_content.encode()).hexdigest()
    result2 = tools.write_file(agent, "research/knowledge-brief.md", "y" * 10000)
    assert "knowledge_brief_locked" in result2
    on_disk = (tmp_path / "research/knowledge-brief.md").read_text("utf-8")
    assert _hl.sha256(on_disk.encode()).hexdigest() == first_hash


def test_patch_after_receipt_locked(tmp_path: Path) -> None:
    """Patch to knowledge-brief is blocked after receipt exists."""

    class ResearchAgent:
        role = "research"
        label = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "A" * 100 + "B" + "C" * 100
    tools.write_file(agent, "research/knowledge-brief.md", content)
    result = tools.patch(
        agent,
        path="research/knowledge-brief.md",
        old_string="B",
        new_string="DDDD",
    )
    assert "knowledge_brief_locked" in result
    assert (tmp_path / "research/knowledge-brief.md").read_text("utf-8") == content


def test_normal_brief_also_locked_after_receipt(tmp_path: Path) -> None:
    """Even a normal (under-recommended) brief is locked after receipt."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "x" * 15000
    tools.write_file(agent, "research/knowledge-brief.md", content)
    result2 = tools.write_file(agent, "research/knowledge-brief.md", "y" * 8000)
    assert "knowledge_brief_locked" in result2


def test_v1_receipt_still_valid(tmp_path: Path) -> None:
    """V1 receipt from legacy code is accepted by research_handoff_is_valid."""
    import json as _json
    import hashlib as _hl
    content = "x" * 10000

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请做 8 页演示"}
        requested_slide_count = 0

    (tmp_path / "research").mkdir(parents=True)
    (tmp_path / "research/knowledge-brief.md").write_text(content, encoding="utf-8")
    encoded = content.encode("utf-8")
    receipt = {
        "version": 1,
        "path": "research/knowledge-brief.md",
        "sha256": _hl.sha256(encoded).hexdigest(),
        "bytes": len(encoded),
        "characters": len(content),
    }
    (tmp_path / "_trace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "_trace/research-handoff.json").write_text(
        _json.dumps(receipt) + "\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(ResearchAgent()) is True


def test_v1_receipt_rejects_over_recommended(tmp_path: Path) -> None:
    """V1 receipt with content exceeding recommended limit is invalid."""
    import json as _json
    import hashlib as _hl
    content = "x" * 13000

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "open_research"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请做 8 页演示"}
        requested_slide_count = 0

    (tmp_path / "research").mkdir(parents=True)
    (tmp_path / "research/knowledge-brief.md").write_text(content, encoding="utf-8")
    encoded = content.encode("utf-8")
    receipt = {
        "version": 1,
        "path": "research/knowledge-brief.md",
        "sha256": _hl.sha256(encoded).hexdigest(),
        "bytes": len(encoded),
        "characters": len(content),
    }
    (tmp_path / "_trace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "_trace/research-handoff.json").write_text(
        _json.dumps(receipt) + "\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(ResearchAgent()) is False


def test_over_budget_warning_triggers_auto_close(tmp_path: Path) -> None:
    """Over-budget write returns '已写入' without '错误' so _research_handoff_written detects it."""

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    content = "x" * 25927
    result = tools.write_file(agent, "research/knowledge-brief.md", content)
    assert "已写入" in result
    assert "错误" not in result

    fake_call = SimpleNamespace(
        id="call_1", name="write_file", type="tool_use",
        input={"path": "research/knowledge-brief.md"},
    )
    tool_result = {"content": result}
    assert _research_handoff_written(agent, [fake_call], [tool_result]) is True


def test_research_role_card_distinguishes_recommended_and_ceiling() -> None:
    """Research role card mentions both recommended and absolute ceiling."""
    role_card = (
        REPO
        / "skills/mural-presenter-v0.4/mural-presenter-v0-4/roles/research.md"
    )
    text = role_card.read_text(encoding="utf-8")
    assert "推荐预算" in text or "recommended" in text.lower()
    assert "绝对" in text or "hard ceiling" in text.lower()


def test_invalid_receipt_does_not_lock_allows_recovery_write(tmp_path: Path) -> None:
    """Corrupted/tampered receipt does not lock the brief; recovery write succeeds."""
    import json as _json

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    (tmp_path / "research").mkdir(parents=True)
    (tmp_path / "_trace").mkdir(parents=True, exist_ok=True)
    old_content = "old brief content"
    (tmp_path / "research/knowledge-brief.md").write_text(old_content, encoding="utf-8")
    bad_receipt = {
        "version": 2,
        "path": "research/knowledge-brief.md",
        "sha256": "0000000000000000000000000000000000000000000000000000000000000000",
        "bytes": 99999,
        "characters": 99999,
        "recommended_characters": 20000,
        "hard_ceiling_characters": 32000,
        "over_budget": True,
    }
    (tmp_path / "_trace/research-handoff.json").write_text(
        _json.dumps(bad_receipt) + "\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(agent) is False
    new_content = "x" * 15000
    result = tools.write_file(agent, "research/knowledge-brief.md", new_content)
    assert "已写入" in result
    assert "错误" not in result
    assert tools.research_handoff_is_valid(agent) is True


def test_tampered_brief_file_does_not_lock(tmp_path: Path) -> None:
    """If brief file is tampered (hash mismatch), receipt is invalid and write allowed."""
    import json as _json, hashlib as _hl

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

        def safe(self, path):
            return str(tmp_path / path)

    agent = ResearchAgent()
    original = "x" * 15000
    result1 = tools.write_file(agent, "research/knowledge-brief.md", original)
    assert "已写入" in result1
    assert tools.research_handoff_is_valid(agent) is True
    (tmp_path / "research/knowledge-brief.md").write_text(
        "tampered content", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(agent) is False
    new_content = "y" * 12000
    result2 = tools.write_file(agent, "research/knowledge-brief.md", new_content)
    assert "已写入" in result2
    assert tools.research_handoff_is_valid(agent) is True


def test_v2_receipt_invalid_when_over_budget_not_bool(tmp_path: Path) -> None:
    """V2 receipt with non-bool over_budget is invalid."""
    import json as _json, hashlib as _hl

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

    content = "x" * 25000
    encoded = content.encode("utf-8")
    (tmp_path / "research").mkdir(parents=True)
    (tmp_path / "research/knowledge-brief.md").write_text(content, encoding="utf-8")
    receipt = {
        "version": 2,
        "path": "research/knowledge-brief.md",
        "sha256": _hl.sha256(encoded).hexdigest(),
        "bytes": len(encoded),
        "characters": len(content),
        "recommended_characters": 20000,
        "hard_ceiling_characters": 32000,
        "over_budget": "true",
    }
    (tmp_path / "_trace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "_trace/research-handoff.json").write_text(
        _json.dumps(receipt) + "\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(ResearchAgent()) is False


def test_v2_receipt_invalid_when_recommended_mismatch(tmp_path: Path) -> None:
    """V2 receipt with wrong recommended_characters is invalid."""
    import json as _json, hashlib as _hl

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

    content = "x" * 25000
    encoded = content.encode("utf-8")
    (tmp_path / "research").mkdir(parents=True)
    (tmp_path / "research/knowledge-brief.md").write_text(content, encoding="utf-8")
    receipt = {
        "version": 2,
        "path": "research/knowledge-brief.md",
        "sha256": _hl.sha256(encoded).hexdigest(),
        "bytes": len(encoded),
        "characters": len(content),
        "recommended_characters": 15000,
        "hard_ceiling_characters": 32000,
        "over_budget": True,
    }
    (tmp_path / "_trace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "_trace/research-handoff.json").write_text(
        _json.dumps(receipt) + "\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(ResearchAgent()) is False


def test_v2_receipt_invalid_when_over_budget_flag_inconsistent(tmp_path: Path) -> None:
    """V2 receipt where over_budget disagrees with actual len vs recommended is invalid."""
    import json as _json, hashlib as _hl

    class ResearchAgent:
        role = "research"
        skill_name = "mural-presenter-v0-4"
        evidence_scope = "verify_external"
        ws = str(tmp_path)
        cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
        requested_slide_count = 0

    content = "x" * 25000
    encoded = content.encode("utf-8")
    (tmp_path / "research").mkdir(parents=True)
    (tmp_path / "research/knowledge-brief.md").write_text(content, encoding="utf-8")
    receipt = {
        "version": 2,
        "path": "research/knowledge-brief.md",
        "sha256": _hl.sha256(encoded).hexdigest(),
        "bytes": len(encoded),
        "characters": len(content),
        "recommended_characters": 20000,
        "hard_ceiling_characters": 32000,
        "over_budget": False,
    }
    (tmp_path / "_trace").mkdir(parents=True, exist_ok=True)
    (tmp_path / "_trace/research-handoff.json").write_text(
        _json.dumps(receipt) + "\n", encoding="utf-8"
    )
    assert tools.research_handoff_is_valid(ResearchAgent()) is False


def test_system_prompt_contains_both_recommended_and_ceiling(tmp_path: Path) -> None:
    """Research agent system prompt contains both recommended and hard ceiling values."""
    cfg = {"_raw_user_query": "请基于附件制作 10 页组会汇报"}
    agent = Agent("sid", str(tmp_path), "research_task", cfg, role="research")
    recommended = agent.research_brief_char_limit
    assert recommended > 0
    ceiling = tools.research_brief_hard_ceiling(agent)
    assert str(recommended) in agent.system
    assert str(ceiling) in agent.system
    assert "推荐预算" in agent.system
    assert "绝对安全上限" in agent.system


# ═══════════════════════════════════════════════════════════════════════════
# Deck 160 root-cause: agent.turn initialization + closure-only logging
# ═══════════════════════════════════════════════════════════════════════════


def test_agent_turn_initialized_at_zero(tmp_path: Path) -> None:
    """Agent.turn must be initialized to 0 at construction."""
    cfg = {"_raw_user_query": "test"}
    agent = Agent("sid", str(tmp_path), "task", cfg, role="review")
    assert hasattr(agent, "turn")
    assert agent.turn == 0


def test_review_turn_is_initialized_for_closure_logging(tmp_path: Path) -> None:
    """Late-turn closure logging can always read the initialized turn counter."""
    cfg = {"_raw_user_query": "test"}
    agent = Agent("sid", str(tmp_path), "task", cfg, role="review")
    agent.review_patches_since_finalize = 8
    assert not _review_should_enter_closure_only(agent, agent.turn)
    agent.turn = int(agent.max_turns * REVIEW_CLOSURE_ONLY_TURN_FRACTION)
    assert _review_should_enter_closure_only(agent, agent.turn)
    log_msg = (
        f"turn={agent.turn}/{agent.max_turns}"
    )
    assert f"turn={agent.turn}/" in log_msg


def test_child_exception_preserves_singleton_identity(tmp_path: Path) -> None:
    """A child crash must produce a durable failed outcome blocking re-delegation."""
    from core.agent_loop import _OPERATIONAL_CHILD_RETRY_REASONS, _operational_child_retry_allowed

    outcome = {
        "label": "review",
        "role": "review",
        "ok": False,
        "status": "child_exception",
        "exit_reason": "child_exception",
        "attempt": 1,
        "summary": "AttributeError: ...",
    }
    assert "child_exception" not in _OPERATIONAL_CHILD_RETRY_REASONS
    assert not _operational_child_retry_allowed(outcome)


def test_no_review_retry2_on_child_exception(tmp_path: Path) -> None:
    """After child_exception outcome, re-delegation of same label must be rejected."""
    cfg = {"_raw_user_query": "test"}
    agent = Agent("sid", str(tmp_path), "task", cfg, role="orchestrator")
    agent.child_outcomes = {
        "review": {
            "role": "review",
            "ok": False,
            "exit_reason": "child_exception",
            "attempt": 1,
            "status": "child_exception",
        }
    }
    from core.agent_loop import _review_verification_allowed
    allowed, reason = _review_verification_allowed(agent, agent.child_outcomes["review"])
    assert not allowed


# ═══════════════════════════════════════════════════════════════════════════
# Review-issues ledger reconciliation after Review completion
# ═══════════════════════════════════════════════════════════════════════════


def test_review_issues_syncs_vision_closure(tmp_path: Path) -> None:
    """After Review closes all vision issues, review-issues.json reflects 0 open."""
    trace = tmp_path / "_trace"
    trace.mkdir()
    (trace / "vision-issues.json").write_text(json.dumps({
        "schema": "mural.vision-issues.v1",
        "issues": [
            {"id": "VIS-AAA", "status": "closed", "page": 2},
            {"id": "VIS-BBB", "status": "closed", "page": 3},
        ],
    }))
    (trace / "review-issues.json").write_text(json.dumps({
        "status": "repair_required",
        "required_review_pages": [2, 3],
        "issues": [
            {"label": "slide_02", "status": "repair_required", "pages": [2],
             "render_budget_exhausted_pages": [2]},
        ],
        "open_vision_issue_ids": ["VIS-AAA", "VIS-BBB"],
    }))
    cfg = {"_raw_user_query": "test"}
    agent = Agent("sid", str(tmp_path), "task", cfg, role="orchestrator")
    _reconcile_review_issues_ledger(agent)
    ledger = json.loads((trace / "review-issues.json").read_text())
    assert ledger["open_vision_issue_ids"] == []
    assert ledger["status"] == "resolved"
    assert ledger["issues"][0]["status"] == "closed_by_review"


def test_review_issues_preserves_open_ids(tmp_path: Path) -> None:
    """Partial closure leaves remaining IDs in open_vision_issue_ids."""
    trace = tmp_path / "_trace"
    trace.mkdir()
    (trace / "vision-issues.json").write_text(json.dumps({
        "schema": "mural.vision-issues.v1",
        "issues": [
            {"id": "VIS-AAA", "status": "closed", "page": 2},
            {"id": "VIS-BBB", "status": "open", "page": 3},
        ],
    }))
    (trace / "review-issues.json").write_text(json.dumps({
        "status": "repair_required",
        "required_review_pages": [2, 3],
        "issues": [{"label": "slide_02", "status": "repair_required", "pages": [2]}],
        "open_vision_issue_ids": ["VIS-AAA", "VIS-BBB"],
    }))
    cfg = {"_raw_user_query": "test"}
    agent = Agent("sid", str(tmp_path), "task", cfg, role="orchestrator")
    _reconcile_review_issues_ledger(agent)
    ledger = json.loads((trace / "review-issues.json").read_text())
    assert ledger["open_vision_issue_ids"] == ["VIS-BBB"]
    assert ledger["status"] == "repair_required"


# ═══════════════════════════════════════════════════════════════════════════
# Console errors surfacing in Review handoff
# ═══════════════════════════════════════════════════════════════════════════


def test_console_errors_in_manifest_quality_findings(tmp_path: Path) -> None:
    """console_errors in render.json must appear in manifest quality findings."""
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(json.dumps({
        "n_pages": 2,
        "console_errors": ["pageerror c.quadTo is not a function"],
        "pages": [],
    }))
    cfg = {"_raw_user_query": "test"}
    agent = Agent("sid", str(tmp_path), "task", cfg, role="review")
    findings = _review_manifest_quality_findings(agent)
    assert any("console_errors" in f for f in findings)


def test_render_console_errors_returns_list(tmp_path: Path) -> None:
    """_render_console_errors extracts errors from render.json."""
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(json.dumps({
        "console_errors": ["err1", "err2"],
        "pages": [],
    }))
    result = _render_console_errors(str(tmp_path))
    assert result == ["err1", "err2"]


def test_render_console_errors_empty_when_none(tmp_path: Path) -> None:
    """_render_console_errors returns [] when no console_errors key."""
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(json.dumps({"pages": []}))
    result = _render_console_errors(str(tmp_path))
    assert result == []


def test_pages_with_active_media_finds_canvas(tmp_path: Path) -> None:
    """_pages_with_active_media returns pages with canvas/echarts/svg media."""
    renders = tmp_path / "renders"
    renders.mkdir()
    (renders / "render.json").write_text(json.dumps({
        "pages": [
            {"page": 3, "geometry": {"media_inventory": {"kinds": ["canvas"]}}},
            {"page": 5, "geometry": {"media_inventory": {"kinds": ["bitmap"]}}},
            {"page": 7, "geometry": {"media_inventory": {"kinds": ["echarts", "bitmap"]}}},
        ],
    }))
    result = _pages_with_active_media(str(tmp_path))
    assert result == {3, 7}


# ═══════════════════════════════════════════════════════════════════════════
# fetch_images replace mode targets only URL-ready entries
# ═══════════════════════════════════════════════════════════════════════════


def test_fetch_images_replace_skips_entries_without_download(tmp_path: Path) -> None:
    """Replace mode calls _download_image_asset only for URL-ready entries."""
    assets = tmp_path / "assets"
    assets.mkdir()
    catalog_content = """# Image Assets

## real-with-url

- id: real-with-url
- kind: real
- source: https://example.com/photo.jpg
- download: https://example.com/photo.jpg
- path: assets/photo.jpg
- slides: [1]
- status: ready

## real-no-url

- id: real-no-url
- kind: real
- source: https://example.com/species.jpg
- download:
- path: assets/species.jpg
- slides: [2]
- status: ready
"""
    (assets / "catalog.md").write_text(catalog_content)
    (assets / "photo.jpg").write_bytes(b"old photo")
    (assets / "species.jpg").write_bytes(b"existing species")

    downloaded_ids: list[str] = []

    def fake_download(root, entry, replace):
        downloaded_ids.append(entry["id"])
        return entry["id"], "downloaded"

    with mock.patch.object(deck_core, "_download_image_asset", side_effect=fake_download):
        deck_core.fetch_images(tmp_path, replace=True)

    assert downloaded_ids == ["real-with-url"], (
        f"only URL-ready entry should be downloaded, got {downloaded_ids}"
    )
    assert (assets / "species.jpg").read_bytes() == b"existing species", (
        "no-url entry's existing file must be left untouched"
    )
