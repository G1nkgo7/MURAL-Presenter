"""验证最终页面并导出 HTML、PNG 引用、PDF、PPTX、result 与 manifest。

导出不调用模型、不修改 Plan、不激活 Owner。CSS、页面容器、PPTX 和 PDF 都读取同一个
实际 CSS 几何，任何 stale PNG、破图或 unresolved deterministic major 都会阻止闭合。
"""
from __future__ import annotations

from dataclasses import dataclass, field
import argparse
from functools import partial
import base64
import copy
from contextlib import contextmanager
from datetime import datetime, timezone
import hashlib
from html import escape, unescape
import json
import mimetypes
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import threading
import time
from typing import Mapping, Sequence
from urllib.parse import quote, unquote, urlsplit
import uuid
import zipfile

from qc.routing import IssueLedger, IssueStatus
from runtime.diagnostics import CommandFailed, run_command
from runtime.shared_objects import SharedSkillStore, load_skill_module, studio_data_root, workspace_skill_module
from runtime.workspace_paths import WorkspacePathError, WorkspacePaths
from state.artifacts import ArtifactStateError, ArtifactStore
from state.locking import workspace_lock
from state.workspace import PageLifecycle, PageSnapshot, WorkspaceState, file_sha256


DELIVERY_MANIFEST_SCHEMA = "mural.delivery.v1"
_ASSEMBLER_LOCK = threading.Lock()


class DeliveryError(RuntimeError):
    """表示最终机械完整性不足以导出。"""


@dataclass(frozen=True)
class DeliveryManifest:
    """保存最终画布、页面、QC 与交付文件哈希。"""

    status: str
    canvas: Mapping[str, object]
    pages: tuple[Mapping[str, object], ...]
    outputs: Mapping[str, Mapping[str, object]]
    image_catalog_sha256: str
    accepted_degradations: tuple[str, ...]
    path: str
    schema: str = DELIVERY_MANIFEST_SCHEMA
    speech: Mapping[str, object] = field(default_factory=dict)

    def __post_init__(self) -> None:
        """深复制嵌套交付字段。"""
        object.__setattr__(self, "canvas", copy.deepcopy(dict(self.canvas)))
        object.__setattr__(
            self, "pages", tuple(copy.deepcopy(dict(item)) for item in self.pages)
        )
        object.__setattr__(
            self,
            "outputs",
            {str(key): copy.deepcopy(dict(value)) for key, value in self.outputs.items()},
        )
        object.__setattr__(
            self, "accepted_degradations", tuple(self.accepted_degradations)
        )
        object.__setattr__(self, "speech", copy.deepcopy(dict(self.speech)))

    def to_dict(self) -> dict[str, object]:
        """把 manifest 转换为可持久 JSON。"""
        return {
            "schema": self.schema,
            "status": self.status,
            "canvas": copy.deepcopy(dict(self.canvas)),
            "pages": [copy.deepcopy(dict(item)) for item in self.pages],
            "outputs": {
                key: copy.deepcopy(dict(value)) for key, value in self.outputs.items()
            },
            "image_catalog_sha256": self.image_catalog_sha256,
            "accepted_degradations": list(self.accepted_degradations),
            "path": self.path,
            "speech": copy.deepcopy(dict(self.speech)),
        }

    @classmethod
    def from_dict(cls, value: Mapping[str, object]) -> "DeliveryManifest":
        """从 JSON 恢复并校验 delivery manifest。"""
        if value.get("schema") != DELIVERY_MANIFEST_SCHEMA:
            raise ValueError(f"unsupported delivery schema {value.get('schema')!r}")
        canvas = value.get("canvas") or {}
        pages = value.get("pages") or []
        outputs = value.get("outputs") or {}
        degradations = value.get("accepted_degradations") or []
        if not isinstance(canvas, dict) or not isinstance(outputs, dict):
            raise ValueError("delivery canvas or outputs are invalid")
        if not isinstance(pages, list) or not all(isinstance(item, dict) for item in pages):
            raise ValueError("delivery pages must be an array of objects")
        if not isinstance(degradations, list):
            raise ValueError("accepted degradations must be an array")
        if not all(isinstance(item, dict) for item in outputs.values()):
            raise ValueError("delivery outputs must be objects")
        return cls(
            status=str(value.get("status") or ""),
            canvas=canvas,
            pages=tuple(pages),
            outputs={str(key): item for key, item in outputs.items()},
            image_catalog_sha256=str(value.get("image_catalog_sha256") or ""),
            accepted_degradations=tuple(str(item) for item in degradations),
            path=str(value.get("path") or ""),
            speech=dict(value.get("speech") or {}),
        )


class DeliveryExporter:
    """从唯一 Workspace/Artifact/Issue 真源构建最终交付。"""

    def __init__(
        self,
        workspace: WorkspaceState,
        artifacts: ArtifactStore,
        issues: IssueLedger,
    ) -> None:
        """绑定只读导出依赖。"""
        self.workspace = workspace
        self.artifacts = artifacts
        self.issues = issues
        self.paths = workspace.paths
        self.manifest_path = self.paths.path("delivery_manifest")
        self.result_path = self.paths.path("result")

    def export(
        self,
        *,
        include_pdf: bool = True,
        include_pptx: bool = True,
        include_zip: bool = True,
        title: str = "MURAL Presentation",
        publish_result: bool = True,
    ) -> DeliveryManifest:
        """验证全册并原子生成选定格式和清单。"""
        try:
            config = self.workspace.creation_config(require_complete=True)
            canvas = self.workspace.require_geometry()
        except (OSError, ValueError) as exc:
            raise DeliveryError(str(exc)) from exc
        pages = self._validate_pages(_canvas=canvas)
        blocking = self.issues.blocking_issues()
        if blocking:
            raise DeliveryError(
                "delivery has unresolved major issues: "
                + ", ".join(issue.issue_id for issue in blocking)
            )
        speech = _sync_speech(self.paths)
        exporting = include_pdf or include_pptx or include_zip
        if exporting:
            self.paths.delivery.mkdir(parents=True, exist_ok=True)
        present_path = self.paths.path("present")
        _atomic_text_write(
            present_path, _present_html(self.paths, canvas, pages, title=str(title), language=config["presentation_language"])
        )
        output_paths: dict[str, Path] = {
            "present_html": present_path,
        }
        if speech.get("notes_count"):
            output_paths["speech"] = self.paths.path("speech")
        output_paths.update(_sync_contact_sheets(self.paths, pages))
        if include_pdf:
            pdf_path = self.paths.delivery / "presentation.pdf"
            self._write_pdf(pdf_path, pages, canvas)
            output_paths["pdf"] = pdf_path
        if include_pptx:
            pptx_path = self.paths.delivery / "presentation.pptx"
            self._write_pptx(pptx_path, pages, canvas)
            output_paths["pptx"] = pptx_path
        package = _build_portable_zip(self.paths, pages) if include_zip else None
        if package is not None:
            output_paths["portable_zip"] = package["path"]
        outputs = {
            name: _output_record(self.paths.root, path)
            for name, path in sorted(output_paths.items())
        }
        for name in ("html", "portable_zip"):
            if name in outputs:
                outputs[name].update(_network_dependencies(self.paths, [page.page for page in pages]))
        if package is not None:
            outputs["portable_zip"].update(
                {key: value for key, value in package.items() if key != "path"}
            )
        accepted = tuple(
            issue.issue_id
            for issue in self.issues.issues()
            if issue.status == IssueStatus.ACCEPTED_DEGRADATION
        )
        manifest = DeliveryManifest(
            status="complete",
            canvas=canvas,
            pages=tuple(page.to_dict() for page in pages),
            outputs=outputs,
            image_catalog_sha256=str(self.artifacts.catalog()["sha256"]),
            accepted_degradations=accepted,
            path=self.paths.relative(self.manifest_path) if exporting else "",
            speech=speech,
        )
        with self.workspace._lock:
            self.verify(manifest)
            if exporting:
                _atomic_json_write(self.manifest_path, manifest.to_dict())
            _publish_shared_delivery_files(
                self.paths.root,
                pages=pages,
                output_paths=tuple(output_paths.values()),
                manifest_path=self.manifest_path,
                result_path=self.result_path,
            )
            if publish_result:
                self._write_result(manifest)
        return manifest

    def write_result(self, manifest: DeliveryManifest, *, training_manifest=None, visual_review=None) -> None:
        """发布制作终态时重新校验，不能沿用导出阶段的旧结论。"""
        with self.workspace._lock:
            self.verify(manifest)
            self._write_result(manifest, training_manifest=training_manifest, visual_review=visual_review)

    def _write_result(self, manifest: DeliveryManifest, *, training_manifest=None, visual_review=None) -> None:
        """仅由同一临界区中刚完成 verify 的发布入口调用。"""
        result = {
            "schema": "mural.result.v1", "status": "completed",
            "artifact_status": manifest.status,
            "generation": {**manifest.to_dict(), "schema": "mural.generation.v1", "path": ""},
            "canvas": manifest.canvas, "page_count": len(manifest.pages),
            "accepted_degradations": list(manifest.accepted_degradations),
        }
        if manifest.path:
            result.update({"delivery_manifest": manifest.path,
                           "delivery_manifest_sha256": file_sha256(self.manifest_path)})
        if visual_review is not None:
            result["external_visual_critic"] = dict(visual_review)
            if visual_review.get("effective") and visual_review.get("critic_incomplete"):
                missing = [row for row in visual_review.get("current_inspections", []) if not row.get("completed")]
                detail = "; ".join(f"P{row['page']}: {row.get('incomplete_reason') or 'missing_current_critic'}" for row in missing)
                result.update(status="error", exit_reason="critic_incomplete", critic_incomplete_pages=missing,
                              error="视觉 Critic 执行未完成；请在同一 workspace 继续。" + detail)
        if training_manifest is not None:
            rejected = training_manifest.get("training_admission") == "rejected"
            result.update({"training_admission": training_manifest.get("training_admission"),
                "training_rejection_reasons": training_manifest.get("rejection_reasons", []),
                "training_manifest": self.paths.relative(self.paths.trace / "training/training-manifest.json"),
                "exit_reason": "training_rejected" if rejected else "completed"})
            if training_manifest.get("training_admission") == "pending":
                critic = training_manifest.get("external_visual_critic") or {}
                missing = [row for row in critic.get("current_inspections", []) if not row.get("completed")]
                detail = "; ".join(f"P{row['page']}: {row.get('incomplete_reason') or 'missing_current_critic'}"
                                   for row in missing)
                reason = "critic_incomplete" if critic.get("critic_incomplete") else "training_pending"
                result.update({"status": "error", "exit_reason": reason,
                    "error": f"训练准入尚未完成：{reason}" + (f" ({detail})" if detail else "")
                        + "；当前页面与轨迹已保留，请在同一 workspace 继续以补齐缺失评审。",
                    "critic_incomplete_pages": missing})
        _atomic_json_write(self.result_path, result)
        self.result_path.chmod(self.result_path.stat().st_mode | 0o044)

    def verify(self, manifest: DeliveryManifest | None = None) -> None:
        """重新计算所有交付哈希并确认画布与页面仍为当前版本。"""
        with self.workspace._lock:
            current = manifest or self.load_manifest()
            canvas = self.workspace.require_geometry()
            if current.canvas != canvas:
                raise DeliveryError("delivery canvas differs from actual CSS geometry")
            for name, value in current.outputs.items():
                relative = str(value.get("path") or "")
                path = self.paths.resolve(relative)
                if not path.is_file():
                    raise DeliveryError(f"delivery output is missing: {name}")
                if file_sha256(path) != str(value.get("sha256") or ""):
                    raise DeliveryError(f"delivery output hash changed: {name}")
                if name == "portable_zip":
                    _verify_portable_zip(
                        path,
                        expected_manifest_sha256=str(
                            value.get("package_manifest_sha256") or ""
                        ),
                    )
            snapshots = {page.page: page for page in self.workspace.pages()}
            if {int(page["page"]) for page in current.pages} != set(snapshots):
                raise DeliveryError("generation page membership changed")
            for value in current.pages:
                page_number = int(value["page"])
                snapshot = snapshots.get(page_number)
                if snapshot is None or snapshot.lifecycle != PageLifecycle.READY:
                    raise DeliveryError(f"delivery page is no longer ready: {page_number}")
                if (
                    snapshot.page_uid != value.get("page_uid")
                    or snapshot.inspection_id != value.get("inspection_id")
                    or snapshot.html_sha256 != str(value.get("html_sha256") or "")
                    or snapshot.png_sha256 != str(value.get("png_sha256") or "")
                ):
                    raise DeliveryError(f"delivery page hash changed: {page_number}")
            self._validate_pages(_canvas=canvas)

    def load_manifest(self) -> DeliveryManifest:
        """读取已有 delivery manifest。"""
        if not self.manifest_path.is_file():
            raise FileNotFoundError(self.manifest_path)
        value = json.loads(self.manifest_path.read_text(encoding="utf-8"))
        if not isinstance(value, dict):
            raise DeliveryError("delivery manifest root must be an object")
        return DeliveryManifest.from_dict(value)

    def _validate_pages(
        self, *, _canvas: Mapping[str, object] | None = None,
    ) -> tuple[PageSnapshot, ...]:
        """确认至少一页且每页 ready、同源并没有破图。"""
        try:
            config = self.workspace.creation_config(require_complete=True)
        except (ValueError, OSError) as exc:
            raise DeliveryError(str(exc)) from exc
        pages = self.workspace.pages()
        if not pages:
            raise DeliveryError("delivery requires at least one registered page")
        if config["target_page_count"] is not None and len(pages) != config["target_page_count"]:
            raise DeliveryError(f"config.json target_page_count={config['target_page_count']}, current pages={len(pages)}")
        canvas = self.workspace.require_geometry() if _canvas is None else _canvas
        for page in pages:
            if page.lifecycle != PageLifecycle.READY:
                raise DeliveryError(
                    f"page {page.page} is {page.lifecycle.value}, not ready"
                )
            try:
                self.artifacts.validate_render_receipt(page.page, _canvas=canvas)
            except ArtifactStateError as exc:
                raise DeliveryError(str(exc)) from exc
            receipt = self.artifacts.render_receipt(page.page)
            loaded = receipt.get("loaded_resources") or []
            if not isinstance(loaded, list):
                raise DeliveryError(
                    f"page {page.page} renderer receipt is invalid"
                )
            validation = self.artifacts.validate_html_assets(
                page.page, loaded_resources=tuple(str(item) for item in loaded)
            )
            if not validation.passed:
                raise DeliveryError(
                    f"page {page.page} has broken resources: {validation.to_dict()}"
                )
        return pages

    def _write_pdf(
        self,
        path: Path,
        pages: Sequence[PageSnapshot],
        canvas: Mapping[str, object],
    ) -> None:
        """显式即时导出也隔离字体转换；与按需导出共用原生 HTML PDF。"""
        started = time.monotonic()
        with workspace_lock(self.paths.root):
            self.paths.runtime_tmp.mkdir(parents=True, exist_ok=True)
            staging = Path(tempfile.mkdtemp(prefix="native-pdf-", dir=self.paths.runtime_tmp))
            try:
                exporter, _ = _validated_export_snapshot(self.paths, staging)
                target = exporter.paths.delivery / path.name
                _write_snapshot_pdf(exporter, target, exporter._validate_pages(), canvas,
                                    _presentation_metadata(self.paths), self.paths.root, started)
                os.replace(target, path)
            finally:
                shutil.rmtree(staging)

    def _write_pptx(
        self,
        path: Path,
        pages: Sequence[PageSnapshot],
        canvas: Mapping[str, object],
    ) -> None:
        """按实际 CSS 物理尺寸把每页当前 PNG 放入 PPTX。"""
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        presentation.slide_width = Inches(canvas["width_in"])
        presentation.slide_height = Inches(canvas["height_in"])
        blank_layout = presentation.slide_layouts[6]
        for page in pages:
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(self.paths.resolve(page.png_path)),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        pending = path.with_name(f".{path.name}.{uuid.uuid4().hex}.pending")
        try:
            presentation.save(pending)
            os.replace(pending, path)
        finally:
            pending.unlink(missing_ok=True)


def _sync_contact_sheets(paths: WorkspacePaths, pages: Sequence[PageSnapshot]) -> dict[str, Path]:
    """收尾从已验证的当前 PNG 派生总览，不重渲染，也不沿用修复前的联系表。"""
    if not paths.skill_ref.is_file():
        return {}
    ref = SharedSkillStore(studio_data_root(paths.root), create=False).load_ref(paths)
    module = load_skill_module(ref.object_root, ref.tree_sha256, "render_deck")
    renderer = module.DeckRenderer(paths.path("present"), paths.renders,
                                   workspace=paths.root, motion_check=False)
    target = renderer._contact_sheet([paths.resolve(page.png_path) for page in pages],
                                    name="contact-sheet.png", labels=[page.page for page in pages])
    if target is None:
        raise DeliveryError("final contact sheet generation failed: contact-sheet.png")
    return {"contact_sheet": target}


def _sync_speech(paths: WorkspacePaths) -> dict:
    """讲稿只从当前作者计划派生；调用者决定作者收尾或隔离导出 staging。"""
    if not paths.skill_ref.is_file() or not paths.path("deck").is_file():
        return {"status": "omitted", "notes_count": 0, "omission_reason": "speaker_note_source_unavailable"}
    ref = SharedSkillStore(studio_data_root(paths.root), create=False).load_ref(paths)
    reports = []

    def collect(line):
        if line.startswith("[speech] "):
            reports.append(json.loads(line[len("[speech] "):]))

    try:
        run_command(
            [sys.executable, str(ref.object_root / "scripts/orchestrator.py"), "sync-speech", str(paths.root)],
            cwd=paths.root, environment=dict(os.environ),
            timeout_s=int(os.environ.get("SKILL_COMMAND_TIMEOUT_S") or 1800), on_stdout_line=collect,
        )
    except (CommandFailed, ValueError) as exc:
        raise DeliveryError(f"speaker notes export failed: {exc}") from exc
    if len(reports) != 1 or not isinstance(reports[0], dict) or reports[0].get("status") not in {"complete", "partial", "omitted"}:
        raise DeliveryError("speaker notes export did not return one coverage receipt")
    return reports[0]


def _present_html(
    paths,
    canvas: Mapping[str, object],
    pages: Sequence[PageSnapshot],
    *,
    title: str,
    language: str,
) -> str:
    """Skill 片段复用固定对象的组装器；完整 HTML 后端仍用独立 frame。"""
    if paths.skill_ref.is_file():
        ref = SharedSkillStore(studio_data_root(paths.root), create=False).load_ref(paths)
        core = load_skill_module(ref.object_root, ref.tree_sha256, "deck_core")
        fragments = [
            (
                page.page,
                core._normalize_fragment_resource_paths(
                    paths.resolve(page.html_path).read_text(encoding="utf-8")
                ),
            )
            for page in pages
        ]
        document = core._document(paths.path("base_css").read_text(encoding="utf-8"), fragments, language)
        return document.replace("<title>Presentation</title>", f"<title>{escape(title)}</title>", 1)
    if any(not re.search(r"<html\b", paths.resolve(page.html_path).read_text(encoding="utf-8"), re.I) for page in pages):
        raise DeliveryError("HTML fragments require the workspace's pinned Skill assembler")
    frames = "\n".join(
        (
            f'<section class="deck-page" data-page="{page.page}">'
            f'<iframe src="{escape(page.html_path)}" '
            f'title="Page {page.page}" loading="eager"></iframe></section>'
        )
        for page in pages
    )
    return f"""<!doctype html>
<html lang="{escape('zh-CN' if language == 'zh' else language)}">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{escape(title)}</title>
  <link rel="stylesheet" href="base.css">
  <style>
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; background: #111; width: auto; height: auto; }}
    .deck-page {{ width: var(--canvas-w); height: var(--canvas-h); margin: 0 auto 24px; }}
    iframe {{ display: block; width: 100%; height: 100%; border: 0; background: white; }}
  </style>
</head>
<body>
{frames}
</body>
</html>
"""


def _output_record(root: Path, path: Path) -> dict[str, object]:
    """记录交付文件的相对路径、哈希和字节数。"""
    return {
        "path": path.resolve().relative_to(root.resolve()).as_posix(),
        "sha256": file_sha256(path),
        "bytes": path.stat().st_size,
    }


_RESOURCE_ATTRIBUTE = re.compile(r"(?is)\b(?:src|poster)\s*=\s*([\"'])(.*?)\1")
_LINK_ATTRIBUTE = re.compile(
    r"(?is)<link\b[^>]*?\bhref\s*=\s*([\"'])(.*?)\1[^>]*>"
)
_CSS_URL = re.compile(r"(?is)url\(\s*([\"']?)(.*?)\1\s*\)")
_CSS_IMPORT = re.compile(r"(?is)@import\s+([\"'])(.*?)\1")
_SVG_HREF = re.compile(r"(?is)<(?:image|use)\b[^>]*?\b(?:xlink:)?href\s*=\s*([\"'])(.*?)\1")
_SRCSET = re.compile(r"(?is)\bsrcset\s*=\s*([\"'])(.*?)\1")


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _json_object(path: Path) -> dict[str, object]:
    if not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DeliveryError(f"delivery provenance is unreadable: {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise DeliveryError(f"delivery provenance must be an object: {path}")
    return value


def _portable_provenance(paths) -> dict[str, object]:
    skill = _json_object(paths.skill_ref)
    font = _json_object(paths.font_manifest)
    source_fonts = []
    for row in [
        *(font.get("source_faces") or []),
        *(font.get("custom_fonts") or []),
    ]:
        if isinstance(row, dict):
            source_fonts.append(
                {
                    "sha256": str(row.get("sha256") or ""),
                    "scope": str(row.get("scope") or ""),
                    "family": str(row.get("family") or ""),
                    "weight": str(row.get("weight") or ""),
                    "style": str(row.get("style") or ""),
                    "license": str(row.get("license") or ""),
                }
            )
    return {
        "layout_version": paths.layout_version,
        "skill": {
            "tree_sha256": str(skill.get("tree_sha256") or ""),
            "release": str(skill.get("release") or ""),
            "object_manifest_sha256": str(skill.get("object_manifest_sha256") or ""),
            "ref_sha256": file_sha256(paths.skill_ref) if paths.skill_ref.is_file() else "",
        },
        "fonts": {
            "manifest_sha256": (
                file_sha256(paths.font_manifest) if paths.font_manifest.is_file() else ""
            ),
            "sources": source_fonts,
        },
    }


def _resource_references(path: Path) -> tuple[str, ...]:
    if path.suffix.lower() not in {".html", ".htm", ".css", ".svg"}:
        return ()
    try:
        text = path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise DeliveryError(f"portable text resource is unreadable: {path}: {exc}") from exc
    values = [match.group(2).strip() for match in _RESOURCE_ATTRIBUTE.finditer(text)]
    values.extend(match.group(2).strip() for match in _LINK_ATTRIBUTE.finditer(text))
    values.extend(match.group(2).strip() for match in _CSS_URL.finditer(text))
    values.extend(match.group(2).strip() for match in _CSS_IMPORT.finditer(text))
    values.extend(match.group(2).strip() for match in _SVG_HREF.finditer(text))
    for match in _SRCSET.finditer(text):
        values.extend(item.strip().split()[0] for item in match.group(2).split(",") if item.strip())
    return tuple(dict.fromkeys(value for value in values if value))


def _pinned_shared_font_file(root: Path, reference: str) -> Path:
    """Authoring CSS keeps file:// URIs to pinned shared objects; portable export inlines them."""
    from runtime.shared_fonts import SharedFontError, font_reference_resolver

    try:
        path, _digest = font_reference_resolver(root)(reference)
    except SharedFontError as exc:
        raise DeliveryError(f"portable package has an external runtime resource: {reference}") from exc
    if not path.is_file() or path.is_symlink():
        raise DeliveryError(f"portable package has an external runtime resource: {reference}")
    return path.resolve()


def _local_reference(base: Path, reference: str, root: Path) -> Path | None:
    parsed = urlsplit(reference)
    if parsed.scheme in {"data", "about"}:
        return None
    if parsed.scheme in {"http", "https"} and parsed.hostname and parsed.username is None and parsed.password is None:
        return None
    if parsed.scheme == "file":
        return _pinned_shared_font_file(root, reference)
    if parsed.scheme or parsed.netloc:
        raise DeliveryError(f"portable package has an external runtime resource: {reference}")
    raw = unquote(parsed.path).strip()
    if not raw:
        return None
    if raw.startswith("/"):
        raise DeliveryError(f"portable package has a root-relative resource: {reference}")
    # BUG-K:与 BUG-I 同源。真实渲染/交付都以 workspace 根为 base(根 present.html),
    # slide 片段里的 `assets/x`(以及归一前的 `../assets/x`)是根相对写法;若按片段所在
    # slides/ 目录解析会得到 slides/assets/x → 文件不存在 → 打包误报缺失。因此非绝对引用
    # 一律按 workspace 根解析,并把片段自然写法的前导 ../assets|fonts 收敛为根相对,
    # 与 state/artifacts._resolve_html_resource 和 deck_core 的 normalize 保持一致。
    fragment = base.suffix.lower() in {".html", ".htm"}
    normalized = re.sub(r"^(?:\.\./)+(assets|fonts)/", r"\1/", raw) if fragment else raw
    if fragment and normalized == "../base.css":
        normalized = "base.css"
    parent = root if fragment else base.parent
    target = parent / normalized
    for candidate in (target, *target.parents):
        if candidate == root:
            break
        if candidate.is_symlink():
            raise DeliveryError("export resource contains a symlink")
    target = Path(os.path.abspath(target))
    _bounded_path(root, target)
    target = target.resolve()
    try:
        target.relative_to(root.resolve())
    except ValueError as exc:
        raise DeliveryError(f"portable resource escapes its root: {reference}") from exc
    return target


def _staging_destination(paths, source: Path) -> Path:
    source = source.resolve()
    runtime_fonts = paths.runtime_fonts.resolve()
    try:
        font_relative = source.relative_to(runtime_fonts)
    except ValueError:
        font_relative = None
    if font_relative is not None:
        return Path("assets/fonts") / font_relative
    try:
        relative = Path(paths.relative(source))
    except WorkspacePathError:
        if source.suffix.lower() in {".ttf", ".otf", ".woff", ".woff2"}:
            return Path("assets/fonts/objects") / source.name
        raise DeliveryError(f"portable resource escapes its root: {source}")
    if relative == Path("assets/catalog.md"):
        raise DeliveryError("assets/catalog.md is authoring metadata, not a portable dependency")
    allowed = (
        relative in {Path("present.html"), Path("base.css"), Path("speech.md")}
        or (relative.parts and relative.parts[0] in {"slides", "assets"})
    )
    if not allowed:
        raise DeliveryError(f"portable resource is outside the delivery allowlist: {relative}")
    return relative


def _copy_portable_resource(paths, source: Path, staging: Path) -> Path:
    if not source.is_file() or source.is_symlink():
        raise DeliveryError(f"portable resource is missing or is a symlink: {source}")
    relative = _staging_destination(paths, source)
    target = staging / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    if source.suffix.lower() in {".html", ".htm", ".css", ".svg", ".json", ".md"}:
        try:
            text = source.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            shutil.copyfile(source, target)
        else:
            runtime_prefix = paths.relative(paths.runtime_fonts)
            if runtime_prefix != "assets/fonts":
                text = text.replace(runtime_prefix, "assets/fonts")
            _atomic_text_write(target, text)
    else:
        shutil.copyfile(source, target)
    return target


def _stage_font_bundle(paths, staging: Path) -> str:
    source_manifest = paths.runtime_fonts / "manifest.json"
    if not source_manifest.is_file():
        return ""
    manifest = _json_object(source_manifest)
    faces = manifest.get("faces") or []
    licenses = manifest.get("license_files") or []
    if not isinstance(faces, list) or not isinstance(licenses, list):
        raise DeliveryError("runtime font bundle manifest is invalid")
    adapted_faces: list[dict[str, object]] = []
    for raw in faces:
        if not isinstance(raw, dict):
            raise DeliveryError("runtime font bundle contains an invalid face")
        relative = str(raw.get("path") or "")
        source = paths.resolve(relative)
        if not source.is_file() or file_sha256(source) != str(raw.get("sha256") or ""):
            raise DeliveryError(f"runtime font subset hash mismatch: {relative}")
        target = _copy_portable_resource(paths, source, staging)
        row = dict(raw)
        row["path"] = target.relative_to(staging).as_posix()
        row.pop("source_path", None)
        row.pop("source_object_ref", None)
        adapted_faces.append(row)
    adapted_licenses: list[str] = []
    for relative in licenses:
        source = paths.resolve(str(relative))
        target = _copy_portable_resource(paths, source, staging)
        adapted_licenses.append(target.relative_to(staging).as_posix())
    adapted = dict(manifest)
    adapted["faces"] = adapted_faces
    adapted["license_files"] = adapted_licenses
    target_manifest = staging / "assets" / "fonts" / "manifest.json"
    _atomic_json_write(target_manifest, adapted)
    return file_sha256(target_manifest)


def _package_files(staging: Path) -> list[Path]:
    files: list[Path] = []
    for path in sorted(staging.rglob("*")):
        if path.is_symlink():
            raise DeliveryError(f"portable package contains a symlink: {path}")
        if not path.is_file():
            continue
        relative = path.relative_to(staging)
        allowed = (
            relative in {
                Path("present.html"),
                Path("base.css"),
                Path("speech.md"),
                Path("package-manifest.json"),
            }
            or (
                relative.parts
                and relative.parts[0] == "slides"
                and relative.suffix == ".html"
            )
            or (relative.parts and relative.parts[0] == "assets")
        )
        if not allowed:
            raise DeliveryError(f"portable package contains a non-whitelisted path: {relative}")
        files.append(path)
    return files


def _verify_staging_references(staging: Path) -> None:
    for path in _package_files(staging):
        for reference in _resource_references(path):
            target = _local_reference(path, reference, staging)
            if target is not None and (not target.is_file() or target.is_symlink()):
                raise DeliveryError(
                    "portable resource reference is unresolved: "
                    f"{path.relative_to(staging)} -> {reference}"
                )


def _build_portable_zip(paths, pages: Sequence[PageSnapshot], *, _shared_snapshot: bool = True) -> dict[str, object]:
    if _shared_snapshot and any(urlsplit(reference).scheme == 'file'
           for source in (paths.path('base_css'), paths.path('present')) if source.is_file()
           for reference in _resource_references(source)):
        return _shared_font_zip_snapshot(paths)
    paths.runtime_tmp.mkdir(parents=True, exist_ok=True)
    paths.delivery.mkdir(parents=True, exist_ok=True)
    export_id = uuid.uuid4().hex
    staging = paths.runtime_tmp / f"export-{export_id}"
    pending = paths.delivery / f".presentation.zip.{export_id}.pending"
    target = paths.path("portable_zip")
    staging.mkdir(parents=False, exist_ok=False)
    try:
        base_css = paths.path("base_css")
        if not base_css.is_file():
            raise DeliveryError(
                "portable ZIP requires the author workspace base.css; refusing to "
                "fabricate empty CSS in staging"
            )
        initial = [paths.path("present"), base_css]
        speech = paths.path("speech")
        if speech.is_file():
            initial.append(speech)
        initial.extend(paths.resolve(page.html_path) for page in pages)
        initial.extend(_observed_dependencies(paths, [page.page for page in pages]))
        queue = list(dict.fromkeys(path.resolve() for path in initial))
        copied: set[Path] = set()
        while queue:
            source = queue.pop(0)
            if source in copied:
                continue
            _copy_portable_resource(paths, source, staging)
            copied.add(source)
            for reference in _resource_references(source):
                dependency = _local_reference(source, reference, paths.root)
                if dependency is not None and dependency not in copied:
                    queue.append(dependency)
        portable_font_manifest_sha256 = _stage_font_bundle(paths, staging)
        _verify_staging_references(staging)
        inventory = [
            {
                "path": path.relative_to(staging).as_posix(),
                "sha256": file_sha256(path),
                "bytes": path.stat().st_size,
            }
            for path in _package_files(staging)
            if path.name != "package-manifest.json"
        ]
        package_manifest = {
            "schema": "mural.portable-package.v1",
            "created_at": _utc_now(),
            "files": inventory,
            "portable_font_manifest_sha256": portable_font_manifest_sha256,
            "provenance": _portable_provenance(paths),
            **_network_dependencies(paths, [page.page for page in pages]),
        }
        package_manifest_path = staging / "package-manifest.json"
        _atomic_json_write(package_manifest_path, package_manifest)
        package_manifest_sha256 = file_sha256(package_manifest_path)
        _verify_staging_references(staging)
        with zipfile.ZipFile(pending, "x", compression=zipfile.ZIP_DEFLATED) as archive:
            for path in _package_files(staging):
                archive.write(path, path.relative_to(staging).as_posix())
        with pending.open("rb") as handle:
            os.fsync(handle.fileno())
        os.replace(pending, target)
        directory_fd = os.open(paths.delivery, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0))
        try:
            os.fsync(directory_fd)
        finally:
            os.close(directory_fd)
        _verify_portable_zip(target, expected_manifest_sha256=package_manifest_sha256)
        provenance = package_manifest["provenance"]
        return {
            "path": target,
            "package_manifest_sha256": package_manifest_sha256,
            "portable_font_manifest_sha256": portable_font_manifest_sha256,
            "skill_tree_sha256": provenance["skill"]["tree_sha256"],
            "font_manifest_sha256": provenance["fonts"]["manifest_sha256"],
            **_network_dependencies(paths, [page.page for page in pages]),
        }
    finally:
        pending.unlink(missing_ok=True)
        if staging.is_dir() and staging.parent == paths.runtime_tmp:
            shutil.rmtree(staging)


def _safe_zip_name(value: str) -> bool:
    candidate = Path(value)
    return (
        bool(value)
        and not candidate.is_absolute()
        and ".." not in candidate.parts
        and value == candidate.as_posix()
    )


def _verify_portable_zip(path: Path, *, expected_manifest_sha256: str) -> None:
    try:
        with zipfile.ZipFile(path, "r") as archive:
            names = archive.namelist()
            if (
                len(names) != len(set(names))
                or not all(_safe_zip_name(name) for name in names)
                or "package-manifest.json" not in names
            ):
                raise DeliveryError("portable ZIP has unsafe/duplicate members or no package manifest")
            manifest_bytes = archive.read("package-manifest.json")
            actual_manifest_sha256 = hashlib.sha256(manifest_bytes).hexdigest()
            if expected_manifest_sha256 and actual_manifest_sha256 != expected_manifest_sha256:
                raise DeliveryError("portable ZIP package manifest hash changed")
            manifest = json.loads(manifest_bytes.decode("utf-8"))
            if not isinstance(manifest, dict) or manifest.get("schema") != "mural.portable-package.v1":
                raise DeliveryError("portable ZIP package manifest schema is invalid")
            declared = manifest.get("files") or []
            if not isinstance(declared, list):
                raise DeliveryError("portable ZIP file inventory is invalid")
            expected_names = {"package-manifest.json"}
            for row in declared:
                if not isinstance(row, dict):
                    raise DeliveryError("portable ZIP file record is invalid")
                name = str(row.get("path") or "")
                if not _safe_zip_name(name):
                    raise DeliveryError(f"portable ZIP contains an unsafe member: {name}")
                expected_names.add(name)
                data = archive.read(name)
                if hashlib.sha256(data).hexdigest() != str(row.get("sha256") or ""):
                    raise DeliveryError(f"portable ZIP member hash mismatch: {name}")
                expected_bytes = row.get("bytes")
                if expected_bytes is None or len(data) != int(expected_bytes):
                    raise DeliveryError(f"portable ZIP member size mismatch: {name}")
            if set(names) != expected_names:
                raise DeliveryError("portable ZIP member inventory differs from its manifest")
    except (OSError, zipfile.BadZipFile, KeyError, json.JSONDecodeError) as exc:
        raise DeliveryError(f"portable ZIP is invalid: {exc}") from exc


def _publish_shared_delivery_files(
    root: Path,
    *,
    pages: Sequence[PageSnapshot],
    output_paths: Sequence[Path],
    manifest_path: Path,
    result_path: Path,
) -> None:
    """让非敏感交付文件可共享读取，并保持输入与内部轨迹权限不变。"""
    workspace = root.resolve()
    directories = (
        workspace,
        workspace / "assets",
        workspace / "slides",
        workspace / "renders",
        workspace / "delivery",
    )
    files = {
        *(path.resolve() for path in output_paths),
        manifest_path.resolve(),
        result_path.resolve(),
        workspace / "base.css",
        workspace / "speech.md",
        *(workspace / page.html_path for page in pages),
        *(workspace / page.png_path for page in pages),
    }
    for directory in directories[1:]:
        if not directory.is_dir() or directory.is_symlink():
            continue
        for path in directory.rglob("*"):
            if path.is_symlink():
                continue
            if path.is_dir():
                path.chmod(path.stat().st_mode | 0o055)
            elif path.is_file():
                files.add(path.resolve())
    for directory in directories:
        if directory.is_dir() and not directory.is_symlink():
            directory.chmod(directory.stat().st_mode | 0o055)
    for path in files:
        if path.is_file() and not path.is_symlink():
            path.chmod(path.stat().st_mode | 0o044)


def _atomic_text_write(path: Path, value: str) -> None:
    """原子写入 UTF-8 交付文本。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    pending = path.with_name(f".{path.name}.{uuid.uuid4().hex}.pending")
    try:
        with pending.open("x", encoding="utf-8") as handle:
            handle.write(str(value))
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(pending, path)
    finally:
        pending.unlink(missing_ok=True)


def _atomic_json_write(path: Path, value: Mapping[str, object]) -> None:
    """原子写入排序后的 JSON 交付文件。"""
    payload = json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    _atomic_text_write(path, payload)


_EXPORT_MEDIA = {
    "html": "text/html", "pdf": "application/pdf", "zip": "application/zip",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pptx-editable": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_EXPORT_SCHEMA = "mural.on-demand-export.v1"


def _bounded_path(root: Path, path: Path) -> Path:
    """导出不跟随目录或文件软链接，即使其最终仍指向工作区内。"""
    try:
        relative = path.relative_to(root)
        path.resolve().relative_to(root.resolve())
    except ValueError as exc:
        raise DeliveryError("export resource escapes workspace") from exc
    if ".." in relative.parts:
        raise DeliveryError("export resource escapes workspace")
    current = root
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            raise DeliveryError("export resource contains a symlink")
    return path


def _snapshot_for_export(paths: WorkspacePaths, staging: Path) -> None:
    # 不在作者目录实例化会初始化/恢复/写入的 Store；只复制明确的验证输入。
    state = _json_object(_bounded_path(paths.root, paths.path("workspace_manifest")))
    records = state.get("pages") or {}
    if not isinstance(records, dict) or not records:
        raise DeliveryError("export requires registered pages")
    required = ["_state/workspace.json", "_state/artifacts.json", "_state/issues.json",
                "base.css", "config.json"]
    optional = ["_state/skill-ref.json", "_state/font-manifest.json", "_state/run.json", "font-intent.json", "deck.md", "speech.md"]
    optional.extend(path.relative_to(paths.root).as_posix() for path in (paths.root / "briefs").glob("slide_*.md"))
    for record in records.values():
        if not isinstance(record, dict):
            raise DeliveryError("export page record is invalid")
        if not record.get("retired"):
            required.extend(str(record[key]) for key in ("html_path", "png_path"))
    initial = [paths.root / relative for relative in required]
    initial.extend(paths.root / relative for relative in optional if (paths.root / relative).exists())
    initial.extend(_observed_dependencies(paths, [int(key) for key, row in records.items() if not row.get("retired")]))
    queue = list(dict.fromkeys(initial))
    copied: set[Path] = set()
    resolve_shared_font = None
    checked_fonts: set[str] = set()
    while queue:
        source = _bounded_path(paths.root, queue.pop(0))
        if source in copied:
            continue
        if not source.is_file():
            raise DeliveryError(f"export input is missing: {paths.relative(source)}")
        relative = source.relative_to(paths.root)
        target = staging / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        copied.add(source)
        for reference in _resource_references(source):
            if urlsplit(reference).scheme == "file":
                if reference not in checked_fonts:
                    from runtime.shared_fonts import font_reference_resolver
                    if resolve_shared_font is None:
                        resolve_shared_font = font_reference_resolver(paths.root)
                    resolve_shared_font(reference)
                    checked_fonts.add(reference)
                continue
            dependency = _local_reference(source, reference, paths.root)
            if dependency is not None:
                _staging_destination(paths, dependency)
                queue.append(dependency)
    # 自定义字体已有 subset manifest 的许可证等不是 HTML 依赖，但仍是验证输入。
    if not any(source.is_relative_to(paths.runtime_fonts) for source in copied):
        return
    manifest_path = _bounded_path(paths.root, paths.runtime_fonts / "manifest.json")
    font_bundle = _json_object(manifest_path)
    bundle_files = [paths.relative(manifest_path)] if manifest_path.is_file() else []
    bundle_files.extend(font_bundle.get("license_files") or [])
    bundle_files.extend(str(row.get('path') or '') for row in font_bundle.get('faces') or [])
    for relative in bundle_files:
        source = _bounded_path(paths.root, paths.root / str(relative))
        _staging_destination(paths, source)
        target = staging / source.relative_to(paths.root)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)


def _network_dependencies(paths, page_numbers) -> dict:
    receipts = _json_object(paths.root / "_state/artifacts.json").get("render_receipts", {})
    references = {ref for number in page_numbers for ref in receipts.get(str(number), {}).get("loaded_resources", [])
                  if urlsplit(ref).scheme in {"http", "https"}}
    for source in (paths.path("base_css"), paths.path("present"), *(paths.page_html(number) for number in page_numbers)):
        if source.is_file():
            references.update(ref for ref in _resource_references(source) if urlsplit(ref).scheme in {"http", "https"})
    return {"requires_network": bool(references), "remote_resources": sorted(references)}


def _observed_dependencies(paths, page_numbers) -> list[Path]:
    receipts = _json_object(_bounded_path(paths.root, paths.root / "_state/artifacts.json")).get("render_receipts", {})
    dependencies = set()
    for number in page_numbers:
        receipt = receipts.get(str(number), {})
        for relative in receipt.get("resource_hashes", {}):
            if urlsplit(relative).scheme in {"file", "http", "https"}:
                continue  # 固定共享字体由既有便携字体打包器提供。
            source = _bounded_path(paths.root, paths.root / relative)
            _staging_destination(paths, source)
            dependencies.add(source)
    return sorted(dependencies)


def _export_resources(paths, pages) -> dict:
    # 本地资源只服务已验证快照中的精确文件；公网走单独的共用读取边界。
    return {
        "https://mural-export.invalid/" + quote(path.relative_to(paths.root).as_posix(), safe="/"):
            (mimetypes.guess_type(path.name)[0] or "application/octet-stream", path.read_bytes())
        for path in _observed_dependencies(paths, [page.page for page in pages])
    }


def _single_file_html(path: Path, root: Path, active: frozenset[Path] = frozenset()) -> str:
    """递归内嵌本地资源，远程依赖保留 URL 并在交付回执中说明需要联网。"""
    if path in active:
        raise DeliveryError("export resource dependency cycle")
    active = active | {path}
    text = path.read_text(encoding="utf-8")

    def embedded(reference: str) -> str:
        target = _local_reference(path, reference, root)
        if target is None:
            return reference
        if not target.is_file():
            raise DeliveryError("export resource is missing")
        _staging_destination(WorkspacePaths.from_root(root), target)
        if target.suffix.lower() in {".html", ".htm", ".css", ".svg"}:
            data = _single_file_html(target, root, active).encode("utf-8")
        else:
            data = target.read_bytes()
        mime = mimetypes.guess_type(target.name)[0] or "application/octet-stream"
        suffix = urlsplit(reference).fragment
        return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}" + (f"#{suffix}" if suffix else "")

    def substitute(match):
        start, end = match.start(2) - match.start(), match.end(2) - match.start()
        value = match.group(0)
        return value[:start] + embedded(match.group(2)) + value[end:]

    for pattern in (_RESOURCE_ATTRIBUTE, _LINK_ATTRIBUTE, _CSS_URL, _CSS_IMPORT, _SVG_HREF):
        text = pattern.sub(substitute, text)

    def srcset(match):
        candidates = []
        for candidate in match.group(2).split(","):
            parts = candidate.strip().split(maxsplit=1)
            if parts:
                candidates.append(embedded(parts[0]) + (" " + parts[1] if len(parts) > 1 else ""))
        return match.group(0).replace(match.group(2), ", ".join(candidates), 1)

    return _SRCSET.sub(srcset, text)


def _presentation_metadata(paths: WorkspacePaths) -> dict[str, str]:
    path = _bounded_path(paths.root, paths.path("present"))
    text = path.read_text(encoding="utf-8") if path.is_file() else ""
    title = re.search(r"(?is)<title\b[^>]*>(.*?)</title>", text)
    return {"title": unescape(title.group(1)) if title else "MURAL Presentation",
            "language": workspace_skill_module(paths.root, "deck_config").read_config(
                paths.root, require_complete=True)["presentation_language"]}


def _embedded_font_licenses(paths: WorkspacePaths) -> str:
    manifest = _json_object(paths.runtime_fonts / 'manifest.json')
    licenses = {}
    for relative in manifest.get('license_files') or []:
        path = _bounded_path(paths.root, paths.root / str(relative))
        licenses[_staging_destination(paths, path).as_posix()] = path.read_text(encoding='utf-8')
    if not licenses:
        return ''
    payload = json.dumps(licenses, ensure_ascii=False, sort_keys=True).replace('<', '\\u003c')
    return f'<script type="application/json" id="mural-font-licenses">{payload}</script>\n'


def _prepare_portable_document(paths: WorkspacePaths, canvas: Mapping[str, object],
                               pages: Sequence[PageSnapshot], metadata: Mapping[str, str]) -> None:
    if paths.skill_ref.is_file():
        if not paths.font_manifest.is_file():
            raise DeliveryError("portable export requires pinned font manifest")
        ref = SharedSkillStore(studio_data_root(paths.root), create=False).load_ref(paths)
        fonts = load_skill_module(ref.object_root, ref.tree_sha256, "font_bundle")
        fonts.bundle_workspace(paths.root, from_plans=False)
    _atomic_text_write(paths.path('present'), _present_html(paths, canvas, pages, **metadata))


_INLINE_STYLE_ATTRIBUTE = re.compile(
    r"(?is)(?P<prefix>\bstyle\s*=\s*)(?P<quote>[\"'])(?P<value>.*?)(?P=quote)"
)
_BASE_STYLESHEET_LINK = re.compile(
    r"(?is)<link\b[^>]*\bhref\s*=\s*[\"'](?:\.\./)?base\.css(?:[?#][^\"']*)?[\"'][^>]*>"
)


def _normalize_page_for_export(path: Path, page: int) -> None:
    """Normalize only the disposable export snapshot, never authored sources."""
    original = path.read_text(encoding="utf-8")

    def terminate_inline_style(match: re.Match[str]) -> str:
        value = match.group("value")
        normalized = value if not value.strip() or value.rstrip().endswith(";") else value.rstrip() + ";"
        return f"{match.group('prefix')}{match.group('quote')}{normalized}{match.group('quote')}"

    text = _INLINE_STYLE_ATTRIBUTE.sub(terminate_inline_style, original)
    if re.search(r"<html\b", text, re.I):
        additions = []
        if not _BASE_STYLESHEET_LINK.search(text):
            additions.append('<link rel="stylesheet" href="base.css">')
        additions.append(
            '<style data-mural-export-normalization>'
            f'.slide[data-slide="{page:02d}"],.slide[data-slide="{page}"]'
            '{visibility:visible!important;opacity:1!important;pointer-events:auto!important}'
            '</style>'
        )
        insertion = "\n".join(additions)
        text, count = re.subn(r"(?i)</head\s*>", insertion + "\n</head>", text, count=1)
        if count != 1:
            raise DeliveryError(f"standalone page {page} has no closing head element")
    if text != original:
        _atomic_text_write(path, text)


def _single_page_documents(paths, canvas, pages, metadata):
    """PDF 和可编辑 PPTX 消费同一份单页内嵌快照，不另设作者格式。"""
    for page in pages:
        _normalize_page_for_export(paths.resolve(page.html_path), page.page)
    _prepare_portable_document(paths, canvas, pages, metadata)
    documents = []
    for page in pages:
        html_path = paths.resolve(page.html_path)
        if not re.search(r"<html\b", html_path.read_text(encoding="utf-8"), re.I):
            html_path = paths.path("present")
            _atomic_text_write(html_path, _present_html(paths, canvas, [page], **metadata))
        documents.append((page.page, _single_file_html(html_path, paths.root)))
    return documents


@contextmanager
def _export_render_slot(root, started):
    from runtime.production import ProductionSettings, ResourceSlots, ResourceSlotTimeout
    settings = ProductionSettings.from_environment({key: os.environ[key] for key in
        ("MURAL_RUNTIME_CONFIG", "SKILL_COMMAND_TIMEOUT_S") if key in os.environ})
    deadline = started + settings.command_timeout_s
    # 顺序与生产渲染一致：工作区锁在外，共享渲染槽在内；缓存命中不占槽。
    try:
        with ResourceSlots(root, settings).acquire("render", deadline_monotonic=deadline):
            yield deadline
    except ResourceSlotTimeout as exc:
        raise DeliveryError("export render queue timeout") from exc


def _write_snapshot_pdf(exporter, target, pages, canvas, metadata, slot_root, started):
    from export.native_pdf import write_native_pdf
    documents = _single_page_documents(exporter.paths, canvas, pages, metadata)
    with _export_render_slot(slot_root, started) as deadline:
        write_native_pdf(documents, target, canvas, scratch=exporter.paths.runtime_tmp,
                         deadline_monotonic=deadline, resources=_export_resources(exporter.paths, pages),
                         remote_fetch=partial(workspace_skill_module(exporter.paths.root, "remote_resources").fetch, root=exporter.paths.root))


def _validated_export_snapshot(paths: WorkspacePaths, staging: Path) -> tuple[DeliveryExporter, dict]:
    _snapshot_for_export(paths, staging)
    state = WorkspaceState(staging)
    if state.paths.font_manifest.is_file():
        from runtime.shared_fonts import SharedFontStore
        SharedFontStore(paths.root, create=False).verify_manifest(state.paths.font_manifest)
    exporter = DeliveryExporter(state, ArtifactStore(staging, workspace_state=state), IssueLedger(staging))
    if exporter.issues.blocking_issues():
        raise DeliveryError("export has unresolved major issues")
    speech = _sync_speech(state.paths)
    return exporter, speech


def _shared_font_zip_snapshot(paths: WorkspacePaths) -> dict[str, object]:
    """显式/训练 ZIP 也走同一便携快照，不让共享字体转换回写作者 CSS。"""
    for path in (paths.runtime_tmp, paths.path('portable_zip')):
        _bounded_path(paths.root, path)
    with workspace_lock(paths.root):
        paths.runtime_tmp.mkdir(parents=True, exist_ok=True)
        staging = Path(tempfile.mkdtemp(prefix='on-demand-', dir=paths.runtime_tmp))
        try:
            exporter, _ = _validated_export_snapshot(paths, staging)
            pages = exporter._validate_pages()
            _prepare_portable_document(exporter.paths, exporter.workspace.require_geometry(), pages,
                                       _presentation_metadata(paths))
            package = _build_portable_zip(exporter.paths, pages, _shared_snapshot=False)
            target = paths.path('portable_zip')
            os.replace(package['path'], target)
            fd = os.open(paths.delivery, os.O_RDONLY | getattr(os, 'O_DIRECTORY', 0))
            try:
                os.fsync(fd)
            finally:
                os.close(fd)
            return {**package, 'path': target}
        finally:
            if staging.is_dir() and staging.parent == paths.runtime_tmp:
                shutil.rmtree(staging)


def export_on_demand(workspace: str | Path, format: str) -> dict[str, object]:
    """离线导出已完成的生成快照；固定交付路径，manifest 按内容哈希复用。"""
    started = time.monotonic()
    format = str(format).lower()
    if format not in _EXPORT_MEDIA:
        raise DeliveryError("unsupported export format")
    paths = WorkspacePaths.from_root(workspace)
    for path in (paths.state, paths.state / 'state.lock', paths.runtime_tmp, paths.delivery):
        _bounded_path(paths.root, path)
    if not paths.path("workspace_manifest").is_file():
        raise DeliveryError("export requires an existing workspace")
    with workspace_lock(paths.root):
        paths.runtime_tmp.mkdir(parents=True, exist_ok=True)
        staging = Path(tempfile.mkdtemp(prefix="on-demand-", dir=paths.runtime_tmp))
        try:
            exporter, speech = _validated_export_snapshot(paths, staging)
            result = _json_object(paths.path("result"))
            generation = result.get("generation")
            if (result.get("status") not in {"completed", "needs_improvement"}
                    or not isinstance(generation, dict)
                    or generation.get("schema") != "mural.generation.v1"
                    or generation.get("status") != "complete"):
                raise DeliveryError("export requires a completed generation receipt")
            # 只读校验作者文件；staging 不复制 present，所以在作者目录验证输出哈希。
            for record in generation.get("outputs", {}).values():
                source = _bounded_path(paths.root, paths.root / str(record.get("path", "")))
                if not source.is_file() or source.stat().st_size != record.get("bytes") or file_sha256(source) != record.get("sha256"):
                    raise DeliveryError("generation output is missing or stale")
            snapshot_proof = DeliveryManifest.from_dict({
                **generation, "schema": DELIVERY_MANIFEST_SCHEMA, "outputs": {}, "path": "", "speech": speech,
            })
            exporter.verify(snapshot_proof)
            metadata = _presentation_metadata(paths)
            state = exporter.workspace
            pages = state.pages()
            canvas = state.require_geometry()
            inputs = {path.relative_to(staging).as_posix(): file_sha256(path)
                      for path in staging.rglob('*') if path.is_file()
                      and path.relative_to(staging).as_posix() not in {
                          '_state/state.lock', '_state/workspace.json', '_state/artifacts.json', '_state/issues.json'}}
            identity = {"schema": _EXPORT_SCHEMA, "format": format,
                        "exporter_sha256": file_sha256(Path(__file__)),
                        "pages": [page.to_dict() for page in pages], "inputs": inputs, "metadata": metadata}
            if format == "pptx-editable":
                from export.editable import source_identity
                identity["editable_converter"] = source_identity()
            elif format in {"pdf", "html"}:
                from export.native_pdf import source_identity
                identity["pdf_converter" if format == "pdf" else "html_validator"] = source_identity()
            source_hash = hashlib.sha256(json.dumps(identity, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
            manifest_path = _bounded_path(paths.root, paths.path("delivery_manifest"))
            previous = _json_object(manifest_path)
            if previous and (previous.get("schema") != DELIVERY_MANIFEST_SCHEMA
                             or previous.get("status") != "complete"):
                raise DeliveryError("export cache manifest is invalid")
            generation_hash = hashlib.sha256(json.dumps(generation, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
            outputs = previous.get("outputs", {}) if previous.get("generation_sha256") == generation_hash else {}
            if not isinstance(outputs, dict):
                raise DeliveryError("export cache outputs are invalid")
            output_key = "portable_zip" if format == "zip" else format
            cached = outputs.get(output_key)
            filename = "presentation.editable.pptx" if format == "pptx-editable" else f"presentation.{format}"
            destination = _bounded_path(paths.root, paths.delivery / filename)
            report_destination = _bounded_path(paths.root, paths.delivery / "presentation.editable.editability.json") if format == "pptx-editable" else None
            if isinstance(cached, dict) and cached.get("source_sha256") == source_hash:
                if (cached.get("path") != paths.relative(destination) or not destination.is_file()
                        or file_sha256(destination) != cached.get("sha256")
                        or destination.stat().st_size != cached.get("bytes")):
                    raise DeliveryError("export cache artifact is missing or damaged")
                if report_destination is not None:
                    report_record = cached.get("editability_report")
                    if (not isinstance(report_record, dict) or report_record.get("path") != paths.relative(report_destination)
                            or not report_destination.is_file() or file_sha256(report_destination) != report_record.get("sha256")
                            or report_destination.stat().st_size != report_record.get("bytes")):
                        raise DeliveryError("export cache editability report is missing or damaged")
                return {**cached, "cached": True}

            stage_paths = state.paths
            stage_paths.delivery.mkdir(exist_ok=True)
            target = stage_paths.delivery / filename
            report = None
            if format == "pdf":
                _write_snapshot_pdf(exporter, target, pages, canvas, metadata, paths.root, started)
            elif format == "pptx":
                exporter._write_pptx(target, pages, canvas)
            elif format == "pptx-editable":
                from export.editable import write_editable_pptx
                documents = _single_page_documents(stage_paths, canvas, pages, metadata)
                bundle = _json_object(stage_paths.runtime_fonts / "manifest.json")
                font_names = {face["delivery_family"]:face["source_family"] for face in bundle.get("faces", [])
                              if face.get("delivery_family") and face.get("source_family")}
                with _export_render_slot(paths.root, started):
                    report = write_editable_pptx(documents, target, canvas, scratch=stage_paths.runtime_tmp, fonts=font_names,
                                                 resources=_export_resources(stage_paths, pages),
                                                 remote_fetch=partial(workspace_skill_module(stage_paths.root, "remote_resources").fetch, root=stage_paths.root))
                report_target = stage_paths.delivery / report_destination.name
                _atomic_json_write(report_target, report)
            else:
                _prepare_portable_document(stage_paths, canvas, pages, metadata)
                if format == "html":
                    html = _single_file_html(stage_paths.path("present"), staging)
                    html = html.replace('</body>', _embedded_font_licenses(stage_paths) + '</body>', 1)
                    from export.native_pdf import validate_single_file_html
                    with _export_render_slot(paths.root, started) as deadline:
                        validate_single_file_html(html, canvas, scratch=stage_paths.runtime_tmp, deadline_monotonic=deadline,
                                                  page_numbers=[page.page for page in pages],
                                                  remote_fetch=partial(workspace_skill_module(stage_paths.root, "remote_resources").fetch, root=stage_paths.root))
                    _atomic_text_write(target, html)
                else:
                    _build_portable_zip(stage_paths, pages, _shared_snapshot=False)
            paths.delivery.mkdir(exist_ok=True)
            with target.open("rb") as handle:
                os.fsync(handle.fileno())
            target.chmod(target.stat().st_mode | 0o044)
            os.replace(target, destination)
            record = {"format": format, **_output_record(paths.root, destination),
                      "filename": filename, "media_type": _EXPORT_MEDIA[format],
                      "source_sha256": source_hash}
            if format in {"html", "zip"}:
                record.update(_network_dependencies(stage_paths, [page.page for page in pages]))
            if report is not None:
                report_target.chmod(report_target.stat().st_mode | 0o044)
                os.replace(report_target, report_destination)
                record["editability_report"] = _output_record(paths.root, report_destination)
                record["summary"] = report["summary"]
            outputs[output_key] = record
            manifest = {**snapshot_proof.to_dict(), "outputs": outputs,
                        "path": paths.relative(manifest_path), "generation_sha256": generation_hash}
            _atomic_json_write(manifest_path, manifest)
            manifest_path.chmod(manifest_path.stat().st_mode | 0o044)
            paths.delivery.chmod(paths.delivery.stat().st_mode | 0o055)
            fd = os.open(paths.delivery, os.O_RDONLY | getattr(os, 'O_DIRECTORY', 0))
            try:
                os.fsync(fd)
            finally:
                os.close(fd)
            return {**record, "cached": False}
        finally:
            if staging.is_dir() and staging.parent == paths.runtime_tmp:
                shutil.rmtree(staging)


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export an inspected MURAL workspace without running agents.")
    parser.add_argument("--workspace", required=True)
    parser.add_argument("--format", required=True, choices=tuple(_EXPORT_MEDIA))
    args = parser.parse_args(argv)
    try:
        result = export_on_demand(args.workspace, args.format)
    except (RuntimeError, ValueError, OSError) as exc:
        detail = str(exc).lower()
        code, category = next(((code, value) for key, code, value in (
            ('searchable pdf text', 'export_pdf_text_unavailable', 'CSS effects prevent searchable text in this PDF; choose another format.'),
            ('native pdf render timeout', 'export_render_timeout', 'Native PDF conversion timed out; retry later.'),
            ('render queue timeout', 'export_render_timeout', 'Timed out waiting for the shared render slot; retry later.'),
            ('cache', 'export_cache_invalid', 'The cached export is invalid.'),
            ('font', 'export_font_invalid', 'Pinned fonts are unavailable or invalid.'),
            ('ready', 'export_not_ready', 'Pages are not ready for export.'),
            ('stale', 'export_source_stale', 'The inspected source has changed.'),
            ('issue', 'export_qc_blocked', 'Unresolved major issues prevent export.'),
        ) if key in detail), ('export_source_invalid', 'The source or its resources failed export validation.'))
        if 'editable' in detail and not any(word in detail for word in ('cache', 'resource', 'snapshot')):
            page_match = re.search(r'editable page (\d+)', detail)
            page_label = f" on page {page_match.group(1)}" if page_match else ""
            code, category = 'export_editability_unsupported', f'Editable objects cannot preserve this content{page_label}. Choose high-fidelity PPTX.'
        diagnostic = f"{type(exc).__name__}: {exc}"[:600]
        print(json.dumps({"error": "delivery_export_failed", "code": code,
                          "message": f"{category} {diagnostic}"}), file=sys.stderr)
        return 1
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
