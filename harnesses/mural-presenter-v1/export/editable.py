"""在已验证的导出快照内读取真实 DOM，写原生 PPTX 与局部栅格覆盖报告。"""
from __future__ import annotations

import base64
import hashlib
import http.client
from contextlib import contextmanager
from io import BytesIO
import os
from pathlib import Path
import tempfile
from urllib.parse import unquote_to_bytes, urlsplit
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


REPORT_SCHEMA = "mural.pptx-editability.v1"
DOM_SOURCE = Path(__file__).with_name("editable_dom.js")
_ORIGIN = "https://mural-export.invalid/slide"


def source_identity() -> dict[str, str]:
    from state.workspace import file_sha256
    from importlib.metadata import version
    return {"adapter": file_sha256(Path(__file__)), "dom": file_sha256(DOM_SOURCE),
            "python-pptx": version("python-pptx"), "playwright": version("playwright")}


def _rgba(target, rgba, opacity=1.0):
    target.rgb = RGBColor(*rgba[:3])
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(round(rgba[3] / 255 * opacity * 100000)))
    target._color._xClr.append(alpha)


def _box(box):
    return tuple(Inches(float(box[key]) / 96) for key in ("x", "y", "w", "h"))


def _font(run, style, fonts):
    run.font.name = fonts.get(style["font"].split(",")[0].strip("\"' "), style["font"].split(",")[0].strip("\"' "))
    run.font.size = Pt(style["size"] * 0.75)
    run.font.bold, run.font.italic = style["bold"], style["italic"]
    run.font.underline = style["underline"]
    _rgba(run.font.color, style["color"], style.get("opacity", 1))
    properties = run._r.get_or_add_rPr()
    properties.set("spc", str(round(style.get("letter_spacing", 0) * 75)))
    if style.get("strike"):
        properties.set("strike", "sngStrike")
    for tag in ("a:ea", "a:cs"):
        face = OxmlElement(tag)
        face.set("typeface", run.font.name)
        properties.append(face)


def _text(slide, row, fonts):
    shape = slide.shapes.add_textbox(*_box(row["box"]))
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = row["text"]
    _font(run, row["style"], fonts)
    return shape


def _shape(slide, row):
    box = row["box"]
    radius = row["radius"]
    kind = MSO_AUTO_SHAPE_TYPE.RECTANGLE
    if row["ellipse"]:
        kind = MSO_AUTO_SHAPE_TYPE.OVAL
    elif radius:
        kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
    shape = slide.shapes.add_shape(kind, *_box(box))
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    for effect in shape._element.xpath("./p:style/a:effectRef"):
        effect.set("idx", "0")
    if kind == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
        shape.adjustments[0] = min(0.5, radius / min(box["w"], box["h"]))
    if row["fill"][3]:
        shape.fill.solid()
        _rgba(shape.fill.fore_color, row["fill"], row["opacity"])
    else:
        shape.fill.background()
    borders = row["borders"]
    if row["uniform"] and borders[0]["width"]:
        shape.line.width = Pt(borders[0]["width"] * 0.75)
        _rgba(shape.line.color, borders[0]["color"], row["opacity"])
    else:
        shape.line.fill.background()
        x, y, w, h = _box(box)
        for border, coords in zip(borders, ((x,y,x+w,y), (x+w,y,x+w,y+h), (x,y+h,x+w,y+h), (x,y,x,y+h))):
            if border["width"]:
                line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *coords)
                line.line.width = Pt(border["width"] * 0.75)
                _rgba(line.line.color, border["color"], row["opacity"])
    return shape


def _edges(cells, axis, dimension):
    values = sorted(value for cell in cells for value in (cell["box"][axis], cell["box"][axis]+cell["box"][dimension]))
    result = []
    for value in values:
        if not result or abs(value-result[-1]) > 0.75:
            result.append(value)
    return result


def _table(slide, row, fonts):
    cells = row["cells"]
    xs, ys = _edges(cells, "x", "w"), _edges(cells, "y", "h")
    shape = slide.shapes.add_table(len(ys)-1, len(xs)-1, Inches(xs[0]/96), Inches(ys[0]/96), Inches((xs[-1]-xs[0])/96), Inches((ys[-1]-ys[0])/96))
    table = shape.table
    table.first_row = table.horz_banding = False
    for column, left, right in zip(table.columns, xs, xs[1:]):
        column.width = Inches((right-left)/96)
    for table_row, top, bottom in zip(table.rows, ys, ys[1:]):
        table_row.height = Inches((bottom-top)/96)
    for source in cells:
        box = source["box"]
        col = min(range(len(xs)), key=lambda i:abs(xs[i]-box["x"]))
        end_col = min(range(len(xs)), key=lambda i:abs(xs[i]-box["x"]-box["w"]))-1
        start = min(range(len(ys)), key=lambda i:abs(ys[i]-box["y"]))
        end = min(range(len(ys)), key=lambda i:abs(ys[i]-box["y"]-box["h"]))-1
        cell = table.cell(start, col)
        if end != start or end_col != col:
            cell.merge(table.cell(end, end_col))
        cell.margin_top, cell.margin_right, cell.margin_bottom, cell.margin_left = [Pt(value*0.75) for value in source["padding"]]
        cell.vertical_anchor = {"middle":MSO_ANCHOR.MIDDLE, "bottom":MSO_ANCHOR.BOTTOM}.get(source["valign"], MSO_ANCHOR.TOP)
        cell.fill.solid()
        _rgba(cell.fill.fore_color, source["fill"])
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.alignment = {"center":PP_ALIGN.CENTER,"right":PP_ALIGN.RIGHT,"end":PP_ALIGN.RIGHT}.get(source["align"],PP_ALIGN.LEFT)
        paragraph.space_before = paragraph.space_after = Pt(0)
        runs = source["runs"]
        if "".join(run["text"] for run in runs).split() != source["text"].split():
            runs = [{"text":source["text"],"style":source["style"]}]
        for value in runs:
            run = paragraph.add_run()
            run.text = value["text"]
            _font(run, value["style"], fonts)
        for tag, border in zip(("lnT","lnR","lnB","lnL"), source["borders"]):
            line = OxmlElement(f"a:{tag}")
            line.set("w", str(round(border["width"] * 9525)))
            fill = OxmlElement("a:solidFill" if border["width"] else "a:noFill")
            if border["width"]:
                rgb = OxmlElement("a:srgbClr")
                rgb.set("val", "".join(f"{v:02X}" for v in border["color"][:3]))
                fill.append(rgb)
            line.append(fill)
            cell._tc.get_or_add_tcPr().append(line)
    return shape


def _image_bytes(src, observed=None):
    if observed is not None and src in observed:
        return observed[src]
    if not src.startswith("data:image/"):
        raise ValueError("editable image is outside the embedded resource snapshot")
    header, payload = src.split(",", 1)
    return base64.b64decode(payload, validate=True) if ";base64" in header else unquote_to_bytes(payload)


def _paint_order_and_fonts(page, pseudo_ids):
    session = page.context.new_cdp_session(page)
    try:
        value = session.send("DOMSnapshot.captureSnapshot", {"computedStyles":[], "includePaintOrder":True})
        document, strings = value["documents"][0], value["strings"]
        attributes = document["nodes"]["attributes"]
        result, backends, families, pseudos = {}, {}, {}, {}
        for index, order in zip(document["layout"]["nodeIndex"], document["layout"]["paintOrders"]):
            attrs = attributes[index]
            for i in range(0, len(attrs), 2):
                if strings[attrs[i]] == "data-mural-export-id":
                    result[strings[attrs[i+1]]] = order
                    backends[document["nodes"]["backendNodeId"][index]] = strings[attrs[i+1]]
        session.send("DOM.enable")
        session.send("CSS.enable")
        session.send("DOM.getDocument")
        nodes = session.send("DOM.pushNodesByBackendIdsToFrontend", {"backendNodeIds":list(backends)})["nodeIds"]
        for backend, node in zip(backends,nodes):
            if not node:
                continue
            fonts = session.send("CSS.getPlatformFontsForNode", {"nodeId":node})["fonts"]
            if fonts:
                families[backends[backend]] = max(fonts,key=lambda font:font["glyphCount"])["familyName"]
            if backends[backend] in pseudo_ids:
                details=session.send("DOM.describeNode",{"nodeId":node})["node"]
                bounds=[]
                for pseudo in details.get("pseudoElements",[]):
                    if pseudo["pseudoType"] not in {"before","after"}:
                        continue
                    model=session.send("DOM.getBoxModel",{"backendNodeId":pseudo["backendNodeId"]})["model"]["border"]
                    bounds.append({"type":pseudo["pseudoType"],"x":min(model[::2]),"y":min(model[1::2]),
                                   "right":max(model[::2]),"bottom":max(model[1::2])})
                pseudos[backends[backend]]=bounds
        return result, families, pseudos
    finally:
        session.detach()


@contextmanager
def _isolate(page, row):
    # 只改浏览器的临时 DOM；保留布局/继承，隐藏其他绘制层，避免截图重复烘焙邻居。
    page.evaluate("""row => {
      const target=document.querySelector(`[data-mural-export-id="${row.dom_id}"]`);
      window.__muralSaved=[];
      for (const el of document.querySelectorAll('*')) {
        window.__muralSaved.push([el,el.getAttribute('style')]);
        if (el!==target && !el.contains(target) && !(target.contains(el)&&!row.decoration))
          el.style.setProperty('visibility','hidden','important');
        if (el.contains(target) && el!==target && !row.keep_ancestor_background) {
          el.style.setProperty('background','transparent','important');
          el.style.setProperty('border-color','transparent','important');
          el.style.setProperty('box-shadow','none','important');
        }
      }
      if (row.decoration) {
        const before=getComputedStyle(target,'::before').color, after=getComputedStyle(target,'::after').color;
        target.style.setProperty('color','transparent','important');
        const rule=document.createElement('style'); rule.id='mural-pseudo-color';
        rule.textContent=`[data-mural-export-id="${row.dom_id}"]::before{color:${before}!important}[data-mural-export-id="${row.dom_id}"]::after{color:${after}!important}`;
        document.head.append(rule);
      }
    }""", row)
    try:
        yield
    finally:
        page.evaluate("""() => {
          document.getElementById('mural-pseudo-color')?.remove();
          for (const [el,style] of window.__muralSaved) style===null ? el.removeAttribute('style') : el.setAttribute('style',style);
          delete window.__muralSaved;
        }""")


def _raster(page, row, root_box, width, height):
    box, inset = row.get("capture_box", row["box"]), row["inset"]
    x, y = box["x"]-inset, box["y"]-inset
    right, bottom = box["x"]+box["w"]+inset, box["y"]+box["h"]+inset
    for pseudo in row.get("pseudo_bounds",[]):
        extent=row["pseudo_insets"].get(pseudo["type"],0)
        x,y=min(x,pseudo["x"]-root_box["x"]-extent),min(y,pseudo["y"]-root_box["y"]-extent)
        right,bottom=max(right,pseudo["right"]-root_box["x"]+extent),max(bottom,pseudo["bottom"]-root_box["y"]+extent)
    x,y,right,bottom=max(0,x),max(0,y),min(width,right),min(height,bottom)
    bounds = {"x":x,"y":y,"w":right-x,"h":bottom-y}
    if bounds["w"]<=0 or bounds["h"]<=0:
        return None, bounds
    with _isolate(page, row):
        data = page.screenshot(clip={"x":x+root_box["x"],"y":y+root_box["y"],"width":right-x,"height":bottom-y}, omit_background=True, animations="disabled")
    return data, bounds


def inspect_pptx(path: Path) -> dict[str, int]:
    ns = {"p":"http://schemas.openxmlformats.org/presentationml/2006/main", "a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
    counts = {"slides":0,"text_runs":0,"pictures":0,"tables":0,"shapes":0}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                document = etree.fromstring(archive.read(name))
                counts["slides"] += 1
                for key, xpath in (("text_runs",".//a:t"),("pictures",".//p:pic"),("tables",".//a:tbl"),("shapes",".//p:sp")):
                    counts[key] += len(document.findall(xpath,ns))
    return counts


def write_editable_pptx(documents, target: Path, canvas, *, scratch: Path, fonts=None, resources=None, remote_fetch=None) -> dict:
    """快照本地资源不越域；公网依赖由固定 Skill 的共同读取边界提供。"""
    from playwright.sync_api import Error as BrowserError, sync_playwright
    from state.workspace import file_sha256
    fonts = fonts or {}
    resources = resources or {}
    width, height = int(canvas["width_px"]), int(canvas["height_px"])
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(canvas["width_in"]), Inches(canvas["height_in"])
    report = {"schema":REPORT_SCHEMA,"mode":"editable","converter":source_identity(),"canvas":dict(canvas),"pages":[],
              "font_policy":"Font family names are preserved, but fonts are not embedded. The recipient needs the listed fonts.","fonts":[]}
    total = {key:0 for key in ("native_text","native_images","native_shapes","native_tables","local_raster_fallbacks","pages_with_warnings")}
    scratch.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="editable-browser-",dir=scratch) as temporary, sync_playwright() as pw:
        env = {**os.environ, **dict.fromkeys(("TMPDIR","TMP","TEMP"),temporary)}
        options = {"headless":True,"env":env,"args":["--no-sandbox","--disable-dev-shm-usage"]}
        if env.get("PPT_SKILL_BROWSER_EXE"):
            options["executable_path"] = env["PPT_SKILL_BROWSER_EXE"]
        browser = pw.chromium.launch(**options)
        try:
            for number, html in documents:
                context = browser.new_context(viewport={"width":width,"height":height},device_scale_factor=2,accept_downloads=False,service_workers="block")
                failures = []
                observed = {}
                context.route_web_socket("**/*", lambda socket: failures.append("external-websocket"))
                def route(request_route):
                    if request_route.request.url == _ORIGIN and request_route.request.is_navigation_request():
                        request_route.fulfill(status=200,content_type="text/html",body=html)
                    elif request_route.request.method == "GET" and request_route.request.url.split("?", 1)[0] in resources:
                        media_type, data = resources[request_route.request.url.split("?", 1)[0]]
                        observed[request_route.request.url] = data
                        request_route.fulfill(status=200, content_type=media_type, body=data)
                    elif (remote_fetch is not None and request_route.request.method == "GET"
                          and urlsplit(request_route.request.url).scheme in {"http", "https"}
                          and urlsplit(request_route.request.url).hostname != "mural-export.invalid"):
                        try:
                            result = remote_fetch(request_route.request.url)
                            if result["status"] >= 400:
                                failures.append(f"HTTP {result['status']}: {request_route.request.url}")
                            observed[request_route.request.url] = result["body"]
                            request_route.fulfill(status=result["status"], headers=result["headers"], body=result["body"])
                        except (OSError, ValueError, http.client.HTTPException) as exc:
                            failures.append(str(exc))
                            request_route.abort()
                    else:
                        failures.append("external-resource")
                        request_route.abort()
                context.route("**/*",route)
                page = context.new_page()
                page.set_default_timeout(30000)
                try:
                    page.goto(_ORIGIN,wait_until="load")
                    page.evaluate("document.fonts.ready")
                    page.evaluate("Promise.all([...document.images].filter(i => i.getClientRects().length).map(i => i.decode()))")
                    if page.evaluate("[...document.fonts].some(f => f.status === 'error')"):
                        raise ValueError(f"editable page {number} has unavailable fonts")
                    page.wait_for_timeout(3000)
                    page.evaluate("() => { for (const a of document.getAnimations()) a.pause(); }")
                    extracted = page.evaluate(DOM_SOURCE.read_text(),{"width":width,"height":height})
                    if failures:
                        raise ValueError(f"editable page {number} requested a resource outside its verified snapshot")
                    orders, actual_fonts, pseudo_bounds = _paint_order_and_fonts(page,{row["dom_id"] for row in extracted["rows"] if row.get("pseudo_insets")})
                    rows = sorted(enumerate(extracted["rows"]),key=lambda pair:(orders.get(pair[1]["dom_id"],0),pair[0]))
                    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                    counts = {key:0 for key in total if key!="pages_with_warnings"}
                    fallbacks, objects, warnings = [], [], []
                    for _, row in rows:
                        kind = row["type"]
                        if kind == "text":
                            row["style"]["font"] = actual_fonts.get(row["dom_id"], row["style"]["font"])
                            shape = _text(slide,row,fonts)
                            counts["native_text"] += 1
                        elif kind == "shape":
                            shape = _shape(slide,row)
                            counts["native_shapes"] += 1
                        elif kind == "table":
                            for cell in row["cells"]:
                                for run in cell["runs"]:
                                    run["style"]["font"] = actual_fonts.get(run["dom_id"],run["style"]["font"])
                                    family = fonts.get(run["style"]["font"],run["style"]["font"])
                                    if family not in report["fonts"]:
                                        report["fonts"].append(family)
                            shape = _table(slide,row,fonts)
                            counts["native_tables"] += 1
                        elif kind == "image":
                            shape = slide.shapes.add_picture(BytesIO(_image_bytes(row["src"], observed)),*_box(row["box"]))
                            counts["native_images"] += 1
                        elif kind == "raster":
                            row["pseudo_bounds"]=pseudo_bounds.get(row["dom_id"],[])
                            data, bounds = _raster(page,row,extracted["root_box"],width,height)
                            if data is None:
                                continue
                            shape = slide.shapes.add_picture(BytesIO(data),*_box(bounds))
                            counts["local_raster_fallbacks"] += 1
                            fallbacks.append({"selector":row["selector"],"box":bounds,"reasons":row["reasons"],
                                              "role":"decoration" if row["decoration"] else "subtree",
                                              "text_not_editable":bool(row["text"].strip()),"text_characters":len(row["text"]),
                                              "sha256":hashlib.sha256(data).hexdigest()})
                        else:
                            raise ValueError(f"editable unsupported object type: {kind}")
                        shape.name = f"{kind}: {row['selector']}"
                        objects.append({"shape_id":shape.shape_id,"type":kind,"selector":row["selector"],"box":row["box"]})
                        if kind == "text":
                            family=shape.text_frame.paragraphs[0].runs[0].font.name
                            if family not in report["fonts"]:
                                report["fonts"].append(family)
                            if row["style"]["weight"] not in (400,700):
                                warnings.append(f"{row['selector']}: CSS weight {row['style']['weight']} maps to Office's regular/bold font switch.")
                    if not counts["native_text"] and not counts["native_tables"] and any(f["text_not_editable"] for f in fallbacks):
                        raise ValueError(f"editable page {number} has no editable text; use the high-fidelity PPTX option")
                    if not counts["native_text"] and not counts["native_tables"] and any(
                            f["role"] == "subtree" and f["box"]["w"]*f["box"]["h"] >= width*height*0.95 for f in fallbacks):
                        raise ValueError(f"editable page {number} would be a full-page raster; use the high-fidelity PPTX option")
                    if not objects:
                        raise ValueError(f"editable page {number} contains no accounted visual objects")
                    if fallbacks:
                        warnings.append("Local effects were rasterized; see each fallback's selector and reason.")
                    if not counts["native_text"] and not counts["native_tables"]:
                        warnings.append("This page has no native text or table objects.")
                    if warnings:
                        total["pages_with_warnings"] += 1
                    for key, value in counts.items():
                        total[key] += value
                    report["pages"].append({"page":number,"counts":counts,"warnings":warnings,"fallbacks":fallbacks,
                                            "objects":objects,"visible_text_nodes":extracted["visible_text_nodes"]})
                except BrowserError as exc:
                    raise ValueError(f"editable page {number} cannot be mapped by Chromium; use the high-fidelity PPTX option") from exc
                finally:
                    context.close()
        finally:
            browser.close()
    target.parent.mkdir(parents=True,exist_ok=True)
    presentation.save(target)
    report["summary"] = total
    report["ooxml"] = inspect_pptx(target)
    report["pptx_sha256"] = file_sha256(target)
    if report["ooxml"]["slides"] != len(documents):
        raise ValueError("editable output slide count differs from the inspected source")
    return report
