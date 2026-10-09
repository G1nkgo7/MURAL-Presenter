from __future__ import annotations

import base64
from io import BytesIO
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
import zipfile

import fitz
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation
import pytest

from export.editable import inspect_pptx, write_editable_pptx
from geometry_fixture import geometry, require_browser


REPO = Path(__file__).resolve().parents[3]


@pytest.fixture
def browser_environment(monkeypatch, tmp_path):
    sys.path.insert(0, str(REPO / "webui/studio"))
    from app.engine import render_env
    environment = render_env(dict(os.environ))
    for key in ("LD_LIBRARY_PATH", "PLAYWRIGHT_BROWSERS_PATH", "PPT_SKILL_BROWSER_EXE", "FONTCONFIG_FILE"):
        if key in environment:
            monkeypatch.setenv(key, environment[key])
    require_browser(tmp_path)
    return environment


@pytest.fixture
def websocket_server():
    requests = []
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            key = self.headers.get("Sec-WebSocket-Key", "")
            accept = base64.b64encode(hashlib.sha1(
                (key + "258EAFA5-E914-47DA-95CA-C5AB0DC85B11").encode()).digest()).decode()
            self.send_response(101)
            self.send_header("Upgrade", "websocket")
            self.send_header("Connection", "Upgrade")
            self.send_header("Sec-WebSocket-Accept", accept)
            self.end_headers()
            self.close_connection = True

        def log_message(self, *_args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"ws://127.0.0.1:{server.server_port}/snapshot-escape", requests
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def sample_html(width, height):
    image = BytesIO()
    Image.new("RGB", (80, 60), (25, 115, 155)).save(image, format="PNG")
    data = base64.b64encode(image.getvalue()).decode()
    return f'''<!doctype html><html><head><style>
      html,body{{margin:0;width:{width}px;height:{height}px;font-family:"DejaVu Sans",sans-serif}}
      .mural-slide{{width:{width}px;height:{height}px;background:#fff;position:relative;overflow:hidden}}
      h1{{position:absolute;left:40px;top:24px;font-size:28px;margin:0;line-height:1.2;font-weight:400}}
      h1 strong{{font-weight:700;color:#186580}}
      .box{{position:absolute;left:40px;top:92px;width:130px;height:66px;background:#186580;border:2px solid #104050;border-radius:12px}}
      img{{position:absolute;left:200px;top:92px;width:88px;height:66px}}
      .badge{{position:absolute;left:322px;top:92px;width:180px;height:66px;background:linear-gradient(90deg,#edf3ff,#cae9f7);border-radius:8px}}
      .badge span{{position:absolute;top:20px;left:15px;font-size:18px}}
      table{{position:absolute;left:40px;top:210px;width:462px;height:132px;border-collapse:collapse;table-layout:fixed;font-size:18px}}
      td,th{{border:1px solid #236580;padding:10px;text-align:left;font-weight:400}}
      th{{background:#e9f3f8;font-weight:700}}
      svg{{position:absolute;left:545px;top:96px;width:60px;height:60px}}
      .rotated{{position:absolute;left:548px;top:240px;font-size:16px;transform:rotate(-9deg)}}
    </style></head><body><main class="mural-slide">
      <h1>Editable <strong>Quarterly Results</strong></h1>
      <div class="box"></div><img src="data:image/png;base64,{data}">
      <div class="badge"><span>Native badge text</span></div>
      <table><tr><th>Metric</th><th>Value</th></tr><tr><td>Revenue</td><td><b>128</b> million</td></tr></table>
      <svg viewBox="0 0 60 60"><circle cx="30" cy="30" r="27" fill="#de8754"/></svg>
      <p class="rotated">Local effect</p>
    </main></body></html>'''


@pytest.mark.parametrize("size", [(960,540),(800,600),(720,960)])
def test_real_editable_objects_and_aspect_ratios(tmp_path, browser_environment, size):
    width,height=size
    output=tmp_path / "editable.pptx"
    report=write_editable_pptx([(1,sample_html(width,height))],output,geometry(width_px=width,height_px=height),scratch=tmp_path / "scratch")
    deck=Presentation(output)
    assert deck.slide_width/914400 == pytest.approx(width/96)
    assert deck.slide_height/914400 == pytest.approx(height/96)
    facts=inspect_pptx(output)
    assert facts["text_runs"]>=6 and facts["pictures"]>=3 and facts["tables"]==1
    assert report["summary"]["native_text"]>=3
    assert report["summary"]["native_images"]==1
    assert report["summary"]["native_shapes"]>=2
    assert report["summary"]["native_tables"]==1
    reasons={reason for item in report["pages"][0]["fallbacks"] for reason in item["reasons"]}
    assert {"svg","transform","background-image-or-gradient"} <= reasons
    with zipfile.ZipFile(output) as archive:
        xml=archive.read("ppt/slides/slide1.xml").decode()
        assert "Quarterly Results" in xml and "Revenue" in xml and "128" in xml
        assert 'b="1"' in xml and "186580" in xml
    assert not list((tmp_path / "scratch").iterdir())


def test_full_page_rasterized_text_is_not_editable_success(tmp_path, browser_environment):
    html=sample_html(960,540).replace('.mural-slide{','.mural-slide{filter:blur(1px);')
    with pytest.raises(ValueError,match="no editable text"):
        write_editable_pptx([(3,html)],tmp_path/'no.pptx',geometry(width_px=960,height_px=540),scratch=tmp_path/'scratch')
    assert not (tmp_path/'no.pptx').exists()


def test_export_browser_cannot_fetch_outside_snapshot(tmp_path, browser_environment):
    html=sample_html(960,540).replace('</body>','<script>fetch("https://outside.invalid/private")</script></body>')
    with pytest.raises(ValueError,match="outside its verified snapshot"):
        write_editable_pptx([(1,html)],tmp_path/'no.pptx',geometry(width_px=960,height_px=540),scratch=tmp_path/'scratch')


def test_export_browser_cannot_connect_websocket(tmp_path, browser_environment, websocket_server):
    endpoint, requests = websocket_server
    html = sample_html(960,540).replace('</body>', f'<script>new WebSocket({json.dumps(endpoint)})</script></body>')
    with pytest.raises(ValueError, match="outside its verified snapshot"):
        write_editable_pptx([(1,html)],tmp_path/'no.pptx',geometry(width_px=960,height_px=540),scratch=tmp_path/'scratch')
    assert requests == []
    assert not (tmp_path/'no.pptx').exists()
    assert not list((tmp_path/'scratch').iterdir())


@pytest.mark.parametrize('hidden_style',['opacity:0','visibility:hidden'])
def test_invisible_full_canvas_pseudo_keeps_native_objects(tmp_path,browser_environment,hidden_style):
    html = sample_html(960,540).replace('</style>', f'''
      .mural-slide::before {{content:"";position:absolute;inset:0;z-index:0;{hidden_style};
        background:radial-gradient(58% 78% at 8% 8%,#176b87,transparent 72%);}}
    </style>''')
    output = tmp_path/'editable.pptx'
    report = write_editable_pptx([(1,html)],output,geometry(width_px=960,height_px=540),scratch=tmp_path/'scratch')
    assert report['summary']['native_text'] >= 3
    assert report['summary']['native_images'] == report['summary']['native_tables'] == 1
    assert not any(row['selector']=='slide' for row in report['pages'][0]['fallbacks'])
    assert 'Quarterly Results' in '\n'.join(shape.text for shape in Presentation(output).slides[0].shapes if shape.has_text_frame)


def test_table_row_paint_and_outside_pseudo_shadow_are_not_omitted(tmp_path,browser_environment):
    html='''<style>html,body{margin:0;width:960px;height:540px}.slide{width:960px;height:540px;position:relative;background:white}
    table{position:absolute;left:30px;top:100px;border-collapse:collapse;width:500px}td{padding:20px;font:24px sans-serif}tr{background:#e00000}
    .decor{position:absolute;left:600px;top:100px;width:200px;height:120px}
    .decor:before{content:"";position:absolute;inset:0;box-shadow:0 0 0 20px blue}</style>
    <main class="slide"><h1>Header</h1><table><tr><td>A</td><td>B</td></tr></table><div class="decor"></div></main>'''
    output=tmp_path/'editable.pptx'
    report=write_editable_pptx([(1,html)],output,geometry(width_px=960,height_px=540),scratch=tmp_path/'scratch')
    fallbacks=report['pages'][0]['fallbacks']
    table=next(row for row in fallbacks if 'complex-table' in row['reasons'])
    shadow=next(row for row in fallbacks if 'pseudo-element-composite' in row['reasons'])
    assert table['text_not_editable'] is True
    assert shadow['box']=={'x':580,'y':80,'w':240,'h':160}
    images=[Image.open(BytesIO(shape.image.blob)).convert('RGBA') for shape in Presentation(output).slides[0].shapes if shape.shape_type==13]
    palettes=[{rgba:count for count,rgba in image.getcolors(image.width*image.height)} for image in images]
    assert any(colors.get((224,0,0,255),0)>100000 for colors in palettes)
    assert any(colors.get((0,0,255,255),0)>50000 for colors in palettes)


def _libreoffice_pdf(pptx, directory, env):
    directory.mkdir()
    profile=directory / 'profile'
    result=subprocess.run(['/mnt/afs/hejiatong/.local/bin/libreoffice',f'-env:UserInstallation={profile.as_uri()}',
        '--headless','--convert-to','pdf','--outdir',str(directory),str(pptx)],env=env,capture_output=True,text=True,timeout=60)
    output=directory / (pptx.stem+'.pdf')
    assert result.returncode==0 and output.is_file(), (result.stdout,result.stderr)
    with fitz.open(output) as document:
        pixmap=document[0].get_pixmap(matrix=fitz.Matrix(96/72,96/72),alpha=False)
        image=Image.frombytes('RGB',(pixmap.width,pixmap.height),pixmap.samples)
        text=document[0].get_text()
    image.save(directory/'render.png')
    return image,text


@pytest.mark.parametrize('size',[(960,540),(800,600),(720,960)])
def test_real_libreoffice_render_and_text_edit(tmp_path, browser_environment,size):
    from playwright.sync_api import sync_playwright
    width,height=size
    html=sample_html(width,height)
    output=tmp_path/'editable.pptx'
    report=write_editable_pptx([(1,html)],output,geometry(width_px=width,height_px=height),scratch=tmp_path/'scratch')
    env={**browser_environment,'TMPDIR':str(tmp_path),'PYTHONDONTWRITEBYTECODE':'1'}
    original,text=_libreoffice_pdf(output,tmp_path/'before',env)
    assert 'Quarterly Results' in text and 'Revenue' in text
    with sync_playwright() as pw:
        browser=pw.chromium.launch(executable_path=env['PPT_SKILL_BROWSER_EXE'],env=env,headless=True,args=['--no-sandbox'])
        try:
            page=browser.new_page(viewport={'width':width,'height':height})
            page.set_content(html)
            page.evaluate('document.fonts.ready')
            reference=Image.open(BytesIO(page.screenshot())).convert('RGB')
            reference.save(tmp_path/'reference.png')
        finally:
            browser.close()
    delta=ImageChops.difference(original,reference)
    mean=sum(ImageStat.Stat(delta).mean)/3
    # Office 字体度量并非浏览器像素等价；原始差异指标保留，不能用整页图遮掩。
    assert mean < 20
    assert all(abs(actual-expected)<3 for actual,expected in zip(original.getpixel((240,120)),(25,115,155)))
    assert all(abs(actual-expected)<3 for actual,expected in zip(original.getpixel((80,120)),(24,101,128)))
    deck=Presentation(output)
    changed=False
    for shape in deck.slides[0].shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if 'Quarterly Results' in run.text:
                        run.text=run.text.replace('Quarterly Results','Edited Results')
                        changed=True
    assert changed
    edited=tmp_path/'edited.pptx'
    deck.save(edited)
    after,edited_text=_libreoffice_pdf(edited,tmp_path/'after',env)
    assert 'Edited Results' in edited_text and 'Quarterly Results' not in edited_text
    edit_box=ImageChops.difference(after,original).getbbox()
    assert edit_box is not None and edit_box[0]>=140 and edit_box[1]>=20 and edit_box[3]<80
    (tmp_path/'editability.json').write_text(json.dumps({**report,'mean_rgb_difference':mean},indent=2))


def test_real_libreoffice_native_table_edit(tmp_path, browser_environment):
    output=tmp_path/'editable.pptx'
    report=write_editable_pptx([(1,sample_html(960,540))],output,
        geometry(width_px=960,height_px=540),scratch=tmp_path/'scratch')
    env={**browser_environment,'TMPDIR':str(tmp_path),'PYTHONDONTWRITEBYTECODE':'1'}
    original,text=_libreoffice_pdf(output,tmp_path/'before',env)
    assert 'Revenue' in text and '128' in text
    deck=Presentation(output)
    table=next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert len(table.rows)==len(table.columns)==2
    table.cell(1,0).text_frame.paragraphs[0].runs[0].text='Growth'
    number=next(run for run in table.cell(1,1).text_frame.paragraphs[0].runs if run.text=='128')
    assert number.font.bold
    number.text='256'
    edited=tmp_path/'table-edited.pptx'
    deck.save(edited)
    after,edited_text=_libreoffice_pdf(edited,tmp_path/'after',env)
    assert 'Growth' in edited_text and '256' in edited_text
    assert 'Revenue' not in edited_text and '128' not in edited_text
    changed=ImageChops.difference(after,original).getbbox()
    assert changed is not None and changed[0]>=40 and changed[1]>=210 and changed[2]<=505 and changed[3]<=345
    assert report['summary']['native_tables']==1
    (tmp_path/'editability.json').write_text(json.dumps(report,indent=2))
