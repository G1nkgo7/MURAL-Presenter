"""字体族变量与字重等排版变量必须在作者/便携链路保持不同类型。"""
from __future__ import annotations

import json
import os
from pathlib import Path
import sys
import zipfile

import pytest

from test_editable_pptx import browser_environment
from test_v1_font_pipeline import font_bundle, pinned_font_sources, _pin_font_manifest


@pytest.mark.parametrize("css", [
    ':root{--font-body:"Manrope",sans-serif;}',
    ':root{--font-body:"Manrope",sans-serif}',
    '--font-body:"Manrope",sans-serif',
])
def test_final_font_declaration_does_not_require_a_semicolon(css):
    assert font_bundle._token_definitions(css) == {'--font-body': '"Manrope",sans-serif'}
    stacks = font_bundle._active_token_stacks(css, ['<h1 style="font:48px var(--font-body)">Title</h1>'])
    assert stacks['--font-body'][0] == 'Manrope'


def test_final_numeric_declaration_remains_a_weight_not_a_family():
    css = ':root{--font-display:"Inter";--font-display-weight:700}'
    assert font_bundle._token_definitions(css)['--font-display-weight'] == '700'
    stacks = font_bundle._active_token_stacks(css, [
        '<h1 style="font:var(--font-display-weight) 48px var(--font-display)">Title</h1>'])
    assert stacks['--font-display'][0] == 'Inter'
    assert '--font-display-weight' not in stacks


@pytest.mark.parametrize("value", ["700", "64px", "1.2", "0.02em", "italic", "calc(2 * 24px)"])
def test_non_family_tokens_are_not_selected_from_page_or_plan_text(value):
    css = ':root{--font-cover:var(--font-heading);--font-heading:"Inter";--font-detail:' + value + ';}'
    markup = '<h1 style="font-family:var(--font-cover)">Cover</h1>'
    markup += '<span style="font-weight:var(--font-detail)">Detail</span>'
    stacks = font_bundle._active_token_stacks(css, [markup, "Use --font-detail for the design."])
    assert "--font-detail" not in stacks
    assert stacks["--font-cover"] == ["Inter", "Noto Sans SC"]


@pytest.mark.parametrize("usage", [
    '.title{font-family:var( --font-cover );font-weight:var(--font-detail)}',
    '<style>.title{font-family:var( --font-cover );font-weight:var(--font-detail)}</style>',
    '<h1 style="font-family:var( --font-cover )" data-note="var(--font-detail)">Cover</h1>',
])
def test_family_usage_in_base_inline_and_embedded_css_preserves_alias_and_cjk(usage):
    css = ':root{--font-cover:var( --font-heading );--font-heading:"Inter";--font-detail:700;}'
    stacks = font_bundle._active_token_stacks(css + (usage if not usage.startswith('<') else ''),
                                               [usage] if usage.startswith('<') else [])
    assert stacks['--font-cover'] == ['Inter', 'Noto Sans SC']
    assert '--font-detail' not in stacks
    assert set(stacks) <= {'--font-cover', '--font-heading'}


@pytest.mark.parametrize('family,expected', [('"Inter"', 'Inter'), ('"InterVariable"', 'Inter'),
                                           ('monospace', 'IBM Plex Mono')])
def test_font_shorthand_classifies_alias_family_without_guessing_numeric_tokens(family, expected):
    css = (':root{--font-cover:var(--font-heading);--font-heading:' + family + ';'
        '--font-weight:700;--font-size:64px;--font-line-height:1.2;}')
    fragment = '<h1 style="font:var(--font-weight) var(--font-size)/var(--font-line-height) var(--font-cover)">Cover</h1>'
    stacks = font_bundle._active_token_stacks(css, [fragment])
    assert stacks['--font-cover'][0] == expected
    assert not {'--font-weight', '--font-size', '--font-line-height'} & stacks.keys()


def test_non_family_cycle_in_shorthand_does_not_override_valid_size_fallback():
    css = ':root{--font-cover:"Inter";--font-size:var(--font-size);}'
    stacks = font_bundle._active_token_stacks(css, ['<p style="font:var(--font-size, 64px) var(--font-cover)">Hi</p>'])
    assert stacks['--font-cover'] == ['Inter', 'Noto Sans SC']
    assert '--font-size' not in stacks


def test_custom_font_variables_preserve_case_nested_fallback_and_other_types(tmp_path, pinned_font_sources):
    from state.workspace import WorkspaceState

    WorkspaceState(tmp_path)
    _pin_font_manifest(tmp_path, pinned_font_sources)
    css = (':root{--Cover-Body:"Inter";--cover-body:"Noto Serif SC";'
           '--Cover-Weight:700;--Cover-Color:#123456;--unused-family:"Missing Unused"}')
    fragment = ('<style>.body{font-family:var(--missing, var(--Cover-Body));'
                'font-weight:var(--Cover-Weight);color:var(--Cover-Color)}'
                '.title{font-family:var(--cover-body)}</style><h1>Title 中文</h1>')
    definitions = font_bundle._token_definitions(css)
    assert definitions['--Cover-Body'] == '"Inter"'
    assert definitions['--cover-body'] == '"Noto Serif SC"'
    identity = font_bundle._bundle_identity(tmp_path, css, [fragment])
    assert identity[2] == {}
    assert identity[5] == {'Inter', 'Noto Serif SC', 'Noto Sans SC'}


@pytest.mark.parametrize('value', ['var(--Missing)', 'var(--custom)', 'var(--CUSTOM)'])
def test_custom_font_variables_still_reject_unavailable_stack(tmp_path, pinned_font_sources, value):
    from state.workspace import WorkspaceState

    WorkspaceState(tmp_path)
    _pin_font_manifest(tmp_path, pinned_font_sources)
    css = ':root{--custom:"Missing Font";--Custom:"Inter"}'
    with pytest.raises(ValueError, match='font_stack_unavailable'):
        font_bundle._bundle_identity(tmp_path, css, [f'<h1 style="font-family:{value}">Title</h1>'])


def test_truncated_font_num_token_aliases_defined_number_role(tmp_path, pinned_font_sources):
    from state.workspace import WorkspaceState

    WorkspaceState(tmp_path)
    _pin_font_manifest(tmp_path, pinned_font_sources)
    css = ':root{--font-number:"Inter";--font-body:"Inter"}'
    identity = font_bundle._bundle_identity(
        tmp_path, css, ['<b style="font-family:var(--font-num)">12%</b>'])
    assert 'Inter' in identity[5]


def test_custom_font_stack_keeps_available_fallback_when_first_family_is_missing(tmp_path, pinned_font_sources):
    from state.workspace import WorkspaceState

    WorkspaceState(tmp_path)
    _pin_font_manifest(tmp_path, pinned_font_sources)
    css = ':root{--Cover-Body:"Missing Primary", var(--Body);--Body:"Inter"}'
    identity = font_bundle._bundle_identity(tmp_path, css, ['<h1 style="font-family:var(--Cover-Body)">Hi 中文</h1>'])
    assert identity[2] == {}
    assert identity[5] == {'Inter', 'Noto Sans SC'}


@pytest.mark.parametrize('value', ['bold', 'italic', 'calc(2 * 24px)', 'clamp(32px, 3vw, 64px)'])
def test_shorthand_non_family_css_values_do_not_become_font_roles(value):
    css = ':root{--font-cover:"Inter";--font-detail:' + value + ';}'
    stacks = font_bundle._active_token_stacks(css, [
        '<h1 style="font:var(--font-detail) var(--font-cover)">Title</h1>'])
    assert stacks['--font-cover'] == ['Inter', 'Noto Sans SC']
    assert '--font-detail' not in stacks


def _write_author_inputs(root):
    (root / 'deck.md').write_text('# Deck\n## Brief\n- language: zh\n\n## Visual Direction\nUse a bold cover.\n')
    (root / 'briefs').mkdir(exist_ok=True)
    (root / 'briefs/slide_01.md').write_text('# P01 — Cover\n- page_type: cover\n- job: Introduce\n- must_show: Bold 中文\n- medium: typography\n')
    (root / 'base.css').write_text('''
    :root{--canvas-w:1600px;--canvas-h:1200px;--font-display:"Inter";
      --font-cover:var(--font-display);--font-display-weight:700;--font-display-size:64px;
      --font-title-leading:1.2;--font-tracking:0.02em}
    html,body{margin:0}.slide,.mural-slide{width:var(--canvas-w);height:var(--canvas-h);background:white}
    h1{position:absolute;left:80px;top:100px;margin:0;font-family:var(--font-cover);
      font-weight:var(--font-display-weight);font-size:var(--font-display-size);
      line-height:var(--font-title-leading);letter-spacing:var(--font-tracking)}
    ''')
    (root / 'slides/slide_01.html').write_text('<!doctype html><html><head><meta charset="utf-8">'
        '<link rel="stylesheet" href="../base.css"></head><body>'
        '<main class="mural-slide"><h1>Bold 中文</h1></main></body></html>')


def _weight(page):
    return page.locator('h1').evaluate('e => ({weight:getComputedStyle(e).fontWeight, '
        'size:getComputedStyle(e).fontSize, family:getComputedStyle(e).fontFamily})')


@pytest.mark.parametrize('authored_stack', [
    'var(--font-cover)', '"Inter"', '"InterVariable", "Inter"',
    'var(--Cover-Body, "Inter")',
])
def test_real_authoring_and_portable_keep_numeric_tokens_and_variable_subset(
    tmp_path, monkeypatch, pinned_font_sources, browser_environment, authored_stack,
):
    from playwright.sync_api import sync_playwright
    from runtime.shared_objects import SharedSkillStore
    from state.workspace import WorkspaceState

    monkeypatch.setenv('PATH', f'{Path(sys.executable).parent}:{os.environ.get("PATH", "")}')
    monkeypatch.setenv('MURAL_STUDIO_DATA_DIR', pinned_font_sources['studio_data_root'])
    state = WorkspaceState(tmp_path)
    _pin_font_manifest(tmp_path, pinned_font_sources)
    _write_author_inputs(tmp_path)
    css_path = tmp_path / 'base.css'
    css_path.write_text(css_path.read_text().replace('font-family:var(--font-cover)', 'font-family:' + authored_stack)
                       + ':root{--Cover-Body:"Missing Primary",var(--font-cover);--cover-body:"Noto Serif SC"}')
    skill = Path(__file__).resolve().parents[3] / 'skills/mural-presenter-v1'
    SharedSkillStore.for_workspace(tmp_path).pin(state.paths, skill, skill_name=skill.name)
    facts, pixels = {}, {}
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=browser_environment['PPT_SKILL_BROWSER_EXE'],
            env=browser_environment, headless=True, args=['--no-sandbox', '--allow-file-access-from-files'])
        try:
            page = browser.new_page()
            for mode in ('plans', 'authoring', 'portable'):
                record = (font_bundle.prepare_authoring_fonts(tmp_path, from_plans=mode == 'plans') if mode != 'portable'
                          else font_bundle.bundle_workspace(tmp_path))
                assert '--font-display-weight' not in record['token_stacks']
                assert record['token_stacks'].get('--font-cover', record['token_stacks']['--font-display']) == ['Inter', 'Noto Sans SC']
                page.goto(state.paths.page_html(1).as_uri())
                page.evaluate('() => document.fonts.ready')
                facts[mode] = _weight(page)
                pixels[mode] = page.locator('h1').screenshot()
                assert facts[mode]['weight'] == '700'
                assert facts[mode]['size'] == '64px'
                if mode != 'portable':
                    assert not state.paths.runtime_fonts.exists()
                else:
                    assert not font_bundle.validate_font_bundle(tmp_path)
                    for family in ('Inter', 'Noto Sans SC'):
                        face = next(row for row in record['faces'] if row['source_family'] == family)
                        assert font_bundle._variable_weight_range(str(tmp_path / face['path'])) == (100.0, 900.0)
        finally:
            browser.close()
    assert pixels['plans'] == pixels['authoring'] == pixels['portable']
    (tmp_path / 'observations.json').write_text(json.dumps(facts, ensure_ascii=False, indent=2))


def test_literal_only_fonts_do_not_require_invented_role_tokens(tmp_path, monkeypatch, pinned_font_sources):
    from state.workspace import WorkspaceState
    monkeypatch.setenv('PATH', f'{Path(sys.executable).parent}:{os.environ.get("PATH", "")}')
    WorkspaceState(tmp_path)
    _pin_font_manifest(tmp_path, pinned_font_sources)
    (tmp_path / 'base.css').write_text('h1{font-family:"Inter", "Noto Sans SC";font-weight:700}')
    (tmp_path / 'slides/slide_01.html').write_text('<h1>Literal font 中文</h1>')
    authored = font_bundle.prepare_authoring_fonts(tmp_path)
    portable = font_bundle.bundle_workspace(tmp_path)
    assert authored['token_stacks'] == portable['token_stacks'] == {}
    assert {'Inter', 'Noto Sans SC'} <= portable['delivery_families'].keys()
    assert not font_bundle.validate_font_bundle(tmp_path)


def test_real_pdf_editable_pptx_and_zip_preserve_cover_weight(
    tmp_path, monkeypatch, pinned_font_sources, browser_environment,
):
    import fitz
    from PIL import Image, ImageChops, ImageStat
    from playwright.sync_api import sync_playwright
    from pptx import Presentation
    from export.delivery import DeliveryExporter, export_on_demand
    from qc.routing import IssueLedger
    from runtime.shared_objects import SharedSkillStore
    from state.artifacts import ArtifactStore
    from state.workspace import WorkspaceState

    monkeypatch.setenv('PATH', f'{Path(sys.executable).parent}:{os.environ.get("PATH", "")}')
    monkeypatch.setenv('MURAL_STUDIO_DATA_DIR', pinned_font_sources['studio_data_root'])
    monkeypatch.setenv('MURAL_RESOURCE_SLOTS_DIR', str(tmp_path / 'render-slots'))
    root = tmp_path / 'deck'
    state = WorkspaceState(root)
    (root / 'config.json').write_text(json.dumps({
        'working_language': 'en', 'presentation_language': 'en', 'target_page_count': 1,
    }))
    _pin_font_manifest(root, pinned_font_sources)
    _write_author_inputs(root)
    skill = Path(__file__).resolve().parents[3] / 'skills/mural-presenter-v1'
    SharedSkillStore.for_workspace(root).pin(state.paths, skill, skill_name=skill.name)
    state.register_pages([1])
    state.register_scaffold(1)
    font_bundle.prepare_authoring_fonts(root)
    store = ArtifactStore(root, workspace_state=state)
    store.declare(artifact_id='cover', kind='page_html', path=state.paths.page_html(1), owner_id='test-owner', page=1)
    store.commit(owner_id='test-owner', artifact_ids=['cover'], source_hashes={})

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=browser_environment['PPT_SKILL_BROWSER_EXE'],
            env=browser_environment, headless=True, args=['--no-sandbox', '--allow-file-access-from-files'])
        try:
            page = browser.new_page(viewport={'width':1600, 'height':1200}, device_scale_factor=1)
            def render(html, png, canvas):
                from urllib.parse import quote, unquote, urlsplit
                loaded = []
                def observed(response):
                    path = Path(unquote(urlsplit(response.url).path)).resolve()
                    if path == html.resolve():
                        return
                    loaded.append(quote(path.relative_to(root).as_posix(), safe='/')
                                  if path.is_relative_to(root) else response.url)
                page.on('response', observed)
                page.goto(html.as_uri())
                page.evaluate('() => document.fonts.ready')
                assert _weight(page)['weight'] == '700'
                page.screenshot(path=str(png))
                return {'geometry':canvas, 'loaded_resources':loaded, 'failed_resources':[]}
            store.render_pages([1], render)
        finally:
            browser.close()
    store.record_inspection(1, inspection_id='offline-fixture-inspection')
    state.mark_ready(1)
    DeliveryExporter(state, store, IssueLedger(root)).export(include_pdf=False, include_pptx=False, include_zip=False)
    authored = {p: p.read_bytes() for p in (root/'base.css', state.paths.page_html(1), state.paths.page_png(1))}
    observations = {}
    for format in ('pdf', 'pptx-editable', 'zip'):
        result = export_on_demand(root, format)
        target = root / result['path']
        observations[format] = {'path': str(target), 'bytes': target.stat().st_size}
        assert all(path.read_bytes() == data for path, data in authored.items())
        if format == 'pdf':
            with fitz.open(target) as document:
                assert len(document) == 1
                assert document[0].rect.width / document[0].rect.height == pytest.approx(4/3)
                assert 'Bold' in document[0].get_text()
                raster = document[0].get_pixmap(matrix=fitz.Matrix(96/72, 96/72), alpha=False)
                actual = Image.frombytes('RGB', [raster.width, raster.height], raster.samples)
                expected = Image.open(state.paths.page_png(1)).convert('RGB')
                bbox = (70, 90, 650, 205)
                error = sum(ImageStat.Stat(ImageChops.difference(actual.crop(bbox), expected.crop(bbox))).mean)/3
                observations['pdf']['cover_mean_pixel_error'] = error
                assert error < 8
        elif format == 'pptx-editable':
            deck = Presentation(target)
            runs = [run for shape in deck.slides[0].shapes if shape.has_text_frame
                    for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if 'Bold' in run.text]
            assert runs and all(run.font.bold for run in runs)
            observations[format]['cover_bold_runs'] = len(runs)
        else:
            extraction = tmp_path / 'zip'
            with zipfile.ZipFile(target) as archive:
                assert 'base.css' in archive.namelist()
                css = archive.read('base.css').decode()
                assert '--font-display-weight:700' in css
                assert 'file://' not in css
                archive.extractall(extraction)
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(executable_path=browser_environment['PPT_SKILL_BROWSER_EXE'],
                    env=browser_environment, headless=True, args=['--no-sandbox', '--allow-file-access-from-files'])
                try:
                    page = browser.new_page()
                    page.goto((extraction / 'slides/slide_01.html').as_uri())
                    page.evaluate('() => document.fonts.ready')
                    observations[format]['css'] = _weight(page)
                    assert observations[format]['css']['weight'] == '700'
                finally:
                    browser.close()
    (tmp_path / 'export-observations.json').write_text(json.dumps(observations, ensure_ascii=False, indent=2))
