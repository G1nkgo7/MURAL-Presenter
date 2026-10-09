"""Office 原生图表、受控公式与同一解析回执的回归；不调用真实 MinerU 或模型。"""
from concurrent.futures import ThreadPoolExecutor
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import shutil
import threading
import zipfile

from PIL import Image
import pytest

from runtime import attachments, mineru, office
from test_attachment_parse import api, corpus, stage


@pytest.mark.parametrize("formula", ["SUM(B2:B3)", "ROUND(A1/2,2)", "IF(A1>0,\"yes\",\"no\")", "'Sheet 2'!A1+1"])
def test_controlled_formula_subset_accepts_local_arithmetic(formula):
    assert office._formula_allowed(formula)


@pytest.mark.parametrize("formula", ["WEBSERVICE(\"http://private.invalid\")", "DDE(\"cmd\",\"/c\",\"calc\")",
                                   "INDIRECT(A1)", "[1]Sheet1!A1", "MY_EXTERNAL_NAME", "NOW()"])
def test_external_or_uncontrolled_formulas_are_not_executed(formula):
    assert not office._formula_allowed(formula)


@pytest.mark.parametrize("part,content,reason", [
    ("xl/vbaProject.bin", b"macro", "office_active_content_disabled"),
    ("xl/_rels/external.rels", b'<Relationships><Relationship Id="r1" TargetMode="External" Target="https://private.invalid" /></Relationships>', "office_external_links_disabled"),
    ("word/fields.xml", b'<root><instrText>DDEAUTO private command</instrText></root>', "office_external_fields_disabled"),
    ("xl/_rels/active.rels", b'<Relationships><Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="unnamed.bin" /></Relationships>', "office_active_content_disabled"),
])
def test_active_office_input_is_rejected_before_native_converter(tmp_path, corpus, monkeypatch, part, content, reason):
    source = tmp_path / "blocked.xlsx"
    with zipfile.ZipFile(corpus["xlsx"]) as original, zipfile.ZipFile(source, "w") as output:
        for item in original.infolist():
            output.writestr(item, original.read(item))
        output.writestr(part, content)
    monkeypatch.setattr(office.subprocess, "Popen", lambda *_a, **_kw: pytest.fail("active document was opened"))
    result = office.derive(source, tmp_path / "temporary", mineru.source_info(source))
    assert result["unresolved_items"] == [{"reason": reason}] and not result["files"]


def test_native_xlsx_charts_and_explicit_calculation_preserve_original(tmp_path, corpus):
    if not shutil.which("libreoffice"):
        pytest.skip("requires installed LibreOffice")
    source = corpus["xlsx"]
    original = source.read_bytes()
    info = mineru.source_info(source)
    assert info["native_charts"] == 2 and info["formulas"][0]["expression"] == "=SUM(B2:B3)"
    result = office.derive(source, tmp_path / "temporary", info)
    assert result["unresolved_items"] == []
    assert [item["worksheet"] for item in result["charts"]] == ["Sales", "Summary"]
    for chart in result["charts"]:
        assert "page" not in chart and chart["derivation"] == "libreoffice_native_chart"
        with Image.open(BytesIO(result["files"][chart["img_path"]])) as image:
            assert image.format == "PNG" and min(image.size) > 100
            assert image.convert("RGB").getextrema() != ((255, 255),) * 3
    for item in result["formulas"]:
        assert item["expression"] == "=SUM(B2:B3)" and item["has_cached_value"] is False
        assert item["calculated_value"] == pytest.approx(802.35)
        assert item["calculation"] == "libreoffice_calc_explicit"
    assert source.read_bytes() == original and not list((tmp_path / "temporary").iterdir())


def test_native_pptx_chart_is_bound_to_slide_not_pdf_page(tmp_path):
    if not shutil.which("libreoffice"):
        pytest.skip("requires installed LibreOffice")
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Alpha", "Beta"]
    data.add_series("Revenue", [100, 200])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data)
    source = tmp_path / "native.pptx"
    deck.save(source)
    before = source.read_bytes()
    result = office.derive(source, tmp_path / "temporary", mineru.source_info(source))
    assert result["unresolved_items"] == [] and len(result["charts"]) == 1
    assert result["charts"][0]["slide"] == 2 and "page" not in result["charts"][0]
    assert source.read_bytes() == before


def test_native_formula_string_and_error_do_not_become_numeric_zero(tmp_path):
    if not shutil.which("libreoffice"):
        pytest.skip("requires installed LibreOffice")
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = "Values"
    book.active["A1"] = '=IF(1>0,"yes","no")'
    book.active["A2"] = "=1/0"
    source = tmp_path / "results.xlsx"
    book.save(source)
    before = source.read_bytes()
    result = office.derive(source, tmp_path / "temporary", mineru.source_info(source))
    string, error = result["formulas"]
    assert string["calculated_value"] == "yes" and string["calculated_text"] == "yes"
    assert error["calculation_error"] > 0 and "calculated_value" not in error
    assert result["unresolved_items"] == [{"reason": "office_formula_calculation_error",
        "worksheet": "Values", "cell": "A2", "error": error["calculation_error"]}]
    assert source.read_bytes() == before


def test_native_docx_chart_retains_original_document_without_fake_physical_page(tmp_path):
    if not shutil.which("libreoffice"):
        pytest.skip("requires installed LibreOffice")
    from docx import Document
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Alpha", "Beta"]
    data.add_series("Revenue", [100, 200])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data)
    presentation = BytesIO()
    deck.save(presentation)
    document = Document()
    document.add_paragraph("Logical introduction")
    paragraph = document.add_paragraph()
    paragraph.add_run()._r.append(etree.fromstring(b'''<w:drawing
      xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
      xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
      xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <wp:inline><wp:extent cx="4572000" cy="2743200"/><wp:docPr id="10" name="Revenue chart"/>
      <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">
      <c:chart r:id="rIdChart"/></a:graphicData></a:graphic></wp:inline></w:drawing>'''))
    template = BytesIO()
    document.save(template)
    source = tmp_path / "native.docx"
    with zipfile.ZipFile(template) as original, zipfile.ZipFile(presentation) as charts, zipfile.ZipFile(source, "w") as output:
        for item in original.infolist():
            raw = original.read(item)
            if item.filename == "word/_rels/document.xml.rels":
                tree = etree.fromstring(raw)
                etree.SubElement(tree, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship",
                    Id="rIdChart", Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart",
                    Target="charts/chart1.xml")
                raw = etree.tostring(tree)
            elif item.filename == "[Content_Types].xml":
                tree = etree.fromstring(raw)
                namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
                etree.SubElement(tree, "{%s}Override" % namespace, PartName="/word/charts/chart1.xml",
                    ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml")
                etree.SubElement(tree, "{%s}Default" % namespace, Extension="xlsx",
                    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                raw = etree.tostring(tree)
            output.writestr(item, raw)
        for name in charts.namelist():
            if name.startswith(("ppt/charts/", "ppt/embeddings/")):
                output.writestr("word/" + name.removeprefix("ppt/"), charts.read(name))
    before = source.read_bytes()
    info = mineru.source_info(source)
    assert info["native_charts"] == 1 and info["granularity"] == "document_block"
    result = office.derive(source, tmp_path / "temporary", info)
    assert result["unresolved_items"] == [] and len(result["charts"]) == 1
    chart = result["charts"][0]
    assert chart["derivation"] == "libreoffice_native_chart" and "page" not in chart
    assert chart["document_object"]
    raw = result["files"][chart["img_path"]]
    with Image.open(BytesIO(raw)) as image:
        assert min(image.size) > 100 and image.convert("RGB").getextrema() != ((255, 255),) * 3
    (tmp_path / "native-chart.png").write_bytes(raw)
    assert source.read_bytes() == before


def test_shared_formula_preserves_master_and_translates_local_cell_expression(tmp_path, corpus):
    from xml.etree import ElementTree as ET
    source = tmp_path / "shared.xlsx"
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(corpus["xlsx"]) as original, zipfile.ZipFile(source, "w") as output:
        for item in original.infolist():
            raw = original.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                tree = ET.fromstring(raw)
                master = tree.find('.//{*}c[@r="B4"]/{*}f')
                master.text = "B2+1"
                master.set("t", "shared")
                master.set("si", "0")
                master.set("ref", "B4:B5")
                row = ET.SubElement(tree.find("{*}sheetData"), "{%s}row" % namespace, r="5")
                cell = ET.SubElement(row, "{%s}c" % namespace, r="B5")
                ET.SubElement(cell, "{%s}f" % namespace, t="shared", si="0")
                raw = ET.tostring(tree)
            output.writestr(item, raw)
    shared = next(item for item in mineru.source_info(source)["formulas"] if item["cell"] == "B5")
    assert shared["expression"] == "=B3+1"
    assert shared["shared_formula_source"] == {"cell": "B4", "expression": "=B2+1"}
    assert shared["formula_attributes"] == {"t": "shared", "si": "0"}


def test_derivatives_share_existing_receipt_hash_cache_and_keep_quality_partial(tmp_path, corpus, api):
    api[0]["derive_office"] = True
    root, record = stage(tmp_path, corpus["xlsx"])
    receipt = attachments.parse_attachment(root, record.attachment_id)
    assert receipt["status"] == "partial" and not receipt["failover"]
    assert len(receipt["office"]["charts"]) == 2
    assert receipt["office"]["source_sha256"] == record.sha256
    manifest = json.loads((root / "_trace/attachment-manifest.json").read_text())
    parse = manifest["records"][0]["parse"]
    for chart in receipt["office"]["charts"]:
        assert chart["img_path"] in receipt["media_paths"]
        assert parse["results"][parse["current"]]["outputs"][chart["img_path"]] == hashlib.sha256((root / chart["img_path"]).read_bytes()).hexdigest()
    resumed = attachments.parse_attachment(root, record.attachment_id)
    assert resumed["cached"] and resumed["office"] == receipt["office"]
    assert api[0]["posts"] == 1 and api[1]["gets"] == 0


def test_office_derivation_version_participates_in_existing_cache_identity(tmp_path, corpus, api, monkeypatch):
    root, record = stage(tmp_path, corpus["xlsx"])
    first = attachments.parse_attachment(root, record.attachment_id)
    monkeypatch.setattr(office, "DERIVATION_VERSION", office.DERIVATION_VERSION + "-test")
    second = attachments.parse_attachment(root, record.attachment_id)
    assert not second["cached"] and first["text_path"] != second["text_path"]
    assert api[0]["posts"] == 2 and api[1]["posts"] == 0


def test_explicit_background_environment_does_not_mutate_process_environment(tmp_path, corpus, api, monkeypatch):
    root, record = stage(tmp_path, corpus["pdf"])
    configured = os.environ["MURAL_RUNTIME_CONFIG"]
    monkeypatch.setenv("MURAL_RUNTIME_CONFIG", str(tmp_path / "must-not-read.json"))
    before = dict(os.environ)
    receipt = attachments.parse_attachment(root, record.attachment_id,
        environment={"MURAL_RUNTIME_CONFIG": configured})
    assert receipt["status"] == "completed" and dict(os.environ) == before


def test_queue_state_is_same_manifest_and_failed_preflight_has_existing_receipt(tmp_path, corpus, api, monkeypatch):
    root, record = stage(tmp_path, corpus["pdf"])
    attachments.queue_attachment_parse(root, record.attachment_id)
    manifest_path = root / "_trace/attachment-manifest.json"
    assert json.loads(manifest_path.read_text())["records"][0]["parse_status"] == "queued"
    assert api[0]["gets"] == 0
    def invalid(_source):
        raise mineru.MinerUError("mineru_invalid_pdf")
    monkeypatch.setattr(mineru, "source_info", invalid)
    receipt = attachments.parse_attachment(root, record.attachment_id)
    detail = json.loads(manifest_path.read_text())["records"][0]
    assert receipt["status"] == detail["parse_status"] == "failed"
    assert detail["parse"]["results"][detail["parse"]["current"]]["receipt"] == receipt
    assert api[0]["gets"] == api[1]["gets"] == 0


def test_duplicate_background_threads_serialize_one_attachment_and_emit_final_state(tmp_path, corpus, api):
    root, record = stage(tmp_path, corpus["pdf"])
    api[0].update(entered=threading.Event(), hold=threading.Event())
    first_events, second_events = [], []
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(attachments.parse_attachment, root, record.attachment_id, event_sink=first_events.append)
        assert api[0]["entered"].wait(10)
        attachments.queue_attachment_parse(root, record.attachment_id)
        second = pool.submit(attachments.parse_attachment, root, record.attachment_id, event_sink=second_events.append)
        try:
            detail = json.loads((root / "_trace/attachment-manifest.json").read_text())["records"][0]
            assert detail["parse_status"] == "processing"
        finally:
            api[0]["hold"].set()
        one, two = first.result(10), second.result(10)
    assert not one["cached"] and two["cached"] and one["text_path"] == two["text_path"]
    assert api[0]["posts"] == 1
    for events in (first_events, second_events):
        assert [row["event_type"] for row in events if row["event_type"].startswith("attachment_")] == [
            "attachment_parse_queued", "attachment_parse_started", "attachment_parse_finished"]
    assert json.loads((root / "_trace/attachment-manifest.json").read_text())["records"][0]["parse_status"] == "completed"
