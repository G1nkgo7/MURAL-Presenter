from __future__ import annotations

from contextlib import contextmanager
import hashlib
import io
import json
import os
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import quote

import pytest
from pptx import Presentation

from canary import real_webui_canary as canary
from runtime.shared_objects import SharedSkillStore
from state.workspace import WorkspaceState


@pytest.mark.parametrize("source", ["dom", "png"])
def test_preview_geometry_uses_actual_author_and_display_dimensions(source):
    preview = {"source": source, "width_px": 1600, "height_px": 900,
               "display_width": 800, "display_height": 450}
    assert canary.validate_preview_geometry(preview, 1600, 900) == preview


@pytest.mark.parametrize("field,value", [
    ("width_px", 1920), ("height_px", 1200), ("display_width", 0),
    ("display_height", 0), ("display_height", 600),
])
def test_preview_geometry_rejects_wrong_or_hidden_canvas(field, value):
    preview = {"source": "dom", "width_px": 1600, "height_px": 900,
               "display_width": 800, "display_height": 450, field: value}
    with pytest.raises(AssertionError):
        canary.validate_preview_geometry(preview, 1600, 900)


@pytest.fixture
def environment(monkeypatch):
    monkeypatch.setattr(os, "environ", {})
    return os.environ


@pytest.mark.parametrize("explicit,configured,expected", [
    ("/selected/alpha", "/configured/v1", "/selected/alpha"),
    (None, "/configured/v1", "/configured/v1"),
    (None, None, str(canary.ROOT / "skills/mural-presenter-v1")),
])
def test_concrete_skill_entry_overrides_dotenv_and_ignores_legacy(environment, monkeypatch, tmp_path,
                                                               explicit, configured, expected):
    legacy = ("MURAL_SKILLS_DIR", "PPTAGENT_MURAL_V10_SKILLS_ROOT", "PPTAGENT_MURAL_V06_SKILLS_ROOT")
    environment.update({key: "/old/parent" for key in legacy})
    if explicit:
        environment["MURAL_SKILL_DIR"] = explicit
    values = {"MURAL_SKILL_DIR": configured} if configured else {}
    monkeypatch.setattr(canary, "read_environment_file", lambda _: values)
    assert canary.configure_skill_environment(tmp_path) == Path(expected)
    assert environment["MURAL_SKILL_DIR"] == expected
    assert not any(key in environment for key in legacy)


def test_resume_identity_uses_verified_pin_not_new_missing_source(environment, monkeypatch, tmp_path):
    source = tmp_path / "original"
    source.mkdir()
    (source / "SKILL.md").write_text("---\nname: original\n---\nPinned instructions.\n")
    workspace = WorkspaceState(tmp_path / "workspace")
    paths = workspace.paths
    store = SharedSkillStore(tmp_path / "studio-data")
    original = store.pin(paths, source, skill_name="original")
    before = paths.skill_ref.read_bytes()
    environment.update(MURAL_SKILL_DIR=str(tmp_path / "missing-new-source"), MURAL_STUDIO_DATA_DIR=str(store.studio_data))
    monkeypatch.setattr(canary, "read_environment_file", lambda _: {})
    canary.configure_skill_environment(tmp_path)
    identity = canary.pinned_skill_identity(paths.root)
    assert identity == {"skill_name": "original", "tree_sha256": original.tree_sha256,
                        "object_root": str(original.object_root)}
    assert paths.skill_ref.read_bytes() == before


@pytest.mark.parametrize("selection", ["pptx", "pptx-editable"])
def test_parser_keeps_no_thinking_and_one_explicit_pptx_format(selection):
    args = canary.argument_parser().parse_args(["--production", str(canary.ROOT), "--output", "new-canary",
                                              "--case", "basic", "--pages", "3", "--model", "opus-5",
                                              "--pptx-format", selection])
    assert (args.pages, args.model, args.thinking, args.pptx_format) == (3, "opus-5", False, selection)
    assert canary.argument_parser().parse_args(["--production", str(canary.ROOT), "--output", "new"]).pptx_format == "pptx"


def test_parser_accepts_existing_attachments_for_synthesis_case():
    args = canary.argument_parser().parse_args(["--production", str(canary.ROOT), "--output", "new",
        "--case", "critic", "--pages", "3", "--attachment", "guide.pdf", "--attachment", "chart.png"])
    assert args.attachment == [Path("guide.pdf"), Path("chart.png")]
    assert args.case == "critic" and args.pages == 3 and args.thinking is False


@pytest.mark.parametrize("damage", [None, "missing", "duplicate", "empty_messages", "wrong_owner"])
def test_synthesis_canary_accepts_raw_segments_without_online_training_admission(tmp_path, damage):
    root = tmp_path / "_trace/raw/segments"
    root.mkdir(parents=True)
    segment = {"owner_id": "orchestrator", "segment_id": "segment-a",
               "agent_context": [{"system": "actual input"}], "messages": [{"role": "assistant", "content": "done"}]}
    events = [{"agent_id": "orchestrator", "segment_id": "segment-a", "event_type": "agent_context"}]
    if damage == "empty_messages":
        segment["messages"] = []
    if damage == "wrong_owner":
        segment["owner_id"] = "another-owner"
    if damage != "missing":
        (root / "segment.json").write_text(json.dumps(segment))
    if damage == "duplicate":
        (root / "duplicate.json").write_text(json.dumps(segment))
    assert canary.raw_segments_complete(tmp_path, events) is (damage is None)
    assert not (tmp_path / "_trace/training/training-manifest.json").exists()


def editable_fixture(tmp_path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = 1600 * 9525, 900 * 9525
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    text = slide.shapes.add_textbox(100, 100, 500000, 500000)
    text.name, text.text = "text: h1", "One conclusion"
    buffer = io.BytesIO()
    deck.save(buffer)
    content = buffer.getvalue()
    delivery = tmp_path / "delivery"
    delivery.mkdir()
    (delivery / "presentation.editable.pptx").write_bytes(content)
    summary = {"native_text": 1, "native_images": 0, "native_shapes": 0, "native_tables": 0,
               "local_raster_fallbacks": 0, "pages_with_warnings": 0}
    report = {"summary": summary, "pptx_sha256": hashlib.sha256(content).hexdigest()}
    (delivery / "presentation.editable.editability.json").write_text(json.dumps(report))
    headers = {"X-Mural-Editability-Report": quote("delivery/presentation.editable.editability.json", safe=""),
               "X-Mural-Editable-Text": "1", "X-Mural-Raster-Fallbacks": "0"}
    return content, headers


def test_editable_validates_native_ooxml_and_corresponding_report(tmp_path):
    content, headers = editable_fixture(tmp_path)
    record = canary.validate_export("pptx-editable", content, headers, tmp_path,
                                    {"width_px": 1600, "height_px": 900}, 1)
    assert record["summary"]["native_text"] == 1
    assert record["object_counts"] == {"text": 1}
    assert not (tmp_path / "delivery/presentation.pptx").exists()


@pytest.mark.parametrize("damage", ["hash", "header", "report_path", "object_count", "bytes", "geometry", "pages"])
def test_editable_rejects_misleading_receipt_or_wrong_download(tmp_path, damage):
    content, headers = editable_fixture(tmp_path)
    report_path = tmp_path / "delivery/presentation.editable.editability.json"
    report = json.loads(report_path.read_text())
    geometry, pages = {"width_px": 1600, "height_px": 900}, 1
    if damage == "hash":
        report["pptx_sha256"] = "wrong"
    elif damage == "header":
        headers["X-Mural-Editable-Text"] = "9"
    elif damage == "report_path":
        headers["X-Mural-Editability-Report"] = "../elsewhere.json"
    elif damage == "object_count":
        report["summary"]["native_text"] = 9
        headers["X-Mural-Editable-Text"] = "9"
    elif damage == "bytes":
        content += b"wrong"
    elif damage == "geometry":
        geometry["width_px"] = 1920
    else:
        pages = 3
    report_path.write_text(json.dumps(report))
    with pytest.raises(AssertionError):
        canary.validate_export("pptx-editable", content, headers, tmp_path, geometry, pages)


@pytest.mark.parametrize("kind", ["pptx", "pptx-editable"])
def test_ui_download_selects_only_requested_format_once(tmp_path, kind):
    actions = []
    content = b"downloaded-pptx"
    response = SimpleNamespace(url=f"http://private/api/decks/7/export/{kind}",
                               request=SimpleNamespace(method="POST"), ok=True, status=200)
    download = SimpleNamespace(failure=lambda: None, save_as=lambda path: path.write_bytes(content))

    @contextmanager
    def expected_response(predicate, **_):
        assert predicate(response)
        yield SimpleNamespace(value=response)

    @contextmanager
    def expected_download(**_):
        yield SimpleNamespace(value=download)

    def locator(selector):
        return SimpleNamespace(click=lambda: actions.append(selector), wait_for=lambda **_: actions.append(selector))

    page = SimpleNamespace(expect_response=expected_response, expect_download=expected_download, locator=locator)
    actual, payload = canary.download_ui_pptx(page, "http://private", 7, tmp_path, kind)
    assert actual is response and payload == content
    assert [action for action in actions if "data-export-format" in action] == [f'[data-export-format="{kind}"]']
    assert len(list(tmp_path.glob("ui-*.pptx"))) == 1
    assert any("file?rel=" in action for action in actions) == (kind == "pptx-editable")


def test_ui_export_http_failure_does_not_wait_for_nonexistent_download(tmp_path):
    actions = []
    response = SimpleNamespace(url="http://private/api/decks/7/export/pptx-editable",
        request=SimpleNamespace(method="POST"), ok=False, status=409,
        text=lambda: '{"detail":"unsupported visible effect"}')

    @contextmanager
    def expected_response(predicate, **_):
        assert predicate(response)
        yield SimpleNamespace(value=response)

    @contextmanager
    def expected_download(**_):
        yield SimpleNamespace()
        actions.append("waited-for-download")

    page = SimpleNamespace(expect_response=expected_response, expect_download=expected_download,
        locator=lambda _: SimpleNamespace(click=lambda: actions.append("clicked")))
    with pytest.raises(AssertionError, match="409.*unsupported visible effect"):
        canary.download_ui_pptx(page, "http://private", 7, tmp_path, "pptx-editable")
    assert actions == ["clicked", "clicked"]
