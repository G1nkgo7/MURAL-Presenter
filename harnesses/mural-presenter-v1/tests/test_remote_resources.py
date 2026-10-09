"""公网依赖允许加载；读取边界、真实字节与字体设计不因放开网络而丢失。"""
import hashlib
import importlib
from io import BytesIO
import json
import time
from urllib.parse import urlsplit

from PIL import Image
import pytest

from test_v1_font_pipeline import font_bundle
from test_editable_pptx import browser_environment

remote = importlib.import_module(font_bundle.__package__ + ".remote_resources")
core = importlib.import_module(font_bundle.__package__ + ".deck_core")


@pytest.mark.parametrize("scheme", ["http", "https"])
def test_remote_capture_hash_is_not_a_local_file_identity(tmp_path, scheme):
    css = tmp_path / "base.css"
    css.write_text("h1{font-weight:700}")
    hashes = {f"{scheme}://cdn.example/font.woff2": hashlib.sha256(b"captured font").hexdigest(),
              "base.css": hashlib.sha256(css.read_bytes()).hexdigest()}
    before = dict(hashes)
    core._validate_render_inputs(tmp_path, hashes)
    assert hashes == before
    css.write_text("h1{font-weight:400}")
    with pytest.raises(ValueError, match="render input changed after capture: base.css"):
        core._validate_render_inputs(tmp_path, hashes)


def test_remote_capture_does_not_bypass_fixed_font_validation(tmp_path, monkeypatch):
    expected = hashlib.sha256(b"fixed font").hexdigest()
    reference = "file:///fixed-font-store/font.woff2"
    checks = []
    def resolve(root, value):
        checks.append((root, value))
        return None, expected if len(checks) == 1 else "changed"
    monkeypatch.setattr(core, "resolve_shared_font_reference", resolve)
    hashes = {"https://cdn.example/remote.woff2": "observed", reference: expected}
    core._validate_render_inputs(tmp_path, hashes)
    with pytest.raises(ValueError, match="render input changed after capture: file:"):
        core._validate_render_inputs(tmp_path, hashes)
    assert checks == [(tmp_path, reference), (tmp_path, reference)]


@pytest.mark.parametrize("url", ["file:///etc/passwd", "https://user:secret@example.com/font", "javascript:alert(1)"])
def test_dependency_url_requires_credential_free_http(url):
    with pytest.raises(ValueError, match="url_invalid"):
        remote.validate_url(url)


@pytest.mark.parametrize("address", ["127.0.0.1", "10.0.0.1", "169.254.169.254", "100.64.0.1", "::1"])
def test_dependencies_cannot_access_private_hosts(monkeypatch, address):
    monkeypatch.setattr(remote.socket, "getaddrinfo", lambda *_args, **_kwargs: [(2, 1, 6, "", (address, 443))])
    with pytest.raises(ValueError, match="scope_denied"):
        remote.validate_url("https://cdn.example/redirect-target")


def test_material_grounded_policy_precedes_network(tmp_path, monkeypatch):
    (tmp_path / "_state").mkdir()
    (tmp_path / "_state/run.json").write_text(json.dumps({"authoring_profile": "material-grounded"}))
    monkeypatch.setattr(remote, "validate_url", lambda *_: pytest.fail("must not resolve a denied dependency"))
    with pytest.raises(ValueError, match="external_source_forbidden"):
        remote.fetch("https://example.com/font", root=tmp_path)


def test_fetch_preserves_bytes_excludes_auth_and_connects_validated_address(monkeypatch):
    import gzip
    from types import SimpleNamespace
    original = b"original EMPTY body token bytes\x00\xff"
    captured = {}
    monkeypatch.setattr(remote, "validate_url", lambda url: (urlsplit(url), "93.184.216.34", 80))
    monkeypatch.setattr(remote.socket, "create_connection", lambda address, **kw: captured.update(address=address) or object())
    def request(connection, method, target, headers):
        connection.connect()
        captured.update(method=method, target=target, headers=headers)
    monkeypatch.setattr(remote.http.client.HTTPConnection, "request", request)
    monkeypatch.setattr(remote.http.client.HTTPConnection, "getresponse", lambda *_: SimpleNamespace(
        status=200, read=lambda: gzip.compress(original), getheaders=lambda: [
            ("Content-Encoding", "gzip"), ("Set-Cookie", "private=value"), ("Content-Length", "0")]))
    monkeypatch.setattr(remote.http.client.HTTPConnection, "close", lambda *_: None)
    value = remote.fetch("http://cdn.example?q=1")
    assert captured["address"] == ("93.184.216.34", 80)
    assert captured["method"] == "GET" and captured["target"] == "/?q=1"
    assert not {"authorization", "cookie", "x-api-key"} & {key.lower() for key in captured["headers"]}
    assert value["body"] == original and value["sha256"] == hashlib.sha256(original).hexdigest()
    assert "set-cookie" not in value["headers"] and "content-encoding" not in value["headers"]


def test_remote_font_definition_keeps_authored_family_weight_and_stack(monkeypatch):
    seen = []
    def fetch(url, **kwargs):
        seen.append(url)
        return {"status": 200, "headers": {}, "body": b'@font-face{font-family:"Actual Font";font-weight:100 900;src:url(face.woff2)}'}
    monkeypatch.setattr(remote, "fetch", fetch)
    css = '@import url("https://cdn.example/fonts.css");:root{--font-title:"Actual Font",sans-serif;--font-title-weight:700}'
    authored = font_bundle._authored_font_families(css, [])
    assert authored == {"Actual Font"} and seen == ["https://cdn.example/fonts.css"]
    resolved = font_bundle.resolve_font_request(role="title", token="--font-title",
        requested_stack=["Actual Font", "sans-serif"], weight="700", authored_families=authored,
        style="normal", custom_registry={}, route_default_stack=[], source_available=lambda *_: False, catalog={})
    assert resolved.resolved_family == "Actual Font" and resolved.resolution_kind == "authored_css"
    assert resolved.weight_resolved == "700"


def test_remote_bytes_work_in_pdf_editable_and_html(tmp_path, browser_environment):
    from export import native_pdf, editable
    from geometry_fixture import geometry
    from pptx import Presentation
    import pymupdf
    png = BytesIO()
    Image.new("RGB", (30, 20), "#bd7937").save(png, format="PNG")
    requested = []
    def fetch(url):
        requested.append(url)
        assert url == "https://cdn.example/image.png"
        return {"status": 200, "headers": {"content-type": "image/png", "access-control-allow-origin": "*"}, "body": png.getvalue()}
    canvas = geometry(width_px=1600, height_px=900)
    html = '<html><style>html,body{margin:0;width:1600px;height:900px;overflow:hidden}</style><body><h1>Remote media</h1><img src="https://cdn.example/image.png" width="300" height="200"></body></html>'
    native_pdf.write_native_pdf([(1, html)], tmp_path / "remote.pdf", canvas, scratch=tmp_path,
        deadline_monotonic=time.monotonic()+60, remote_fetch=fetch)
    with pymupdf.open(tmp_path / "remote.pdf") as pdf:
        assert len(pdf) == 1 and pdf[0].get_images()
    editable.write_editable_pptx([(1, html)], tmp_path / "remote.pptx", canvas, scratch=tmp_path, remote_fetch=fetch)
    assert any(shape.shape_type == 13 for shape in Presentation(tmp_path / "remote.pptx").slides[0].shapes)
    native_pdf.validate_single_file_html(html, canvas, scratch=tmp_path,
        deadline_monotonic=time.monotonic()+60, page_numbers=[1], remote_fetch=fetch)
    assert len(requested) == 3
